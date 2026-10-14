import sys as _rb_sys
from application.dependencies import Bindings as _rb_bind

import io
import csv

from report_edit_control import (
    assert_role_allowed,
    assert_not_final,
    assert_lease_owner,
    assert_no_foreign_active_lease,
    acquire_or_renew_lease,
    release_lease,
    ReportEditControlError,
    ReportEditLocked,
    ReportRoleDenied,
    FinalReportLocked,
)

APP_VERSION = '5.46.9'
BUILD_NAME = 'fresh-schema-runtime-hardening'

from urllib.parse import quote, urlparse
"""RadBridge RIS/PACS main application.

This historical entry point composes owned models, policies, web components and
integration services. Independent instances are constructed through web.factory.
Before changing contracts, read AI_DEVELOPER_GUIDE.md and docs/ARCHITECTURE.md.

Safety invariants: preserve patient/accession/DICOM identity, server-side
permissions, branch scope, auditability, migration history and non-destructive
PACS routing semantics. New large features should prefer extracted service
modules while keeping compatibility with existing routes/templates.
"""

import json
import hmac
import shutil
from pypdf import PdfReader
import tempfile
import zipfile
from docx.shared import Inches

import os, smtplib, secrets, io, json, re, csv, tempfile, hashlib
from infrastructure.composition_environment import composition_os
os = composition_os()
from datetime import datetime, timedelta, date
from functools import wraps
from workflow_engine import (
    EXAM_FORWARD_STATES,
    JOURNEY_STAGES as CANONICAL_JOURNEY_STAGES,
    blockers_text,
    can_transition,
    canonical_display_state,
    derive_journey_stage,
    exam_transition_gate,
    TransitionDecision,
    normalize_state,
    order_status_for_exam_transition,
    report_source_allowed,
    resend_mwl_gate,
    work_order_gate,
)
from patient_demographics import resolve_patient_dob
import uuid
from flask import Flask, request, redirect, url_for, render_template, flash, send_file, jsonify, session, abort, Response
from infrastructure.extensions import db
from procedure_catalog import PROCEDURE_CATALOG
from sqlalchemy import or_, text, inspect, and_, func
from sqlalchemy.exc import IntegrityError
import requests
from dicom_node_control import (policy_ready as dicom_ingress_policy_ready,
                                node_payload, modality_key, delete_modality,
                                legacy_endpoint_payload, legacy_endpoint_matches)
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from werkzeug.security import generate_password_hash, check_password_hash
from itsdangerous import URLSafeTimedSerializer, BadSignature, SignatureExpired
from email.message import EmailMessage
import qrcode
from reportlab.graphics.barcode import code128
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from radbridge_ocr import (
    OCRProviderError,
    SUPPORTED_PROVIDERS,
    configuration_hash as ocr_configuration_hash,
    extract_file as ocr_extract_file,
    normalize_config as normalize_ocr_config,
    validate_local as validate_ocr_local,
    validate_remote as validate_ocr_remote,
)
from radbridge_intake import (
    classify_document_text,
    CORE_FIELDS as OCR_INTAKE_CORE_FIELDS,
    digits as intake_digits,
    normalize_identity as normalize_intake_identity,
    normalize_name as normalize_intake_name,
    normalize_phone as normalize_intake_phone,
    registration_fingerprint as intake_registration_fingerprint,
    score_patient_candidates as score_intake_patient_candidates,
    validate_identity as validate_intake_identity,
)

from infrastructure.settings import WebSettings
from web.factory import create_web_app

app = create_web_app(
    __name__,
    root_path=os.path.dirname(os.path.abspath(__file__)),
    settings=WebSettings.from_environment(os.environ),
)
app.extensions['radbridge_runtime'] = _rb_sys.modules[__name__]

REPORT_DIR = os.getenv('REPORT_DIR', '/app/report-files')
ATTACHMENT_DIR = os.getenv('ATTACHMENT_DIR', '/app/attachments')
os.makedirs(REPORT_DIR, exist_ok=True)
os.makedirs(ATTACHMENT_DIR, exist_ok=True)
db.init_app(app)


# RB_BIND:0000
from web.routes.shared import install_0000 as _rb_install_0000
_same_origin_unsafe_request_guard = _rb_install_0000(_rb_bind(_rb_sys.modules[__name__], ('RIS_PUBLIC_URL', 'abort', 'app', 'os', 'request', 'session')))

ORTHANC = os.getenv('ORTHANC_URL', 'http://ris-pacs:8042')
ORTHANC_USER = os.getenv('ORTHANC_USER', 'admin')
ORTHANC_PASSWORD = os.getenv('ORTHANC_PASSWORD', 'ChangeMe_PACS_2026')
ORTHANC_PUBLIC_URL = os.getenv('ORTHANC_PUBLIC_URL', 'https://ris.dynuddns.net/pacs/')
RIS_PUBLIC_URL = os.getenv('RIS_PUBLIC_URL', 'https://ris.dynuddns.net').rstrip('/')

RADBRIDGE_ENV = os.getenv('RADBRIDGE_ENV', 'development').strip().lower()


# RB_BIND:0001
from application.components.shared import install_0001 as _rb_install_0001
_validate_production_secrets = _rb_install_0001(_rb_bind(_rb_sys.modules[__name__], ('ORTHANC_PASSWORD', 'RADBRIDGE_ENV', 'app')))


_validate_production_secrets()

ERP_SYNC_ENABLED = os.getenv('ERP_SYNC_ENABLED','0') == '1'
ERP_BASE_URL = os.getenv('ERP_BASE_URL','').rstrip('/')
ERP_API_TOKEN = os.getenv('ERP_API_TOKEN','')
ERP_TIMEOUT = int(os.getenv('ERP_TIMEOUT','15') or 15)
ERP_PROVIDER = os.getenv('ERP_PROVIDER','GENERIC').upper()

PORTAL_REQUIRE_PHONE_VERIFY = os.getenv('PORTAL_REQUIRE_PHONE_VERIFY','1') == '1'
STICKER_DEFAULT_WIDTH_MM = float(os.getenv('STICKER_DEFAULT_WIDTH_MM','70') or 70)
STICKER_DEFAULT_HEIGHT_MM = float(os.getenv('STICKER_DEFAULT_HEIGHT_MM','45') or 45)


PACS_IMAGES_STABLE_SECONDS = int(os.getenv('PACS_IMAGES_STABLE_SECONDS','120') or 120)
PACS_ORTHANC_URL = os.getenv('PACS_ORTHANC_URL','http://ris-pacs:8042').rstrip('/')
PACS_ORTHANC_USER = os.getenv('PACS_ORTHANC_USER','')
PACS_ORTHANC_PASSWORD = os.getenv('PACS_ORTHANC_PASSWORD','')


PATIENT_PORTAL_TOKEN_DAYS = int(os.getenv('PATIENT_PORTAL_TOKEN_DAYS','7') or 7)

SMTP_ENABLED = os.getenv('SMTP_ENABLED','0') == '1'
SMTP_HOST = os.getenv('SMTP_HOST','')
SMTP_PORT = int(os.getenv('SMTP_PORT','587') or 587)
SMTP_USER = os.getenv('SMTP_USER','')
SMTP_PASSWORD = os.getenv('SMTP_PASSWORD','')
SMTP_FROM = os.getenv('SMTP_FROM',SMTP_USER)
SMTP_TLS = os.getenv('SMTP_TLS','1') == '1'

TELEGRAM_ENABLED = os.getenv('TELEGRAM_ENABLED','0') == '1'
TELEGRAM_BOT_TOKEN = os.getenv('TELEGRAM_BOT_TOKEN','')
TELEGRAM_STAFF_CHAT_ID = os.getenv('TELEGRAM_STAFF_CHAT_ID','')

WHATSAPP_ENABLED = os.getenv('WHATSAPP_ENABLED','0') == '1'
WHATSAPP_API_URL = os.getenv('WHATSAPP_API_URL','')
WHATSAPP_API_KEY = os.getenv('WHATSAPP_API_KEY','')
WHATSAPP_INSTANCE = os.getenv('WHATSAPP_INSTANCE','')

AUTH = (ORTHANC_USER, ORTHANC_PASSWORD)

ROLE_PERMS = {
    'SUPER_ADMIN': ['*'],
    'IT': ['dashboard.view','patients.view','orders.view','devices.manage','users.view','audit.view','pacs.view'],
    'RECEPTION': ['dashboard.view','patients.view','patients.create','patients.edit','orders.view','orders.create','appointments.manage','payments.create','attachments.view','attachments.upload','delivery.manage','finance.discount.request','correction.request'],
    'CALL_CENTER': ['dashboard.view','patients.view','patients.create','orders.view','appointments.manage'],
    'NURSING': ['dashboard.view','patients.view','orders.view','nursing.worklist','consumables.use','attachments.view','attachments.upload'],
    'TECHNICIAN': ['dashboard.view','patients.view','orders.view','technician.worklist','consumables.use','orders.status','attachments.view','pacs.transfer','correction.request'],
    'SERVICE': ['dashboard.view','patients.view','orders.view','technician.worklist','consumables.use','orders.status'],
    'RADIOLOGIST': ['dashboard.view','patients.view','orders.view','reports.view','reports.edit','pacs.view','attachments.view','workflow.command.view','pacs.open_pool.take'],
    'CASE_COORDINATOR': ['dashboard.view','patients.view','orders.view','reports.view','attachments.view','reports.coordinate','reports.version.view','assignment.manage','correction.request','correction.review'],
    'BRANCH_MANAGER': ['dashboard.view','patients.view','orders.view','reports.view','attachments.view','correction.request'],
    'CASHIER': ['dashboard.view','patients.view','orders.view','payments.create','payments.refund','receipts.print','cashier.workstation','finance.discount.request'],
    'CONTRACTS': ['dashboard.view','contracts.manage','procedures.view','patients.view','orders.view'],
    'INVENTORY': ['dashboard.view','inventory.manage','orders.view'],
    'MANAGEMENT': ['dashboard.view','patients.view','orders.view','reports.view','financial.view','analytics.view','analytics.executive.view','export.view','receivables.view','delivery.manage','audit.view','finance.discount.approve','finance.doctor_share.view','inventory.view','inventory.transfer.approve','inventory.reports'],

}

ALL_PERMISSIONS = [
 'dashboard.view','patients.view','patients.create','patients.edit',
 'orders.view','orders.create','orders.status',
 'payments.create','payments.refund','financial.view',
 'nursing.worklist','technician.worklist','consumables.use',
 'reports.view','reports.edit','pacs.view',
 'devices.manage','procedures.view','contracts.manage','inventory.manage',
 'users.view','users.manage','roles.manage','branches.manage','audit.view','consumables.correct','workflow.override',
 'analytics.view','export.view','import.manage','referrers.manage','appointments.manage','receivables.view','delivery.manage','attachments.view','attachments.upload',
 'finance.discount.request','finance.discount.apply','finance.discount.approve','finance.discount.manage',
 'finance.doctor_share.manage','finance.doctor_share.view',
 'inventory.view','inventory.consume','inventory.add_extra_consumable','inventory.adjust_quantity',
 'inventory.waste','inventory.return','inventory.transfer.request','inventory.transfer.approve',
 'inventory.transfer.dispatch','inventory.transfer.receive','inventory.manage','inventory.reports',

    'backup.view',
    'backup.create',
    'backup.restore',
    'storage.view',
    'storage.manage',
    'data.export',
    'data.archive',
    'data.anonymize',
    'data.delete.patient',
    'data.delete.module',
    'data.reset',
    'data.purge',
    'data.cancel.case',

    'backup.download',
    'backup.all_patients',
    'data.delete.all_patients',

    'analytics.executive.view',
    'workflow.command.view','pacs.intake.view','pacs.intake.manage',
    'pacs.destinations.manage','pacs.transfer','pacs.archive.manage','pacs.discovery.view','pacs.discovery.manage',
    'storage.cloud.view','storage.cloud.manage','storage.targets.view','storage.targets.manage','dr.view','dr.manage','dr.failover',
    'reports.coordinate','reports.version.view','pacs.open_pool.take','assignment.manage',
    'correction.request','correction.review','correction.execute',]

from infrastructure.models import Branch

from infrastructure.models import Role

from infrastructure.models import User

from infrastructure.models import AuditLog

from infrastructure.models import Patient

from infrastructure.models import DicomEndpoint


from infrastructure.models import Device

from infrastructure.models import Procedure


from infrastructure.models import ProcedureBranch


from infrastructure.models import ClinicalNote

from infrastructure.models import UserProfile

from infrastructure.models import DictionaryEntry

from infrastructure.models import RibbonAction

from infrastructure.models import UserRibbonAction

from infrastructure.models import BookingBlock

from infrastructure.models import JourneyEvent

from infrastructure.models import ChatThread

from infrastructure.models import ChatThreadMember

from infrastructure.models import ChatMessage


from infrastructure.models import PatientDocument

from infrastructure.models import DocumentMatch

from infrastructure.models import TelephonyExtension

from infrastructure.models import CallLog


from infrastructure.models import FieldRule

from infrastructure.models import UserRegistrationPreference






from infrastructure.models import Warehouse

from infrastructure.models import InventoryItem

from infrastructure.models import InventoryLot

from infrastructure.models import WarehouseStockPolicy

from infrastructure.models import ProcedureConsumableDefault

from infrastructure.models import OrderConsumableUsage

from infrastructure.models import WarehouseStockMovement

from infrastructure.models import StockTransfer

from infrastructure.models import StockTransferLine

from infrastructure.models import PerformingDoctorShareRule

from infrastructure.models import OrderPerformingDoctor

from infrastructure.models import DiscountPolicy

from infrastructure.models import OrderDiscount

from infrastructure.models import DiscountApprovalRequest

from infrastructure.models import ReferralCommissionRule

from infrastructure.models import DailyClosing

from infrastructure.models import StudioShortcut

from infrastructure.models import RegistrationDefault

from infrastructure.models import PatientAlert

from infrastructure.models import CompanyLegalProfile

from infrastructure.models import StudioTemplate

from infrastructure.models import StudioTemplateVersion

from infrastructure.models import DoctorWordTemplate


from infrastructure.models import PatientDemographicsExtra


from infrastructure.models import AIProviderSetting

from infrastructure.models import AIDraftReport

from infrastructure.models import ExternalPriorStudy

from infrastructure.models import PriorStudySelection


from infrastructure.models import CommunicationChannelSetting

from infrastructure.models import CommunicationTemplate

from infrastructure.models import PatientCommunication


from infrastructure.models import ImagingAIProvider


# ============================================================
# RadBridge Automation & QA Control Center
# ============================================================

from infrastructure.models import QARun


from infrastructure.models import QATestResult


from infrastructure.models import QARemediationAction


QA_N8N_BASE_URL=os.getenv('QA_N8N_BASE_URL','').rstrip('/')
QA_N8N_RUN_WEBHOOK=os.getenv('QA_N8N_RUN_WEBHOOK','').strip()
QA_N8N_RETRY_WEBHOOK=os.getenv('QA_N8N_RETRY_WEBHOOK','').strip()
QA_N8N_ENABLE_WEBHOOK=os.getenv('QA_N8N_ENABLE_WEBHOOK','').strip()
QA_N8N_DISABLE_WEBHOOK=os.getenv('QA_N8N_DISABLE_WEBHOOK','').strip()
QA_N8N_STATUS_WEBHOOK=os.getenv('QA_N8N_STATUS_WEBHOOK','').strip()
QA_CALLBACK_TOKEN=os.getenv('QA_CALLBACK_TOKEN','').strip()


# RB_BIND:0002
from application.components.operations import install_0002 as _rb_install_0002
qa_n8n_call = _rb_install_0002(_rb_bind(_rb_sys.modules[__name__], ('requests',)))


# RB_BIND:0003
from application.components.operations import install_0003 as _rb_install_0003
qa_run_summary = _rb_install_0003(_rb_bind(_rb_sys.modules[__name__], ()))


from infrastructure.models import ImagingAIAnalysis

from infrastructure.models import DemographicOverrideAudit


from infrastructure.models import DoctorAvailabilityState

from infrastructure.models import AssignmentOfferPolicy


from infrastructure.models import AssignmentStrategy


from infrastructure.models import AssignmentStrategyVersion


from infrastructure.models import AssignmentVisit


from infrastructure.models import AssignmentDistributionPolicy


from infrastructure.models import AssignmentOffer

from infrastructure.models import AssignmentEscalationEvent


from infrastructure.models import ProcedureQualityRule

from infrastructure.models import CaseReadiness


from infrastructure.models import ReadinessSeriesAlias


from infrastructure.models import CaseReadinessOverride


from infrastructure.models import CaseExceptionResolutionEvent

from infrastructure.models import CaseWorkflowEvent

from infrastructure.models import CaseReturnReason

from infrastructure.models import TATPolicy

from infrastructure.models import RISNotification

from infrastructure.models import NotificationEscalationRule


from infrastructure.models import SavedSearchFilter

from infrastructure.models import LabelTemplate


from infrastructure.models import SmartIntakeSession

from infrastructure.models import SmartIntakeDocument

from infrastructure.models import SmartIntakeFieldReview

from infrastructure.models import OCRProviderSetting



from infrastructure.models import PrinterProfile

from infrastructure.models import ClinicalAssessment

from infrastructure.models import VitalSign

from infrastructure.models import ClinicalDiagnosis

from infrastructure.models import ClinicalMedication

from infrastructure.models import ClinicalLabItem

from infrastructure.models import Contract

from infrastructure.models import ContractPrice

from infrastructure.models import Consumable

from infrastructure.models import ProcedureConsumable


from infrastructure.models import ReferringDoctor


from infrastructure.models import Receipt

from infrastructure.models import DeliveryEvent

from infrastructure.models import Attachment


from infrastructure.models import DeviceSchedule

from infrastructure.models import DeviceBlock


from infrastructure.models import DoctorAssignmentProfile

from infrastructure.models import CaseAssignment

from infrastructure.models import Appointment

from infrastructure.models import Order





# ============================================================
# Insurance / Medical Approval Module - v5.34
# Isolated from Order.status / PACS / MWL / Auto Assignment.
# ============================================================

from infrastructure.models import InsuranceApprovalRequest


from infrastructure.models import InsuranceApprovalEvent


from infrastructure.models import InsurancePortalAccount


from infrastructure.models import EventOutbox

from infrastructure.models import CompanyProfile

from infrastructure.models import StickerTemplate

from infrastructure.models import PacsStudyMatch


from infrastructure.models import PacsExamLink


from infrastructure.models import PacsResolutionEvent


from infrastructure.models import DicomCorrection



from infrastructure.models import PacsStudyInbox


# External reporting is deliberately separated from normal RIS registration.
# A received study cannot enter assignment until an administrator queues it.
from infrastructure.models import ExternalReportingSetting


from infrastructure.models import ExternalReportingCenter


from infrastructure.models import ExternalCenterDoctor


from infrastructure.models import ExternalIntakeCase


from infrastructure.models import DiscoveredDicomNode

from infrastructure.models import DicomDestination

from infrastructure.models import PacsRoutingRule

from infrastructure.models import PacsTransferJob

from infrastructure.models import CoordinatorTask



# RB_CORRECTION_REQUEST_WORKFLOW_V1
from infrastructure.models import CorrectionRequest


from infrastructure.models import PatientDemographicsCorrection

# RB_CORRECTION_EXECUTION_LEDGER_PHASE2A_V1
from infrastructure.models import CorrectionExecution


from infrastructure.models import StorageLocation

from infrastructure.models import ReplicationNode

from infrastructure.models import ReplicationPolicy

from infrastructure.models import ReplicationState

from infrastructure.models import FailoverEvent

from infrastructure.models import CloudBackupConfig

from infrastructure.models import BackupRun

from infrastructure.models import InventoryBatch

from infrastructure.models import StockMovement

from infrastructure.models import ERPItemMapping

from infrastructure.models import ERPSyncQueue

from infrastructure.models import FinancialEntry

from infrastructure.models import Payment


from infrastructure.models import PaymentRefund

from infrastructure.models import OrderConsumable


from infrastructure.models import NursingConsumableEntry

from infrastructure.models import PortalAccessToken

from infrastructure.models import NotificationEvent

from infrastructure.models import Report

from infrastructure.models import ReportVersion


# RB_REPORT_EDIT_LEASE_FRESH_BOOTSTRAP_MODEL_V1
# Keep this SQLAlchemy metadata contract aligned with migration
# 20260901_001_report_edit_lease.py. Fresh bootstrap uses db.create_all(),
# so this table must be registered in application metadata as well.
from infrastructure.models import ReportEditLease


from infrastructure.models import DoctorSigningKey



I18N = {
 'en': {
   'dashboard':'Dashboard','patients':'Patients','orders':'Orders / Visits','technician':'Technician Worklist',
   'doctor':'Doctor Worklist','analytics':'Analytics','appointments':'Appointments','admin':'Admin Panel',
   'users':'Users','roles':'Roles & Permissions','branches':'Branches','devices':'Devices / AE Titles',
   'procedures':'Procedure Catalog','contracts':'Contracts & Pricing','inventory':'Inventory',
   'referrers':'Referring Doctors','import_export':'Import / Export','audit':'Audit Log',
   'logout':'Logout','change_password':'Change password','pacs':'PACS'
 },
 'ar': {
   'dashboard':'لوحة المتابعة','patients':'المرضى','orders':'الطلبات / الزيارات','technician':'قائمة عمل الفنيين',
   'doctor':'قائمة عمل الأطباء','analytics':'الإحصائيات','appointments':'المواعيد','admin':'لوحة الإدارة',
   'users':'المستخدمون','roles':'الأدوار والصلاحيات','branches':'الفروع','devices':'الأجهزة / AE Titles',
   'procedures':'دليل الفحوصات','contracts':'جهات التعاقد والأسعار','inventory':'المخزون',
   'referrers':'الأطباء المحولون','import_export':'استيراد / تصدير','audit':'سجل المراجعة',
   'logout':'خروج','change_password':'تغيير كلمة المرور','pacs':'PACS'
 }
}
# RB_BIND:0004
from web.components.shared import install_0004 as _rb_install_0004
lang = _rb_install_0004(_rb_bind(_rb_sys.modules[__name__], ('current_ui_language', 'session')))
# RB_BIND:0005
from application.components.shared import install_0005 as _rb_install_0005
tr = _rb_install_0005(_rb_bind(_rb_sys.modules[__name__], ('I18N', 'lang')))



# RB_BIND:0006
from application.components.shared import install_0006 as _rb_install_0006
_active_portal_token = _rb_install_0006(_rb_bind(_rb_sys.modules[__name__], ('PATIENT_PORTAL_TOKEN_DAYS', 'PortalAccessToken', 'datetime', 'db', 'secrets', 'timedelta')))

# RB_BIND:0007
from application.components.communications import install_0007 as _rb_install_0007
_notification_exists = _rb_install_0007(_rb_bind(_rb_sys.modules[__name__], ('NotificationEvent',)))

# RB_BIND:0008
from web.components.shared import install_0008 as _rb_install_0008
queue_final_delivery = _rb_install_0008(_rb_bind(_rb_sys.modules[__name__], ('NotificationEvent', 'RIS_PUBLIC_URL', 'SMTP_ENABLED', 'TELEGRAM_ENABLED', 'TELEGRAM_STAFF_CHAT_ID', 'WHATSAPP_ENABLED', '_active_portal_token', '_notification_exists', 'audit', 'db', 'url_for')))

# RB_BIND:0009
from web.components.shared import install_0009 as _rb_install_0009
_final_portal_url = _rb_install_0009(_rb_bind(_rb_sys.modules[__name__], ('RIS_PUBLIC_URL', '_active_portal_token', 'url_for')))

# RB_BIND:0010
from application.components.communications import install_0010 as _rb_install_0010
send_notification_job = _rb_install_0010(_rb_bind(_rb_sys.modules[__name__], ('EmailMessage', 'RIS_PUBLIC_URL', 'SMTP_ENABLED', 'SMTP_FROM', 'SMTP_HOST', 'SMTP_PASSWORD', 'SMTP_PORT', 'SMTP_TLS', 'SMTP_USER', 'TELEGRAM_BOT_TOKEN', 'TELEGRAM_ENABLED', 'WHATSAPP_API_KEY', 'WHATSAPP_API_URL', 'WHATSAPP_ENABLED', 'WHATSAPP_INSTANCE', '_final_portal_url', 'audit', 'datetime', 'requests', 'smtplib')))

# RB_BIND:0011
from application.components.communications import install_0011 as _rb_install_0011
process_notification_queue = _rb_install_0011(_rb_bind(_rb_sys.modules[__name__], ('NotificationEvent', 'db', 'send_notification_job')))


# Examination workflow is separate from report workflow.
# VERIFIED / SECOND_VERIFIED / FINAL belong to REPORT_WORKFLOW_STATUSES
# and must never be set by the technician exam-status endpoint.
EXAM_STATUS_ORDER=list(EXAM_FORWARD_STATES)

# RB_BIND:0012
from application.components.clinical import install_0012 as _rb_install_0012
exam_status_rank = _rb_install_0012(_rb_bind(_rb_sys.modules[__name__], ('EXAM_STATUS_ORDER', 'normalize_state')))

# RB_BIND:0013
from application.components.shared import install_0013 as _rb_install_0013
exam_transition_allowed = _rb_install_0013(_rb_bind(_rb_sys.modules[__name__], ('can_transition',)))

# RB_BIND:0014
from application.components.imaging import install_0014 as _rb_install_0014
_orthanc_auth = _rb_install_0014(_rb_bind(_rb_sys.modules[__name__], ('PACS_ORTHANC_PASSWORD', 'PACS_ORTHANC_USER')))

# RB_BIND:0015
from application.components.imaging import install_0015 as _rb_install_0015
pacs_lookup_order = _rb_install_0015(_rb_bind(_rb_sys.modules[__name__], ('DicomCorrection', 'PACS_IMAGES_STABLE_SECONDS', 'PACS_ORTHANC_URL', '_orthanc_auth', 'datetime', 'db', 'requests')))


FINANCIAL_REPORTS={
 'summary':'Executive Financial Summary',
 'profit_loss':'Profit & Loss',
 'balance_sheet':'Management Balance Sheet',
 'cash_flow':'Cash Flow',
 'revenue_procedure':'Revenue by Procedure / Modality',
 'revenue_payer':'Revenue by Payer / Contract',
 'aging':'Receivables Aging',
 'exam_costing':'Exam Costing / COGS',
 'maintenance':'Medical Equipment Maintenance',
 'payroll_doctor':'Payroll & Doctor Fees',
 'inventory':'Inventory & Consumable Movement',
 'device_profitability':'Device / Modality Profitability',
 'branch_profitability':'Branch Profitability',
 'collections':'Collections by Payment Method'
}
FINANCE_ENTRY_TYPES=['EXPENSE','MAINTENANCE','PAYROLL','DOCTOR_FEE','ASSET','LIABILITY','EQUITY','INVENTORY_PURCHASE']

# RB_BIND:0016
from application.components.shared import install_0016 as _rb_install_0016
_money = _rb_install_0016(_rb_bind(_rb_sys.modules[__name__], ()))

# RB_BIND:0017
from web.components.finance import install_0017 as _rb_install_0017
_finance_filters = _rb_install_0017(_rb_bind(_rb_sys.modules[__name__], ('date', 'request')))

# RB_BIND:0018
from application.components.finance import install_0018 as _rb_install_0018
_orders_for_finance = _rb_install_0018(_rb_bind(_rb_sys.modules[__name__], ('Order', 'datetime', 'timedelta')))

# RB_BIND:0019
from application.components.finance import install_0019 as _rb_install_0019
_entries_for_finance = _rb_install_0019(_rb_bind(_rb_sys.modules[__name__], ('FinancialEntry',)))

# RB_BIND:0020
from application.components.inventory import install_0020 as _rb_install_0020
_order_direct_consumable_cost = _rb_install_0020(_rb_bind(_rb_sys.modules[__name__], ('NursingConsumableEntry', 'OrderConsumable', '_money')))

# RB_BIND:0021
from application.components.registration import install_0021 as _rb_install_0021
_order_revenue = _rb_install_0021(_rb_bind(_rb_sys.modules[__name__], ('_money',)))

# RB_BIND:0022
from application.components.finance import install_0022 as _rb_install_0022
_financial_dataset = _rb_install_0022(_rb_bind(_rb_sys.modules[__name__], ('Consumable', 'FINANCIAL_REPORTS', 'NursingConsumableEntry', 'Receipt', '_entries_for_finance', '_money', '_order_direct_consumable_cost', '_order_revenue', '_orders_for_finance', 'date', 'order_balance')))

# RB_BIND:0023
from web.components.shared import install_0023 as _rb_install_0023
_export_csv = _rb_install_0023(_rb_bind(_rb_sys.modules[__name__], ('csv', 'io', 'send_file')))

# RB_BIND:0024
from web.components.shared import install_0024 as _rb_install_0024
_export_xlsx = _rb_install_0024(_rb_bind(_rb_sys.modules[__name__], ('Alignment', 'Font', 'PatternFill', 'Workbook', 'get_column_letter', 'io', 'send_file')))

# RB_BIND:0025
from web.components.shared import install_0025 as _rb_install_0025
_export_pdf = _rb_install_0025(_rb_bind(_rb_sys.modules[__name__], ('A4', 'Paragraph', 'SimpleDocTemplate', 'Spacer', 'Table', 'TableStyle', 'colors', 'getSampleStyleSheet', 'io', 'send_file')))


DEFAULT_STICKER_FIELDS=['patient_name','patient_id','accession','exam','datetime','branch']
DEFAULT_STICKER_LABELS={
'patient_name':{'en':'Patient Name','ar':'اسم المريض'},
'patient_id':{'en':'Patient ID','ar':'رقم المريض'},
'accession':{'en':'Accession No.','ar':'رقم الحالة'},
'exam':{'en':'Exam','ar':'نوع الفحص'},
'datetime':{'en':'Date & Time','ar':'التاريخ والوقت'},
'branch':{'en':'Branch','ar':'الفرع'},
'modality':{'en':'Modality','ar':'نوع الجهاز'},
'phone':{'en':'Mobile','ar':'الموبايل'},
'worklist_name':{'en':'Worklist Name','ar':'الاسم في قائمة العمل'}
}

# RB_BIND:0026
from application.components.shared import install_0026 as _rb_install_0026
get_company_profile = _rb_install_0026(_rb_bind(_rb_sys.modules[__name__], ('CompanyProfile', 'db')))

# RB_BIND:0027
from application.components.shared import install_0027 as _rb_install_0027
get_default_sticker_template = _rb_install_0027(_rb_bind(_rb_sys.modules[__name__], ('DEFAULT_STICKER_FIELDS', 'DEFAULT_STICKER_LABELS', 'STICKER_DEFAULT_HEIGHT_MM', 'STICKER_DEFAULT_WIDTH_MM', 'StickerTemplate', 'db', 'json')))

# RB_BIND:0028
from application.components.shared import install_0028 as _rb_install_0028
_sticker_fields = _rb_install_0028(_rb_bind(_rb_sys.modules[__name__], ('DEFAULT_STICKER_FIELDS', 'json')))

# RB_BIND:0029
from application.components.shared import install_0029 as _rb_install_0029
_sticker_labels = _rb_install_0029(_rb_bind(_rb_sys.modules[__name__], ('DEFAULT_STICKER_LABELS', 'json')))

# RB_BIND:0030
from application.components.shared import install_0030 as _rb_install_0030
_sticker_value = _rb_install_0030(_rb_bind(_rb_sys.modules[__name__], ()))

# RB_BIND:0031
from web.components.shared import install_0031 as _rb_install_0031
_portal_url_for_order = _rb_install_0031(_rb_bind(_rb_sys.modules[__name__], ('RIS_PUBLIC_URL', '_active_portal_token', 'db', 'url_for')))

# RB_BIND:0032
from application.components.registration import install_0032 as _rb_install_0032
find_duplicate_patients = _rb_install_0032(_rb_bind(_rb_sys.modules[__name__], ('Patient',)))

# RB_BIND:0033
from application.components.imaging import install_0033 as _rb_install_0033
_orthanc_recent_studies = _rb_install_0033(_rb_bind(_rb_sys.modules[__name__], ('PACS_ORTHANC_URL', '_orthanc_auth', 'requests')))

# RB_BIND:0034
from application.components.shared import install_0034 as _rb_install_0034
_study_tags = _rb_install_0034(_rb_bind(_rb_sys.modules[__name__], ()))


# RB_BIND:0035
from application.components.imaging import install_0035 as _rb_install_0035
strict_pacs_demographic_match = _rb_install_0035(_rb_bind(_rb_sys.modules[__name__], ('Order', 'Patient')))



# RB_BIND:0036
from application.components.imaging import install_0036 as _rb_install_0036
pacs_demographic_gate = _rb_install_0036(_rb_bind(_rb_sys.modules[__name__], ('CaseReadiness', 'PacsStudyInbox', 'PacsStudyMatch')))


# RB_BIND:0037
from application.components.registration import install_0037 as _rb_install_0037
suggest_order_for_study = _rb_install_0037(_rb_bind(_rb_sys.modules[__name__], ('Order', 'Patient')))



# RB_BIND:0038
from application.components.imaging import install_0038 as _rb_install_0038
_orthanc_study_detail = _rb_install_0038(_rb_bind(_rb_sys.modules[__name__], ('PACS_ORTHANC_URL', '_orthanc_auth', 'requests')))

# RB_BIND:0039
from application.components.shared import install_0039 as _rb_install_0039
_study_tags_extended = _rb_install_0039(_rb_bind(_rb_sys.modules[__name__], ('_orthanc_instance_remote_aet', '_orthanc_series_detail', '_orthanc_study_detail', '_study_tags')))

# RB_BIND:0040
from application.components.imaging import install_0040 as _rb_install_0040
_orthanc_series_detail = _rb_install_0040(_rb_bind(_rb_sys.modules[__name__], ('PACS_ORTHANC_URL', '_orthanc_auth', 'requests')))

# RB_BIND:0041
from application.components.imaging import install_0041 as _rb_install_0041
_orthanc_instance_remote_aet = _rb_install_0041(_rb_bind(_rb_sys.modules[__name__], ('PACS_ORTHANC_URL', '_orthanc_auth', 'requests')))

# RB_BIND:0042
from application.components.imaging import install_0042 as _rb_install_0042
_parse_dicom_date = _rb_install_0042(_rb_bind(_rb_sys.modules[__name__], ('datetime',)))

# RB_BIND:0043
from application.components.imaging import install_0043 as _rb_install_0043
provisional_patient_from_pacs = _rb_install_0043(_rb_bind(_rb_sys.modules[__name__], ('Patient', 'db', 'next_patient_id')))

# RB_BIND:0044
from application.components.imaging import install_0044 as _rb_install_0044
create_order_from_pacs = _rb_install_0044(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Order', '_parse_dicom_date', 'datetime', 'db', 'new_accession', 'provisional_patient_from_pacs')))


# RB_BIND:0045
from application.components.imaging import install_0045 as _rb_install_0045
pacs_auto_intake_once = _rb_install_0045(_rb_bind(_rb_sys.modules[__name__], ('DicomCorrection', 'ExternalIntakeCase', 'ExternalReportingCenter', 'Order', 'PacsResolutionEvent', 'PacsStudyInbox', 'PacsStudyMatch', 'Report', '_orthanc_recent_studies', '_study_tags_extended', 'app', 'datetime', 'db', 'discover_dicom_node_from_tags', 'enqueue_auto_routes_for_inbox', 'evaluate_case_readiness', 'external_reporting_setting', 'strict_pacs_demographic_match')))


# RB_BIND:0046
from application.components.imaging import install_0046 as _rb_install_0046
discover_dicom_node_from_tags = _rb_install_0046(_rb_bind(_rb_sys.modules[__name__], ('DiscoveredDicomNode', 'datetime', 'db')))

# RB_BIND:0047
from application.components.shared import install_0047 as _rb_install_0047
approve_discovered_node = _rb_install_0047(_rb_bind(_rb_sys.modules[__name__], ('DicomDestination', 'current_user', 'datetime', 'db', 're', 'sync_destination_to_orthanc')))

# RB_BIND:0048
from application.components.imaging import install_0048 as _rb_install_0048
sync_destination_to_orthanc = _rb_install_0048(_rb_bind(_rb_sys.modules[__name__], ('PACS_ORTHANC_URL', '_orthanc_auth', 'requests')))

# RB_BIND:0049
from application.components.shared import install_0049 as _rb_install_0049
enqueue_transfer = _rb_install_0049(_rb_bind(_rb_sys.modules[__name__], ('PacsTransferJob', 'db')))

# RB_BIND:0050
from application.components.shared import install_0050 as _rb_install_0050
enqueue_auto_routes_for_inbox = _rb_install_0050(_rb_bind(_rb_sys.modules[__name__], ('Order', 'PacsRoutingRule', 'enqueue_transfer')))

# RB_BIND:0051
from application.components.imaging import install_0051 as _rb_install_0051
test_dicom_destination = _rb_install_0051(_rb_bind(_rb_sys.modules[__name__], ('PACS_ORTHANC_URL', '_orthanc_auth', 'requests', 'sync_destination_to_orthanc')))

# RB_BIND:0052
from application.components.imaging import install_0052 as _rb_install_0052
process_pacs_transfer_jobs = _rb_install_0052(_rb_bind(_rb_sys.modules[__name__], ('PACS_ORTHANC_URL', 'PacsTransferJob', '_orthanc_auth', 'datetime', 'db', 'requests', 'test_dicom_destination')))

# RB_BIND:0053
from application.components.shared import install_0053 as _rb_install_0053
workflow_command_snapshot = _rb_install_0053(_rb_bind(_rb_sys.modules[__name__], ('Order', 'datetime', 'orthanc_ok', 'timedelta')))

# RB_BIND:0054
from application.components.shared import install_0054 as _rb_install_0054
doctor_can_take_open_case = _rb_install_0054(_rb_bind(_rb_sys.modules[__name__], ('assignment_candidate_users', 'assignment_offer_policy', 'readiness_allows_auto_assign')))

# RB_BIND:0055
from application.components.operations import install_0055 as _rb_install_0055
storage_defaults = _rb_install_0055(_rb_bind(_rb_sys.modules[__name__], ('ATTACHMENT_DIR', 'REPORT_DIR', 'StorageLocation', 'db', 'os')))

# RB_BIND:0056
from application.components.shared import install_0056 as _rb_install_0056
enqueue_event = _rb_install_0056(_rb_bind(_rb_sys.modules[__name__], ('EventOutbox', 'db', 'json')))

# RB_BIND:0057
from application.components.shared import install_0057 as _rb_install_0057
pending_outbox_events = _rb_install_0057(_rb_bind(_rb_sys.modules[__name__], ('EventOutbox',)))


DEFAULT_CONSUMABLE_CATALOG=[
 {'code':'CON-001','barcode':'CON001','name_en':'CT Contrast','name_ar':'صبغة الأشعة المقطعية','category':'Contrast','unit':'ml'},
 {'code':'CON-002','barcode':'CON002','name_en':'MRI Contrast','name_ar':'صبغة الرنين المغناطيسي','category':'Contrast','unit':'ml'},
 {'code':'CON-003','barcode':'CON003','name_en':'Normal Saline','name_ar':'محلول ملحي','category':'Fluids','unit':'ml'},
 {'code':'CON-004','barcode':'CON004','name_en':'Syringe 10ml','name_ar':'سرنجة 10 مل','category':'Syringes','unit':'pcs'},
 {'code':'CON-005','barcode':'CON005','name_en':'Syringe 20ml','name_ar':'سرنجة 20 مل','category':'Syringes','unit':'pcs'},
 {'code':'CON-006','barcode':'CON006','name_en':'Syringe 50ml','name_ar':'سرنجة 50 مل','category':'Syringes','unit':'pcs'},
 {'code':'CON-007','barcode':'CON007','name_en':'IV Cannula 18G','name_ar':'كانيولا 18G','category':'IV Access','unit':'pcs'},
 {'code':'CON-008','barcode':'CON008','name_en':'IV Cannula 20G','name_ar':'كانيولا 20G','category':'IV Access','unit':'pcs'},
 {'code':'CON-009','barcode':'CON009','name_en':'IV Cannula 22G','name_ar':'كانيولا 22G','category':'IV Access','unit':'pcs'},
 {'code':'CON-010','barcode':'CON010','name_en':'Examination Gloves','name_ar':'جوانتي فحص','category':'PPE','unit':'pair'},
 {'code':'CON-011','barcode':'CON011','name_en':'Alcohol Swab','name_ar':'مسحة كحول','category':'Disinfection','unit':'pcs'},
 {'code':'CON-012','barcode':'CON012','name_en':'Medical Tape','name_ar':'لاصق طبي','category':'Dressing','unit':'pcs'},
 {'code':'US-001','barcode':'US001','name_en':'Ultrasound Gel','name_ar':'جل سونار','category':'Ultrasound','unit':'ml'},
 {'code':'US-002','barcode':'US002','name_en':'Probe Cover','name_ar':'غطاء بروب','category':'Ultrasound','unit':'pcs'},
 {'code':'XR-001','barcode':'XR001','name_en':'X-Ray Film','name_ar':'فيلم أشعة','category':'Printing','unit':'sheet'},
 {'code':'GEN-001','barcode':'GEN001','name_en':'Printing Paper','name_ar':'ورق طباعة','category':'Printing','unit':'sheet'},
 {'code':'GEN-002','barcode':'GEN002','name_en':'Result Envelope','name_ar':'ظرف نتيجة','category':'Printing','unit':'pcs'}
]

INVENTORY_IMPORT_COLUMNS=[
 'code','barcode','name_en','name_ar','category','unit','stock_qty',
 'reorder_level','unit_cost','selling_price','expiry_tracking','active'
]

# RB_BIND:0058
from application.components.inventory import install_0058 as _rb_install_0058
_record_stock_movement = _rb_install_0058(_rb_bind(_rb_sys.modules[__name__], ('ERP_SYNC_ENABLED', 'StockMovement', 'db', 'queue_erp_event')))

# RB_BIND:0059
from application.components.inventory import install_0059 as _rb_install_0059
seed_default_consumables = _rb_install_0059(_rb_bind(_rb_sys.modules[__name__], ('Consumable', 'DEFAULT_CONSUMABLE_CATALOG', 'db')))

# RB_BIND:0060
from application.components.inventory import install_0060 as _rb_install_0060
inventory_row = _rb_install_0060(_rb_bind(_rb_sys.modules[__name__], ()))

# RB_BIND:0061
from application.components.inventory import install_0061 as _rb_install_0061
queue_erp_event = _rb_install_0061(_rb_bind(_rb_sys.modules[__name__], ('ERPSyncQueue', 'db', 'json')))

# RB_BIND:0062
from application.components.inventory import install_0062 as _rb_install_0062
_erp_headers = _rb_install_0062(_rb_bind(_rb_sys.modules[__name__], ('ERP_API_TOKEN',)))

# RB_BIND:0063
from application.components.inventory import install_0063 as _rb_install_0063
send_erp_queue_item = _rb_install_0063(_rb_bind(_rb_sys.modules[__name__], ('ERP_BASE_URL', 'ERP_PROVIDER', 'ERP_SYNC_ENABLED', 'ERP_TIMEOUT', '_erp_headers', 'datetime', 'json', 'requests')))

# RB_BIND:0064
from application.components.inventory import install_0064 as _rb_install_0064
process_erp_queue = _rb_install_0064(_rb_bind(_rb_sys.modules[__name__], ('ERPSyncQueue', 'db', 'send_erp_queue_item')))

# RB_BIND:0065
from application.components.finance import install_0065 as _rb_install_0065
next_receipt_no = _rb_install_0065(_rb_bind(_rb_sys.modules[__name__], ('Receipt', 'datetime')))

# RB_BIND:0066
from application.components.registration import install_0066 as _rb_install_0066
order_balance = _rb_install_0066(_rb_bind(_rb_sys.modules[__name__], ('effective_patient_due',)))

# RB_BIND:0067
from application.components.registration import install_0067 as _rb_install_0067
procedure_display_name = _rb_install_0067(_rb_bind(_rb_sys.modules[__name__], ('lang',)))


# RB_BIND:0068
from application.components.shared import install_0068 as _rb_install_0068
_hhmm_minutes = _rb_install_0068(_rb_bind(_rb_sys.modules[__name__], ()))

# RB_BIND:0069
from application.components.shared import install_0069 as _rb_install_0069
generate_slots = _rb_install_0069(_rb_bind(_rb_sys.modules[__name__], ('Appointment', 'DeviceBlock', 'DeviceSchedule', '_hhmm_minutes', 'datetime', 'timedelta')))


# RB_BIND:0070
from application.components.shared import install_0070 as _rb_install_0070
_safe_token = _rb_install_0070(_rb_bind(_rb_sys.modules[__name__], ('re',)))

# RB_BIND:0071
from application.components.reporting import install_0071 as _rb_install_0071
report_filename = _rb_install_0071(_rb_bind(_rb_sys.modules[__name__], ('_safe_token',)))

# RB_BIND:0072
from application.components.reporting import install_0072 as _rb_install_0072
latest_final_report = _rb_install_0072(_rb_bind(_rb_sys.modules[__name__], ('Report',)))


# RB_BIND:0073
from application.components.shared import install_0073 as _rb_install_0073
_error_id = _rb_install_0073(_rb_bind(_rb_sys.modules[__name__], ('datetime', 'secrets')))

# RB_BIND:0074
from application.components.imaging import install_0074 as _rb_install_0074
_public_pacs_base = _rb_install_0074(_rb_bind(_rb_sys.modules[__name__], ('ORTHANC_PUBLIC_URL', 're')))



# ============================================================
# OHIF INTEGRATION
# Resolve StudyInstanceUID through DICOMweb using Accession.
# Order.pacs_study_id contains the Orthanc internal study ID,
# not the DICOM StudyInstanceUID.
# ============================================================

OHIF_INTERNAL_URL=os.getenv(
    'OHIF_INTERNAL_URL',
    'http://radbridge-ohif'
).rstrip('/')

OHIF_PUBLIC_URL=os.getenv(
    'OHIF_PUBLIC_URL',
    'http://51.89.248.57:8092'
).rstrip('/')


# RB_BIND:0075
from application.components.imaging import install_0075 as _rb_install_0075
ohif_study_uid = _rb_install_0075(_rb_bind(_rb_sys.modules[__name__], ('OHIF_INTERNAL_URL', 'app', 'requests')))



# ============================================================
# RadBridge Doctor Identity Bridge
# ============================================================

OHIF_DOCTOR_IDENTITY_MAX_AGE = int(
    os.getenv("OHIF_DOCTOR_IDENTITY_MAX_AGE", "300")
)


# RB_BIND:0076
from application.components.imaging import install_0076 as _rb_install_0076
_ohif_identity_serializer = _rb_install_0076(_rb_bind(_rb_sys.modules[__name__], ('app', 'os')))


# RB_BIND:0077
from application.components.imaging import install_0077 as _rb_install_0077
create_ohif_doctor_identity = _rb_install_0077(_rb_bind(_rb_sys.modules[__name__], ('_ohif_identity_serializer',)))


# RB_BIND:0078
from application.components.imaging import install_0078 as _rb_install_0078
verify_ohif_doctor_identity = _rb_install_0078(_rb_bind(_rb_sys.modules[__name__], ('OHIF_DOCTOR_IDENTITY_MAX_AGE', 'Order', 'User', '_ohif_identity_serializer', 'app', 'db', 'doctor_case_access_allowed')))


# RB_BIND:0079
from web.components.imaging import install_0079 as _rb_install_0079
get_ohif_doctor_identity = _rb_install_0079(_rb_bind(_rb_sys.modules[__name__], ('request', 'verify_ohif_doctor_identity')))



# RB_BIND:0080
from application.components.shared import install_0080 as _rb_install_0080
audit_as_user = _rb_install_0080(_rb_bind(_rb_sys.modules[__name__], ('AuditLog', '_security_audit_ip', 'db')))


# RB_BIND:0081
from web.components.imaging import install_0081 as _rb_install_0081
require_ohif_doctor_user = _rb_install_0081(_rb_bind(_rb_sys.modules[__name__], ('Order', '__name__', '_reporting_support_sys', 'db', 'get_ohif_doctor_identity', 'jsonify')))


# RB_BIND:0082
from application.components.imaging import install_0082 as _rb_install_0082
ohif_audit_user = _rb_install_0082(_rb_bind(_rb_sys.modules[__name__], ('get_ohif_doctor_identity',)))


# RB_BIND:0083
from application.components.imaging import install_0083 as _rb_install_0083
ohif_url_for_order = _rb_install_0083(_rb_bind(_rb_sys.modules[__name__], ('OHIF_PUBLIC_URL', 'app', 'create_ohif_doctor_identity', 'current_user', 'ohif_study_uid', 'os')))


# RB_BIND:0084
from web.components.imaging import install_0084 as _rb_install_0084
case_pacs_url = _rb_install_0084(_rb_bind(_rb_sys.modules[__name__], ('current_user', 'doctor_case_access_allowed', 'ohif_url_for_order', 'url_for')))


# RB_CLINICAL_IMAGE_ONLY_WORKSPACE_V1
# RB_BIND:0085
from application.components.clinical import install_0085 as _rb_install_0085
clinical_image_only_url = _rb_install_0085(_rb_bind(_rb_sys.modules[__name__], ('OHIF_PUBLIC_URL', 'ohif_study_uid')))




# RB_BIND:0086
from application.components.reporting import install_0086 as _rb_install_0086
_safe_report_token = _rb_install_0086(_rb_bind(_rb_sys.modules[__name__], ('re',)))

# RB_BIND:0087
from application.components.reporting import install_0087 as _rb_install_0087
_next_report_version = _rb_install_0087(_rb_bind(_rb_sys.modules[__name__], ('ReportVersion',)))

# RB_BIND:0088
from application.components.reporting import install_0088 as _rb_install_0088
build_report_filename = _rb_install_0088(_rb_bind(_rb_sys.modules[__name__], ('_safe_report_token',)))

# RB_BIND:0089
from application.components.shared import install_0089 as _rb_install_0089
latest_final_version = _rb_install_0089(_rb_bind(_rb_sys.modules[__name__], ('ReportVersion',)))


REPORT_WORKFLOW_STATUSES=['DRAFT','PARTIAL','READY_FORMATTING','FORMATTING','READY_APPROVAL','VERIFIED','SECOND_VERIFIED','FINAL','ADDENDUM']

# RB_BIND:0090
from application.components.reporting import install_0090 as _rb_install_0090
allowed_report_transitions = _rb_install_0090(_rb_bind(_rb_sys.modules[__name__], ()))


# RB_BIND:0091
from application.components.reporting import install_0091 as _rb_install_0091
report_transition_allowed = _rb_install_0091(_rb_bind(_rb_sys.modules[__name__], ('allowed_report_transitions',)))


# RB_BIND:0092
from application.components.reporting import install_0092 as _rb_install_0092
extract_report_text = _rb_install_0092(_rb_bind(_rb_sys.modules[__name__], ('PdfReader', 're', 'zipfile')))

# RB_BIND:0093
from application.components.shared import install_0093 as _rb_install_0093
_normalize_compare = _rb_install_0093(_rb_bind(_rb_sys.modules[__name__], ('re',)))

# RB_BIND:0094
from application.components.shared import install_0094 as _rb_install_0094
detect_labeled_value = _rb_install_0094(_rb_bind(_rb_sys.modules[__name__], ('re',)))

# RB_BIND:0095
from application.components.reporting import install_0095 as _rb_install_0095
validate_report_patient_file = _rb_install_0095(_rb_bind(_rb_sys.modules[__name__], ('_normalize_compare', 'arabic_to_latin_name', 'detect_labeled_value', 'extract_report_text', 'normalized_person_name', 're')))



# RB_BIND:0096
from application.components.shared import install_0096 as _rb_install_0096
_csv_values = _rb_install_0096(_rb_bind(_rb_sys.modules[__name__], ()))


# RB_BIND:0097
from application.components.assignment import install_0097 as _rb_install_0097
assignable_doctors = _rb_install_0097(_rb_bind(_rb_sys.modules[__name__], ('Role', 'User')))

# RB_BIND:0098
from application.components.shared import install_0098 as _rb_install_0098
doctor_active_case_count = _rb_install_0098(_rb_bind(_rb_sys.modules[__name__], ('Order',)))

# RB_BIND:0099
from application.components.shared import install_0099 as _rb_install_0099
eligible_doctors_for_order = _rb_install_0099(_rb_bind(_rb_sys.modules[__name__], ('DoctorAssignmentProfile', '__name__', '_csv_values', '_reporting_support_sys', 'doctor_active_case_count')))

# RB_BIND:0100
from application.components.assignment import install_0100 as _rb_install_0100
auto_assign_order = _rb_install_0100(_rb_bind(_rb_sys.modules[__name__], ('AssignmentOffer', 'CaseAssignment', 'Order', 'Report', 'User', 'audit', 'db', 'eligible_doctors_for_order')))

# RB_BIND:0101
from application.components.assignment import install_0101 as _rb_install_0101
manual_assign_order = _rb_install_0101(_rb_bind(_rb_sys.modules[__name__], ('AssignmentOffer', 'CaseAssignment', 'Order', 'Report', 'User', 'audit', 'db')))


ARABIC_NAME_MAP = {
    'ا':'A','أ':'A','إ':'E','آ':'A','ب':'B','ت':'T','ث':'TH','ج':'G','ح':'H','خ':'KH',
    'د':'D','ذ':'Z','ر':'R','ز':'Z','س':'S','ش':'SH','ص':'S','ض':'D','ط':'T','ظ':'Z',
    'ع':'A','غ':'GH','ف':'F','ق':'Q','ك':'K','ل':'L','م':'M','ن':'N','ه':'H','ة':'A',
    'و':'W','ؤ':'O','ي':'Y','ى':'A','ئ':'E','ء':'A'
}
# RB_BIND:0102
from application.components.shared import install_0102 as _rb_install_0102
arabic_to_latin_name = _rb_install_0102(_rb_bind(_rb_sys.modules[__name__], ('ARABIC_NAME_MAP', 're')))

# RB_BIND:0103
from application.components.shared import install_0103 as _rb_install_0103
normalized_worklist_name = _rb_install_0103(_rb_bind(_rb_sys.modules[__name__], ('arabic_to_latin_name', 're')))

# RB_BIND:0104
from application.components.registration import install_0104 as _rb_install_0104
patient_worklist_display_name = _rb_install_0104(_rb_bind(_rb_sys.modules[__name__], ('normalized_worklist_name',)))

# RB_BIND:0105
from application.components.imaging import install_0105 as _rb_install_0105
dicom_worklist_patient_name = _rb_install_0105(_rb_bind(_rb_sys.modules[__name__], ('ARABIC_NAME_MAP', 'patient_worklist_display_name', 're')))


# RB_BIND:0106
from application.components.imaging import install_0106 as _rb_install_0106
dicom_person_name = _rb_install_0106(_rb_bind(_rb_sys.modules[__name__], ('re',)))

# RB_BIND:0107
from application.components.shared import install_0107 as _rb_install_0107
normalized_person_name = _rb_install_0107(_rb_bind(_rb_sys.modules[__name__], ('re',)))

# RB_BIND:0108
from web.components.identity import install_0108 as _rb_install_0108
current_user = _rb_install_0108(_rb_bind(_rb_sys.modules[__name__], ('User', 'session')))

# RB_BIND:0109
from application.components.identity import install_0109 as _rb_install_0109
has_perm = _rb_install_0109(_rb_bind(_rb_sys.modules[__name__], ('current_user',)))

# RB_BIND:0110
from web.components.identity import install_0110 as _rb_install_0110
login_required = _rb_install_0110(_rb_bind(_rb_sys.modules[__name__], ('current_user', 'redirect', 'request', 'url_for', 'wraps')))

# RB_BIND:0111
from web.components.shared import install_0111 as _rb_install_0111
require_perm = _rb_install_0111(_rb_bind(_rb_sys.modules[__name__], ('abort', 'current_user', 'has_perm', 'redirect', 'request', 'url_for', 'wraps')))

# RB_BIND:0112
from web.components.identity import install_0112 as _rb_install_0112
_security_audit_ip = _rb_install_0112(_rb_bind(_rb_sys.modules[__name__], ('app', 'os', 'request')))


# RB_BIND:0113
from application.components.shared import install_0113 as _rb_install_0113
audit = _rb_install_0113(_rb_bind(_rb_sys.modules[__name__], ('AuditLog', '_security_audit_ip', 'current_user', 'db')))


# RB_BIND:0114
from application.components.registration import install_0114 as _rb_install_0114
next_patient_id = _rb_install_0114(_rb_bind(_rb_sys.modules[__name__], ('Patient', 'db', 'text')))


# ---------------------------------------------------------------------
# RadBridge v5.45.0 - Unified Registration V2 helpers
# ---------------------------------------------------------------------
# RB_BIND:0115
from application.components.registration import install_0115 as _rb_install_0115
_rb_registration_v2_name_parts = _rb_install_0115(_rb_bind(_rb_sys.modules[__name__], ()))

# RB_BIND:0116
from application.components.registration import install_0116 as _rb_install_0116
_rb_registration_v2_patient_name = _rb_install_0116(_rb_bind(_rb_sys.modules[__name__], ('_rb_registration_v2_name_parts',)))

# RB_BIND:0117
from application.components.registration import install_0117 as _rb_install_0117
_rb_registration_v2_worklist_name = _rb_install_0117(_rb_bind(_rb_sys.modules[__name__], ('_rb_registration_v2_patient_name', 'normalized_worklist_name')))

# RB_BIND:0118
from application.components.registration import install_0118 as _rb_install_0118
_rb_registration_v2_dob = _rb_install_0118(_rb_bind(_rb_sys.modules[__name__], ('resolve_patient_dob',)))

# RB_BIND:0119
from application.components.registration import install_0119 as _rb_install_0119
_rb_registration_v2_create_patient = _rb_install_0119(_rb_bind(_rb_sys.modules[__name__], ('Patient', 'PatientDemographicsExtra', '_rb_registration_v2_name_parts', '_rb_registration_v2_patient_name', '_rb_registration_v2_worklist_name', 'audit', 'db', 'find_duplicate_patients', 'generate_patient_id', 'next_patient_id', 'resolve_patient_dob', 'secrets')))

# RB_BIND:0120
from web.components.shared import install_0120 as _rb_install_0120
_rb_calendar_redirect = _rb_install_0120(_rb_bind(_rb_sys.modules[__name__], ('redirect', 'url_for')))

# RB_BIND:0121
from application.components.shared import install_0121 as _rb_install_0121
new_accession = _rb_install_0121(_rb_bind(_rb_sys.modules[__name__], ('datetime', 'secrets')))

# RB_BIND:0122
from application.components.registration import install_0122 as _rb_install_0122
generate_patient_id = _rb_install_0122(_rb_bind(_rb_sys.modules[__name__], ()))

# RB_BIND:0123
from application.components.imaging import install_0123 as _rb_install_0123
orthanc_ok = _rb_install_0123(_rb_bind(_rb_sys.modules[__name__], ('AUTH', 'ORTHANC', 'requests')))

# RB_BIND:0124
from application.components.shared import install_0124 as _rb_install_0124
route_device = _rb_install_0124(_rb_bind(_rb_sys.modules[__name__], ('Device',)))

# RB_BIND:0125
from application.components.shared import install_0125 as _rb_install_0125
price_order = _rb_install_0125(_rb_bind(_rb_sys.modules[__name__], ('ContractPrice',)))

# RB_BIND:0126
from application.components.imaging import install_0126 as _rb_install_0126
send_mwl = _rb_install_0126(_rb_bind(_rb_sys.modules[__name__], ('AUTH', 'ORTHANC', 'dicom_worklist_patient_name', 'requests')))


# RB_BIND:0127
from application.components.reporting import install_0127 as _rb_install_0127
get_report = _rb_install_0127(_rb_bind(_rb_sys.modules[__name__], ('Report', 'db')))

# RB_BIND:0128
from application.components.reporting import install_0128 as _rb_install_0128
report_history = _rb_install_0128(_rb_bind(_rb_sys.modules[__name__], ('ReportVersion',)))

# RB_BIND:0129
from application.components.shared import install_0129 as _rb_install_0129
next_version = _rb_install_0129(_rb_bind(_rb_sys.modules[__name__], ('ReportVersion',)))

# RB_BIND:0130
from application.components.shared import install_0130 as _rb_install_0130
add_version = _rb_install_0130(_rb_bind(_rb_sys.modules[__name__], ('ReportVersion', 'current_user', 'db', 'next_version', 'os')))

# RB_BIND:0131
from application.components.shared import install_0131 as _rb_install_0131
make_docx = _rb_install_0131(_rb_bind(_rb_sys.modules[__name__], ('Document', 'io')))

# RB_BIND:0132
from application.components.shared import install_0132 as _rb_install_0132
make_pdf = _rb_install_0132(_rb_bind(_rb_sys.modules[__name__], ('A4', 'canvas')))

# RB_BIND:0133
from application.components.shared import install_0133 as _rb_install_0133
extract_docx_sections = _rb_install_0133(_rb_bind(_rb_sys.modules[__name__], ('Document',)))

# RB_BIND:0134
from application.components.shared import install_0134 as _rb_install_0134
save_uploaded_word = _rb_install_0134(_rb_bind(_rb_sys.modules[__name__], ()))

# RB_BIND:0135
from application.components.shared import install_0135 as _rb_install_0135
case_json = _rb_install_0135(_rb_bind(_rb_sys.modules[__name__], ('ORTHANC_PUBLIC_URL', 'RIS_PUBLIC_URL', '_rb_latest_report_version')))


# RB_BIND:0136
from web.routes.imaging import install_0136 as _rb_install_0136
inject_pacs_helpers = _rb_install_0136(_rb_bind(_rb_sys.modules[__name__], ('app', 'case_pacs_url')))

# RB_BIND:0137
from application.components.shared import install_0137 as _rb_install_0137
ui_phrase = _rb_install_0137(_rb_bind(_rb_sys.modules[__name__], ('current_ui_language',)))


_flask_flash = flash

# RB_BIND:0138
from application.components.shared import install_0138 as _rb_install_0138
flash = _rb_install_0138(_rb_bind(_rb_sys.modules[__name__], ('_flask_flash', 'ui_phrase')))


# Imported Jinja macros also need the explicit UI translator.
app.jinja_env.globals['ui_phrase'] = ui_phrase

# RB_BIND:0139
from web.routes.shared import install_0139 as _rb_install_0139
template_globals = _rb_install_0139(_rb_bind(_rb_sys.modules[__name__], ('ARABIC_NAME_MAP', 'RIS_PUBLIC_URL', 'app', 'current_ui_language', 'current_user', 'get_company_profile', 'get_user_profile', 'has_perm', 'orthanc_ok', 'procedure_display_name', 'quick_action_role', 'ribbon_actions_for_user', 'role_quick_actions', 'safe_url_for_endpoint', 'status_text', 'ui_text', 'unread_chat_count')))


# RB_BIND:0140
from application.components.shared import install_0140 as _rb_install_0140
inject_language = _rb_install_0140(_rb_bind(_rb_sys.modules[__name__], ('lang', 'tr')))

# RB_BIND:0141
from web.routes.shared import install_0141 as _rb_install_0141
set_language = _rb_install_0141(_rb_bind(_rb_sys.modules[__name__], ('UserProfile', 'app', 'current_user', 'db', 'login_required', 'redirect', 'request', 'session', 'url_for')))


# RB_BIND:0142
from web.routes.shared import install_0142 as _rb_install_0142
friendly_401 = _rb_install_0142(_rb_bind(_rb_sys.modules[__name__], ('_error_id', 'app', 'render_template')))


# RB_BIND:0143
from web.routes.shared import install_0143 as _rb_install_0143
friendly_403 = _rb_install_0143(_rb_bind(_rb_sys.modules[__name__], ('_error_id', 'app', 'render_template')))


# RB_BIND:0144
from web.routes.shared import install_0144 as _rb_install_0144
friendly_404 = _rb_install_0144(_rb_bind(_rb_sys.modules[__name__], ('_error_id', 'app', 'render_template')))

# RB_BIND:0145
from web.routes.shared import install_0145 as _rb_install_0145
friendly_500 = _rb_install_0145(_rb_bind(_rb_sys.modules[__name__], ('_error_id', 'app', 'db', 'render_template')))


# RB_BIND:0146
from web.components.identity import install_0146 as _rb_install_0146
_safe_login_next = _rb_install_0146(_rb_bind(_rb_sys.modules[__name__], ('url_for',)))


# RB_BIND:0147
from web.routes.identity import install_0147 as _rb_install_0147
login = _rb_install_0147(_rb_bind(_rb_sys.modules[__name__], ('AuditLog', 'User', '_safe_login_next', '_security_audit_ip', 'app', 'audit', 'check_password_hash', 'db', 'flash', 'redirect', 'render_template', 'request', 'session')))

# RB_BIND:0148
from web.routes.identity import install_0148 as _rb_install_0148
logout = _rb_install_0148(_rb_bind(_rb_sys.modules[__name__], ('app', 'redirect', 'session', 'url_for')))

# RB_BIND:0149
from web.routes.shared import install_0149 as _rb_install_0149
dashboard = _rb_install_0149(_rb_bind(_rb_sys.modules[__name__], ('Order', 'Patient', 'app', 'date', 'login_required', 'render_template')))

# RB_BIND:0150
from web.routes.registration import install_0150 as _rb_install_0150
patients = _rb_install_0150(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Patient', 'Procedure', 'ProcedureBranch', 'app', 'or_', 'render_template', 'request', 'require_perm')))

# RB_BIND:0151
from web.routes.registration import install_0151 as _rb_install_0151
new_patient = _rb_install_0151(_rb_bind(_rb_sys.modules[__name__], ('Patient', 'PatientDemographicsExtra', 'app', 'audit', 'db', 'find_duplicate_patients', 'flash', 'generate_patient_id', 'normalized_worklist_name', 'redirect', 'render_template', 'request', 'require_perm', 'resolve_patient_dob', 'secrets', 'url_for')))

# RB_BIND:0152
from web.routes.shared import install_0152 as _rb_install_0152
orders = _rb_install_0152(_rb_bind(_rb_sys.modules[__name__], ('Order', 'Patient', 'app', 'or_', 'render_template', 'request', 'require_perm')))

# RB_BIND:0153
from web.routes.shared import install_0153 as _rb_install_0153
new_order = _rb_install_0153(_rb_bind(_rb_sys.modules[__name__], ('ATTACHMENT_DIR', 'Attachment', 'Branch', 'ClinicalNote', 'Contract', 'InsuranceApprovalEvent', 'InsuranceApprovalRequest', 'Order', 'OrderConsumable', 'Patient', 'PatientDemographicsExtra', 'Payment', 'Procedure', 'ProcedureConsumable', 'ReferringDoctor', 'SmartIntakeSession', 'abort', 'app', 'audit', 'booking_restriction_reason', 'current_user', 'datetime', 'db', 'ensure_journey', 'find_duplicate_patients', 'flash', 'generated_attachment_filename', 'has_perm', 'json', 'new_accession', 'next_patient_id', 'normalized_worklist_name', 'os', 'price_order', 'procedure_available_in_branch', 'redirect', 'registration_rule_map', 'registration_validation_error', 'render_registration_form', 'request', 'require_perm', 'resolve_patient_dob', 'route_device', 'secrets', 'url_for', 'validate_registration_value')))



# RB_BIND:0154
from web.components.registration import install_0154 as _rb_install_0154
api_registration_procedures = _rb_install_0154(_rb_bind(_rb_sys.modules[__name__], ('allowed_procedures_for_branch', 'jsonify', 'login_required', 'request')))

# RB_BIND:0155
from web.components.shared import install_0155 as _rb_install_0155
api_pricing = _rb_install_0155(_rb_bind(_rb_sys.modules[__name__], ('Contract', 'Procedure', 'jsonify', 'login_required', 'price_order', 'request')))

# RB_BIND:0156
from web.routes.identity import install_0156 as _rb_install_0156
change_password = _rb_install_0156(_rb_bind(_rb_sys.modules[__name__], ('app', 'audit', 'check_password_hash', 'current_user', 'db', 'flash', 'generate_password_hash', 'login_required', 'redirect', 'render_template', 'request', 'url_for')))



# RB_BIND:0157
from web.routes.registration import install_0157 as _rb_install_0157
admin_insurance_approvals = _rb_install_0157(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Contract', 'InsurancePortalAccount', 'app', 'audit', 'datetime', 'db', 'flash', 'json', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))


# RB_BIND:0158
from web.routes.registration import install_0158 as _rb_install_0158
admin_insurance_approval_edit = _rb_install_0158(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Contract', 'InsurancePortalAccount', 'app', 'render_template', 'require_perm')))


# RB_BIND:0159
from web.routes.registration import install_0159 as _rb_install_0159
admin_insurance_approval_toggle = _rb_install_0159(_rb_bind(_rb_sys.modules[__name__], ('InsurancePortalAccount', 'app', 'audit', 'datetime', 'db', 'redirect', 'require_perm', 'url_for')))


# RB_BIND:0160
from web.routes.registration import install_0160 as _rb_install_0160
admin_insurance_approval_test = _rb_install_0160(_rb_bind(_rb_sys.modules[__name__], ('InsurancePortalAccount', 'app', 'datetime', 'db', 'flash', 'redirect', 'require_perm', 'url_for')))


# RB_BIND:0161
from web.routes.registration import install_0161 as _rb_install_0161
insurance_approvals_dashboard = _rb_install_0161(_rb_bind(_rb_sys.modules[__name__], ('InsuranceApprovalRequest', 'Order', 'Patient', 'app', 'datetime', 'render_template', 'request', 'require_perm')))


# RB_BIND:0162
from web.routes.registration import install_0162 as _rb_install_0162
order_detail = _rb_install_0162(_rb_bind(_rb_sys.modules[__name__], ('Attachment', 'DeliveryEvent', 'InsuranceApprovalRequest', 'Order', 'OrderConsumable', 'Payment', 'Receipt', 'User', 'app', 'approved_order_discounts', 'default_warehouse_for_order', 'discount_policy_for', 'effective_patient_due', 'order_balance', 'order_consumable_cost', 'performing_doctor_for_order', 'performing_doctor_share_for_order', 'render_template', 'require_perm', 'work_order_gate')))


# RB_BIND:0163
from web.routes.registration import install_0163 as _rb_install_0163
insurance_approval_start = _rb_install_0163(_rb_bind(_rb_sys.modules[__name__], ('InsuranceApprovalEvent', 'InsuranceApprovalRequest', 'Order', 'app', 'audit', 'current_user', 'datetime', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_BIND:0164
from web.routes.registration import install_0164 as _rb_install_0164
insurance_approval_update = _rb_install_0164(_rb_bind(_rb_sys.modules[__name__], ('InsuranceApprovalEvent', 'InsuranceApprovalRequest', 'Order', 'app', 'audit', 'current_user', 'datetime', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_WORK_ORDER_GATE_V5_CANONICAL
# RB_BIND:0165
from web.routes.registration import install_0165 as _rb_install_0165
convert_to_work_order = _rb_install_0165(_rb_bind(_rb_sys.modules[__name__], ('EventOutbox', 'Order', 'Patient', 'app', 'audit', 'blockers_text', 'datetime', 'db', 'enqueue_event', 'flash', 'order_balance', 'redirect', 'require_perm', 'send_mwl', 'url_for', 'work_order_gate')))


# RB_BIND:0166
from web.routes.imaging import install_0166 as _rb_install_0166
resend_mwl = _rb_install_0166(_rb_bind(_rb_sys.modules[__name__], ('Order', 'app', 'audit', 'blockers_text', 'db', 'flash', 'redirect', 'require_perm', 'resend_mwl_gate', 'route_device', 'send_mwl', 'url_for')))

# RB_BIND:0167
from web.routes.finance import install_0167 as _rb_install_0167
add_payment = _rb_install_0167(_rb_bind(_rb_sys.modules[__name__], ('app', 'receive_payment', 'require_perm')))


# RB_BIND:0168
from web.routes.clinical import install_0168 as _rb_install_0168
nursing_worklist = _rb_install_0168(_rb_bind(_rb_sys.modules[__name__], ('NursingConsumableEntry', 'Order', 'Patient', 'ProcedureConsumable', '__name__', '_order_attachment_count', '_reporting_support_sys', '_saved_views_for', 'app', 'clinical_image_only_url', 'datetime', 'db', 'render_template', 'request', 'require_perm', 'timedelta')))

# RB_BIND:0169
from web.routes.inventory import install_0169 as _rb_install_0169
nursing_case = _rb_install_0169(_rb_bind(_rb_sys.modules[__name__], ('Attachment', 'ClinicalNote', 'Consumable', 'InventoryItem', 'NursingConsumableEntry', 'Order', 'OrderConsumable', 'ProcedureConsumable', '_record_stock_movement', 'app', 'audit', 'clinical_image_only_url', 'current_user', 'datetime', 'db', 'default_warehouse_for_order', 'flash', 'inventory_available_qty', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0170
from web.routes.inventory import install_0170 as _rb_install_0170
nursing_consumable_correct = _rb_install_0170(_rb_bind(_rb_sys.modules[__name__], ('Consumable', 'NursingConsumableEntry', 'Order', 'OrderConsumable', '_record_stock_movement', 'app', 'audit', 'current_user', 'datetime', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0171
from web.routes.inventory import install_0171 as _rb_install_0171
procedure_consumables = _rb_install_0171(_rb_bind(_rb_sys.modules[__name__], ('Consumable', 'Procedure', 'ProcedureConsumable', 'app', 'audit', 'db', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0172
from web.routes.inventory import install_0172 as _rb_install_0172
procedure_consumable_delete = _rb_install_0172(_rb_bind(_rb_sys.modules[__name__], ('ProcedureConsumable', 'app', 'audit', 'db', 'flash', 'redirect', 'require_perm', 'url_for')))


# RB_BIND:0173
from web.routes.registration import install_0173 as _rb_install_0173
patient_portal = _rb_install_0173(_rb_bind(_rb_sys.modules[__name__], ('PORTAL_REQUIRE_PHONE_VERIFY', 'PortalAccessToken', 'abort', 'app', 'case_pacs_url', 'datetime', 'flash', 're', 'redirect', 'render_template', 'request', 'session', 'url_for')))

# RB_BIND:0174
from web.routes.reporting import install_0174 as _rb_install_0174
patient_portal_report = _rb_install_0174(_rb_bind(_rb_sys.modules[__name__], ('PortalAccessToken', '_rb_final_revision_document', '_rb_revision_pdf_download', 'abort', 'app', 'datetime', 'os', 'send_file')))

# RB_BIND:0175
from web.routes.communications import install_0175 as _rb_install_0175
admin_notifications = _rb_install_0175(_rb_bind(_rb_sys.modules[__name__], ('NotificationEvent', 'app', 'render_template', 'require_perm')))


# RB_BIND:0176
from web.routes.communications import install_0176 as _rb_install_0176
admin_notifications_process = _rb_install_0176(_rb_bind(_rb_sys.modules[__name__], ('app', 'flash', 'process_notification_queue', 'redirect', 'require_perm', 'url_for')))

# RB_BIND:0177
from web.routes.communications import install_0177 as _rb_install_0177
admin_notification_retry = _rb_install_0177(_rb_bind(_rb_sys.modules[__name__], ('NotificationEvent', 'app', 'db', 'flash', 'redirect', 'require_perm', 'url_for')))

# RB_BIND:0178
from web.routes.imaging import install_0178 as _rb_install_0178
order_study_thumbnail = _rb_install_0178(_rb_bind(_rb_sys.modules[__name__], ('Order', 'PACS_ORTHANC_URL', '_orthanc_auth', '_orthanc_series_detail', '_orthanc_study_detail', 'abort', 'app', 'requests', 'require_perm')))


# RB_UNIFIED_TECHNICIAN_CASE_WORKSPACE_V1
# RB_BIND:0179
from web.routes.clinical import install_0179 as _rb_install_0179
technician_case_workspace = _rb_install_0179(_rb_bind(_rb_sys.modules[__name__], ('Attachment', 'AuditLog', 'CaseReadiness', 'ClinicalNote', 'CorrectionRequest', 'InventoryItem', 'Order', 'OrderConsumableUsage', '__name__', '_reporting_support_sys', 'app', 'case_return_snapshot', 'clinical_image_only_url', 'default_warehouse_for_order', 'inventory_available_qty', 'order_consumable_cost', 'quality_rule_for', 'render_template', 'require_perm', 'technician_acquisition_notes', 'technician_minimum_requirement')))


# RB_BIND:0180
from web.routes.clinical import install_0180 as _rb_install_0180
technician_case_note = _rb_install_0180(_rb_bind(_rb_sys.modules[__name__], ('ClinicalNote', 'Order', 'app', 'audit', 'current_user', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_BIND:0181
from web.routes.inventory import install_0181 as _rb_install_0181
technician_case_consume = _rb_install_0181(_rb_bind(_rb_sys.modules[__name__], ('InventoryItem', 'Order', '__name__', '_reporting_support_sys', 'app', 'audit', 'consume_inventory', 'current_user', 'db', 'default_warehouse_for_order', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_BIND:0182
from web.routes.clinical import install_0182 as _rb_install_0182
technician_worklist = _rb_install_0182(_rb_bind(_rb_sys.modules[__name__], ('__name__', '_reporting_support_sys', 'app', 'require_perm')))


# RB_BIND:0183
from application.components.imaging import install_0183 as _rb_install_0183
technician_acquisition_notes = _rb_install_0183(_rb_bind(_rb_sys.modules[__name__], ('StudioTemplate',)))


# RB_BIND:0184
from application.components.clinical import install_0184 as _rb_install_0184
technician_minimum_requirement = _rb_install_0184(_rb_bind(_rb_sys.modules[__name__], ('CaseReadinessOverride',)))


# RB_BIND:0185
from application.components.clinical import install_0185 as _rb_install_0185
technician_completion_evidence = _rb_install_0185(_rb_bind(_rb_sys.modules[__name__], ('ReadinessSeriesAlias', '__name__', '_reporting_support_sys', 'evaluate_case_readiness', 'pacs_demographic_gate', 'pacs_readiness_evidence', 'quality_rule_for', 'technician_minimum_requirement')))


# RB_BIND:0186
from web.routes.clinical import install_0186 as _rb_install_0186
technician_check_images = _rb_install_0186(_rb_bind(_rb_sys.modules[__name__], ('Order', '__name__', '_reporting_support_sys', 'app', 'audit', 'db', 'evaluate_case_readiness', 'flash', 'json', 'redirect', 'require_perm', 'technician_completion_evidence', 'url_for')))


# RB_BIND:0187
from web.components.shared import install_0187 as _rb_install_0187
flash_exam_transition_blocked = _rb_install_0187(_rb_bind(_rb_sys.modules[__name__], ('flash',)))


# RB_BIND:0188
from web.components.shared import install_0188 as _rb_install_0188
_apply_exam_transition = _rb_install_0188(_rb_bind(_rb_sys.modules[__name__], ('__name__', '_reporting_support_sys', 'audit', 'current_user', 'datetime', 'exam_transition_gate', 'flash', 'json', 'order_status_for_exam_transition', 'pacs_demographic_gate', 'technician_completion_evidence')))


# RB_BIND:0189
from web.routes.clinical import install_0189 as _rb_install_0189
technician_status = _rb_install_0189(_rb_bind(_rb_sys.modules[__name__], ('AuditLog', 'NursingConsumableEntry', 'Order', '__name__', '_apply_exam_transition', '_reporting_support_sys', 'abort', 'app', 'audit', 'current_user', 'db', 'flash', 'flash_exam_transition_blocked', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_BIND:0190
from web.components.reporting import install_0190 as _rb_install_0190
reporting_agent_download = _rb_install_0190(_rb_bind(_rb_sys.modules[__name__], ('__file__', 'abort', 'login_required', 'os', 'send_file')))



# RB_BIND:0191
from web.components.reporting import install_0191 as _rb_install_0191
api_reporting_health = _rb_install_0191(_rb_bind(_rb_sys.modules[__name__], ('APP_VERSION', 'ATTACHMENT_DIR', 'REPORT_DIR', '__file__', 'db', 'jsonify', 'orthanc_ok', 'os', 'text')))



# RB_BIND:0192
from web.routes.clinical import install_0192 as _rb_install_0192
order_exam_status = _rb_install_0192(_rb_bind(_rb_sys.modules[__name__], ('Order', '__name__', '_apply_exam_transition', '_reporting_support_sys', 'app', 'audit', 'coordinate_post_exam_assignment', 'current_user', 'db', 'enqueue_event', 'evaluate_case_readiness', 'flash', 'flash_exam_transition_blocked', 'pacs_lookup_order', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0193
from web.routes.clinical import install_0193 as _rb_install_0193
order_exam_status_override = _rb_install_0193(_rb_bind(_rb_sys.modules[__name__], ('EXAM_STATUS_ORDER', 'Order', '__name__', '_apply_exam_transition', '_reporting_support_sys', 'abort', 'app', 'audit', 'current_user', 'db', 'enqueue_event', 'flash', 'flash_exam_transition_blocked', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_BIND:0194
from web.routes.imaging import install_0194 as _rb_install_0194
pacs_home = _rb_install_0194(_rb_bind(_rb_sys.modules[__name__], ('app', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0195
from web.routes.imaging import install_0195 as _rb_install_0195
pacs_tracking = _rb_install_0195(_rb_bind(_rb_sys.modules[__name__], ('Order', 'app', 'render_template', 'require_perm')))

# RB_BIND:0196
from web.routes.imaging import install_0196 as _rb_install_0196
pacs_tracking_refresh = _rb_install_0196(_rb_bind(_rb_sys.modules[__name__], ('Order', 'app', 'coordinate_post_exam_assignment', 'current_user', 'db', 'evaluate_case_readiness', 'jsonify', 'pacs_lookup_order', 'require_perm')))


# RB_BIND:0197
from web.routes.assignment import install_0197 as _rb_install_0197
doctor_presence_heartbeat = _rb_install_0197(_rb_bind(_rb_sys.modules[__name__], ('abort', 'app', 'current_user', 'datetime', 'db', 'doctor_availability', 'jsonify', 'login_required')))


# RB_BIND:0198
from web.routes.reporting import install_0198 as _rb_install_0198
doctor_workspace = _rb_install_0198(_rb_bind(_rb_sys.modules[__name__], ('_saved_views_for', 'abort', 'app', 'current_user', 'doctor_workspace_user', 'render_template', 'require_perm')))


# RB_BIND:0199
from web.routes.reporting import install_0199 as _rb_install_0199
doctor_workspace_data = _rb_install_0199(_rb_bind(_rb_sys.modules[__name__], ('AssignmentOffer', 'ClinicalAssessment', 'ExternalIntakeCase', 'Order', 'OrderPerformingDoctor', 'PatientDocument', 'PriorStudySelection', 'UserProfile', 'abort', 'app', 'assignment_visit_summary', 'current_user', 'datetime', 'doctor_workspace_user', 'expire_assignment_offers', 'jsonify', 'patient_worklist_display_name', 'report_sla_snapshot', 'require_perm', 'url_for', 'user_profile_settings')))


DOCTOR_WORKLIST_COLUMN_KEYS=(
    'priority','accession','patient_id','patient','patient_worklist','dob','sex',
    'procedure','modality','body_part','contrast','branch','device','room',
    'station_ae','referring_doctor','assigned_doctor','performing_doctor',
    'clinical_indication','clinical_history','clinical_diagnosis','allergy_notes',
    'status','exam_status','pacs_status','images','attachment_count','prior_count',
    'study_instance_uid','report_status','reported_by','verified_by','verified_at',
    'second_verified_by','second_verified_at','finalized_by','finalized_at',
    'report_updated_at','report_delivered','images_delivered','sla_status',
    'sla_elapsed_minutes','sla_remaining_minutes','scheduled_at','assigned_at',
    'exam_started_at','exam_completed_at','pacs_last_seen_at','center','action'
)


# RB_BIND:0200
from application.components.shared import install_0200 as _rb_install_0200
normalize_doctor_worklist_layout = _rb_install_0200(_rb_bind(_rb_sys.modules[__name__], ('DOCTOR_WORKLIST_COLUMN_KEYS',)))


# RB_BIND:0201
from web.routes.reporting import install_0201 as _rb_install_0201
doctor_workspace_column_layout_save = _rb_install_0201(_rb_bind(_rb_sys.modules[__name__], ('abort', 'app', 'current_user', 'db', 'doctor_workspace_user', 'get_user_profile', 'json', 'jsonify', 'normalize_doctor_worklist_layout', 'request', 'require_perm', 'user_profile_settings')))


# RB_BIND:0202
from application.components.shared import install_0202 as _rb_install_0202
doctor_workspace_user = _rb_install_0202(_rb_bind(_rb_sys.modules[__name__], ()))


# RB_BIND:0203
from application.components.shared import install_0203 as _rb_install_0203
doctor_case_access_allowed = _rb_install_0203(_rb_bind(_rb_sys.modules[__name__], ('__name__', '_reporting_support_sys', 'doctor_workspace_user')))


# RB_BIND:0204
from web.routes.reporting import install_0204 as _rb_install_0204
doctor_report_ownership_guard = _rb_install_0204(_rb_bind(_rb_sys.modules[__name__], ('Order', 'abort', 'app', 'current_user', 'db', 'doctor_case_access_allowed', 'request')))


# RB_BIND:0205
from web.routes.reporting import install_0205 as _rb_install_0205
doctor_workspace_open = _rb_install_0205(_rb_bind(_rb_sys.modules[__name__], ('Order', 'abort', 'app', 'current_user', 'doctor_case_access_allowed', 'open_image_viewer_case', 'require_perm')))


# RB_BIND:0206
from web.routes.assignment import install_0206 as _rb_install_0206
admin_assignment_visits = _rb_install_0206(_rb_bind(_rb_sys.modules[__name__], ('AssignmentVisit', 'Order', 'abort', 'app', 'assignable_doctors', 'audit', 'db', 'flash', 'func', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))


# RB_BIND:0207
from web.routes.registration import install_0207 as _rb_install_0207
api_registration_visits = _rb_install_0207(_rb_bind(_rb_sys.modules[__name__], ('AssignmentVisit', '_rb_v5462_allow_registration', 'app', 'datetime', 'jsonify', 'login_required', 'request')))

# RB_BIND:0208
from web.routes.assignment import install_0208 as _rb_install_0208
complete_visit_registration = _rb_install_0208(_rb_bind(_rb_sys.modules[__name__], ('AssignmentVisit', 'Order', '_rb_v5462_allow_registration', 'app', 'audit', 'current_user', 'db', 'flash', 'login_required', 'offer_next_doctor', 'redirect', 'text', 'url_for')))

# RB_BIND:0209
from web.routes.assignment import install_0209 as _rb_install_0209
admin_assignment_monitor = _rb_install_0209(_rb_bind(_rb_sys.modules[__name__], ('app', 'render_template', 'require_perm')))


# RB_BIND:0210
from web.routes.assignment import install_0210 as _rb_install_0210
coordinator_assignments = _rb_install_0210(_rb_bind(_rb_sys.modules[__name__], ('abort', 'app', 'current_user', 'has_perm', 'redirect', 'render_template', 'request', 'url_for')))



# RB_BIND:0211
from web.routes.assignment import install_0211 as _rb_install_0211
admin_assignment_distribution = _rb_install_0211(_rb_bind(_rb_sys.modules[__name__], ('AssignmentDistributionPolicy', 'abort', 'app', 'assignable_doctors', 'audit', 'current_user', 'db', 'flash', 'json', 're', 'redirect', 'render_template', 'request', 'require_perm', 'text', 'url_for')))


# RB_BIND:0212
from web.routes.reporting import install_0212 as _rb_install_0212
doctor_reading_start = _rb_install_0212(_rb_bind(_rb_sys.modules[__name__], ('Order', 'abort', 'app', 'current_user', 'db', 'doctor_case_access_allowed', 'doctor_workspace_user', 'redirect', 'require_perm', 'url_for')))


# RB_BIND:0213
from web.routes.reporting import install_0213 as _rb_install_0213
reports = _rb_install_0213(_rb_bind(_rb_sys.modules[__name__], ('Order', 'app', 'render_template', 'require_perm')))

# RB_BIND:0214
from web.routes.reporting import install_0214 as _rb_install_0214
admin_report_reopen = _rb_install_0214(_rb_bind(_rb_sys.modules[__name__], ('Order', 'abort', 'app', 'audit_as_user', 'current_user', 'db', 'flash', 'get_report', 'redirect', 'require_perm', 'url_for')))




# RB_BIND:0215
from web.routes.reporting import install_0215 as _rb_install_0215
report_lease_acquire = _rb_install_0215(_rb_bind(_rb_sys.modules[__name__], ('Order', 'ReportEditControlError', 'acquire_or_renew_lease', 'app', 'audit', 'current_user', 'db', 'get_report', 'jsonify', 'request', 'require_perm')))



# RB_BIND:0216
from web.routes.reporting import install_0216 as _rb_install_0216
report_lease_release = _rb_install_0216(_rb_bind(_rb_sys.modules[__name__], ('Order', 'app', 'audit', 'current_user', 'db', 'jsonify', 'release_lease', 'request', 'require_perm')))


# RB_BIND:0217
from web.routes.reporting import install_0217 as _rb_install_0217
report_editor = _rb_install_0217(_rb_bind(_rb_sys.modules[__name__], ('DoctorWordTemplate', 'Order', 'REPORT_DIR', 'Report', 'ReportEditControlError', 'ReportVersion', '_error_id', 'abort', 'allowed_report_transitions', 'app', 'assert_lease_owner', 'assert_no_foreign_active_lease', 'assert_not_final', 'assert_role_allowed', 'audit', 'case_pacs_url', 'current_user', 'db', 'doctor_case_access_allowed', 'enqueue_event', 'flash', 'get_report', 'make_docx', 'or_', 'os', 'pacs_demographic_gate', 'queue_final_delivery', 'redirect', 'render_template', 'request', 'require_perm', 'shutil', 'tempfile', 'url_for')))


# RB_BIND:0218
from web.routes.reporting import install_0218 as _rb_install_0218
word_template = _rb_install_0218(_rb_bind(_rb_sys.modules[__name__], ('Document', 'Inches', 'Order', '_error_id', '_next_report_version', '_rb_authoritative_docx_download', 'app', 'build_report_filename', 'flash', 'get_report', 'io', 'redirect', 'require_perm', 'send_file', 'url_for')))




RADBRIDGE_REPORTING_TOKEN = os.getenv('RADBRIDGE_REPORTING_TOKEN', '').strip()
REPORTING_CASE_TOKEN_TTL_SECONDS = int(os.getenv('REPORTING_CASE_TOKEN_TTL_SECONDS', '28800') or 28800)


# RB_BIND:0219
from application.components.reporting import install_0219 as _rb_install_0219
_reporting_serializer = _rb_install_0219(_rb_bind(_rb_sys.modules[__name__], ('URLSafeTimedSerializer', 'app')))


# RB_BIND:0220
from application.components.reporting import install_0220 as _rb_install_0220
reporting_case_token = _rb_install_0220(_rb_bind(_rb_sys.modules[__name__], ('_reporting_serializer', 'current_user', 'has_perm')))


app.jinja_env.globals['reporting_case_token'] = reporting_case_token


# RB_BIND:0221
from application.components.reporting import install_0221 as _rb_install_0221
_valid_reporting_case_token = _rb_install_0221(_rb_bind(_rb_sys.modules[__name__], ('BadSignature', 'REPORTING_CASE_TOKEN_TTL_SECONDS', 'SignatureExpired', '_reporting_serializer')))



# RB_BIND:0222
from web.components.reporting import install_0222 as _rb_install_0222
reporting_actor = _rb_install_0222(_rb_bind(_rb_sys.modules[__name__], ('BadSignature', 'REPORTING_CASE_TOKEN_TTL_SECONDS', 'SignatureExpired', 'User', '_reporting_serializer', 'current_user', 'db', 'request')))


# RB_BIND:0223
from web.components.reporting import install_0223 as _rb_install_0223
_reporting_api_authorized = _rb_install_0223(_rb_bind(_rb_sys.modules[__name__], ('RADBRIDGE_REPORTING_TOKEN', '_valid_reporting_case_token', 'current_user', 'has_perm', 'hmac', 'request')))


# RB_BIND:0224
from web.components.reporting import install_0224 as _rb_install_0224
require_reporting_api = _rb_install_0224(_rb_bind(_rb_sys.modules[__name__], ('_reporting_api_authorized', 'current_user', 'jsonify', 'wraps')))

# RB_BIND:0225
from web.components.reporting import install_0225 as _rb_install_0225
word_addin_taskpane = _rb_install_0225(_rb_bind(_rb_sys.modules[__name__], ('render_template',)))


# RB_BIND:0226
from web.components.reporting import install_0226 as _rb_install_0226
word_addin_manifest = _rb_install_0226(_rb_bind(_rb_sys.modules[__name__], ('RIS_PUBLIC_URL', 'Response', 'render_template')))


# RB_BIND:0227
from application.components.reporting import install_0227 as _rb_install_0227
word_addin_taskpane_js = _rb_install_0227(_rb_bind(_rb_sys.modules[__name__], ('app',)))



# RB_BIND:0228
from web.routes.reporting import install_0228 as _rb_install_0228
api_reporting_case = _rb_install_0228(_rb_bind(_rb_sys.modules[__name__], ('Order', 'app', 'case_json', 'get_report', 'jsonify', 'require_reporting_api')))


# RB_BIND:0229
from web.routes.reporting import install_0229 as _rb_install_0229
api_reporting_lease_acquire = _rb_install_0229(_rb_bind(_rb_sys.modules[__name__], ('Order', 'ReportEditControlError', 'acquire_or_renew_lease', 'app', 'audit_as_user', 'db', 'get_report', 'jsonify', 'reporting_actor', 'request', 'require_reporting_api')))



# RB_BIND:0230
from web.routes.reporting import install_0230 as _rb_install_0230
api_reporting_lease_release = _rb_install_0230(_rb_bind(_rb_sys.modules[__name__], ('Order', 'app', 'audit_as_user', 'db', 'jsonify', 'release_lease', 'reporting_actor', 'request', 'require_reporting_api')))


# RB_BIND:0231
from web.routes.reporting import install_0231 as _rb_install_0231
api_reporting_draft = _rb_install_0231(_rb_bind(_rb_sys.modules[__name__], ('Order', 'REPORT_DIR', 'Report', 'ReportEditControlError', 'ReportVersion', '_ohif_validate_word_upload', 'app', 'assert_lease_owner', 'assert_not_final', 'assert_role_allowed', 'audit', 'audit_as_user', 'db', 'extract_docx_sections', 'get_report', 'jsonify', 'reporting_actor', 'request', 'require_reporting_api')))

# RB_BIND:0232
from web.routes.reporting import install_0232 as _rb_install_0232
api_reporting_verify = _rb_install_0232(_rb_bind(_rb_sys.modules[__name__], ('Order', 'Report', 'ReportEditControlError', 'ReportVersion', '_rb_latest_report_version', 'app', 'assert_no_foreign_active_lease', 'assert_not_final', 'assert_role_allowed', 'audit', 'audit_as_user', 'db', 'get_report', 'jsonify', 'pacs_demographic_gate', 'reporting_actor', 'request', 'require_reporting_api')))


# RB_BIND:0233
from web.routes.reporting import install_0233 as _rb_install_0233
api_reporting_finalize = _rb_install_0233(_rb_bind(_rb_sys.modules[__name__], ('Order', 'Report', 'ReportEditControlError', 'ReportVersion', '_rb_latest_report_version', 'app', 'assert_no_foreign_active_lease', 'assert_not_final', 'assert_role_allowed', 'audit', 'audit_as_user', 'db', 'enqueue_event', 'get_report', 'jsonify', 'pacs_demographic_gate', 'queue_final_delivery', 'reporting_actor', 'request', 'require_reporting_api')))


# ---------------- Appointments / Referrers / Analytics / Import-Export ----------------



APPOINTMENT_STATUS_COLORS={
 'WAITING':'gray','BOOKED':'blue','CONFIRMED':'cyan','ARRIVED':'yellow',
 'ORDERED':'purple','IN_PROGRESS':'orange','EXAM_COMPLETED':'green',
 'REPORTING':'navy','VERIFIED':'teal','FINAL':'darkgreen',
 'CANCELLED':'red','NO_SHOW':'salmon'
}

# RB_BIND:0234
from application.components.registration import install_0234 as _rb_install_0234
appointment_overlap = _rb_install_0234(_rb_bind(_rb_sys.modules[__name__], ('Appointment', 'timedelta')))


DEVICE_BOOKING_STATUSES={
    'ACTIVE':'Available',
    'OUT_OF_SERVICE':'Out of Service',
    'MAINTENANCE':'Maintenance',
    'DISABLED':'Booking Disabled'
}

# RB_BIND:0235
from application.components.shared import install_0235 as _rb_install_0235
_hhmm_to_minutes = _rb_install_0235(_rb_bind(_rb_sys.modules[__name__], ()))

# RB_BIND:0236
from application.components.registration import install_0236 as _rb_install_0236
device_booking_reason = _rb_install_0236(_rb_bind(_rb_sys.modules[__name__], ('DeviceBlock', 'datetime')))

# RB_BIND:0237
from application.components.registration import install_0237 as _rb_install_0237
device_available_for_booking = _rb_install_0237(_rb_bind(_rb_sys.modules[__name__], ('DEVICE_BOOKING_STATUSES', '_hhmm_to_minutes', 'appointment_overlap', 'device_booking_reason', 'timedelta')))

# RB_BIND:0238
from application.components.registration import install_0238 as _rb_install_0238
booking_v2_slots = _rb_install_0238(_rb_bind(_rb_sys.modules[__name__], ('_hhmm_to_minutes', 'datetime', 'timedelta')))

# RB_BIND:0239
from application.components.registration import install_0239 as _rb_install_0239
booking_v2_slot_state = _rb_install_0239(_rb_bind(_rb_sys.modules[__name__], ('device_available_for_booking',)))


# RB_BIND:0240
from application.components.registration import install_0240 as _rb_install_0240
procedure_available_in_branch = _rb_install_0240(_rb_bind(_rb_sys.modules[__name__], ('ProcedureBranch',)))

# RB_BIND:0241
from application.components.registration import install_0241 as _rb_install_0241
allowed_procedures_for_branch = _rb_install_0241(_rb_bind(_rb_sys.modules[__name__], ('Procedure', 'procedure_available_in_branch')))

# RB_BIND:0242
from application.components.registration import install_0242 as _rb_install_0242
booking_restriction_reason = _rb_install_0242(_rb_bind(_rb_sys.modules[__name__], ('booking_control_reason', 'device_available_for_booking', 'procedure_available_in_branch', 'timedelta')))

# RB_BIND:0243
from application.components.registration import install_0243 as _rb_install_0243
ensure_default_procedure_branch_mappings = _rb_install_0243(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Procedure', 'ProcedureBranch', 'db')))


DEFAULT_DICTIONARY=[
 ('patient','GLOBAL','Patient','المريض'),
 ('booking','SCHEDULING','Booking','الحجز'),
 ('procedure','GLOBAL','Examination','الفحص'),
 ('clinical_note','CLINICAL','Clinical Note','ملاحظات الفحص'),
 ('attachment','GLOBAL','Attachment','المرفقات'),
 ('technician','GLOBAL','Technician','الفني'),
 ('radiologist','GLOBAL','Radiologist','الطبيب'),
 ('branch','GLOBAL','Branch','الفرع'),
 ('referrer','GLOBAL','Referring Doctor','الطبيب المحول')
]

DEFAULT_RIBBON_ACTIONS=[
 ('new_patient','New Patient','مريض جديد','👤','new_patient','patients.manage','GLOBAL',10),
 ('new_visit','Patient & Visit','المريض والزيارة','➕','new_order','orders.create','GLOBAL',20),
 ('booking_v2','New Booking','حجز جديد','📅','schedule_calendar','appointments.manage','GLOBAL',30),
 ('cashier','Cashier','الكاشير','💳','cashier_workstation','cashier.workstation','GLOBAL',40),
 ('pacs_tracking','PACS Tracking','متابعة PACS','🩻','pacs_tracking','orders.view','GLOBAL',50),
 ('pacs_reconciliation','PACS Reconciliation','مطابقة PACS','🔗','pacs_reconciliation','orders.view','GLOBAL',60),
 ('doctor_workspace','Doctor Workspace','مساحة الطبيب','📄','doctor_workspace','reports.view','GLOBAL',70),
 ('booking_control','Booking Control','التحكم في الحجز','⛔','admin_booking_control','devices.manage','GLOBAL',80),
 ('dictionary','Dictionary','القاموس','Aع','admin_dictionary','users.view','GLOBAL',90)
]

BOOKING_BLOCK_REASONS=['MAINTENANCE','DEVICE_FAILURE','CALIBRATION','DOCTOR_UNAVAILABLE','TECHNICIAN_UNAVAILABLE','HOLIDAY','EMERGENCY_BLOCK','OTHER']

# RB_BIND:0244
from application.components.identity import install_0244 as _rb_install_0244
get_user_profile = _rb_install_0244(_rb_bind(_rb_sys.modules[__name__], ('UserProfile', 'current_user', 'db')))


# RB_BIND:0245
from application.components.identity import install_0245 as _rb_install_0245
user_profile_settings = _rb_install_0245(_rb_bind(_rb_sys.modules[__name__], ('json',)))


# RB_BIND:0246
from application.components.reporting import install_0246 as _rb_install_0246
user_reporting_preferences = _rb_install_0246(_rb_bind(_rb_sys.modules[__name__], ('user_profile_settings',)))


# RB_BIND:0247
from application.components.reporting import install_0247 as _rb_install_0247
set_user_reporting_preferences = _rb_install_0247(_rb_bind(_rb_sys.modules[__name__], ('json', 'user_profile_settings')))


# RB_BIND:0248
from application.components.shared import install_0248 as _rb_install_0248
dictionary_label = _rb_install_0248(_rb_bind(_rb_sys.modules[__name__], ('DictionaryEntry',)))

# RB_BIND:0249
from application.components.shared import install_0249 as _rb_install_0249
seed_dictionary_and_ribbon = _rb_install_0249(_rb_bind(_rb_sys.modules[__name__], ('DEFAULT_DICTIONARY', 'DEFAULT_RIBBON_ACTIONS', 'DictionaryEntry', 'RibbonAction', 'db')))

# RB_BIND:0250
from application.components.registration import install_0250 as _rb_install_0250
booking_control_reason = _rb_install_0250(_rb_bind(_rb_sys.modules[__name__], ('BookingBlock',)))


# RB_BIND:0251
from application.components.registration import install_0251 as _rb_install_0251
validate_booking_block_scope = _rb_install_0251(_rb_bind(_rb_sys.modules[__name__], ()))


# RB_BIND:0252
from application.components.registration import install_0252 as _rb_install_0252
acquire_booking_slot_lock = _rb_install_0252(_rb_bind(_rb_sys.modules[__name__], ('db', 'text')))


DEFAULT_UI_DICTIONARY=[
 ('nav.dashboard','GLOBAL','Dashboard','لوحة التحكم'),
 ('nav.patients','PATIENT','Patients','المرضى'),
 ('nav.orders','ORDERS','Orders / Visits','الطلبات / الزيارات'),
 ('nav.appointments','BOOKING','Appointments','المواعيد'),
 ('nav.booking','BOOKING','Booking','الحجز'),
 ('nav.booking_control','BOOKING','Booking Control','التحكم في الحجز'),
 ('nav.booking_availability','BOOKING','Booking Availability','إتاحة الحجز'),
 ('nav.technician','CLINICAL','Technician Worklist','قائمة عمل الفني'),
 ('nav.nursing','CLINICAL','Nursing','التمريض'),
 ('nav.doctor','REPORTING','Doctor Workspace','مساحة الطبيب'),
 ('nav.pacs_tracking','PACS','PACS Tracking','متابعة PACS'),
 ('nav.cashier','CASHIER','Cashier','الكاشير'),
 ('nav.dictionary','USERS','Dictionary','القاموس'),
 ('nav.profile','USERS','My Profile','ملفي الشخصي'),

 ('common.save','GLOBAL','Save','حفظ'),
 ('common.cancel','GLOBAL','Cancel','إلغاء'),
 ('common.back','GLOBAL','Back','رجوع'),
 ('common.search','GLOBAL','Search','بحث'),
 ('common.edit','GLOBAL','Edit','تعديل'),
 ('common.delete','GLOBAL','Delete','حذف'),
 ('common.add','GLOBAL','Add','إضافة'),
 ('common.print','GLOBAL','Print','طباعة'),
 ('common.upload','GLOBAL','Upload','رفع'),
 ('common.refresh','GLOBAL','Refresh','تحديث'),
 ('common.actions','GLOBAL','Actions','إجراءات'),
 ('common.status','GLOBAL','Status','الحالة'),
 ('common.date','GLOBAL','Date','التاريخ'),
 ('common.time','GLOBAL','Time','الوقت'),
 ('common.from','GLOBAL','From','من'),
 ('common.to','GLOBAL','To','إلى'),
 ('common.reason','GLOBAL','Reason','السبب'),
 ('common.notes','GLOBAL','Notes','ملاحظات'),
 ('common.active','GLOBAL','Active','نشط'),
 ('common.inactive','GLOBAL','Inactive','غير نشط'),
 ('common.yes','GLOBAL','Yes','نعم'),
 ('common.no','GLOBAL','No','لا'),
 ('common.close','GLOBAL','Close','إغلاق'),
 ('common.select','GLOBAL','Select','اختيار'),
 ('common.name','GLOBAL','Name','الاسم'),
 ('common.code','GLOBAL','Code','الكود'),
 ('common.branch','GLOBAL','Branch','الفرع'),
 ('common.device','GLOBAL','Device','الجهاز'),
 ('common.room','GLOBAL','Room','الغرفة'),
 ('common.priority','GLOBAL','Priority','الأولوية'),

 ('patient.title','PATIENT','Patient','المريض'),
 ('patient.name','PATIENT','Patient Name','اسم المريض'),
 ('patient.id','PATIENT','Patient ID','كود المريض'),
 ('patient.new','PATIENT','New Patient','مريض جديد'),
 ('patient.existing','PATIENT','Existing Patient','مريض مسجل'),
 ('patient.mobile','PATIENT','Mobile','الموبايل'),
 ('patient.dob','PATIENT','Date of Birth','تاريخ الميلاد'),
 ('patient.sex','PATIENT','Sex','النوع'),
 ('patient.national_id','PATIENT','National ID','الرقم القومي'),
 ('patient.address','PATIENT','Address','العنوان'),
 ('patient.email','PATIENT','Email','البريد الإلكتروني'),
 ('patient.worklist_name','PATIENT','Worklist Name','اسم قائمة العمل'),
 ('patient.search_hint','PATIENT','Patient ID, name or mobile','كود المريض أو الاسم أو الموبايل'),

 ('visit.title','ORDERS','Patient & Visit','المريض والزيارة'),
 ('visit.save','ORDERS','Save Patient & Visit','حفظ المريض والزيارة'),
 ('visit.date_time','ORDERS','Visit Date / Time','تاريخ ووقت الزيارة'),
 ('visit.accession','ORDERS','Accession Number','رقم الزيارة'),

 ('exam.title','BOOKING','Examination','الفحص'),
 ('exam.select','BOOKING','Select examination','اختر الفحص'),
 ('exam.modality','BOOKING','Modality','نوع الجهاز'),
 ('exam.duration','BOOKING','Duration','المدة'),
 ('exam.not_available','BOOKING','Examination is not available','الفحص غير متاح'),

 ('booking.title','BOOKING','Booking','الحجز'),
 ('booking.new','BOOKING','New Booking','حجز جديد'),
 ('booking.book_slot','BOOKING','Book Slot','حجز موعد'),
 ('booking.waiting_list','BOOKING','Waiting List','قائمة الانتظار'),
 ('booking.available','BOOKING','Available','متاح'),
 ('booking.unavailable','BOOKING','Unavailable','غير متاح'),
 ('booking.maintenance','BOOKING','Maintenance','صيانة'),
 ('booking.device_failure','BOOKING','Device Failure','عطل بالجهاز'),
 ('booking.calibration','BOOKING','Calibration','معايرة'),
 ('booking.doctor_unavailable','BOOKING','Doctor Unavailable','الطبيب غير متاح'),
 ('booking.technician_unavailable','BOOKING','Technician Unavailable','الفني غير متاح'),
 ('booking.holiday','BOOKING','Holiday','إجازة'),
 ('booking.emergency_block','BOOKING','Emergency Block','إيقاف طارئ'),
 ('booking.other','BOOKING','Other','أخرى'),

 ('clinical.note','CLINICAL','Clinical Note','ملاحظات الفحص'),
 ('clinical.notes','CLINICAL','Clinical Notes','ملاحظات الفحص'),
 ('clinical.attachment','CLINICAL','Attachment','مرفق'),
 ('clinical.attachments','CLINICAL','Attachments','المرفقات'),
 ('clinical.reception','CLINICAL','Reception','الاستقبال'),
 ('clinical.technician','CLINICAL','Technician','الفني'),
 ('clinical.radiologist','CLINICAL','Radiologist','طبيب الأشعة'),
 ('clinical.nursing','CLINICAL','Nursing','التمريض'),
 ('clinical.visibility','CLINICAL','Visibility','الظهور'),

 ('referrer.title','CLINICAL','Referring Doctor','الطبيب المحول'),
 ('payer.title','CASHIER','Payer','جهة التعاقد'),
 ('payer.cash','CASHIER','Cash / Private','نقدي / خاص'),
 ('approval.number','CASHIER','Approval No.','رقم الموافقة'),

 ('cashier.title','CASHIER','Cashier Workstation','شاشة الكاشير'),
 ('cashier.payment','CASHIER','Payment','التحصيل'),
 ('cashier.collect','CASHIER','Collect Payment','تحصيل'),
 ('cashier.amount','CASHIER','Amount','المبلغ'),
 ('cashier.method','CASHIER','Payment Method','طريقة الدفع'),
 ('cashier.reference','CASHIER','Reference / POS','المرجع / نقطة البيع'),
 ('cashier.receipt','CASHIER','Receipt','إيصال'),
 ('cashier.refund','CASHIER','Refund','استرداد'),
 ('cashier.refund_reason','CASHIER','Refund Reason','سبب الاسترداد'),
 ('cashier.confirm_refund','CASHIER','Confirm Refund','تأكيد الاسترداد'),
 ('cashier.paid','CASHIER','Paid','المدفوع'),
 ('cashier.due','CASHIER','Due','المستحق'),

 ('profile.title','USERS','My Profile & Workspace','ملفي الشخصي وواجهة العمل'),
 ('profile.language','USERS','Language','اللغة'),
 ('profile.theme','USERS','Theme','المظهر'),
 ('profile.accent','USERS','Accent','اللون الرئيسي'),
 ('profile.density','USERS','Density','كثافة العرض'),
 ('profile.font_scale','USERS','Font Size','حجم الخط'),
 ('profile.sidebar','USERS','Sidebar','القائمة الجانبية'),

 ('dictionary.title','USERS','System Dictionary','قاموس النظام'),
 ('dictionary.key','USERS','Key','المفتاح'),
 ('dictionary.module','USERS','Module','القسم'),
 ('dictionary.english','USERS','English Label','الاسم الإنجليزي'),
 ('dictionary.arabic','USERS','Arabic Label','الاسم العربي'),
 ('dictionary.reset','USERS','Reset to Default','استرجاع الافتراضي'),

 ('status.BOOKED','STATUS','Booked','محجوز'),
 ('status.WAITING','STATUS','Waiting','انتظار'),
 ('status.CONFIRMED','STATUS','Confirmed','مؤكد'),
 ('status.ARRIVED','STATUS','Arrived','وصل'),
 ('status.REGISTERED','STATUS','Registered','مسجل'),
 ('status.SCHEDULED','STATUS','Scheduled','مجدول'),
 ('status.MWL_ERROR','STATUS','MWL Error','خطأ في قائمة العمل'),
 ('status.ORDERED','STATUS','Ordered','تم الطلب'),
 ('status.IN_PROGRESS','STATUS','In Progress','قيد التنفيذ'),
 ('status.EXAM_COMPLETED','STATUS','Exam Completed','اكتمل الفحص'),
 ('status.REPORTING','STATUS','Reporting','جاري كتابة التقرير'),
 ('status.VERIFIED','STATUS','Verified','تمت المراجعة'),
 ('status.FINAL','STATUS','Final','نهائي'),
 ('status.CANCELLED','STATUS','Cancelled','ملغي'),
 ('status.NO_SHOW','STATUS','No Show','لم يحضر'),
 ('status.REFUNDED','STATUS','Refunded','تم الاسترداد'),
 ('status.PARTIAL_REFUND','STATUS','Partial Refund','استرداد جزئي')
]


EXTRA_UI_DICTIONARY=[
 ('nav.assignment','USERS','Assignment','التوزيع'),
 ('nav.unassigned','REPORTING','Unassigned','غير موزع'),
 ('nav.calendar','BOOKING','Calendar','التقويم'),
 ('nav.reconciliation','PACS','PACS Reconciliation','مطابقة PACS'),
 ('nav.logout','GLOBAL','Logout','تسجيل الخروج'),
 ('nav.online','GLOBAL','Online','متصل'),
 ('nav.offline','GLOBAL','Offline','غير متصل'),

 ('patient.select','PATIENT','Select patient...','اختر المريض...'),
 ('patient.duplicate','PATIENT','Possible duplicate patients','مرضى محتمل تكرارهم'),
 ('patient.worklist_name','PATIENT','Worklist Name','الاسم في قائمة العمل'),
 ('patient.email','PATIENT','Email','البريد الإلكتروني'),
 ('patient.male','PATIENT','Male','ذكر'),
 ('patient.female','PATIENT','Female','أنثى'),

 ('branch.select','BOOKING','Select branch','اختر الفرع'),
 ('exam.select_branch_first','BOOKING','Select branch first...','اختر الفرع أولاً...'),
 ('priority.ROUTINE','ORDERS','Routine','عادي'),
 ('priority.URGENT','ORDERS','Urgent','عاجل'),
 ('priority.STAT','ORDERS','STAT','طارئ فوري'),

 ('clinical.information','CLINICAL','Clinical Information','المعلومات الإكلينيكية'),
 ('clinical.note_placeholder','CLINICAL','Reason for exam, symptoms, special instructions...','سبب الفحص، الأعراض، والتعليمات الخاصة...'),
 ('clinical.attachment_category','CLINICAL','Attachment Category','تصنيف المرفق'),
 ('clinical.attachment_notes','CLINICAL','Attachment Notes','ملاحظات المرفق'),
 ('clinical.category.CLINICAL','CLINICAL','Clinical','إكلينيكي'),
 ('clinical.category.REFERRAL','CLINICAL','Referral','تحويل'),
 ('clinical.category.APPROVAL','CLINICAL','Approval','موافقة'),
 ('clinical.category.OTHER','CLINICAL','Other','أخرى'),

 ('payment.no_payment','CASHIER','No payment now','بدون تحصيل الآن'),
 ('payment.CASH','CASHIER','Cash','نقدي'),
 ('payment.VISA','CASHIER','Visa','فيزا'),
 ('payment.MASTERCARD','CASHIER','Mastercard','ماستركارد'),
 ('payment.WALLET','CASHIER','Wallet','محفظة إلكترونية'),
 ('payment.BANK_TRANSFER','CASHIER','Bank Transfer','تحويل بنكي'),

 ('registration.validation_notice','ORDERS',
  'The system automatically checks branch availability, examination availability, device status, maintenance, staffing blocks and overlapping appointments before Save.',
  'يتحقق النظام تلقائيًا قبل الحفظ من إتاحة الفرع والفحص وحالة الجهاز والصيانة وتوفر الطاقم وتعارض المواعيد.'),

 ('booking.control_center','BOOKING','Booking Control Center','مركز التحكم في الحجز'),
 ('booking.block','BOOKING','Block Booking','إيقاف الحجز'),
 ('booking.block_from','BOOKING','Block From','إيقاف من'),
 ('booking.block_to','BOOKING','Block To','إيقاف إلى'),
 ('booking.staffing','BOOKING','Staffing','توفر الطاقم'),

 ('dictionary.search_hint','USERS','Search key, English or Arabic','ابحث بالمفتاح أو الإنجليزي أو العربي'),
 ('profile.workspace','USERS','Workspace','واجهة العمل'),
]
for _row in EXTRA_UI_DICTIONARY:
    if _row[0] not in {x[0] for x in DEFAULT_UI_DICTIONARY}:
        DEFAULT_UI_DICTIONARY.append(_row)


ORGANIZED_UI_DICTIONARY=[
 ('group.workflow','GLOBAL','Workflow','التشغيل'),
 ('group.clinical','CLINICAL','Clinical','الطبي'),
 ('group.pacs','PACS','PACS & Imaging','PACS والصور'),
 ('group.finance','CASHIER','Finance','المالية'),
 ('group.management','GLOBAL','Management','الإدارة التشغيلية'),
 ('group.system','USERS','System Administration','إدارة النظام'),
 ('nav.new_visit','ORDERS','Patient & Visit','المريض والزيارة'),
 ('nav.calendar','BOOKING','Calendar','التقويم'),
 ('nav.reconciliation','PACS','PACS Reconciliation','مطابقة PACS'),
 ('nav.receivables','CASHIER','Receivables','المديونيات والبواقي'),
 ('nav.accounting_reports','CASHIER','Accounting Reports','التقارير المحاسبية'),
 ('nav.analytics','GLOBAL','Analytics','التحليلات'),
 ('nav.assignment','CLINICAL','Assignment','التوزيع'),
 ('nav.unassigned','REPORTING','Unassigned Cases','الحالات غير الموزعة'),
 ('nav.reports','REPORTING','Reports','التقارير'),
 ('nav.admin_home','USERS','Administration','الإدارة'),
 ('nav.roles','USERS','Roles','الأدوار والصلاحيات'),
 ('nav.branches','USERS','Branches','الفروع'),
 ('nav.devices','USERS','Devices','الأجهزة'),
 ('nav.procedures','USERS','Procedures','الفحوصات'),
 ('nav.contracts','CASHIER','Contracts','التعاقدات'),
 ('nav.inventory','INVENTORY','Inventory','المخزون والمستهلكات'),
 ('nav.import_export','USERS','Import / Export','استيراد / تصدير'),
 ('nav.diagnostics','USERS','Diagnostics','التشخيص الفني'),
 ('nav.audit','USERS','Audit Log','سجل المراجعة'),
 ('nav.ribbon','USERS','My Ribbon','شريط الاختصارات'),
 ('nav.ribbon_designer','USERS','Ribbon Designer','مصمم شريط الاختصارات')
]
for _row in ORGANIZED_UI_DICTIONARY:
    if _row[0] not in {x[0] for x in DEFAULT_UI_DICTIONARY}:
        DEFAULT_UI_DICTIONARY.append(_row)

DEFAULT_UI_MAP={k:{'module':m,'en':en,'ar':ar} for k,m,en,ar in DEFAULT_UI_DICTIONARY}

# RB_BIND:0253
from web.components.shared import install_0253 as _rb_install_0253
current_ui_language = _rb_install_0253(_rb_bind(_rb_sys.modules[__name__], ('UserProfile', 'current_user', 'session')))

# RB_BIND:0254
from application.components.shared import install_0254 as _rb_install_0254
ui_text = _rb_install_0254(_rb_bind(_rb_sys.modules[__name__], ('DEFAULT_UI_MAP', 'DictionaryEntry', 'current_ui_language')))

# RB_BIND:0255
from application.components.shared import install_0255 as _rb_install_0255
status_text = _rb_install_0255(_rb_bind(_rb_sys.modules[__name__], ('ui_text',)))

# RB_BIND:0256
from application.components.shared import install_0256 as _rb_install_0256
seed_full_ui_dictionary = _rb_install_0256(_rb_bind(_rb_sys.modules[__name__], ('DEFAULT_UI_DICTIONARY', 'DictionaryEntry', 'db')))


# RB_BIND:0257
from application.components.shared import install_0257 as _rb_install_0257
ribbon_actions_for_user = _rb_install_0257(_rb_bind(_rb_sys.modules[__name__], ('RibbonAction', 'UserRibbonAction', 'current_user', 'has_perm')))

JOURNEY_STAGES=list(CANONICAL_JOURNEY_STAGES)

# RB_BIND:0258
from application.components.shared import install_0258 as _rb_install_0258
ensure_journey = _rb_install_0258(_rb_bind(_rb_sys.modules[__name__], ('JourneyEvent', 'current_user', 'datetime', 'db')))

# RB_BIND:0259
from application.components.shared import install_0259 as _rb_install_0259
journey_snapshot = _rb_install_0259(_rb_bind(_rb_sys.modules[__name__], ('JourneyEvent', 'derive_journey_stage')))

# RB_BIND:0260
from application.components.shared import install_0260 as _rb_install_0260
role_dashboard_kind = _rb_install_0260(_rb_bind(_rb_sys.modules[__name__], ('current_user',)))

# RB_BIND:0261
from application.components.communications import install_0261 as _rb_install_0261
ensure_chat_member = _rb_install_0261(_rb_bind(_rb_sys.modules[__name__], ('ChatThreadMember', 'db')))

# RB_BIND:0262
from application.components.shared import install_0262 as _rb_install_0262
get_or_create_case_thread = _rb_install_0262(_rb_bind(_rb_sys.modules[__name__], ('ChatThread', 'current_user', 'db', 'ensure_chat_member')))

# RB_BIND:0263
from application.components.communications import install_0263 as _rb_install_0263
unread_chat_count = _rb_install_0263(_rb_bind(_rb_sys.modules[__name__], ('ChatMessage', 'ChatThreadMember', 'current_user')))

# RB_BIND:0264
from application.components.registration import install_0264 as _rb_install_0264
smart_booking_cards = _rb_install_0264(_rb_bind(_rb_sys.modules[__name__], ('Appointment', 'BookingBlock', 'Device', 'and_', 'datetime', 'or_', 'timedelta')))


DOCUMENT_TYPES={
 'PRESCRIPTION':('Prescription','روشتة','CLINICAL'),
 'MEDICAL_APPROVAL':('Medical Approval','موافقة طبية','INSURANCE'),
 'NATIONAL_ID':('National ID','بطاقة رقم قومي','IDENTITY'),
 'INSURANCE_CARD':('Insurance Card','بطاقة التأمين','INSURANCE'),
 'REFERRAL':('Referral','تحويل','CLINICAL'),
 'QUESTIONNAIRE':('Questionnaire','استبيان','CLINICAL'),
 'CONSENT':('Consent Form','إقرار / موافقة','CONSENT'),
 'PREVIOUS_REPORT':('Previous Report','تقرير سابق','HISTORY'),
 'LAB_RESULT':('Lab Result','نتيجة معمل','CLINICAL'),
 'PATHOLOGY':('Pathology','باثولوجي','CLINICAL'),
 'PREVIOUS_IMAGING':('Previous Imaging','أشعة سابقة','HISTORY'),
 'OTHER':('Other','أخرى','OTHER')
}
DOCUMENT_REVIEW_STATUSES=['NEW','REVIEWED','APPROVED','REJECTED','NEEDS_CLARIFICATION']

# RB_BIND:0265
from application.components.shared import install_0265 as _rb_install_0265
normalize_phone = _rb_install_0265(_rb_bind(_rb_sys.modules[__name__], ('re',)))

# RB_BIND:0266
from application.components.registration import install_0266 as _rb_install_0266
patient_by_phone = _rb_install_0266(_rb_bind(_rb_sys.modules[__name__], ('Patient', 'normalize_phone')))

# RB_BIND:0267
from application.components.shared import install_0267 as _rb_install_0267
document_timeline = _rb_install_0267(_rb_bind(_rb_sys.modules[__name__], ('DOCUMENT_TYPES', 'Order', 'PatientDocument', 'current_ui_language', 'datetime')))

# RB_BIND:0268
from application.components.shared import install_0268 as _rb_install_0268
extract_document_text_basic = _rb_install_0268(_rb_bind(_rb_sys.modules[__name__], ('os',)))

# RB_BIND:0269
from application.components.shared import install_0269 as _rb_install_0269
build_document_matches = _rb_install_0269(_rb_bind(_rb_sys.modules[__name__], ('DocumentMatch', 'db')))



DEFAULT_V521_FIELD_RULES=[
 ('patient.dob','Date of Birth','تاريخ الميلاد','ANY',0,None,None,True,True),
 ('patient.sex','Sex','النوع','ANY',0,None,None,True,True),
 ('patient.name','Patient Name','اسم المريض','LETTERS',3,100,None,True,True),
 ('patient.phone','Mobile','الموبايل','NUMBERS',10,11,None,True,True),
 ('patient.national_id','National ID','الرقم القومي','NUMBERS',14,14,14,False,True),
 ('patient.nationality','Nationality','الجنسية','LETTERS',2,100,None,False,True),
 ('patient.email','Email','البريد الإلكتروني','ANY',0,180,None,False,True),
 ('patient.address','Address','العنوان','ANY',0,300,None,False,True),
]

# RB_BIND:0270
from application.components.shared import install_0270 as _rb_install_0270
seed_v521_functional_defaults = _rb_install_0270(_rb_bind(_rb_sys.modules[__name__], ('DEFAULT_V521_FIELD_RULES', 'FieldRule', 'db')))



STUDIO_ICON_LIBRARY = [
 ('home','⌂','Home'),('patient','👤','Patient'),('patients','👥','Patients'),
 ('calendar','📅','Booking'),('search','🔎','Search'),('pacs','🩻','PACS'),
 ('report','📄','Report'),('print','🖨️','Print'),('receipt','🧾','Receipt'),
 ('label','🏷️','Label'),('attachment','📎','Attachments'),('doctor','🩺','Doctor'),
 ('lab','🧪','Lab'),('finance','💳','Finance'),('chart','📊','Analytics'),
 ('alert','⚠️','Alerts'),('message','💬','Messages'),('phone','☎️','Phone'),
 ('ai','🤖','AI'),('settings','⚙️','Settings'),('shield','🛡️','Security'),
 ('inventory','📦','Inventory'),('branch','🏥','Branch'),('link','🔗','Link'),
 ('star','★','Favorite'),('workflow','⇄','Workflow'),('camera','📷','Imaging'),
 ('upload','⬆️','Upload'),('download','⬇️','Download'),('check','✓','Complete')
]
STUDIO_ICON_MAP={k:g for k,g,_ in STUDIO_ICON_LIBRARY}

# RB_BIND:0271
from application.components.shared import install_0271 as _rb_install_0271
studio_shortcut_context = _rb_install_0271(_rb_bind(_rb_sys.modules[__name__], ('current_user',)))







# RB_BIND:0272
from application.components.identity import install_0272 as _rb_install_0272
verify_current_user_password = _rb_install_0272(_rb_bind(_rb_sys.modules[__name__], ('check_password_hash', 'current_user')))

# RB_BIND:0273
from application.components.registration import install_0273 as _rb_install_0273
_collect_patient_file_paths = _rb_install_0273(_rb_bind(_rb_sys.modules[__name__], ('Attachment', 'Report', 'os')))

# RB_BIND:0274
from application.components.registration import install_0274 as _rb_install_0274
create_all_patient_backup = _rb_install_0274(_rb_bind(_rb_sys.modules[__name__], ('APP_VERSION', 'Attachment', 'Order', 'Patient', 'Report', '_collect_patient_file_paths', 'admin_backup_dir', 'csv', 'datetime', 'hashlib', 'io', 'json', 'os', 'safe_export_name', 'zipfile')))

# RB_BIND:0275
from application.components.registration import install_0275 as _rb_install_0275
quarantine_patient_files = _rb_install_0275(_rb_bind(_rb_sys.modules[__name__], ('_collect_patient_file_paths', 'datetime', 'hashlib', 'os', 'safe_export_name', 'shutil')))

# RB_BIND:0276
from application.components.registration import install_0276 as _rb_install_0276
operational_patient_table_order = _rb_install_0276(_rb_bind(_rb_sys.modules[__name__], ('db',)))

# RB_BIND:0277
from application.components.registration import install_0277 as _rb_install_0277
delete_all_patient_operational_data = _rb_install_0277(_rb_bind(_rb_sys.modules[__name__], ('db', 'operational_patient_table_order', 'quarantine_patient_files')))

# RB_BIND:0278
from application.components.operations import install_0278 as _rb_install_0278
admin_storage_paths = _rb_install_0278(_rb_bind(_rb_sys.modules[__name__], ('os',)))

# RB_BIND:0279
from application.components.shared import install_0279 as _rb_install_0279
safe_export_name = _rb_install_0279(_rb_bind(_rb_sys.modules[__name__], ('re',)))

# RB_BIND:0280
from application.components.shared import install_0280 as _rb_install_0280
destructive_confirmation_token = _rb_install_0280(_rb_bind(_rb_sys.modules[__name__], ()))

# RB_BIND:0281
from application.components.registration import install_0281 as _rb_install_0281
patient_export_manifest = _rb_install_0281(_rb_bind(_rb_sys.modules[__name__], ('Order', 'datetime')))

# RB_BIND:0282
from application.components.operations import install_0282 as _rb_install_0282
admin_backup_dir = _rb_install_0282(_rb_bind(_rb_sys.modules[__name__], ('os',)))

# RB_BIND:0283
from application.components.inventory import install_0283 as _rb_install_0283
_warehouse_code_part = _rb_install_0283(_rb_bind(_rb_sys.modules[__name__], ()))

# RB_BIND:0284
from application.components.inventory import install_0284 as _rb_install_0284
ensure_main_warehouse = _rb_install_0284(_rb_bind(_rb_sys.modules[__name__], ('Warehouse', 'db')))

# RB_BIND:0285
from application.components.inventory import install_0285 as _rb_install_0285
ensure_branch_warehouses = _rb_install_0285(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Warehouse', '_warehouse_code_part', 'db', 'ensure_main_warehouse')))

# RB_BIND:0286
from application.components.shared import install_0286 as _rb_install_0286
next_transfer_no = _rb_install_0286(_rb_bind(_rb_sys.modules[__name__], ('StockTransfer', 'datetime')))

# RB_BIND:0287
from application.components.inventory import install_0287 as _rb_install_0287
default_warehouse_for_order = _rb_install_0287(_rb_bind(_rb_sys.modules[__name__], ('Warehouse',)))

# RB_BIND:0288
from application.components.inventory import install_0288 as _rb_install_0288
stock_lots = _rb_install_0288(_rb_bind(_rb_sys.modules[__name__], ('InventoryLot', 'db')))

# RB_BIND:0289
from application.components.inventory import install_0289 as _rb_install_0289
inventory_available_qty = _rb_install_0289(_rb_bind(_rb_sys.modules[__name__], ('InventoryLot', 'db')))

# RB_BIND:0290
from application.components.inventory import install_0290 as _rb_install_0290
inventory_on_hand_qty = _rb_install_0290(_rb_bind(_rb_sys.modules[__name__], ('InventoryLot',)))

# RB_BIND:0291
from application.components.inventory import install_0291 as _rb_install_0291
inventory_move = _rb_install_0291(_rb_bind(_rb_sys.modules[__name__], ('InventoryLot', 'WarehouseStockMovement', 'current_user', 'db')))

# RB_BIND:0292
from application.components.inventory import install_0292 as _rb_install_0292
consume_inventory = _rb_install_0292(_rb_bind(_rb_sys.modules[__name__], ('OrderConsumableUsage', 'current_user', 'db', 'default_warehouse_for_order', 'inventory_move', 'stock_lots')))

# RB_BIND:0293
from application.components.inventory import install_0293 as _rb_install_0293
return_inventory = _rb_install_0293(_rb_bind(_rb_sys.modules[__name__], ()))

# RB_BIND:0294
from application.components.inventory import install_0294 as _rb_install_0294
order_consumable_cost = _rb_install_0294(_rb_bind(_rb_sys.modules[__name__], ('OrderConsumableUsage',)))

# RB_BIND:0295
from application.components.inventory import install_0295 as _rb_install_0295
order_default_consumables = _rb_install_0295(_rb_bind(_rb_sys.modules[__name__], ('OrderConsumableUsage', 'ProcedureConsumableDefault')))

# RB_BIND:0296
from application.components.inventory import install_0296 as _rb_install_0296
low_stock_rows = _rb_install_0296(_rb_bind(_rb_sys.modules[__name__], ('WarehouseStockPolicy', 'inventory_available_qty')))

# RB_BIND:0297
from application.components.finance import install_0297 as _rb_install_0297
approved_order_discounts = _rb_install_0297(_rb_bind(_rb_sys.modules[__name__], ('OrderDiscount',)))

# RB_BIND:0298
from application.components.finance import install_0298 as _rb_install_0298
total_order_discount = _rb_install_0298(_rb_bind(_rb_sys.modules[__name__], ('approved_order_discounts',)))

# RB_BIND:0299
from application.components.registration import install_0299 as _rb_install_0299
base_patient_due = _rb_install_0299(_rb_bind(_rb_sys.modules[__name__], ()))

# RB_BIND:0300
from application.components.registration import install_0300 as _rb_install_0300
effective_patient_due = _rb_install_0300(_rb_bind(_rb_sys.modules[__name__], ('base_patient_due', 'total_order_discount')))

# RB_BIND:0301
from application.components.finance import install_0301 as _rb_install_0301
discount_policy_for = _rb_install_0301(_rb_bind(_rb_sys.modules[__name__], ('DiscountPolicy', 'current_user')))

# RB_BIND:0302
from application.components.finance import install_0302 as _rb_install_0302
apply_or_request_discount = _rb_install_0302(_rb_bind(_rb_sys.modules[__name__], ('DiscountApprovalRequest', 'OrderDiscount', 'audit', 'base_patient_due', 'current_user', 'datetime', 'db', 'discount_policy_for', 'has_perm')))

# RB_BIND:0303
from application.components.shared import install_0303 as _rb_install_0303
performing_doctor_for_order = _rb_install_0303(_rb_bind(_rb_sys.modules[__name__], ('CaseAssignment', 'OrderPerformingDoctor')))

# RB_BIND:0304
from application.components.shared import install_0304 as _rb_install_0304
performing_doctor_share_for_order = _rb_install_0304(_rb_bind(_rb_sys.modules[__name__], ('PerformingDoctorShareRule', '__name__', '_reporting_support_sys', 'effective_patient_due', 'performing_doctor_for_order')))


# RB_BIND:0305
from web.components.shared import install_0305 as _rb_install_0305
executive_analytics_filters = _rb_install_0305(_rb_bind(_rb_sys.modules[__name__], ('request',)))

# RB_BIND:0306
from application.components.shared import install_0306 as _rb_install_0306
executive_orders = _rb_install_0306(_rb_bind(_rb_sys.modules[__name__], ('Order',)))

# RB_BIND:0307
from application.components.shared import install_0307 as _rb_install_0307
_chart_rows = _rb_install_0307(_rb_bind(_rb_sys.modules[__name__], ()))

# RB_BIND:0308
from application.components.shared import install_0308 as _rb_install_0308
executive_analytics_payload = _rb_install_0308(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Device', 'Patient', 'Report', 'User', '_chart_rows', 'executive_orders')))

# RB_BIND:0309
from web.components.shared import install_0309 as _rb_install_0309
management_date_range = _rb_install_0309(_rb_bind(_rb_sys.modules[__name__], ('request',)))

# RB_BIND:0310
from application.components.shared import install_0310 as _rb_install_0310
management_orders_query = _rb_install_0310(_rb_bind(_rb_sys.modules[__name__], ('Order',)))

# RB_BIND:0311
from application.components.shared import install_0311 as _rb_install_0311
management_summary = _rb_install_0311(_rb_bind(_rb_sys.modules[__name__], ('CaseAssignment', 'Report', 'case_delay_snapshot', 'management_orders_query')))

# RB_BIND:0312
from application.components.finance import install_0312 as _rb_install_0312
referral_commission_for_order = _rb_install_0312(_rb_bind(_rb_sys.modules[__name__], ('ReferralCommissionRule',)))

# RB_BIND:0313
from application.components.finance import install_0313 as _rb_install_0313
daily_payment_totals = _rb_install_0313(_rb_bind(_rb_sys.modules[__name__], ('db', 'Order', 'Payment', 'PaymentRefund')))

# RB_BIND:0314
from application.components.shared import install_0314 as _rb_install_0314
studio_shortcuts_for = _rb_install_0314(_rb_bind(_rb_sys.modules[__name__], ('STUDIO_ICON_MAP', 'StudioShortcut', 'current_user', 'has_perm')))

# RB_BIND:0315
from application.components.registration import install_0315 as _rb_install_0315
registration_defaults_for = _rb_install_0315(_rb_bind(_rb_sys.modules[__name__], ('RegistrationDefault', 'current_user')))

# RB_BIND:0316
from application.components.registration import install_0316 as _rb_install_0316
registration_form_context = _rb_install_0316(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Contract', 'Patient', 'Procedure', 'ReferringDoctor', 'registration_defaults_for', 'registration_rule_map', 'registration_style_for')))

# RB_BIND:0317
from web.components.registration import install_0317 as _rb_install_0317
render_registration_form = _rb_install_0317(_rb_bind(_rb_sys.modules[__name__], ('registration_form_context', 'registration_template_for', 'render_template')))

# RB_BIND:0318
from web.components.registration import install_0318 as _rb_install_0318
registration_validation_error = _rb_install_0318(_rb_bind(_rb_sys.modules[__name__], ('db', 'flash', 'render_registration_form', 'request')))

# RB_BIND:0319
from application.components.registration import install_0319 as _rb_install_0319
registration_style_for = _rb_install_0319(_rb_bind(_rb_sys.modules[__name__], ('UserRegistrationPreference', 'current_user')))

# RB_BIND:0320
from application.components.registration import install_0320 as _rb_install_0320
registration_template_for = _rb_install_0320(_rb_bind(_rb_sys.modules[__name__], ('registration_style_for',)))

# RB_BIND:0321
from application.components.registration import install_0321 as _rb_install_0321
registration_rule_map = _rb_install_0321(_rb_bind(_rb_sys.modules[__name__], ('FieldRule',)))

# RB_BIND:0322
from application.components.registration import install_0322 as _rb_install_0322
validate_registration_value = _rb_install_0322(_rb_bind(_rb_sys.modules[__name__], ('current_ui_language', 're')))

# RB_BIND:0323
from application.components.registration import install_0323 as _rb_install_0323
patient_extra = _rb_install_0323(_rb_bind(_rb_sys.modules[__name__], ('PatientDemographicsExtra',)))

# RB_BIND:0324
from application.components.finance import install_0324 as _rb_install_0324
receipt_studio_template = _rb_install_0324(_rb_bind(_rb_sys.modules[__name__], ('StudioTemplate', 'or_')))

# RB_BIND:0325
from application.components.shared import install_0325 as _rb_install_0325
company_legal_for_order = _rb_install_0325(_rb_bind(_rb_sys.modules[__name__], ('CompanyLegalProfile',)))

# RB_BIND:0326
from application.components.shared import install_0326 as _rb_install_0326
replace_docx_placeholders = _rb_install_0326(_rb_bind(_rb_sys.modules[__name__], ()))



# RB_BIND:0327
from application.components.shared import install_0327 as _rb_install_0327
ai_provider_setting = _rb_install_0327(_rb_bind(_rb_sys.modules[__name__], ('AIProviderSetting',)))

# RB_BIND:0328
from application.components.shared import install_0328 as _rb_install_0328
related_prior_orders = _rb_install_0328(_rb_bind(_rb_sys.modules[__name__], ('Order',)))

# RB_BIND:0329
from application.components.shared import install_0329 as _rb_install_0329
external_priors_for_order = _rb_install_0329(_rb_bind(_rb_sys.modules[__name__], ('ExternalPriorStudy', 'or_')))

# RB_BIND:0330
from application.components.imaging import install_0330 as _rb_install_0330
safe_dicom_metadata = _rb_install_0330(_rb_bind(_rb_sys.modules[__name__], ()))

# RB_BIND:0331
from application.components.shared import install_0331 as _rb_install_0331
score_prior_match = _rb_install_0331(_rb_bind(_rb_sys.modules[__name__], ('re',)))

# RB_BIND:0332
from application.components.shared import install_0332 as _rb_install_0332
selected_prior_context = _rb_install_0332(_rb_bind(_rb_sys.modules[__name__], ('ExternalPriorStudy', 'Order', 'PriorStudySelection', 'get_report')))


# RB_BIND:0333
from application.components.shared import install_0333 as _rb_install_0333
generate_ai_draft_payload = _rb_install_0333(_rb_bind(_rb_sys.modules[__name__], ('ai_case_context',)))

# RB_BIND:0334
from application.components.communications import install_0334 as _rb_install_0334
call_configured_ai_provider = _rb_install_0334(_rb_bind(_rb_sys.modules[__name__], ('ai_provider_setting', 'generate_ai_draft_payload', 'os')))



ROLE_QUICK_ACTION_PRESETS = {
    'RECEPTION': [
        ('new_visit','New Patient / Visit','مريض / زيارة جديدة','➕','new_order'),
        ('booking','Booking','الحجز','📅','schedule_v2'),
        ('patients','Patients','المرضى','👤','patients'),
        ('documents','Documents','المستندات','📎','patients'),
        ('history','Patient History','تاريخ المريض','🕘','patients'),
        ('label','Print Label','طباعة ملصق','🏷️','patients'),
        ('receipt','Print Receipt','طباعة إيصال','🧾','cashier'),
    ],
    'CASHIER': [
        ('cashier','Cashier','الكاشير','💳','cashier'),
        ('receipt','Print Receipt','طباعة إيصال','🧾','cashier'),
        ('refund','Refund','استرداد','↩️','cashier'),
        ('outstanding','Outstanding','البواقي','💰','cashier'),
        ('financial','Financial Reports','التقارير المالية','📊','accounting_reports'),
    ],
    'TECHNICIAN': [
        ('worklist','Technician Worklist','قائمة الفني','🩻','technician'),
        ('label','Print Label','طباعة ملصق','🏷️','technician'),
        ('pacs','PACS Tracking','متابعة PACS','🖥️','pacs_tracking'),
        ('priors','Previous Studies','فحوصات سابقة','🕘','patients'),
        ('documents','Documents','المستندات','📎','patients'),
    ],
    'NURSING': [
        ('nursing','Nursing','التمريض','✚','nursing'),
        ('journey','Patient Journey','رحلة المريض','🧭','operations_journey'),
        ('documents','Documents','المستندات','📎','patients'),
        ('history','Previous History','تاريخ سابق','🕘','patients'),
        ('label','Print Label','طباعة ملصق','🏷️','patients'),
    ],
    'DOCTOR': [
        ('doctor','Doctor Workspace','مساحة الطبيب','🩺','doctor'),
        ('report','Reports','التقارير','📄','reports'),
        ('ai','AI Draft','مسودة AI','🤖','doctor'),
        ('priors','Previous Studies','فحوصات سابقة','🕘','doctor'),
        ('documents','Documents','المستندات','📎','doctor'),
        ('word','Word Templates','قوالب Word','📝','doctor_templates'),
    ],
    'MANAGEMENT': [
        ('operations','Operations','التشغيل','📊','operations_dashboard'),
        ('studio','RIS Studio','RIS Studio','🎨','studio_home'),
        ('reports','Financial Reports','التقارير المالية','📊','accounting_reports'),
        ('pacs','PACS','PACS','🖥️','pacs_tracking'),
    ],
}

# RB_BIND:0335
from application.components.shared import install_0335 as _rb_install_0335
quick_action_role = _rb_install_0335(_rb_bind(_rb_sys.modules[__name__], ('current_user',)))

# RB_BIND:0336
from application.components.shared import install_0336 as _rb_install_0336
role_quick_actions = _rb_install_0336(_rb_bind(_rb_sys.modules[__name__], ('ROLE_QUICK_ACTION_PRESETS', 'quick_action_role')))

# RB_BIND:0337
from web.components.shared import install_0337 as _rb_install_0337
safe_url_for_endpoint = _rb_install_0337(_rb_bind(_rb_sys.modules[__name__], ('url_for',)))

# RB_BIND:0338
from web.components.shared import install_0338 as _rb_install_0338
common_search_filters = _rb_install_0338(_rb_bind(_rb_sys.modules[__name__], ('request',)))



DEFAULT_COMM_TEMPLATES=[
 ('BOOKING_CONFIRM','Booking Confirmation','تأكيد الحجز','ANY','BOOKING',
  'Booking confirmation - {{patient_name}}',
  'Dear {{patient_name}}, your appointment for {{procedure}} is scheduled on {{scheduled}}. Accession: {{accession}}.'),
 ('APPOINTMENT_REMINDER','Appointment Reminder','تذكير بالموعد','ANY','REMINDER',
  'Appointment reminder - {{patient_name}}',
  'Reminder: {{patient_name}} has an appointment for {{procedure}} on {{scheduled}}.'),
 ('REPORT_READY','Report Ready','التقرير جاهز','ANY','REPORT_FINAL',
  'Report ready - {{patient_name}}',
  'Dear {{patient_name}}, your report for {{procedure}} is ready. Please use the secure patient portal/link provided by the center.'),
 ('RECEIPT','Payment Receipt','إيصال دفع','ANY','PAYMENT',
  'Payment receipt - {{patient_name}}',
  'Payment received for {{procedure}}. Receipt: {{receipt_no}}. Amount: {{amount}}.'),
 ('DOCUMENT_REQUEST','Document Request','طلب مستند','ANY','DOCUMENT',
  'Required document - {{patient_name}}',
  'Please provide the required document for your visit: {{document_name}}.')
]

# RB_BIND:0339
from application.components.communications import install_0339 as _rb_install_0339
seed_communication_defaults = _rb_install_0339(_rb_bind(_rb_sys.modules[__name__], ('CommunicationChannelSetting', 'CommunicationTemplate', 'DEFAULT_COMM_TEMPLATES', 'db')))

# RB_BIND:0340
from application.components.communications import install_0340 as _rb_install_0340
communication_context = _rb_install_0340(_rb_bind(_rb_sys.modules[__name__], ()))

# RB_BIND:0341
from application.components.shared import install_0341 as _rb_install_0341
render_comm_text = _rb_install_0341(_rb_bind(_rb_sys.modules[__name__], ()))

# RB_BIND:0342
from application.components.registration import install_0342 as _rb_install_0342
patient_default_recipient = _rb_install_0342(_rb_bind(_rb_sys.modules[__name__], ()))

# RB_BIND:0343
from application.components.communications import install_0343 as _rb_install_0343
communication_setting = _rb_install_0343(_rb_bind(_rb_sys.modules[__name__], ('CommunicationChannelSetting',)))

# RB_BIND:0344
from application.components.communications import install_0344 as _rb_install_0344
send_patient_communication = _rb_install_0344(_rb_bind(_rb_sys.modules[__name__], ('communication_setting', 'datetime', 'db', 'os')))



# RB_BIND:0345
from application.components.imaging import install_0345 as _rb_install_0345
imaging_ai_provider_for_modality = _rb_install_0345(_rb_bind(_rb_sys.modules[__name__], ('ImagingAIProvider',)))


# RB_BIND:0346
from application.components.imaging import install_0346 as _rb_install_0346
pacs_study_lookup = _rb_install_0346(_rb_bind(_rb_sys.modules[__name__], ('app', 'os')))


# RB_BIND:0347
from application.components.shared import install_0347 as _rb_install_0347
normalize_person_name = _rb_install_0347(_rb_bind(_rb_sys.modules[__name__], ('re',)))

# RB_BIND:0348
from application.components.shared import install_0348 as _rb_install_0348
normalize_date = _rb_install_0348(_rb_bind(_rb_sys.modules[__name__], ('re',)))

# RB_BIND:0349
from application.components.shared import install_0349 as _rb_install_0349
demographic_verification = _rb_install_0349(_rb_bind(_rb_sys.modules[__name__], ('normalize_date', 'normalize_person_name')))

# RB_BIND:0350
from application.components.imaging import install_0350 as _rb_install_0350
build_imaging_ai_request = _rb_install_0350(_rb_bind(_rb_sys.modules[__name__], ()))

# RB_BIND:0351
from application.components.imaging import install_0351 as _rb_install_0351
call_imaging_ai = _rb_install_0351(_rb_bind(_rb_sys.modules[__name__], ('build_imaging_ai_request', 'imaging_ai_provider_for_modality', 'os')))

# RB_BIND:0352
from application.components.imaging import install_0352 as _rb_install_0352
latest_verified_imaging_analysis = _rb_install_0352(_rb_bind(_rb_sys.modules[__name__], ('ImagingAIAnalysis',)))

# RB_BIND:0353
from application.components.shared import install_0353 as _rb_install_0353
ai_case_context = _rb_install_0353(_rb_bind(_rb_sys.modules[__name__], ('ClinicalNote', 'json', 'latest_verified_imaging_analysis', 'selected_prior_context')))



# RB_BIND:0354
from application.components.assignment import install_0354 as _rb_install_0354
seed_assignment_offer_policy = _rb_install_0354(_rb_bind(_rb_sys.modules[__name__], ('AssignmentOfferPolicy', 'db')))

# RB_BIND:0355
from application.components.assignment import install_0355 as _rb_install_0355
assignment_offer_policy = _rb_install_0355(_rb_bind(_rb_sys.modules[__name__], ('AssignmentOfferPolicy',)))

# RB_BIND:0356
from application.components.assignment import install_0356 as _rb_install_0356
doctor_availability = _rb_install_0356(_rb_bind(_rb_sys.modules[__name__], ('DoctorAvailabilityState', 'datetime', 'db')))

# RB_BIND:0357
from application.components.shared import install_0357 as _rb_install_0357
doctor_is_online = _rb_install_0357(_rb_bind(_rb_sys.modules[__name__], ('datetime', 'doctor_availability', 'timedelta')))

# RB_BIND:0358
from application.components.assignment import install_0358 as _rb_install_0358
assignment_timeout_for = _rb_install_0358(_rb_bind(_rb_sys.modules[__name__], ()))

# RB_BIND:0359
from application.components.assignment import install_0359 as _rb_install_0359
assignment_candidate_users = _rb_install_0359(_rb_bind(_rb_sys.modules[__name__], ('DoctorAssignmentProfile', 'ExternalCenterDoctor', 'ExternalIntakeCase', 'User', 'doctor_availability', 'doctor_is_online')))

# RB_BIND:0360
from application.components.assignment import install_0360 as _rb_install_0360
assignment_offer_channels = _rb_install_0360(_rb_bind(_rb_sys.modules[__name__], ()))

# RB_BIND:0361
from application.components.assignment import install_0361 as _rb_install_0361
create_assignment_offer = _rb_install_0361(_rb_bind(_rb_sys.modules[__name__], ('AssignmentEscalationEvent', 'AssignmentOffer', 'PatientCommunication', 'assignment_offer_channels', 'assignment_offer_policy', 'assignment_strategy_timeout', 'datetime', 'db', 'send_patient_communication', 'timedelta')))

# RB_BIND:0362
from application.components.shared import install_0362 as _rb_install_0362
case_return_snapshot = _rb_install_0362(_rb_bind(_rb_sys.modules[__name__], ('CaseWorkflowEvent',)))


# RB_BIND:0363
from application.components.shared import install_0363 as _rb_install_0363
case_return_is_open = _rb_install_0363(_rb_bind(_rb_sys.modules[__name__], ('case_return_snapshot',)))


# RB_BIND:0364
from application.components.assignment import install_0364 as _rb_install_0364
offer_next_doctor = _rb_install_0364(_rb_bind(_rb_sys.modules[__name__], ('AssignmentEscalationEvent', 'AssignmentOffer', 'CaseAssignment', 'Order', 'Report', 'User', 'assignment_candidate_users', 'assignment_offer_policy', 'assignment_readiness_decision', 'assignment_strategy_max_attempts', 'case_return_is_open', 'create_assignment_offer', 'create_ris_notification', 'datetime', 'db', 'doctor_is_online')))

# RB_BIND:0365
from application.components.assignment import install_0365 as _rb_install_0365
assignment_decision_phase1 = _rb_install_0365(_rb_bind(_rb_sys.modules[__name__], ('CaseReadiness', 'app', 'os')))


# RB_BIND:0366
from application.components.assignment import install_0366 as _rb_install_0366
coordinate_post_exam_assignment = _rb_install_0366(_rb_bind(_rb_sys.modules[__name__], ('AssignmentOffer', 'CaseAssignment', 'Order', 'app', 'assignment_decision_phase1', 'case_return_is_open', 'offer_next_doctor')))


# RB_BIND:0367
from application.components.assignment import install_0367 as _rb_install_0367
expire_assignment_offers = _rb_install_0367(_rb_bind(_rb_sys.modules[__name__], ('AssignmentEscalationEvent', 'AssignmentOffer', 'datetime', 'db', 'offer_next_doctor')))



QUALITY_DEFAULTS=[
 ('CT','BRAIN',80,60,'AXIAL'),
 ('CT','CHEST',250,180,'AXIAL'),
 ('CT','ABDOMEN',300,220,'AXIAL,CORONAL'),
 ('CT','PELVIS',250,180,'AXIAL,CORONAL'),
 ('CT','KUB',180,120,'AXIAL,CORONAL'),
 ('CT','ANGIO',400,300,'THIN,MPR,MIP'),
 ('MR','BRAIN',120,80,'T1,T2,FLAIR,DWI,ADC'),
 ('MR','C-SPINE',100,70,'SAG T1,SAG T2,AX T2'),
 ('MR','L-SPINE',120,80,'SAG T1,SAG T2,AX'),
 ('MR','KNEE',120,80,'SAG,COR,AX'),
 ('MR','SHOULDER',100,70,'AX,COR,SAG'),
 ('US','ABDOMEN',15,8,''),
 ('US','PELVIS',12,6,''),
 ('US','DOPPLER VENOUS',20,12,''),
 ('CR','CHEST',1,1,'PA/AP'),
 ('DX','CHEST',1,1,'PA/AP'),
 ('CR','KNEE',2,1,'AP,LATERAL'),
 ('DX','KNEE',2,1,'AP,LATERAL'),
 ('MG','MAMMOGRAPHY',4,4,'CC,MLO'),
]

# RB_BIND:0368
from application.components.shared import install_0368 as _rb_install_0368
seed_quality_defaults = _rb_install_0368(_rb_bind(_rb_sys.modules[__name__], ('CaseReturnReason', 'NotificationEscalationRule', 'ProcedureQualityRule', 'QUALITY_DEFAULTS', 'TATPolicy', 'db')))

# RB_BIND:0369
from application.components.readiness import install_0369 as _rb_install_0369
quality_rule_for = _rb_install_0369(_rb_bind(_rb_sys.modules[__name__], ('ProcedureQualityRule',)))

# RB_BIND:0370
from application.components.registration import install_0370 as _rb_install_0370
_order_history_present = _rb_install_0370(_rb_bind(_rb_sys.modules[__name__], ('ClinicalNote',)))

# RB_BIND:0371
from application.components.registration import install_0371 as _rb_install_0371
_order_attachment_count = _rb_install_0371(_rb_bind(_rb_sys.modules[__name__], ('PatientDocument',)))

# RB_BIND:0372
from application.components.imaging import install_0372 as _rb_install_0372
pacs_image_count = _rb_install_0372(_rb_bind(_rb_sys.modules[__name__], ('pacs_readiness_evidence',)))

# RB_BIND:0373
from application.components.readiness import install_0373 as _rb_install_0373
pacs_readiness_evidence = _rb_install_0373(_rb_bind(_rb_sys.modules[__name__], ('app', 'pacs_study_lookup')))

# RB_BIND:0374
from application.components.readiness import install_0374 as _rb_install_0374
evaluate_case_readiness = _rb_install_0374(_rb_bind(_rb_sys.modules[__name__], ('CaseReadiness', 'CaseReadinessOverride', 'ReadinessSeriesAlias', '_order_attachment_count', '_order_history_present', 'datetime', 'db', 'json', 'pacs_readiness_evidence', 'quality_rule_for')))

# RB_BIND:0375
from application.components.communications import install_0375 as _rb_install_0375
create_ris_notification = _rb_install_0375(_rb_bind(_rb_sys.modules[__name__], ('RISNotification', 'db')))


# RB_BIND:0376
from application.components.assignment import install_0376 as _rb_install_0376
assignment_strategy_for_order = _rb_install_0376(_rb_bind(_rb_sys.modules[__name__], ('AssignmentStrategy',)))


# RB_BIND:0377
from application.components.assignment import install_0377 as _rb_install_0377
assignment_strategy_timeout = _rb_install_0377(_rb_bind(_rb_sys.modules[__name__], ('assignment_timeout_for',)))


# RB_BIND:0378
from application.components.assignment import install_0378 as _rb_install_0378
assignment_strategy_max_attempts = _rb_install_0378(_rb_bind(_rb_sys.modules[__name__], ()))



# RB_BIND:0379
from application.components.assignment import install_0379 as _rb_install_0379
assignment_readiness_decision = _rb_install_0379(_rb_bind(_rb_sys.modules[__name__], ('CaseReadiness', 'assignment_strategy_for_order', 'datetime', 'evaluate_case_readiness')))


# RB_BIND:0380
from application.components.assignment import install_0380 as _rb_install_0380
readiness_allows_auto_assign = _rb_install_0380(_rb_bind(_rb_sys.modules[__name__], ('assignment_readiness_decision',)))


# RB_TAT_POLICY_RESOLVER_V1
# RB_BIND:0381
from application.components.readiness import install_0381 as _rb_install_0381
resolve_tat_policy = _rb_install_0381(_rb_bind(_rb_sys.modules[__name__], ('TATPolicy',)))


# RB_BIND:0382
from application.components.shared import install_0382 as _rb_install_0382
case_delay_snapshot = _rb_install_0382(_rb_bind(_rb_sys.modules[__name__], ('CaseAssignment', 'CaseReadiness', 'CaseWorkflowEvent', 'Report', 'datetime', 'resolve_tat_policy')))



# RB_BIND:0383
from web.components.registration import install_0383 as _rb_install_0383
advanced_order_query = _rb_install_0383(_rb_bind(_rb_sys.modules[__name__], ('CaseAssignment', 'CaseReadiness', 'CaseWorkflowEvent', 'Order', 'Patient', 'PatientDocument', 'Report', 'case_delay_snapshot', 'external_priors_for_order', 'or_', 'related_prior_orders', 'request')))

# RB_BIND:0384
from application.components.registration import install_0384 as _rb_install_0384
order_action_context = _rb_install_0384(_rb_bind(_rb_sys.modules[__name__], ('Receipt', 'Report')))

# RB_BIND:0385
from application.components.shared import install_0385 as _rb_install_0385
default_label_template = _rb_install_0385(_rb_bind(_rb_sys.modules[__name__], ('LabelTemplate',)))

# RB_BIND:0386
from application.components.reporting import install_0386 as _rb_install_0386
assignment_report_rows = _rb_install_0386(_rb_bind(_rb_sys.modules[__name__], ('AssignmentOffer', 'CaseAssignment', 'CaseReadiness', 'Order', 'Report', 'case_delay_snapshot', 'case_return_snapshot')))



SMART_DOCUMENT_TYPES=[
 ('IDENTITY','National ID / Passport','بطاقة / جواز سفر'),
 ('INSURANCE_CARD','Insurance / Contract Card','كارنيه التأمين / التعاقد'),
 ('PRESCRIPTION','Prescription / Referral','روشتة / تحويل'),
 ('APPROVAL','Medical Approval','موافقة طبية'),
 ('PREVIOUS_REPORT','Previous Report','تقرير سابق'),
 ('PREVIOUS_IMAGES','Previous Images / DICOM','صور / DICOM سابق'),
 ('CONSENT','Consent','إقرار / موافقة'),
 ('QUESTIONNAIRE','Questionnaire','استبيان'),
 ('OTHER','Other','أخرى'),
]

# RB_BIND:0387
from application.components.shared import install_0387 as _rb_install_0387
enabled_ocr_provider = _rb_install_0387(_rb_bind(_rb_sys.modules[__name__], ('OCRProviderSetting',)))



CASE_LIFECYCLE = [
    ('REGISTERED','Registered'),
    ('ARRIVED','Arrived'),
    ('IN_EXAM','In Exam'),
    ('IMAGES_PENDING','Images Pending'),
    ('READY_REPORT','Ready for Reporting'),
    ('ASSIGNED','Assigned'),
    ('REPORTING','Reporting'),
    ('FINAL','Finalized'),
    ('DELIVERED','Delivered'),
]

# RB_BIND:0388
from application.components.registration import install_0388 as _rb_install_0388
order_visual_state = _rb_install_0388(_rb_bind(_rb_sys.modules[__name__], ('Attachment', 'CaseAssignment', 'CaseReadiness', 'Report')))

# RB_BIND:0389
from application.components.shared import install_0389 as _rb_install_0389
resolve_printer_profile = _rb_install_0389(_rb_bind(_rb_sys.modules[__name__], ('PrinterProfile', 'current_user')))

# RB_BIND:0390
from application.components.shared import install_0390 as _rb_install_0390
generated_attachment_filename = _rb_install_0390(_rb_bind(_rb_sys.modules[__name__], ('os', 're', 'uuid')))

# RB_BIND:0391
from application.components.registration import install_0391 as _rb_install_0391
smart_intake_dir = _rb_install_0391(_rb_bind(_rb_sys.modules[__name__], ('ATTACHMENT_DIR', 'os')))

# RB_BIND:0392
from application.components.shared import install_0392 as _rb_install_0392
_simple_identity_extract = _rb_install_0392(_rb_bind(_rb_sys.modules[__name__], ('normalize_intake_identity', 're', 'validate_intake_identity')))


# RB_BIND:0393
from application.components.registration import install_0393 as _rb_install_0393
_suggest_procedure_from_text = _rb_install_0393(_rb_bind(_rb_sys.modules[__name__], ('Procedure', 're')))

# RB_BIND:0394
from application.components.registration import install_0394 as _rb_install_0394
_suggest_contract_from_text = _rb_install_0394(_rb_bind(_rb_sys.modules[__name__], ('Contract', 're')))


INTAKE_IDENTITY_SOURCE_PRIORITY={
    'IDENTITY':100,
    'INSURANCE_CARD':80,
    'PRESCRIPTION':60,
    'APPROVAL':50,
    'QUESTIONNAIRE':40,
    'CONSENT':30,
    'OTHER':20,
}


# RB_BIND:0395
from application.components.registration import install_0395 as _rb_install_0395
_intake_source_field = _rb_install_0395(_rb_bind(_rb_sys.modules[__name__], ()))


# RB_BIND:0396
from application.components.registration import install_0396 as _rb_install_0396
_intake_normalized_field = _rb_install_0396(_rb_bind(_rb_sys.modules[__name__], ('normalize_intake_identity',)))


# RB_BIND:0397
from web.components.registration import install_0397 as _rb_install_0397
_intake_redirect = _rb_install_0397(_rb_bind(_rb_sys.modules[__name__], ('re', 'redirect', 'url_for')))

# RB_BIND:0398
from application.components.communications import install_0398 as _rb_install_0398
call_ocr_provider = _rb_install_0398(_rb_bind(_rb_sys.modules[__name__], ('OCRProviderError', '_simple_identity_extract', 'app', 'enabled_ocr_provider', 'normalize_intake_identity', 'ocr_configuration_hash', 'ocr_extract_file', 'os')))

# RB_BIND:0399
from application.components.registration import install_0399 as _rb_install_0399
_intake_file_sha256 = _rb_install_0399(_rb_bind(_rb_sys.modules[__name__], ('hashlib',)))


# RB_BIND:0400
from application.components.registration import install_0400 as _rb_install_0400
_intake_final_identity = _rb_install_0400(_rb_bind(_rb_sys.modules[__name__], ('SmartIntakeFieldReview', 'json', 'normalize_intake_identity')))


# RB_BIND:0401
from application.components.registration import install_0401 as _rb_install_0401
_intake_candidate_patients = _rb_install_0401(_rb_bind(_rb_sys.modules[__name__], ('Patient', 'db', 'func', 'intake_digits', 'normalize_intake_identity', 'or_')))


# RB_BIND:0402
from application.components.registration import install_0402 as _rb_install_0402
_intake_match_context = _rb_install_0402(_rb_bind(_rb_sys.modules[__name__], ('_intake_candidate_patients', 'score_intake_patient_candidates')))


# RB_BIND:0403
from application.components.registration import install_0403 as _rb_install_0403
_sync_intake_field_reviews = _rb_install_0403(_rb_bind(_rb_sys.modules[__name__], ('OCR_INTAKE_CORE_FIELDS', 'SmartIntakeFieldReview', '_intake_source_field', 'db', 'normalize_intake_identity', 'validate_intake_identity')))


# RB_BIND:0404
from application.components.registration import install_0404 as _rb_install_0404
rebuild_intake_summary = _rb_install_0404(_rb_bind(_rb_sys.modules[__name__], ('INTAKE_IDENTITY_SOURCE_PRIORITY', 'OCR_INTAKE_CORE_FIELDS', 'SmartIntakeDocument', 'SmartIntakeFieldReview', '_intake_final_identity', '_intake_match_context', '_intake_normalized_field', '_intake_source_field', '_suggest_contract_from_text', '_suggest_procedure_from_text', '_sync_intake_field_reviews', 'db', 'json', 'validate_intake_identity')))


# RB_BIND:0405
from web.components.registration import install_0405 as _rb_install_0405
api_slots = _rb_install_0405(_rb_bind(_rb_sys.modules[__name__], ('Device', 'Procedure', 'abort', 'datetime', 'generate_slots', 'jsonify', 'request', 'require_perm')))

# RB_BIND:0406
from web.routes.registration import install_0406 as _rb_install_0406
schedule_calendar = _rb_install_0406(_rb_bind(_rb_sys.modules[__name__], ('Appointment', 'Branch', 'Contract', 'Device', 'Patient', 'Procedure', 'ReferringDoctor', 'SMART_DOCUMENT_TYPES', 'abort', 'app', 'assignable_doctors', 'build_calendar_cells', 'has_perm', 'or_', 'render_template', 'request', 'require_perm')))


# RB_BIND:0407
from application.components.shared import install_0407 as _rb_install_0407
build_calendar_cells = _rb_install_0407(_rb_bind(_rb_sys.modules[__name__], ('timedelta',)))



# ============================================================
# v5.46.2 Unified Scheduler + Registration Workspace
# ============================================================

# RB_BIND:0408
from web.components.registration import install_0408 as _rb_install_0408
_rb_v5462_allow_registration = _rb_install_0408(_rb_bind(_rb_sys.modules[__name__], ('abort', 'current_user', 'has_perm')))

# RB_BIND:0409
from application.components.shared import install_0409 as _rb_install_0409
_rb_v5462_proc_ids = _rb_install_0409(_rb_bind(_rb_sys.modules[__name__], ()))

# RB_BIND:0410
from application.components.registration import install_0410 as _rb_install_0410
_rb_v5462_patient = _rb_install_0410(_rb_bind(_rb_sys.modules[__name__], ('Patient', '_rb_registration_v2_create_patient')))

# RB_BIND:0411
from web.components.shared import install_0411 as _rb_install_0411
_rb_v5462_direct_attachments = _rb_install_0411(_rb_bind(_rb_sys.modules[__name__], ('ATTACHMENT_DIR', 'Attachment', 'current_user', 'db', 'generated_attachment_filename', 'os', 'request')))

# RB_BIND:0412
from web.components.shared import install_0412 as _rb_install_0412
_rb_v5462_attach_ocr = _rb_install_0412(_rb_bind(_rb_sys.modules[__name__], ('Attachment', 'PatientDocument', 'SmartIntakeDocument', 'SmartIntakeSession', 'current_user', 'datetime', 'db', 'request')))

# RB_BIND:0413
from application.components.registration import install_0413 as _rb_install_0413
_rb_v5462_insurance = _rb_install_0413(_rb_bind(_rb_sys.modules[__name__], ('InsuranceApprovalEvent', 'InsuranceApprovalRequest', 'current_user', 'datetime', 'db')))

# RB_BIND:0414
from application.components.shared import install_0414 as _rb_install_0414
_rb_v5462_notes = _rb_install_0414(_rb_bind(_rb_sys.modules[__name__], ('ClinicalNote', 'current_user', 'db')))

# RB_BIND:0415
from web.components.finance import install_0415 as _rb_install_0415
_rb_v5462_payment = _rb_install_0415(_rb_bind(_rb_sys.modules[__name__], ('Payment', 'Receipt', 'abort', 'audit', 'current_user', 'db', 'has_perm', 'next_receipt_no')))

# RB_BIND:0416
from web.routes.registration import install_0416 as _rb_install_0416
schedule_ocr_preview = _rb_install_0416(_rb_bind(_rb_sys.modules[__name__], ('Contract', 'Procedure', 'SMART_DOCUMENT_TYPES', 'SmartIntakeDocument', 'SmartIntakeSession', '_apply_smart_intake_ocr_result', '_intake_file_sha256', '_rb_v5462_allow_registration', 'app', 'call_ocr_provider', 'current_user', 'db', 'json', 'jsonify', 'login_required', 'os', 're', 'request', 'secrets', 'smart_intake_dir', 'uuid')))

# RB_BIND:0417
from web.routes.registration import install_0417 as _rb_install_0417
schedule_unified_submit = _rb_install_0417(_rb_bind(_rb_sys.modules[__name__], ('Appointment', 'Branch', 'Contract', 'Device', 'Order', 'OrderConsumable', 'Procedure', 'ProcedureConsumable', '_rb_calendar_redirect', '_rb_v5462_allow_registration', '_rb_v5462_attach_ocr', '_rb_v5462_direct_attachments', '_rb_v5462_insurance', '_rb_v5462_notes', '_rb_v5462_patient', '_rb_v5462_payment', '_rb_v5462_proc_ids', 'abort', 'acquire_booking_slot_lock', 'app', 'audit', 'booking_restriction_reason', 'current_user', 'datetime', 'db', 'ensure_journey', 'flash', 'has_perm', 'login_required', 'new_accession', 'os', 'price_order', 'procedure_available_in_branch', 'redirect', 'request', 'timedelta', 'url_for')))


# RB_BIND:0418
from web.routes.registration import install_0418 as _rb_install_0418
schedule_book = _rb_install_0418(_rb_bind(_rb_sys.modules[__name__], ('Appointment', 'Branch', 'Device', 'Patient', 'Procedure', '_rb_calendar_redirect', '_rb_registration_v2_create_patient', 'acquire_booking_slot_lock', 'app', 'audit', 'booking_restriction_reason', 'current_user', 'datetime', 'db', 'flash', 'procedure_available_in_branch', 'redirect', 'request', 'require_perm', 'timedelta', 'url_for')))

# RB_BIND:0419
from web.routes.registration import install_0419 as _rb_install_0419
order_workflow_status = _rb_install_0419(_rb_bind(_rb_sys.modules[__name__], ('Order', 'app', 'canonical_display_state', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_BIND:0420
from web.routes.registration import install_0420 as _rb_install_0420
appointments = _rb_install_0420(_rb_bind(_rb_sys.modules[__name__], ('Appointment', 'Branch', 'Device', 'Patient', 'Procedure', 'app', 'render_template', 'request', 'require_perm')))


# RB_BIND:0421
from web.routes.registration import install_0421 as _rb_install_0421
appointment_action = _rb_install_0421(_rb_bind(_rb_sys.modules[__name__], ('Appointment', 'Order', 'abort', 'app', 'audit', 'current_user', 'datetime', 'db', 'flash', 'new_accession', 'price_order', 'redirect', 'require_perm', 'route_device', 'url_for')))




# RB_BIND:0422
from web.routes.registration import install_0422 as _rb_install_0422
appointment_edit = _rb_install_0422(_rb_bind(_rb_sys.modules[__name__], ('Appointment', 'AssignmentVisit', 'Device', 'DoctorAssignmentProfile', 'Order', 'Patient', 'PatientDemographicsExtra', 'Procedure', 'User', 'abort', 'acquire_booking_slot_lock', 'app', 'appointment_overlap', 'arabic_to_latin_name', 'audit', 'booking_restriction_reason', 'datetime', 'db', 'flash', 'procedure_available_in_branch', 'redirect', 'render_template', 'request', 'require_perm', 'resolve_patient_dob', 'url_for')))

# RB_BIND:0423
from web.routes.registration import install_0423 as _rb_install_0423
appointment_reschedule = _rb_install_0423(_rb_bind(_rb_sys.modules[__name__], ('Appointment', 'AssignmentVisit', 'Device', 'acquire_booking_slot_lock', 'app', 'appointment_overlap', 'audit', 'booking_restriction_reason', 'datetime', 'db', 'jsonify', 'request', 'require_perm')))

# RB_BIND:0424
from web.routes.registration import install_0424 as _rb_install_0424
appointment_status = _rb_install_0424(_rb_bind(_rb_sys.modules[__name__], ('Appointment', 'abort', 'app', 'audit', 'db', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0425
from web.routes.registration import install_0425 as _rb_install_0425
schedule_maintenance = _rb_install_0425(_rb_bind(_rb_sys.modules[__name__], ('Device', 'DeviceBlock', 'abort', 'app', 'audit', 'datetime', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_BIND:0426
from web.routes.registration import install_0426 as _rb_install_0426
schedule_v2 = _rb_install_0426(_rb_bind(_rb_sys.modules[__name__], ('app', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0427
from web.routes.registration import install_0427 as _rb_install_0427
schedule_v2_book_form = _rb_install_0427(_rb_bind(_rb_sys.modules[__name__], ('app', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0428
from web.routes.registration import install_0428 as _rb_install_0428
schedule_v2_book_submit = _rb_install_0428(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Device', 'Procedure', 'app', 'booking_restriction_reason', 'datetime', 'flash', 'redirect', 'request', 'require_perm', 'schedule_book', 'url_for')))

# RB_BIND:0429
from web.routes.registration import install_0429 as _rb_install_0429
device_booking_status = _rb_install_0429(_rb_bind(_rb_sys.modules[__name__], ('DEVICE_BOOKING_STATUSES', 'Device', '_hhmm_to_minutes', 'abort', 'app', 'audit', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_BIND:0430
from web.routes.registration import install_0430 as _rb_install_0430
admin_booking_availability = _rb_install_0430(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Procedure', 'ProcedureBranch', 'app', 'audit', 'db', 'flash', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))


# RB_SMART_BOOKING_V2_AVAILABILITY_API_V1
# RB_BIND:0431
from web.routes.registration import install_0431 as _rb_install_0431
api_smart_booking_availability = _rb_install_0431(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Device', 'Procedure', 'app', 'booking_restriction_reason', 'booking_v2_slots', 'datetime', 'db', 'jsonify', 'procedure_available_in_branch', 'request', 'require_perm')))


# RB_BIND:0432
from web.routes.registration import install_0432 as _rb_install_0432
admin_booking_availability_seed = _rb_install_0432(_rb_bind(_rb_sys.modules[__name__], ('app', 'audit', 'db', 'ensure_default_procedure_branch_mappings', 'flash', 'redirect', 'require_perm', 'url_for')))

# RB_BIND:0433
from web.components.registration import install_0433 as _rb_install_0433
api_booking_procedures = _rb_install_0433(_rb_bind(_rb_sys.modules[__name__], ('Device', 'allowed_procedures_for_branch', 'jsonify', 'request', 'require_perm')))


# RB_BIND:0434
from web.routes.clinical import install_0434 as _rb_install_0434
order_clinical_notes = _rb_install_0434(_rb_bind(_rb_sys.modules[__name__], ('ClinicalNote', 'Order', 'app', 'current_user', 'db', 'flash', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0435
from web.routes.shared import install_0435 as _rb_install_0435
my_profile = _rb_install_0435(_rb_bind(_rb_sys.modules[__name__], ('abort', 'app', 'audit', 'current_user', 'db', 'doctor_workspace_user', 'flash', 'get_user_profile', 'image_viewer_preference', 'login_required', 'redirect', 'render_template', 'request', 'save_image_viewer_preference', 'set_user_reporting_preferences', 'url_for', 'user_reporting_preferences')))

# RB_BIND:0436
from web.routes.operations import install_0436 as _rb_install_0436
admin_dictionary = _rb_install_0436(_rb_bind(_rb_sys.modules[__name__], ('DictionaryEntry', 'app', 'db', 'or_', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))


# RB_BIND:0437
from web.routes.registration import install_0437 as _rb_install_0437
admin_booking_control_center = _rb_install_0437(_rb_bind(_rb_sys.modules[__name__], ('Appointment', 'BookingBlock', 'Device', 'DeviceBlock', 'app', 'datetime', 'render_template', 'require_perm', 'timedelta')))


# RB_BIND:0438
from web.routes.registration import install_0438 as _rb_install_0438
admin_booking_control = _rb_install_0438(_rb_bind(_rb_sys.modules[__name__], ('BOOKING_BLOCK_REASONS', 'BookingBlock', 'Branch', 'Device', 'Procedure', 'User', 'app', 'current_user', 'datetime', 'db', 'flash', 'redirect', 'render_template', 'request', 'require_perm', 'url_for', 'validate_booking_block_scope')))


# RB_BIND:0439
from web.routes.registration import install_0439 as _rb_install_0439
admin_booking_control_deactivate = _rb_install_0439(_rb_bind(_rb_sys.modules[__name__], ('BookingBlock', 'app', 'audit', 'db', 'flash', 'redirect', 'require_perm', 'url_for')))


# RB_BIND:0440
from web.routes.shared import install_0440 as _rb_install_0440
ui_language_switch = _rb_install_0440(_rb_bind(_rb_sys.modules[__name__], ('UserProfile', 'abort', 'app', 'current_user', 'db', 'login_required', 'redirect', 'request', 'session', 'url_for')))

# RB_BIND:0441
from web.routes.operations import install_0441 as _rb_install_0441
admin_dictionary_reset = _rb_install_0441(_rb_bind(_rb_sys.modules[__name__], ('DEFAULT_UI_MAP', 'DictionaryEntry', 'app', 'db', 'flash', 'redirect', 'require_perm', 'url_for')))


# RB_BIND:0442
from web.routes.shared import install_0442 as _rb_install_0442
my_ribbon = _rb_install_0442(_rb_bind(_rb_sys.modules[__name__], ('RibbonAction', 'UserRibbonAction', 'abort', 'app', 'current_user', 'db', 'flash', 'has_perm', 'login_required', 'redirect', 'render_template', 'request', 'url_for')))

# RB_BIND:0443
from web.routes.operations import install_0443 as _rb_install_0443
admin_ribbon_designer = _rb_install_0443(_rb_bind(_rb_sys.modules[__name__], ('RibbonAction', 'app', 'db', 'flash', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0444
from web.routes.shared import install_0444 as _rb_install_0444
operations_journey = _rb_install_0444(_rb_bind(_rb_sys.modules[__name__], ('JOURNEY_STAGES', 'Order', 'Patient', 'app', 'journey_snapshot', 'or_', 'render_template', 'request', 'require_perm')))

# RB_BIND:0445
from web.routes.registration import install_0445 as _rb_install_0445
order_journey_stage = _rb_install_0445(_rb_bind(_rb_sys.modules[__name__], ('JOURNEY_STAGES', 'Order', 'abort', 'app', 'flash', 'journey_snapshot', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0446
from web.routes.registration import install_0446 as _rb_install_0446
operations_booking_board = _rb_install_0446(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'app', 'datetime', 'render_template', 'request', 'require_perm', 'smart_booking_cards')))

# RB_BIND:0447
from web.routes.shared import install_0447 as _rb_install_0447
operations_dashboard = _rb_install_0447(_rb_bind(_rb_sys.modules[__name__], ('Order', 'Payment', 'PaymentRefund', 'app', 'datetime', 'journey_snapshot', 'render_template', 'require_perm', 'role_dashboard_kind', 'timedelta', 'unread_chat_count')))

# RB_BIND:0448
from web.routes.communications import install_0448 as _rb_install_0448
chat_home = _rb_install_0448(_rb_bind(_rb_sys.modules[__name__], ('ChatThreadMember', 'User', 'app', 'current_user', 'login_required', 'render_template')))

# RB_BIND:0449
from web.routes.communications import install_0449 as _rb_install_0449
chat_direct_create = _rb_install_0449(_rb_bind(_rb_sys.modules[__name__], ('ChatThread', 'ChatThreadMember', 'User', 'abort', 'app', 'current_user', 'db', 'ensure_chat_member', 'login_required', 'redirect', 'request', 'url_for')))

# RB_BIND:0450
from web.routes.communications import install_0450 as _rb_install_0450
chat_group_create = _rb_install_0450(_rb_bind(_rb_sys.modules[__name__], ('ChatThread', 'app', 'current_user', 'db', 'ensure_chat_member', 'login_required', 'redirect', 'request', 'url_for')))

# RB_BIND:0451
from web.routes.communications import install_0451 as _rb_install_0451
chat_case = _rb_install_0451(_rb_bind(_rb_sys.modules[__name__], ('Order', 'app', 'current_user', 'db', 'ensure_chat_member', 'get_or_create_case_thread', 'login_required', 'redirect', 'url_for')))

# RB_BIND:0452
from web.routes.communications import install_0452 as _rb_install_0452
chat_thread = _rb_install_0452(_rb_bind(_rb_sys.modules[__name__], ('ATTACHMENT_DIR', 'ChatMessage', 'ChatThread', 'ChatThreadMember', 'abort', 'app', 'current_user', 'datetime', 'db', 'login_required', 'os', 're', 'redirect', 'render_template', 'request', 'url_for')))

# RB_BIND:0453
from web.routes.clinical import install_0453 as _rb_install_0453
chat_message_to_clinical_note = _rb_install_0453(_rb_bind(_rb_sys.modules[__name__], ('ChatMessage', 'ChatThread', 'ClinicalNote', 'abort', 'app', 'audit', 'current_user', 'db', 'redirect', 'require_perm', 'url_for')))


# RB_BIND:0454
from web.routes.registration import install_0454 as _rb_install_0454
patient_documents = _rb_install_0454(_rb_bind(_rb_sys.modules[__name__], ('ATTACHMENT_DIR', 'DOCUMENT_TYPES', 'Order', 'Patient', 'PatientDocument', 'app', 'audit', 'build_document_matches', 'current_user', 'db', 'extract_document_text_basic', 'flash', 'generated_attachment_filename', 'intake', 'login_required', 'os', 'redirect', 'render_template', 'request', 'url_for')))

# RB_BIND:0455
from web.routes.registration import install_0455 as _rb_install_0455
patient_timeline = _rb_install_0455(_rb_bind(_rb_sys.modules[__name__], ('Patient', 'app', 'document_timeline', 'login_required', 'render_template')))

# RB_BIND:0456
from web.routes.shared import install_0456 as _rb_install_0456
document_review = _rb_install_0456(_rb_bind(_rb_sys.modules[__name__], ('DOCUMENT_REVIEW_STATUSES', 'DocumentMatch', 'PatientDocument', 'app', 'login_required', 'render_template')))

# RB_BIND:0457
from web.routes.shared import install_0457 as _rb_install_0457
document_review_save = _rb_install_0457(_rb_bind(_rb_sys.modules[__name__], ('DOCUMENT_REVIEW_STATUSES', 'PatientDocument', 'abort', 'app', 'audit', 'current_user', 'datetime', 'db', 'login_required', 'redirect', 'request', 'url_for')))

# RB_BIND:0458
from web.routes.communications import install_0458 as _rb_install_0458
telephony_incoming = _rb_install_0458(_rb_bind(_rb_sys.modules[__name__], ('CallLog', 'TelephonyExtension', 'app', 'db', 'jsonify', 'patient_by_phone', 'request')))

# RB_BIND:0459
from web.routes.communications import install_0459 as _rb_install_0459
telephony_calls = _rb_install_0459(_rb_bind(_rb_sys.modules[__name__], ('CallLog', 'app', 'login_required', 'render_template')))

# RB_BIND:0460
from web.routes.communications import install_0460 as _rb_install_0460
telephony_caller = _rb_install_0460(_rb_bind(_rb_sys.modules[__name__], ('CallLog', 'Order', 'PatientDocument', 'app', 'login_required', 'render_template')))

# RB_BIND:0461
from web.routes.communications import install_0461 as _rb_install_0461
telephony_call_notes = _rb_install_0461(_rb_bind(_rb_sys.modules[__name__], ('CallLog', 'app', 'db', 'login_required', 'redirect', 'request', 'url_for')))


# RB_BIND:0462
from web.routes.operations import install_0462 as _rb_install_0462
admin_referrers = _rb_install_0462(_rb_bind(_rb_sys.modules[__name__], ('ReferringDoctor', 'app', 'audit', 'db', 'flash', 'render_template', 'request', 'require_perm')))

# RB_BIND:0463
from web.components.shared import install_0463 as _rb_install_0463
analytics_query = _rb_install_0463(_rb_bind(_rb_sys.modules[__name__], ('Order', 'datetime', 'request')))

# RB_BIND:0464
from web.components.shared import install_0464 as _rb_install_0464
analytics = _rb_install_0464(_rb_bind(_rb_sys.modules[__name__], ('Appointment', 'Branch', 'Contract', 'Order', 'OrderConsumable', 'analytics_query', 'datetime', 'order_balance', 'render_template', 'request', 'require_perm', 'timedelta')))

# RB_BIND:0465
from web.components.shared import install_0465 as _rb_install_0465
analytics_export_csv = _rb_install_0465(_rb_bind(_rb_sys.modules[__name__], ('Order', 'analytics_query', 'csv', 'io', 'require_perm', 'send_file')))

# RB_BIND:0466
from web.routes.operations import install_0466 as _rb_install_0466
admin_import_export = _rb_install_0466(_rb_bind(_rb_sys.modules[__name__], ('Consumable', 'Contract', 'Procedure', 'abort', 'app', 'audit', 'csv', 'db', 'flash', 'has_perm', 'io', 'login_required', 'redirect', 'render_template', 'request', 'url_for')))

# RB_BIND:0467
from web.routes.operations import install_0467 as _rb_install_0467
admin_export_entity = _rb_install_0467(_rb_bind(_rb_sys.modules[__name__], ('Consumable', 'Contract', 'Device', 'Procedure', 'abort', 'app', 'csv', 'has_perm', 'io', 'login_required', 'send_file')))

# RB_BIND:0468
from web.routes.reporting import install_0468 as _rb_install_0468
report_upload = _rb_install_0468(_rb_bind(_rb_sys.modules[__name__], ('Order', 'REPORT_DIR', 'Report', 'ReportEditControlError', 'ReportVersion', '_error_id', 'abort', 'app', 'assert_lease_owner', 'assert_not_final', 'assert_role_allowed', 'audit', 'current_user', 'db', 'extract_docx_sections', 'flash', 'get_report', 'os', 'redirect', 'request', 'require_perm', 'shutil', 'tempfile', 'url_for', 'validate_report_patient_file')))


# RB_BIND:0469
from web.routes.reporting import install_0469 as _rb_install_0469
report_source = _rb_install_0469(_rb_bind(_rb_sys.modules[__name__], ('Order', 'Report', 'ReportVersion', 'abort', 'app', 'os', 'report_filename', 'report_source_allowed', 'require_perm', 'send_file')))


# RB_BIND:0470
from web.routes.shared import install_0470 as _rb_install_0470
attachment_upload = _rb_install_0470(_rb_bind(_rb_sys.modules[__name__], ('ATTACHMENT_DIR', 'Attachment', 'app', 'audit', 'current_user', 'db', 'flash', 'generated_attachment_filename', 'os', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0471
from web.routes.shared import install_0471 as _rb_install_0471
attachment_download = _rb_install_0471(_rb_bind(_rb_sys.modules[__name__], ('Attachment', 'app', 'require_perm', 'send_file')))


# RB_BIND:0472
from web.routes.shared import install_0472 as _rb_install_0472
attachment_preview = _rb_install_0472(_rb_bind(_rb_sys.modules[__name__], ('Attachment', 'abort', 'app', 'os', 'require_perm', 'send_file')))

# RB_BIND:0473
from web.routes.shared import install_0473 as _rb_install_0473
attachment_delete = _rb_install_0473(_rb_bind(_rb_sys.modules[__name__], ('Attachment', 'app', 'audit', 'db', 'flash', 'os', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_BIND:0474
from web.routes.registration import install_0474 as _rb_install_0474
order_attachments = _rb_install_0474(_rb_bind(_rb_sys.modules[__name__], ('Attachment', 'Order', 'app', 'render_template', 'require_perm')))

# ---------------- Accounting / Receivables / Delivery ----------------



# RB_BIND:0475
from web.routes.registration import install_0475 as _rb_install_0475
order_sticker = _rb_install_0475(_rb_bind(_rb_sys.modules[__name__], ('Order', '_portal_url_for_order', '_sticker_fields', '_sticker_labels', '_sticker_value', 'app', 'canvas', 'code128', 'db', 'get_company_profile', 'get_default_sticker_template', 'io', 'qrcode', 'require_perm', 'send_file')))

# RB_BIND:0476
from web.routes.registration import install_0476 as _rb_install_0476
order_sticker_preview = _rb_install_0476(_rb_bind(_rb_sys.modules[__name__], ('Order', '_portal_url_for_order', '_sticker_fields', '_sticker_labels', 'app', 'db', 'get_company_profile', 'get_default_sticker_template', 'render_template', 'require_perm')))

# RB_BIND:0477
from web.routes.shared import install_0477 as _rb_install_0477
barcode_scan = _rb_install_0477(_rb_bind(_rb_sys.modules[__name__], ('Order', 'Patient', 'app', 'flash', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0478
from web.routes.operations import install_0478 as _rb_install_0478
admin_company = _rb_install_0478(_rb_bind(_rb_sys.modules[__name__], ('ATTACHMENT_DIR', 'Path', 'app', 'audit', 'db', 'flash', 'get_company_profile', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0479
from web.routes.shared import install_0479 as _rb_install_0479
company_logo = _rb_install_0479(_rb_bind(_rb_sys.modules[__name__], ('abort', 'app', 'get_company_profile', 'os', 'send_file')))

# RB_BIND:0480
from web.routes.operations import install_0480 as _rb_install_0480
admin_sticker_templates = _rb_install_0480(_rb_bind(_rb_sys.modules[__name__], ('DEFAULT_STICKER_FIELDS', 'DEFAULT_STICKER_LABELS', '_sticker_fields', '_sticker_labels', 'app', 'audit', 'db', 'flash', 'get_default_sticker_template', 'json', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))


# RB_BIND:0481
from web.routes.imaging import install_0481 as _rb_install_0481
pacs_reconciliation_preview = _rb_install_0481(_rb_bind(_rb_sys.modules[__name__], ('Order', 'PacsStudyInbox', '__name__', 'app', 'jsonify', 'request', 'requests', 'require_perm')))


# RB_BIND:0482
from web.routes.imaging import install_0482 as _rb_install_0482
pacs_reconciliation_correct_dicom = _rb_install_0482(_rb_bind(_rb_sys.modules[__name__], ('Order', 'PacsStudyInbox', '__name__', 'abort', 'app', 'db', 'flash', 'has_perm', 'redirect', 'request', 'requests', 'require_perm', 'url_for')))


# RB_BIND:0483
from web.routes.imaging import install_0483 as _rb_install_0483
pacs_reconciliation_audit = _rb_install_0483(_rb_bind(_rb_sys.modules[__name__], ('DicomCorrection', 'Order', 'abort', 'app', 'current_user', 'has_perm', 'json', 'render_template', 'require_perm')))


# RB_BIND:0484
from web.routes.imaging import install_0484 as _rb_install_0484
pacs_reconciliation_example = _rb_install_0484(_rb_bind(_rb_sys.modules[__name__], ('app', 'render_template', 'require_perm')))

# RB_BIND:0485
from web.routes.imaging import install_0485 as _rb_install_0485
pacs_reconciliation = _rb_install_0485(_rb_bind(_rb_sys.modules[__name__], ('DicomCorrection', 'Order', 'PacsExamLink', 'PacsStudyInbox', 'Patient', 'Report', '__name__', 'app', 'correction_search_orders', 'current_user', 'db', 'flash', 'has_perm', 'render_template', 'request', 'requests', 'require_perm')))


# RB_BIND:0486
from web.routes.imaging import install_0486 as _rb_install_0486
pacs_exam_link_create = _rb_install_0486(_rb_bind(_rb_sys.modules[__name__], ('Order', 'PacsExamLink', 'Report', 'app', 'audit', 'current_user', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_BIND:0487
from web.routes.imaging import install_0487 as _rb_install_0487
pacs_exam_link_remove = _rb_install_0487(_rb_bind(_rb_sys.modules[__name__], ('Order', 'PacsExamLink', 'Report', 'app', 'audit', 'current_user', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_BIND:0488
from web.routes.imaging import install_0488 as _rb_install_0488
pacs_reconciliation_match = _rb_install_0488(_rb_bind(_rb_sys.modules[__name__], ('Order', 'PacsResolutionEvent', 'PacsStudyInbox', 'PacsStudyMatch', 'Report', '__name__', '_orthanc_recent_studies', '_study_tags_extended', 'abort', 'app', 'audit', 'current_user', 'db', 'enqueue_auto_routes_for_inbox', 'evaluate_case_readiness', 'flash', 'json', 'redirect', 'request', 'requests', 'require_perm', 'url_for')))


# RB_BIND:0489
from web.routes.imaging import install_0489 as _rb_install_0489
pacs_reconciliation_quarantine = _rb_install_0489(_rb_bind(_rb_sys.modules[__name__], ('PacsStudyInbox', 'app', 'audit', 'current_user', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_BIND:0490
from web.routes.imaging import install_0490 as _rb_install_0490
pacs_reconciliation_auto_safe = _rb_install_0490(_rb_bind(_rb_sys.modules[__name__], ('DicomCorrection', 'Order', 'PacsResolutionEvent', 'PacsStudyInbox', 'PacsStudyMatch', 'Report', '_orthanc_recent_studies', '_study_tags_extended', 'app', 'audit', 'current_user', 'db', 'evaluate_case_readiness', 'flash', 'redirect', 'require_perm', 'strict_pacs_demographic_match', 'url_for')))


# RB_BIND:0491
from web.components.operations import install_0491 as _rb_install_0491
admin_qa = _rb_install_0491(_rb_bind(_rb_sys.modules[__name__], ('EventOutbox', 'Order', 'Patient', 'User', 'render_template')))


# ============================================================
# RadBridge Live Command Center
# Read-only observability API.
# RB_LIVE_COMMAND_CENTER_SNAPSHOT_V1
# ============================================================

# RB_LIVE_COMMAND_CENTER_UI_V1

# ============================================================

# RB_COMMAND_CENTER_CASE_GRAPH_V3
# Read-only per-case workflow contract for the Live Command Center.
# RB_BIND:0492
from web.routes.registration import install_0492 as _rb_install_0492
api_command_center_case_graph_v3 = _rb_install_0492(_rb_bind(_rb_sys.modules[__name__], ('CaseAssignment', 'CaseReadiness', 'CaseWorkflowEvent', 'DoctorAvailabilityState', 'Order', 'PacsStudyInbox', 'PacsStudyMatch', 'app', 'datetime', 'jsonify', 'require_perm', 'timedelta', 'url_for')))



# ============================================================
# RB_COMMAND_CENTER_MONITORING_CONTROL_V1
# Persistent multi-worker runtime control.
#
# OFF:
#   No SSE monitoring loop.
#
# NORMAL:
#   Low-frequency SSE refresh.
#
# LIVE:
#   Real-time SSE refresh.
#
# Storage is intentionally outside the application image.
# A persistent bind mount is required before writes are enabled.
# ============================================================

# RB_BIND:0493
from application.components.shared import install_0493 as _rb_install_0493
_command_center_control_defaults = _rb_install_0493(_rb_bind(_rb_sys.modules[__name__], ('os',)))


# RB_BIND:0494
from application.components.shared import install_0494 as _rb_install_0494
_command_center_control_path = _rb_install_0494(_rb_bind(_rb_sys.modules[__name__], ('os',)))


# RB_BIND:0495
from application.components.shared import install_0495 as _rb_install_0495
command_center_control_state = _rb_install_0495(_rb_bind(_rb_sys.modules[__name__], ('_command_center_control_defaults', '_command_center_control_path', 'app', 'json')))


# RB_BIND:0496
from application.components.shared import install_0496 as _rb_install_0496
_command_center_write_control_state = _rb_install_0496(_rb_bind(_rb_sys.modules[__name__], ('_command_center_control_path', 'app', 'command_center_control_state', 'datetime', 'json', 'os')))


# RB_BIND:0497
from web.routes.shared import install_0497 as _rb_install_0497
api_command_center_control_get = _rb_install_0497(_rb_bind(_rb_sys.modules[__name__], ('app', 'command_center_control_state', 'jsonify', 'require_perm')))


# RB_BIND:0498
from web.routes.shared import install_0498 as _rb_install_0498
api_command_center_control_set = _rb_install_0498(_rb_bind(_rb_sys.modules[__name__], ('_command_center_write_control_state', 'abort', 'app', 'audit', 'current_user', 'db', 'jsonify', 'login_required', 'request')))



# ============================================================

# ============================================================
# RB_COMMAND_CENTER_AUTO_PROTECTION_V1
# Protects Command Center monitoring only.
#
# LIVE -> NORMAL when sustained resource pressure is detected.
# Never changes RIS/PACS/DB/clinical workflow behavior.
# ============================================================

# RB_BIND:0499
from application.components.shared import install_0499 as _rb_install_0499
_rb_auto_protection_state_path = _rb_install_0499(_rb_bind(_rb_sys.modules[__name__], ('_command_center_control_path',)))


# RB_BIND:0500
from application.components.shared import install_0500 as _rb_install_0500
_rb_auto_protection_defaults = _rb_install_0500(_rb_bind(_rb_sys.modules[__name__], ()))


# RB_BIND:0501
from application.components.shared import install_0501 as _rb_install_0501
_rb_auto_protection_read_state = _rb_install_0501(_rb_bind(_rb_sys.modules[__name__], ('_rb_auto_protection_defaults', '_rb_auto_protection_state_path', 'app', 'json')))


# RB_BIND:0502
from application.components.shared import install_0502 as _rb_install_0502
_rb_auto_protection_write_state = _rb_install_0502(_rb_bind(_rb_sys.modules[__name__], ('_rb_auto_protection_state_path', 'app', 'json', 'os')))


# RB_BIND:0503
from application.components.shared import install_0503 as _rb_install_0503
_rb_auto_protection_evaluate = _rb_install_0503(_rb_bind(_rb_sys.modules[__name__], ('_command_center_write_control_state', '_rb_auto_protection_read_state', '_rb_auto_protection_write_state', 'command_center_control_state', 'datetime')))


# RB_COMMAND_CENTER_RESOURCE_MONITOR_V1
# Read-only resource and service health telemetry.
# ============================================================

# RB_BIND:0504
from application.components.shared import install_0504 as _rb_install_0504
_rb_read_text = _rb_install_0504(_rb_bind(_rb_sys.modules[__name__], ()))


# RB_BIND:0505
from application.components.shared import install_0505 as _rb_install_0505
_rb_int = _rb_install_0505(_rb_bind(_rb_sys.modules[__name__], ()))


# RB_BIND:0506
from application.components.shared import install_0506 as _rb_install_0506
_rb_cgroup_metrics = _rb_install_0506(_rb_bind(_rb_sys.modules[__name__], ('_rb_int', '_rb_read_text')))


# RB_BIND:0507
from application.components.shared import install_0507 as _rb_install_0507
_rb_host_metrics = _rb_install_0507(_rb_bind(_rb_sys.modules[__name__], ('_rb_int', '_rb_read_text')))


# RB_BIND:0508
from application.components.operations import install_0508 as _rb_install_0508
_rb_db_health = _rb_install_0508(_rb_bind(_rb_sys.modules[__name__], ('db', 'text')))


# RB_BIND:0509
from application.components.imaging import install_0509 as _rb_install_0509
_rb_pacs_health = _rb_install_0509(_rb_bind(_rb_sys.modules[__name__], ('PACS_ORTHANC_URL', '_orthanc_auth', 'requests')))


# RB_BIND:0510
from web.routes.shared import install_0510 as _rb_install_0510
api_command_center_resources = _rb_install_0510(_rb_bind(_rb_sys.modules[__name__], ('_rb_auto_protection_evaluate', '_rb_cgroup_metrics', '_rb_db_health', '_rb_host_metrics', '_rb_pacs_health', 'app', 'command_center_control_state', 'datetime', 'jsonify', 'require_perm')))


# RB_LIVE_COMMAND_CENTER_SSE_V1
# Read-only Server-Sent Events transport for Command Center.
# It reuses the normalized event API and never writes clinical
# workflow state.
# ============================================================

# RB_BIND:0511
from web.routes.shared import install_0511 as _rb_install_0511
command_center_events_stream = _rb_install_0511(_rb_bind(_rb_sys.modules[__name__], ('app', 'command_center_control_state', 'json', 'require_perm')))


# RB_BIND:0512
from web.routes.shared import install_0512 as _rb_install_0512
live_command_center = _rb_install_0512(_rb_bind(_rb_sys.modules[__name__], ('app', 'render_template', 'require_perm')))



# ============================================================
# RadBridge Live Command Center
# Normalized read-only event feed.
# RB_LIVE_COMMAND_CENTER_EVENTS_API_V1
# ============================================================

# RB_BIND:0513
from web.routes.shared import install_0513 as _rb_install_0513
command_center_events = _rb_install_0513(_rb_bind(_rb_sys.modules[__name__], ('AssignmentEscalationEvent', 'AuditLog', 'CaseWorkflowEvent', 'EventOutbox', 'Order', 'PacsResolutionEvent', 'PacsStudyInbox', 'PacsStudyMatch', 'app', 'datetime', 'db', 'json', 'jsonify', 'request', 'require_perm')))


# RB_BIND:0514
from web.routes.shared import install_0514 as _rb_install_0514
command_center_snapshot = _rb_install_0514(_rb_bind(_rb_sys.modules[__name__], ('AssignmentOffer', 'CaseAssignment', 'CaseReadiness', 'EventOutbox', 'Order', 'PacsStudyInbox', 'PacsStudyMatch', 'Report', 'app', 'datetime', 'db', 'jsonify', 'require_perm')))


# RB_BIND:0515
from web.routes.finance import install_0515 as _rb_install_0515
financial_center = _rb_install_0515(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Contract', 'Device', 'FINANCIAL_REPORTS', 'Procedure', '_finance_filters', '_financial_dataset', 'app', 'db', 'render_template', 'request', 'require_perm')))

# RB_BIND:0516
from web.routes.finance import install_0516 as _rb_install_0516
financial_center_export = _rb_install_0516(_rb_bind(_rb_sys.modules[__name__], ('FINANCIAL_REPORTS', '_export_csv', '_export_pdf', '_export_xlsx', '_finance_filters', '_financial_dataset', 'abort', 'app', 'request', 'require_perm')))

# RB_BIND:0517
from web.routes.finance import install_0517 as _rb_install_0517
financial_entries = _rb_install_0517(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Consumable', 'Device', 'FINANCE_ENTRY_TYPES', 'FinancialEntry', 'User', '__name__', '_reporting_support_sys', 'app', 'audit', 'current_user', 'date', 'db', 'flash', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0518
from web.routes.finance import install_0518 as _rb_install_0518
financial_entry_void = _rb_install_0518(_rb_bind(_rb_sys.modules[__name__], ('FinancialEntry', 'app', 'audit', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0519
from web.routes.shared import install_0519 as _rb_install_0519
receivables = _rb_install_0519(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Contract', 'Order', 'app', 'order_balance', 'render_template', 'request', 'require_perm')))

# RB_BIND:0520
from web.routes.shared import install_0520 as _rb_install_0520
receivables_csv = _rb_install_0520(_rb_bind(_rb_sys.modules[__name__], ('Order', 'app', 'csv', 'io', 'order_balance', 'require_perm', 'send_file')))

# RB_BIND:0521
from web.routes.finance import install_0521 as _rb_install_0521
receive_payment = _rb_install_0521(_rb_bind(_rb_sys.modules[__name__], ('Order', 'Payment', 'Receipt', 'abort', 'app', 'apply_or_request_discount', 'audit', 'current_user', 'db', 'flash', 'next_receipt_no', 'order_balance', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0522
from web.routes.finance import install_0522 as _rb_install_0522
receipt_view = _rb_install_0522(_rb_bind(_rb_sys.modules[__name__], ('Receipt', 'abort', 'app', 'has_perm', 'login_required', 'render_template')))

# RB_BIND:0523
from web.routes.registration import install_0523 as _rb_install_0523
mark_delivery = _rb_install_0523(_rb_bind(_rb_sys.modules[__name__], ('DeliveryEvent', 'Order', 'Report', 'abort', 'app', 'audit', 'current_user', 'datetime', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_BIND:0524
from web.routes.finance import install_0524 as _rb_install_0524
cashier_workstation = _rb_install_0524(_rb_bind(_rb_sys.modules[__name__], ('Order', 'Patient', 'Payment', 'PaymentRefund', 'abort', 'app', 'current_user', 'discount_policy_for', 'effective_patient_due', 'or_', 'render_template', 'request', 'require_perm', 'total_order_discount')))

# RB_BIND:0525
from web.routes.finance import install_0525 as _rb_install_0525
refund_payment = _rb_install_0525(_rb_bind(_rb_sys.modules[__name__], ('Payment', 'PaymentRefund', 'app', 'audit', 'current_user', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0526
from web.routes.reporting import install_0526 as _rb_install_0526
accounting_reports = _rb_install_0526(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Order', 'Receipt', 'app', 'datetime', 'order_balance', 'render_template', 'request', 'require_perm', 'timedelta')))

# RB_BIND:0527
from web.routes.reporting import install_0527 as _rb_install_0527
accounting_reports_csv = _rb_install_0527(_rb_bind(_rb_sys.modules[__name__], ('Order', 'app', 'csv', 'io', 'order_balance', 'require_perm', 'send_file')))


# RB_BIND:0528
from web.routes.clinical import install_0528 as _rb_install_0528
admin_diagnostics = _rb_install_0528(_rb_bind(_rb_sys.modules[__name__], ('APP_VERSION', 'ATTACHMENT_DIR', 'REPORT_DIR', '__file__', 'app', 'db', 'orthanc_ok', 'os', 'render_template', 'require_perm', 'text')))


# RB_BIND:0529
from application.components.reporting import install_0529 as _rb_install_0529
create_reporting_demo_cases = _rb_install_0529(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Order', 'Patient', 'Procedure', 'Report', 'ReportVersion', 'current_user', 'datetime', 'db', 'route_device')))

# RB_BIND:0530
from web.routes.reporting import install_0530 as _rb_install_0530
admin_reporting_demo = _rb_install_0530(_rb_bind(_rb_sys.modules[__name__], ('app', 'create_reporting_demo_cases', 'flash', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))



# RB_BIND:0531
from web.routes.assignment import install_0531 as _rb_install_0531
admin_assignment = _rb_install_0531(_rb_bind(_rb_sys.modules[__name__], ('DoctorAssignmentProfile', 'app', 'assignable_doctors', 'db', 'flash', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))


# RB_BIND:0532
from application.components.shared import install_0532 as _rb_install_0532
_strategy_action = _rb_install_0532(_rb_bind(_rb_sys.modules[__name__], ()))


# RB_BIND:0533
from application.components.shared import install_0533 as _rb_install_0533
_strategy_snapshot = _rb_install_0533(_rb_bind(_rb_sys.modules[__name__], ()))

# RB_BIND:0534
from application.components.shared import install_0534 as _rb_install_0534
_strategy_change_summary = _rb_install_0534(_rb_bind(_rb_sys.modules[__name__], ()))





# RB_BIND:0535
from web.routes.assignment import install_0535 as _rb_install_0535
admin_assignment_strategies = _rb_install_0535(_rb_bind(_rb_sys.modules[__name__], ('AssignmentStrategy', 'AssignmentStrategyVersion', 'Branch', '_strategy_action', '_strategy_change_summary', '_strategy_snapshot', 'abort', 'app', 'audit', 'current_user', 'db', 'flash', 'json', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))


# RB_BIND:0536
from application.components.assignment import install_0536 as _rb_install_0536
assignment_candidate_diagnostics = _rb_install_0536(_rb_bind(_rb_sys.modules[__name__], ('CaseAssignment', 'DoctorAssignmentProfile', 'Role', 'User', 'db', 'doctor_availability', 'doctor_is_online')))


# RB_BIND:0537
from web.routes.assignment import install_0537 as _rb_install_0537
admin_assignment_strategy_simulate = _rb_install_0537(_rb_bind(_rb_sys.modules[__name__], ('Order', 'app', 'assignment_candidate_diagnostics', 'assignment_candidate_users', 'assignment_offer_policy', 'assignment_readiness_decision', 'assignment_strategy_for_order', 'jsonify', 'request', 'require_perm')))



# RB_BIND:0538
from web.routes.assignment import install_0538 as _rb_install_0538
admin_assignment_strategy_toggle = _rb_install_0538(_rb_bind(_rb_sys.modules[__name__], ('AssignmentStrategy', 'AssignmentStrategyVersion', '_strategy_snapshot', 'app', 'audit', 'current_user', 'db', 'flash', 'json', 'redirect', 'require_perm', 'url_for')))



# RB_BIND:0539
from web.routes.assignment import install_0539 as _rb_install_0539
admin_assignment_strategy_simulate_ui = _rb_install_0539(_rb_bind(_rb_sys.modules[__name__], ('Order', 'app', 'assignment_candidate_diagnostics', 'assignment_offer_policy', 'assignment_readiness_decision', 'assignment_strategy_for_order', 'flash', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))


# RB_BIND:0540
from web.routes.assignment import install_0540 as _rb_install_0540
admin_assignment_strategy_clone = _rb_install_0540(_rb_bind(_rb_sys.modules[__name__], ('AssignmentStrategy', 'AssignmentStrategyVersion', '_strategy_snapshot', 'app', 'audit', 'current_user', 'db', 'flash', 'json', 'redirect', 'require_perm', 'url_for')))



# RB_BIND:0541
from web.routes.assignment import install_0541 as _rb_install_0541
admin_assignment_strategy_history = _rb_install_0541(_rb_bind(_rb_sys.modules[__name__], ('AssignmentStrategy', 'AssignmentStrategyVersion', 'app', 'render_template', 'require_perm')))


# RB_BIND:0542
from web.routes.assignment import install_0542 as _rb_install_0542
admin_assignment_strategy_rollback = _rb_install_0542(_rb_bind(_rb_sys.modules[__name__], ('AssignmentStrategy', 'AssignmentStrategyVersion', '_strategy_snapshot', 'app', 'audit', 'current_user', 'db', 'flash', 'json', 'redirect', 'require_perm', 'url_for')))


# RB_BIND:0543
from web.routes.assignment import install_0543 as _rb_install_0543
assign_case = _rb_install_0543(_rb_bind(_rb_sys.modules[__name__], ('Order', '__name__', 'abort', 'app', 'assignment_readiness_decision', 'current_user', 'db', 'flash', 'manual_assign_order', 'offer_next_doctor', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0544
from web.routes.imaging import install_0544 as _rb_install_0544
take_case = _rb_install_0544(_rb_bind(_rb_sys.modules[__name__], ('Order', 'app', 'current_user', 'db', 'doctor_can_take_open_case', 'flash', 'manual_assign_order', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0545
from web.routes.reporting import install_0545 as _rb_install_0545
unassigned_cases = _rb_install_0545(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'CaseReadiness', 'Order', 'Patient', 'Procedure', 'TATPolicy', 'abort', 'app', 'assignable_doctors', 'datetime', 'db', 'or_', 'render_template', 'request', 'require_perm', 'unassigned_case_snapshot')))


# RB_BIND:0546
from application.components.assignment import install_0546 as _rb_install_0546
unassigned_case_snapshot = _rb_install_0546(_rb_bind(_rb_sys.modules[__name__], ('assignment_readiness_decision', 'datetime')))


# ---------------- Admin ----------------

# RB_BIND:0547
from web.routes.operations import install_0547 as _rb_install_0547
admin_home = _rb_install_0547(_rb_bind(_rb_sys.modules[__name__], ('abort', 'app', 'has_perm', 'login_required', 'render_template', 'url_for')))

# RB_BIND:0548
from web.routes.operations import install_0548 as _rb_install_0548
admin_users = _rb_install_0548(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Role', 'User', 'UserProfile', '__name__', '_reporting_support_sys', 'abort', 'app', 'audit', 'db', 'flash', 'generate_password_hash', 'has_perm', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0549
from web.routes.operations import install_0549 as _rb_install_0549
admin_user_toggle = _rb_install_0549(_rb_bind(_rb_sys.modules[__name__], ('User', 'app', 'audit', 'current_user', 'db', 'flash', 'redirect', 'require_perm', 'url_for')))

# RB_BIND:0550
from web.routes.operations import install_0550 as _rb_install_0550
admin_user_edit = _rb_install_0550(_rb_bind(_rb_sys.modules[__name__], ('Role', 'User', '__name__', '_reporting_support_sys', 'app', 'audit', 'current_user', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0551
from web.routes.identity import install_0551 as _rb_install_0551
admin_user_reset_password = _rb_install_0551(_rb_bind(_rb_sys.modules[__name__], ('User', 'app', 'audit', 'datetime', 'db', 'flash', 'generate_password_hash', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0552
from web.routes.registration import install_0552 as _rb_install_0552
admin_device_schedules = _rb_install_0552(_rb_bind(_rb_sys.modules[__name__], ('Device', 'DeviceBlock', 'DeviceSchedule', 'app', 'audit', 'db', 'flash', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0553
from web.routes.registration import install_0553 as _rb_install_0553
admin_device_blocks = _rb_install_0553(_rb_bind(_rb_sys.modules[__name__], ('DeviceBlock', 'app', 'audit', 'datetime', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0554
from web.routes.imaging import install_0554 as _rb_install_0554
dicom_connections_page = _rb_install_0554(_rb_bind(_rb_sys.modules[__name__], ('AUTH', 'Branch', 'Device', 'DicomEndpoint', 'ORTHANC', 'abort', 'app', 'current_user', 'has_perm', 'render_template', 'request', 'requests', 'require_perm')))

# RB_BIND:0555
from web.routes.imaging import install_0555 as _rb_install_0555
dicom_connection_edit = _rb_install_0555(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'DicomEndpoint', 'abort', 'app', 'audit', 'current_user', 'db', 'flash', 'has_perm', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0556
from web.routes.imaging import install_0556 as _rb_install_0556
dicom_connection_echo = _rb_install_0556(_rb_bind(_rb_sys.modules[__name__], ('DicomEndpoint', 'abort', 'app', 'audit', 'current_user', 'datetime', 'db', 'flash', 'has_perm', 'redirect', 'require_perm', 'url_for')))

# RB_BIND:0557
from web.routes.operations import install_0557 as _rb_install_0557
admin_branch_edit = _rb_install_0557(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'app', 'audit', 'db', 'flash', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))


# RB_BIND:0558
from web.routes.operations import install_0558 as _rb_install_0558
admin_devices = _rb_install_0558(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Device', 'app', 'audit', 'db', 'flash', 'render_template', 'request', 'require_perm')))

# RB_BIND:0559
from web.routes.registration import install_0559 as _rb_install_0559
admin_device_edit = _rb_install_0559(_rb_bind(_rb_sys.modules[__name__], ('Device', 'app', 'audit', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0560
from web.routes.registration import install_0560 as _rb_install_0560
admin_procedures = _rb_install_0560(_rb_bind(_rb_sys.modules[__name__], ('Procedure', 'abort', 'app', 'audit', 'db', 'flash', 'has_perm', 'render_template', 'request', 'require_perm')))

# RB_BIND:0561
from web.routes.registration import install_0561 as _rb_install_0561
admin_contracts = _rb_install_0561(_rb_bind(_rb_sys.modules[__name__], ('Contract', 'app', 'audit', 'db', 'flash', 'render_template', 'request', 'require_perm')))

# RB_BIND:0562
from web.routes.registration import install_0562 as _rb_install_0562
admin_contract_prices = _rb_install_0562(_rb_bind(_rb_sys.modules[__name__], ('Contract', 'ContractPrice', 'Procedure', 'app', 'audit', 'db', 'flash', 'render_template', 'request', 'require_perm')))

# RB_BIND:0563
from web.routes.inventory import install_0563 as _rb_install_0563
admin_inventory = _rb_install_0563(_rb_bind(_rb_sys.modules[__name__], ('Consumable', 'app', 'audit', 'db', 'flash', 'render_template', 'request', 'require_perm')))

# RB_BIND:0564
from web.routes.inventory import install_0564 as _rb_install_0564
admin_inventory_edit = _rb_install_0564(_rb_bind(_rb_sys.modules[__name__], ('Consumable', 'app', 'audit', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0565
from web.routes.inventory import install_0565 as _rb_install_0565
admin_inventory_adjust = _rb_install_0565(_rb_bind(_rb_sys.modules[__name__], ('Consumable', '_record_stock_movement', 'app', 'audit', 'current_user', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_BIND:0566
from web.routes.inventory import install_0566 as _rb_install_0566
admin_inventory_seed_defaults = _rb_install_0566(_rb_bind(_rb_sys.modules[__name__], ('app', 'audit', 'db', 'flash', 'redirect', 'require_perm', 'seed_default_consumables', 'url_for')))

# RB_BIND:0567
from web.routes.inventory import install_0567 as _rb_install_0567
admin_inventory_export = _rb_install_0567(_rb_bind(_rb_sys.modules[__name__], ('Consumable', 'Font', 'INVENTORY_IMPORT_COLUMNS', 'PatternFill', 'Workbook', 'abort', 'app', 'csv', 'get_column_letter', 'inventory_row', 'io', 'require_perm', 'send_file')))

# RB_BIND:0568
from web.routes.inventory import install_0568 as _rb_install_0568
admin_inventory_import_template = _rb_install_0568(_rb_bind(_rb_sys.modules[__name__], ('DEFAULT_CONSUMABLE_CATALOG', 'Font', 'INVENTORY_IMPORT_COLUMNS', 'PatternFill', 'Workbook', 'app', 'get_column_letter', 'io', 'require_perm', 'send_file')))

# RB_BIND:0569
from web.routes.inventory import install_0569 as _rb_install_0569
admin_inventory_import = _rb_install_0569(_rb_bind(_rb_sys.modules[__name__], ('Consumable', 'app', 'audit', 'csv', 'db', 'flash', 'io', 'load_workbook', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0570
from web.routes.inventory import install_0570 as _rb_install_0570
admin_inventory_movements = _rb_install_0570(_rb_bind(_rb_sys.modules[__name__], ('StockMovement', 'app', 'render_template', 'require_perm')))

# RB_BIND:0571
from web.routes.inventory import install_0571 as _rb_install_0571
admin_inventory_batches = _rb_install_0571(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Consumable', 'InventoryBatch', '_record_stock_movement', 'app', 'audit', 'current_user', 'date', 'db', 'flash', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0572
from web.routes.inventory import install_0572 as _rb_install_0572
admin_erp = _rb_install_0572(_rb_bind(_rb_sys.modules[__name__], ('Consumable', 'ERPItemMapping', 'ERPSyncQueue', 'ERP_BASE_URL', 'ERP_PROVIDER', 'ERP_SYNC_ENABLED', 'app', 'audit', 'db', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0573
from web.routes.inventory import install_0573 as _rb_install_0573
admin_erp_process = _rb_install_0573(_rb_bind(_rb_sys.modules[__name__], ('app', 'flash', 'process_erp_queue', 'redirect', 'require_perm', 'url_for')))

# RB_BIND:0574
from web.routes.inventory import install_0574 as _rb_install_0574
admin_erp_sync_items = _rb_install_0574(_rb_bind(_rb_sys.modules[__name__], ('Consumable', 'ERPItemMapping', 'app', 'db', 'flash', 'queue_erp_event', 'redirect', 'require_perm', 'url_for')))

# RB_BIND:0575
from web.routes.operations import install_0575 as _rb_install_0575
admin_audit = _rb_install_0575(_rb_bind(_rb_sys.modules[__name__], ('AuditLog', 'app', 'datetime', 'render_template', 'request', 'require_perm', 'timedelta')))

# RB_BIND:0576
from web.components.shared import install_0576 as _rb_install_0576
cancellation_reason = _rb_install_0576(_rb_bind(_rb_sys.modules[__name__], ('abort', 'request')))

# RB_BIND:0577
from application.components.shared import install_0577 as _rb_install_0577
_cancellation_literal_search_pattern = _rb_install_0577(_rb_bind(_rb_sys.modules[__name__], ()))


# RB_BIND:0578
from application.components.shared import install_0578 as _rb_install_0578
cancellation_search_rows = _rb_install_0578(_rb_bind(_rb_sys.modules[__name__], ('Appointment', 'Order', 'Patient', '_cancellation_literal_search_pattern', 'db')))




# RB_GRANULAR_DELETE_IMAGES_ONLY_V1
# RB_BIND:0579
from application.components.shared import install_0579 as _rb_install_0579
granular_delete_images_only_preflight = _rb_install_0579(_rb_bind(_rb_sys.modules[__name__], ('Order', 'PACS_ORTHANC_URL', 'PacsTransferJob', 'Patient', '_orthanc_auth', 'current_user', 'db', 'requests')))


# RB_BIND:0580
from application.components.shared import install_0580 as _rb_install_0580
granular_delete_images_only_execute = _rb_install_0580(_rb_bind(_rb_sys.modules[__name__], ('PACS_ORTHANC_URL', '_orthanc_auth', 'audit', 'datetime', 'db', 'granular_delete_images_only_preflight', 'requests')))


# RB_BIND:0581
from web.routes.registration import install_0581 as _rb_install_0581
granular_delete_order_images_only = _rb_install_0581(_rb_bind(_rb_sys.modules[__name__], ('Order', 'app', 'granular_delete_images_only_execute', 'request', 'require_perm')))


# RB_DELETE_REPORTS_ONLY_PREVIEW_V1
# RB_BIND:0582
from application.components.reporting import install_0582 as _rb_install_0582
granular_delete_reports_only_preview_snapshot = _rb_install_0582(_rb_bind(_rb_sys.modules[__name__], ('CoordinatorTask', 'REPORT_DIR', 'ReportVersion', 'db')))


# RB_DELETE_REPORTS_ONLY_EXECUTION_V1
# RB_BIND:0583
from application.components.reporting import install_0583 as _rb_install_0583
_delete_reports_only_is_super_admin = _rb_install_0583(_rb_bind(_rb_sys.modules[__name__], ()))


# RB_BIND:0584
from application.components.reporting import install_0584 as _rb_install_0584
_delete_reports_only_quarantine_target = _rb_install_0584(_rb_bind(_rb_sys.modules[__name__], ('REPORT_DIR',)))


# RB_BIND:0585
from application.components.reporting import install_0585 as _rb_install_0585
_delete_reports_only_blocked_response = _rb_install_0585(_rb_bind(_rb_sys.modules[__name__], ('audit', 'db')))


# RB_BIND:0586
from web.components.reporting import install_0586 as _rb_install_0586
granular_delete_reports_only_execute = _rb_install_0586(_rb_bind(_rb_sys.modules[__name__], ('Order', 'Patient', 'REPORT_DIR', 'Report', 'ReportVersion', '_delete_reports_only_blocked_response', '_delete_reports_only_is_super_admin', '_delete_reports_only_quarantine_target', 'abort', 'audit', 'current_user', 'datetime', 'db', 'granular_delete_reports_only_preview_snapshot', 'os', 'secrets')))


# RB_RIS_CASE_DATA_ONLY_PREFLIGHT_V1
_RIS_CASE_ONLY_DEPENDENCY_POLICY = {
    ('appointment','order_id'): ('PRESERVE','NULLIFY_ORDER_LINK','APPOINTMENT_PRESERVED'),
    ('attachment','order_id'): ('PRESERVE','NULLIFY_ORDER_LINK','ATTACHMENT_FILE_PRESERVED'),
    ('patient_document','order_id'): ('PRESERVE','NULLIFY_ORDER_LINK','PATIENT_DOCUMENT_PRESERVED'),
    ('chat_thread','order_id'): ('PRESERVE','NULLIFY_ORDER_LINK','CHAT_THREAD_PRESERVED'),
    ('external_prior_study','linked_order_id'): ('PRESERVE','NULLIFY_ORDER_LINK','EXTERNAL_PRIOR_PRESERVED'),
    ('patient_communication','order_id'): ('PRESERVE','NULLIFY_ORDER_LINK','COMMUNICATION_PRESERVED'),
    ('ris_notification','order_id'): ('PRESERVE','NULLIFY_ORDER_LINK','RIS_NOTIFICATION_PRESERVED'),
    ('pacs_study_inbox','order_id'): ('PRESERVE','NULLIFY_ORDER_LINK','PACS_INBOX_PRESERVED'),
    ('external_intake_case','order_id'): ('PRESERVE','NULLIFY_ORDER_LINK','EXTERNAL_INTAKE_PRESERVED'),

    ('clinical_note','order_id'): ('CASE_DATA','DELETE_CASE_ROW','CASE_OWNED'),
    ('journey_event','order_id'): ('CASE_DATA','DELETE_CASE_ROW','CASE_OWNED'),
    ('order_performing_doctor','order_id'): ('CASE_DATA','DELETE_CASE_ROW','CASE_OWNED'),
    ('ai_draft_report','order_id'): ('CASE_DATA','DELETE_CASE_ROW','CASE_OWNED'),
    ('prior_study_selection','order_id'): ('CASE_DATA','DELETE_CASE_ROW','CASE_OWNED'),
    ('imaging_ai_analysis','order_id'): ('CASE_DATA','DELETE_CASE_ROW','CASE_OWNED'),
    ('case_readiness','order_id'): ('CASE_DATA','DELETE_CASE_ROW','CASE_OWNED'),
    ('case_readiness_override','order_id'): ('CASE_DATA','DELETE_CASE_ROW','CASE_OWNED'),
    ('case_exception_resolution_event','order_id'): ('CASE_DATA','DELETE_CASE_ROW','CASE_OWNED'),
    ('case_workflow_event','order_id'): ('CASE_DATA','DELETE_CASE_ROW','CASE_OWNED'),
    ('clinical_assessment','order_id'): ('CASE_DATA','DELETE_CASE_ROW','CASE_OWNED'),
    ('vital_sign','order_id'): ('CASE_DATA','DELETE_CASE_ROW','CASE_OWNED'),
    ('clinical_diagnosis','order_id'): ('CASE_DATA','DELETE_CASE_ROW','CASE_OWNED'),
    ('clinical_medication','order_id'): ('CASE_DATA','DELETE_CASE_ROW','CASE_OWNED'),
    ('clinical_lab_item','order_id'): ('CASE_DATA','DELETE_CASE_ROW','CASE_OWNED'),
    ('pacs_study_match','order_id'): ('RIS_PACS_METADATA','DELETE_CASE_ROW','PACS_IMAGES_PRESERVED'),
    ('pacs_exam_link','order_id'): ('RIS_PACS_METADATA','DELETE_CASE_ROW','PACS_IMAGES_PRESERVED'),
    ('pacs_resolution_event','order_id'): ('RIS_PACS_METADATA','DELETE_CASE_ROW','PACS_IMAGES_PRESERVED'),
    ('portal_access_token','order_id'): ('CASE_DATA','DELETE_CASE_ROW','TOKEN_REVOKE_DELETE'),
    ('notification_event','order_id'): ('CASE_DATA','DELETE_CASE_ROW','CASE_OWNED'),

    ('report','order_id'): ('REPORT_DOMAIN','BLOCK','DELETE_REPORTS_FIRST'),
    ('coordinator_task','order_id'): ('WORKFLOW_GOVERNANCE','BLOCK','COORDINATOR_TASK_PRESENT'),
    ('correction_request','order_id'): ('GOVERNANCE_HISTORY','BLOCK','CORRECTION_HISTORY_PRESENT'),
    ('correction_execution','order_id'): ('GOVERNANCE_HISTORY','BLOCK','CORRECTION_HISTORY_PRESENT'),
    ('demographic_override_audit','order_id'): ('AUDIT_HISTORY','BLOCK','AUDIT_HISTORY_PRESENT'),
    ('payment','order_id'): ('FINANCIAL_HISTORY','BLOCK','FINANCIAL_HISTORY_PRESENT'),
    ('payment_refund','order_id'): ('FINANCIAL_HISTORY','BLOCK','FINANCIAL_HISTORY_PRESENT'),
    ('receipt','order_id'): ('FINANCIAL_HISTORY','BLOCK','FINANCIAL_HISTORY_PRESENT'),
    ('insurance_approval_request','order_id'): ('FINANCIAL_HISTORY','BLOCK','INSURANCE_HISTORY_PRESENT'),
    ('order_discount','order_id'): ('FINANCIAL_HISTORY','BLOCK','FINANCIAL_HISTORY_PRESENT'),
    ('discount_approval_request','order_id'): ('FINANCIAL_HISTORY','BLOCK','FINANCIAL_HISTORY_PRESENT'),
    ('order_consumable_usage','order_id'): ('INVENTORY_HISTORY','BLOCK','INVENTORY_HISTORY_PRESENT'),
    ('order_consumable','order_id'): ('INVENTORY_HISTORY','BLOCK','INVENTORY_HISTORY_PRESENT'),
    ('nursing_consumable_entry','order_id'): ('INVENTORY_HISTORY','BLOCK','INVENTORY_HISTORY_PRESENT'),
    ('delivery_event','order_id'): ('DELIVERY_HISTORY','BLOCK','DELIVERY_HISTORY_PRESENT'),
    ('case_assignment','order_id'): ('ACTIVE_WORKFLOW','BLOCK','ASSIGNMENT_HISTORY_REVIEW'),
    ('assignment_offer','order_id'): ('ACTIVE_WORKFLOW','BLOCK','ASSIGNMENT_OFFER_REVIEW'),
    ('assignment_escalation_event','order_id'): ('ACTIVE_WORKFLOW','BLOCK','ASSIGNMENT_ESCALATION_REVIEW'),
    ('pacs_transfer_job','order_id'): ('ACTIVE_WORKFLOW','BLOCK','PACS_TRANSFER_REVIEW'),
}


# RB_BIND:0587
from application.components.shared import install_0587 as _rb_install_0587
granular_delete_ris_case_only_preview_snapshot = _rb_install_0587(_rb_bind(_rb_sys.modules[__name__], ('_RIS_CASE_ONLY_DEPENDENCY_POLICY', 'db')))


# RB_RIS_CASE_DATA_ONLY_EXECUTION_V1
# RB_BIND:0588
from application.components.shared import install_0588 as _rb_install_0588
_ris_case_only_dependency_rows = _rb_install_0588(_rb_bind(_rb_sys.modules[__name__], ('Order', 'db')))


# RB_BIND:0589
from application.components.registration import install_0589 as _rb_install_0589
_ris_case_only_second_order_checks = _rb_install_0589(_rb_bind(_rb_sys.modules[__name__], ('db',)))


# RB_BIND:0590
from application.components.shared import install_0590 as _rb_install_0590
granular_delete_ris_case_only_execution_preview_snapshot = _rb_install_0590(_rb_bind(_rb_sys.modules[__name__], ('Report', '_ris_case_only_dependency_rows', '_ris_case_only_second_order_checks', 'db', 'granular_delete_ris_case_only_preview_snapshot')))


# RB_BIND:0591
from application.components.shared import install_0591 as _rb_install_0591
_ris_case_only_blocked_response = _rb_install_0591(_rb_bind(_rb_sys.modules[__name__], ('audit', 'db')))


# RB_BIND:0592
from web.components.shared import install_0592 as _rb_install_0592
granular_delete_ris_case_only_execute = _rb_install_0592(_rb_bind(_rb_sys.modules[__name__], ('Order', 'Patient', 'Report', '_RIS_CASE_ONLY_DEPENDENCY_POLICY', '_delete_reports_only_is_super_admin', '_ris_case_only_blocked_response', '_ris_case_only_dependency_rows', 'abort', 'audit', 'current_user', 'datetime', 'db', 'granular_delete_ris_case_only_execution_preview_snapshot')))


# RB_GRANULAR_DELETE_READONLY_PREVIEW_V1
# RB_BIND:0593
from application.components.registration import install_0593 as _rb_install_0593
granular_delete_order_preview_snapshot = _rb_install_0593(_rb_bind(_rb_sys.modules[__name__], ('Order', 'PACS_ORTHANC_URL', 'Patient', 'Report', '_orthanc_auth', 'db', 'granular_delete_full_case_orchestration_preview_snapshot', 'granular_delete_reports_only_preview_snapshot', 'granular_delete_ris_case_only_execution_preview_snapshot', 'requests')))


# RB_FULL_CASE_ORCHESTRATION_PREVIEW_V1
# RB_BIND:0594
from application.components.shared import install_0594 as _rb_install_0594
granular_delete_full_case_orchestration_preview_snapshot = _rb_install_0594(_rb_bind(_rb_sys.modules[__name__], ('_delete_reports_only_is_super_admin', '_full_case_orchestration_execution_enabled', 'current_user', 'granular_delete_images_only_preflight', 'granular_delete_ris_case_only_execution_preview_snapshot')))


# RB_BIND:0595
from application.components.registration import install_0595 as _rb_install_0595
granular_delete_patient_preview_snapshot = _rb_install_0595(_rb_bind(_rb_sys.modules[__name__], ('Appointment', 'Order', 'Report')))


# RB_BIND:0596
from web.routes.operations import install_0596 as _rb_install_0596
granular_delete_preview_home = _rb_install_0596(_rb_bind(_rb_sys.modules[__name__], ('app', 'render_template', 'require_perm')))


# RB_BIND:0597
from web.routes.registration import install_0597 as _rb_install_0597
granular_delete_order_preview = _rb_install_0597(_rb_bind(_rb_sys.modules[__name__], ('Order', 'app', 'granular_delete_order_preview_snapshot', 'jsonify', 'render_template', 'request', 'require_perm')))


# RB_BIND:0598
from web.routes.reporting import install_0598 as _rb_install_0598
granular_delete_order_reports_only = _rb_install_0598(_rb_bind(_rb_sys.modules[__name__], ('Order', '_delete_reports_only_is_super_admin', 'abort', 'app', 'current_user', 'granular_delete_reports_only_execute', 'request', 'require_perm')))


# RB_BIND:0599
from web.routes.registration import install_0599 as _rb_install_0599
granular_delete_order_ris_case_data_only = _rb_install_0599(_rb_bind(_rb_sys.modules[__name__], ('Order', '_delete_reports_only_is_super_admin', 'abort', 'app', 'current_user', 'granular_delete_ris_case_only_execute', 'request', 'require_perm')))


# RB_FULL_CASE_ORCHESTRATION_EXECUTOR_V1
# RB_BIND:0600
from application.components.shared import install_0600 as _rb_install_0600
_full_case_orchestration_execution_enabled = _rb_install_0600(_rb_bind(_rb_sys.modules[__name__], ('os',)))


# RB_BIND:0601
from application.components.shared import install_0601 as _rb_install_0601
_full_case_orchestration_failure = _rb_install_0601(_rb_bind(_rb_sys.modules[__name__], ('audit', 'db')))


# RB_BIND:0602
from application.components.shared import install_0602 as _rb_install_0602
_full_case_orchestration_final_postcheck = _rb_install_0602(_rb_bind(_rb_sys.modules[__name__], ('Order', 'Patient', 'Report', '_ris_case_only_dependency_rows', 'db')))


# RB_BIND:0603
from web.components.shared import install_0603 as _rb_install_0603
granular_delete_full_case_orchestration_execute = _rb_install_0603(_rb_bind(_rb_sys.modules[__name__], ('Order', 'Patient', 'Report', '_delete_reports_only_is_super_admin', '_full_case_orchestration_execution_enabled', '_full_case_orchestration_failure', '_full_case_orchestration_final_postcheck', 'abort', 'audit', 'current_user', 'db', 'granular_delete_images_only_execute', 'granular_delete_order_preview_snapshot', 'granular_delete_reports_only_execute', 'granular_delete_ris_case_only_execute')))


# RB_BIND:0604
from web.routes.registration import install_0604 as _rb_install_0604
granular_delete_order_full_case = _rb_install_0604(_rb_bind(_rb_sys.modules[__name__], ('Order', '_delete_reports_only_is_super_admin', 'abort', 'app', 'current_user', 'granular_delete_full_case_orchestration_execute', 'request', 'require_perm')))


# RB_BIND:0605
from web.routes.registration import install_0605 as _rb_install_0605
granular_delete_patient_preview = _rb_install_0605(_rb_bind(_rb_sys.modules[__name__], ('Patient', 'app', 'granular_delete_patient_preview_snapshot', 'jsonify', 'render_template', 'request', 'require_perm')))



# RB_CORRECTION_CANCELLATION_CENTER_V1
CORRECTION_REQUEST_TYPES=(
    ('WRONG_IMAGES','Wrong images / study sent'),
    ('WRONG_DEMOGRAPHICS','Wrong patient demographics'),
    ('WRONG_ORDER','Wrong order / examination'),
    ('WRONG_ACCESSION','Wrong accession'),
    ('DUPLICATE_CASE','Duplicate case'),
    ('CANCEL_ORDER','Cancel order request'),
    ('WRONG_REPORT','Wrong report / report correction'),
    ('OTHER','Other correction'),
)
CORRECTION_REQUEST_TYPE_CODES={code for code,_ in CORRECTION_REQUEST_TYPES}
CORRECTION_OPEN_STATUSES=('NEW','UNDER_REVIEW','NEED_INFO','APPROVED')


# RB_CORRECTION_EXECUTION_STATE_MACHINE_PHASE2A_V1

CORRECTION_EXECUTION_STATUSES=(
    'READY_FOR_EXECUTION',
    'CLAIMED_FOR_EXECUTION',
    'ON_HOLD',
    'PREFLIGHT_PASSED',
    'PREFLIGHT_BLOCKED',
    'EXECUTING',
    'VERIFYING',
    'COMPLETED',
    'EXECUTION_FAILED',
    'VERIFICATION_FAILED',
    'EXPIRED',
    'CANCELLED_BEFORE_EXECUTION',
)

CORRECTION_EXECUTION_TERMINAL_STATUSES={
    'COMPLETED',
    'EXECUTION_FAILED',
    'VERIFICATION_FAILED',
    'EXPIRED',
    'CANCELLED_BEFORE_EXECUTION',
}

CORRECTION_EXECUTION_TRANSITIONS={
    'READY_FOR_EXECUTION':{
        'CLAIMED_FOR_EXECUTION',
        'ON_HOLD',
        'CANCELLED_BEFORE_EXECUTION',
        'EXPIRED',
    },
    'CLAIMED_FOR_EXECUTION':{
        'READY_FOR_EXECUTION',
        'ON_HOLD',
        'PREFLIGHT_PASSED',
        'PREFLIGHT_BLOCKED',
        'CANCELLED_BEFORE_EXECUTION',
        'EXPIRED',
    },
    'ON_HOLD':{
        'READY_FOR_EXECUTION',
        'CLAIMED_FOR_EXECUTION',
        'CANCELLED_BEFORE_EXECUTION',
        'EXPIRED',
    },
    'PREFLIGHT_PASSED':{
        'EXECUTING',
        'ON_HOLD',
        'CANCELLED_BEFORE_EXECUTION',
        'EXPIRED',
    },
    'PREFLIGHT_BLOCKED':{
        'READY_FOR_EXECUTION',
        'ON_HOLD',
        'CANCELLED_BEFORE_EXECUTION',
        'EXPIRED',
    },
    'EXECUTING':{
        'VERIFYING',
        'EXECUTION_FAILED',
    },
    'VERIFYING':{
        'COMPLETED',
        'VERIFICATION_FAILED',
    },
}


# RB_BIND:0606
from application.components.identity import install_0606 as _rb_install_0606
correction_execution_user_has_perm = _rb_install_0606(_rb_bind(_rb_sys.modules[__name__], ()))


# RB_BIND:0607
from application.components.shared import install_0607 as _rb_install_0607
correction_execution_snapshot = _rb_install_0607(_rb_bind(_rb_sys.modules[__name__], ()))



# RB_CORRECTION_EXECUTION_CONCURRENCY_HARDENING_V1

CORRECTION_EXECUTION_SNAPSHOT_KEYS=(
    'correction_request_id',
    'correction_request_status',
    'request_type',
    'requester_id',
    'reviewer_id',
    'order_id',
    'branch_id',
    'order_status',
    'exam_status',
    'accession',
    'pacs_status',
    'pacs_image_count',
    'pacs_study_id',
    'study_instance_uid',
)


# RB_BIND:0608
from application.components.shared import install_0608 as _rb_install_0608
correction_execution_request_lock = _rb_install_0608(_rb_bind(_rb_sys.modules[__name__], ('CorrectionRequest',)))


# RB_BIND:0609
from application.components.shared import install_0609 as _rb_install_0609
correction_execution_row_lock = _rb_install_0609(_rb_bind(_rb_sys.modules[__name__], ('CorrectionExecution',)))


# RB_BIND:0610
from application.components.shared import install_0610 as _rb_install_0610
correction_execution_saved_snapshot = _rb_install_0610(_rb_bind(_rb_sys.modules[__name__], ('json',)))


# RB_BIND:0611
from application.components.shared import install_0611 as _rb_install_0611
correction_execution_assert_snapshot_current = _rb_install_0611(_rb_bind(_rb_sys.modules[__name__], ('CORRECTION_EXECUTION_SNAPSHOT_KEYS', 'correction_execution_saved_snapshot', 'correction_execution_snapshot')))


# RB_BIND:0612
from application.components.shared import install_0612 as _rb_install_0612
correction_execution_prepare = _rb_install_0612(_rb_bind(_rb_sys.modules[__name__], ('CorrectionExecution', 'correction_execution_request_lock', 'correction_execution_snapshot', 'db', 'json')))


# RB_BIND:0613
from application.components.shared import install_0613 as _rb_install_0613
correction_execution_assert_actor = _rb_install_0613(_rb_bind(_rb_sys.modules[__name__], ('Order', 'correction_branch_allowed', 'correction_execution_request_lock', 'correction_execution_user_has_perm', 'db')))


# RB_BIND:0614
from application.components.shared import install_0614 as _rb_install_0614
correction_execution_transition = _rb_install_0614(_rb_bind(_rb_sys.modules[__name__], ('CORRECTION_EXECUTION_STATUSES', 'CORRECTION_EXECUTION_TRANSITIONS', 'correction_execution_assert_actor', 'correction_execution_assert_snapshot_current', 'correction_execution_row_lock', 'datetime', 'db')))


# RB_BIND:0615
from application.components.shared import install_0615 as _rb_install_0615
correction_execution_hold = _rb_install_0615(_rb_bind(_rb_sys.modules[__name__], ('correction_execution_transition',)))


# RB_BIND:0616
from application.components.shared import install_0616 as _rb_install_0616
correction_execution_release = _rb_install_0616(_rb_bind(_rb_sys.modules[__name__], ('correction_execution_transition',)))










CORRECTION_REVIEW_STATUSES=('ALL','NEW','UNDER_REVIEW','NEED_INFO','APPROVED','REJECTED','COMPLETED')


# RB_BIND:0617
from application.components.shared import install_0617 as _rb_install_0617
correction_can_request = _rb_install_0617(_rb_bind(_rb_sys.modules[__name__], ('has_perm',)))


# RB_BIND:0618
from application.components.shared import install_0618 as _rb_install_0618
correction_can_review = _rb_install_0618(_rb_bind(_rb_sys.modules[__name__], ('has_perm',)))


# RB_BIND:0619
from application.components.shared import install_0619 as _rb_install_0619
correction_branch_allowed = _rb_install_0619(_rb_bind(_rb_sys.modules[__name__], ('current_user', 'has_perm')))


# RB_BIND:0620
from application.components.shared import install_0620 as _rb_install_0620
correction_search_orders = _rb_install_0620(_rb_bind(_rb_sys.modules[__name__], ('Order', 'Patient', 'current_user', 'db', 'has_perm')))


# RB_BIND:0621
from web.routes.shared import install_0621 as _rb_install_0621
correction_center = _rb_install_0621(_rb_bind(_rb_sys.modules[__name__], ('CORRECTION_REQUEST_TYPES', 'CorrectionRequest', 'abort', 'app', 'correction_can_request', 'correction_can_review', 'correction_search_orders', 'current_user', 'db', 'login_required', 'render_template', 'request')))


# RB_BIND:0622
from web.routes.registration import install_0622 as _rb_install_0622
patient_demographics_correction = _rb_install_0622(_rb_bind(_rb_sys.modules[__name__], ('Order', 'Patient', 'abort', 'app', 'correction_branch_allowed', 'db', 'flash', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0623
from web.routes.registration import install_0623 as _rb_install_0623
patient_demographics_correction_view = _rb_install_0623(_rb_bind(_rb_sys.modules[__name__], ('CorrectionRequest', 'Patient', 'abort', 'app', 'correction_branch_allowed', 'correction_can_request', 'correction_can_review', 'current_user', 'json', 'login_required', 'render_template')))

# RB_BIND:0624
from web.routes.shared import install_0624 as _rb_install_0624
correction_request_submit = _rb_install_0624(_rb_bind(_rb_sys.modules[__name__], ('CORRECTION_OPEN_STATUSES', 'CORRECTION_REQUEST_TYPE_CODES', 'CorrectionRequest', 'IntegrityError', 'Order', 'abort', 'app', 'audit', 'correction_branch_allowed', 'current_user', 'datetime', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))



# RB_CORRECTION_EXECUTION_PREVIEW_ROUTES_PHASE2A_V1

# RB_BIND:0625
from application.components.shared import install_0625 as _rb_install_0625
correction_execution_assert_request_actor = _rb_install_0625(_rb_bind(_rb_sys.modules[__name__], ('correction_branch_allowed', 'correction_execution_user_has_perm')))


# RB_BIND:0626
from web.components.shared import install_0626 as _rb_install_0626
correction_execution_http_error = _rb_install_0626(_rb_bind(_rb_sys.modules[__name__], ('abort',)))


# RB_BIND:0627
from web.routes.shared import install_0627 as _rb_install_0627
correction_execution_queue = _rb_install_0627(_rb_bind(_rb_sys.modules[__name__], ('CorrectionExecution', 'CorrectionRequest', 'Order', 'app', 'correction_branch_allowed', 'current_user', 'db', 'render_template', 'require_perm')))


# RB_BIND:0628
from web.routes.shared import install_0628 as _rb_install_0628
correction_execution_prepare_route = _rb_install_0628(_rb_bind(_rb_sys.modules[__name__], ('CorrectionRequest', 'app', 'audit', 'correction_execution_assert_request_actor', 'correction_execution_http_error', 'correction_execution_prepare', 'current_user', 'db', 'flash', 'redirect', 'require_perm', 'url_for')))


# RB_BIND:0629
from web.routes.shared import install_0629 as _rb_install_0629
correction_execution_claim_route = _rb_install_0629(_rb_bind(_rb_sys.modules[__name__], ('CorrectionExecution', 'abort', 'app', 'audit', 'correction_execution_http_error', 'correction_execution_transition', 'current_user', 'db', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_BIND:0630
from web.routes.shared import install_0630 as _rb_install_0630
correction_execution_hold_route = _rb_install_0630(_rb_bind(_rb_sys.modules[__name__], ('CorrectionExecution', 'abort', 'app', 'audit', 'correction_execution_hold', 'correction_execution_http_error', 'current_user', 'db', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_BIND:0631
from web.routes.shared import install_0631 as _rb_install_0631
correction_execution_release_route = _rb_install_0631(_rb_bind(_rb_sys.modules[__name__], ('CorrectionExecution', 'abort', 'app', 'audit', 'correction_execution_http_error', 'correction_execution_release', 'current_user', 'db', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_CORRECTION_WRONG_IMAGES_EXECUTION_PHASE2B_V1
# RB_CORRECTION_PHASE2B_TRUE_STATE_MACHINE_V1
# RB_BIND:0632
from application.components.shared import install_0632 as _rb_install_0632
correction_wrong_images_execute = _rb_install_0632(_rb_bind(_rb_sys.modules[__name__], ('CorrectionExecution', 'CorrectionRequest', 'Order', 'audit', 'correction_execution_assert_snapshot_current', 'correction_execution_transition', 'datetime', 'db', 'granular_delete_images_only_execute', 'granular_delete_images_only_preflight', 'json')))




# RB_CORRECTION_PHASE2B_HTTP_EXECUTE_ROUTE_V1
# RB_BIND:0633
from web.routes.shared import install_0633 as _rb_install_0633
correction_execution_execute_route = _rb_install_0633(_rb_bind(_rb_sys.modules[__name__], ('abort', 'app', 'correction_wrong_images_execute', 'current_user', 'flash', 'jsonify', 'os', 'request', 'require_perm')))



# RB_BIND:0634
from web.routes.shared import install_0634 as _rb_install_0634
correction_request_queue = _rb_install_0634(_rb_bind(_rb_sys.modules[__name__], ('CORRECTION_REVIEW_STATUSES', 'CorrectionRequest', 'app', 'current_user', 'db', 'render_template', 'request', 'require_perm')))


# RB_BIND:0635
from web.routes.shared import install_0635 as _rb_install_0635
correction_request_action = _rb_install_0635(_rb_bind(_rb_sys.modules[__name__], ('CorrectionRequest', 'abort', 'app', 'audit', 'correction_branch_allowed', 'current_user', 'datetime', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_BIND:0636
from web.routes.operations import install_0636 as _rb_install_0636
admin_data_deletion_center = _rb_install_0636(_rb_bind(_rb_sys.modules[__name__], ('app', 'cancellation_search_rows', 'render_template', 'request', 'require_perm')))


# RB_BIND:0637
from web.routes.operations import install_0637 as _rb_install_0637
admin_cancellation_center = _rb_install_0637(_rb_bind(_rb_sys.modules[__name__], ('app', 'cancellation_search_rows', 'render_template', 'request', 'require_perm')))

# RB_BIND:0638
from web.routes.registration import install_0638 as _rb_install_0638
admin_cancel_appointment = _rb_install_0638(_rb_bind(_rb_sys.modules[__name__], ('Appointment', 'app', 'audit', 'cancellation_reason', 'datetime', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0639
from web.routes.registration import install_0639 as _rb_install_0639
admin_cancel_order = _rb_install_0639(_rb_bind(_rb_sys.modules[__name__], ('AssignmentOffer', 'CaseAssignment', 'Order', 'abort', 'app', 'audit', 'cancellation_reason', 'datetime', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0640
from web.routes.registration import install_0640 as _rb_install_0640
admin_remove_order_referral = _rb_install_0640(_rb_bind(_rb_sys.modules[__name__], ('Order', 'abort', 'app', 'audit', 'cancellation_reason', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_BIND:0641
from web.routes.operations import install_0641 as _rb_install_0641
admin_roles = _rb_install_0641(_rb_bind(_rb_sys.modules[__name__], ('Role', 'abort', 'app', 'audit', 'current_user', 'db', 'flash', 'json', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0642
from web.routes.operations import install_0642 as _rb_install_0642
admin_access_review = _rb_install_0642(_rb_bind(_rb_sys.modules[__name__], ('Role', 'app', 'render_template', 'request', 'require_perm')))

# RB_BIND:0643
from web.routes.operations import install_0643 as _rb_install_0643
admin_branches = _rb_install_0643(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'app', 'audit', 'db', 'flash', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0644
from web.routes.registration import install_0644 as _rb_install_0644
admin_device_page = _rb_install_0644(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Device', 'app', 'audit', 'db', 'flash', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))


# RB_BIND:0645
from application.components.shared import install_0645 as _rb_install_0645
migrate_existing_schema = _rb_install_0645(_rb_bind(_rb_sys.modules[__name__], ('db', 'inspect', 'text')))

# RB_BIND:0646
from application.components.shared import install_0646 as _rb_install_0646
seed_master_data = _rb_install_0646(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Consumable', 'Contract', 'PROCEDURE_CATALOG', 'Procedure', 'ProcedureConsumable', 'ROLE_PERMS', 'Role', 'User', 'db', 'generate_password_hash', 'json', 'os', 'storage_defaults')))


# RB_BIND:0647
from web.routes.registration import install_0647 as _rb_install_0647
my_registration_style = _rb_install_0647(_rb_bind(_rb_sys.modules[__name__], ('UserRegistrationPreference', 'abort', 'app', 'current_user', 'db', 'login_required', 'redirect', 'request', 'url_for')))

# RB_BIND:0648
from web.routes.operations import install_0648 as _rb_install_0648
admin_field_rules = _rb_install_0648(_rb_bind(_rb_sys.modules[__name__], ('FieldRule', 'app', 'audit', 'current_user', 'db', 'flash', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0649
from web.routes.registration import install_0649 as _rb_install_0649
patient_alerts = _rb_install_0649(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Patient', 'PatientAlert', 'app', 'current_user', 'db', 'login_required', 'redirect', 'render_template', 'request', 'url_for')))

# RB_BIND:0650
from web.routes.registration import install_0650 as _rb_install_0650
patient_alert_resolve = _rb_install_0650(_rb_bind(_rb_sys.modules[__name__], ('PatientAlert', 'app', 'current_user', 'datetime', 'db', 'login_required', 'redirect', 'request', 'url_for')))

# RB_BIND:0651
from web.routes.operations import install_0651 as _rb_install_0651
admin_company_legal = _rb_install_0651(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'CompanyLegalProfile', 'app', 'db', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0652
from web.routes.shared import install_0652 as _rb_install_0652
studio_home = _rb_install_0652(_rb_bind(_rb_sys.modules[__name__], ('app', 'render_template', 'require_perm')))

# RB_BIND:0653
from web.routes.shared import install_0653 as _rb_install_0653
studio_templates = _rb_install_0653(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'StudioTemplate', 'abort', 'app', 'current_user', 'db', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0654
from web.routes.shared import install_0654 as _rb_install_0654
studio_template_publish = _rb_install_0654(_rb_bind(_rb_sys.modules[__name__], ('StudioTemplate', 'StudioTemplateVersion', 'abort', 'app', 'current_user', 'datetime', 'db', 'json', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0655
from web.routes.shared import install_0655 as _rb_install_0655
doctor_templates = _rb_install_0655(_rb_bind(_rb_sys.modules[__name__], ('DoctorWordTemplate', 'Procedure', 'REPORT_DIR', 'abort', 'app', 'current_user', 'datetime', 'db', 'login_required', 'or_', 'os', 're', 'redirect', 'render_template', 'request', 'url_for')))


# RB_BIND:0656
from web.routes.finance import install_0656 as _rb_install_0656
receipt_studio_print = _rb_install_0656(_rb_bind(_rb_sys.modules[__name__], ('Receipt', 'app', 'company_legal_for_order', 'json', 'login_required', 'receipt_studio_template', 'render_template')))

# RB_BIND:0657
from web.routes.shared import install_0657 as _rb_install_0657
studio_template_preview = _rb_install_0657(_rb_bind(_rb_sys.modules[__name__], ('StudioTemplate', 'app', 'json', 'render_template', 'require_perm')))

# RB_BIND:0658
from web.routes.reporting import install_0658 as _rb_install_0658
report_word_from_doctor_template = _rb_install_0658(_rb_bind(_rb_sys.modules[__name__], ('DoctorWordTemplate', 'Document', 'Order', '_next_report_version', '_rb_authoritative_docx_download', 'abort', 'app', 'build_report_filename', 'current_user', 'get_report', 'has_perm', 'io', 'replace_docx_placeholders', 'require_perm', 'send_file')))


# RB_BIND:0659
from web.routes.reporting import install_0659 as _rb_install_0659
admin_ai_reporting = _rb_install_0659(_rb_bind(_rb_sys.modules[__name__], ('AIProviderSetting', 'app', 'current_user', 'db', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0660
from web.routes.reporting import install_0660 as _rb_install_0660
ai_report_assistant = _rb_install_0660(_rb_bind(_rb_sys.modules[__name__], ('AIDraftReport', 'Order', 'PriorStudySelection', 'ai_provider_setting', 'app', 'external_priors_for_order', 'related_prior_orders', 'render_template', 'require_perm')))

# RB_BIND:0661
from web.routes.reporting import install_0661 as _rb_install_0661
ai_report_generate = _rb_install_0661(_rb_bind(_rb_sys.modules[__name__], ('AIDraftReport', 'Order', 'ai_provider_setting', 'app', 'audit', 'call_configured_ai_provider', 'current_user', 'db', 'flash', 'generate_ai_draft_payload', 'json', 'redirect', 'require_perm', 'url_for')))

# RB_BIND:0662
from web.routes.reporting import install_0662 as _rb_install_0662
ai_report_apply = _rb_install_0662(_rb_bind(_rb_sys.modules[__name__], ('AIDraftReport', 'Order', 'app', 'audit', 'current_user', 'datetime', 'db', 'flash', 'get_report', 'redirect', 'require_perm', 'url_for')))

# RB_BIND:0663
from web.routes.reporting import install_0663 as _rb_install_0663
ai_report_reject = _rb_install_0663(_rb_bind(_rb_sys.modules[__name__], ('AIDraftReport', 'app', 'current_user', 'datetime', 'db', 'redirect', 'require_perm', 'url_for')))

# RB_BIND:0664
from web.routes.reporting import install_0664 as _rb_install_0664
report_prior_select = _rb_install_0664(_rb_bind(_rb_sys.modules[__name__], ('PriorStudySelection', 'abort', 'app', 'current_user', 'db', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0665
from web.routes.registration import install_0665 as _rb_install_0665
patient_external_priors = _rb_install_0665(_rb_bind(_rb_sys.modules[__name__], ('ATTACHMENT_DIR', 'ExternalPriorStudy', 'Order', 'Patient', 'abort', 'app', 'current_user', 'datetime', 'db', 'json', 'login_required', 'os', 're', 'redirect', 'render_template', 'request', 'safe_dicom_metadata', 'score_prior_match', 'url_for')))

# RB_BIND:0666
from web.routes.reporting import install_0666 as _rb_install_0666
external_prior_confirm_match = _rb_install_0666(_rb_bind(_rb_sys.modules[__name__], ('ExternalPriorStudy', 'app', 'db', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0667
from web.routes.reporting import install_0667 as _rb_install_0667
external_prior_mark_mismatch = _rb_install_0667(_rb_bind(_rb_sys.modules[__name__], ('ExternalPriorStudy', 'app', 'db', 'redirect', 'request', 'require_perm', 'url_for')))



# RB_BIND:0668
from web.routes.shared import install_0668 as _rb_install_0668
my_quick_actions = _rb_install_0668(_rb_bind(_rb_sys.modules[__name__], ('app', 'login_required', 'quick_action_role', 'render_template', 'role_quick_actions')))

# RB_BIND:0669
from web.routes.shared import install_0669 as _rb_install_0669
universal_search = _rb_install_0669(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Order', 'Patient', 'PatientDocument', 'app', 'common_search_filters', 'external_priors_for_order', 'login_required', 'or_', 'related_prior_orders', 'render_template')))



# RB_BIND:0670
from web.routes.communications import install_0670 as _rb_install_0670
admin_communications = _rb_install_0670(_rb_bind(_rb_sys.modules[__name__], ('CommunicationChannelSetting', 'CommunicationTemplate', 'app', 'current_user', 'db', 'flash', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0671
from web.routes.communications import install_0671 as _rb_install_0671
admin_communication_template_save = _rb_install_0671(_rb_bind(_rb_sys.modules[__name__], ('CommunicationTemplate', 'app', 'current_user', 'db', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0672
from web.routes.communications import install_0672 as _rb_install_0672
patient_communications = _rb_install_0672(_rb_bind(_rb_sys.modules[__name__], ('CommunicationTemplate', 'Order', 'Patient', 'PatientCommunication', 'Receipt', 'app', 'audit', 'communication_context', 'current_user', 'db', 'flash', 'login_required', 'patient_default_recipient', 'redirect', 'render_comm_text', 'render_template', 'request', 'send_patient_communication', 'url_for')))

# RB_BIND:0673
from web.routes.communications import install_0673 as _rb_install_0673
communication_retry = _rb_install_0673(_rb_bind(_rb_sys.modules[__name__], ('PatientCommunication', 'app', 'flash', 'login_required', 'redirect', 'request', 'send_patient_communication', 'url_for')))



# RB_BIND:0674
from web.routes.imaging import install_0674 as _rb_install_0674
admin_imaging_ai = _rb_install_0674(_rb_bind(_rb_sys.modules[__name__], ('ImagingAIProvider', 'app', 'current_user', 'db', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0675
from web.routes.reporting import install_0675 as _rb_install_0675
imaging_ai_workspace = _rb_install_0675(_rb_bind(_rb_sys.modules[__name__], ('ImagingAIAnalysis', 'Order', 'app', 'imaging_ai_provider_for_modality', 'render_template', 'require_perm')))

# RB_BIND:0676
from web.routes.reporting import install_0676 as _rb_install_0676
imaging_ai_verify_demographics = _rb_install_0676(_rb_bind(_rb_sys.modules[__name__], ('ImagingAIAnalysis', 'Order', 'app', 'audit', 'current_user', 'db', 'demographic_verification', 'flash', 'imaging_ai_provider_for_modality', 'json', 'pacs_demographic_gate', 'pacs_study_lookup', 'redirect', 'require_perm', 'url_for')))

# RB_BIND:0677
from web.routes.reporting import install_0677 as _rb_install_0677
imaging_ai_demographic_override = _rb_install_0677(_rb_bind(_rb_sys.modules[__name__], ('DemographicOverrideAudit', 'ImagingAIAnalysis', 'app', 'audit', 'current_user', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0678
from web.routes.reporting import install_0678 as _rb_install_0678
imaging_ai_analyze = _rb_install_0678(_rb_bind(_rb_sys.modules[__name__], ('ImagingAIAnalysis', 'Order', 'app', 'audit', 'call_imaging_ai', 'datetime', 'db', 'flash', 'json', 'pacs_demographic_gate', 'redirect', 'require_perm', 'url_for')))

# RB_BIND:0679
from web.routes.reporting import install_0679 as _rb_install_0679
imaging_ai_mark_reviewed = _rb_install_0679(_rb_bind(_rb_sys.modules[__name__], ('ImagingAIAnalysis', 'app', 'current_user', 'datetime', 'db', 'redirect', 'require_perm', 'url_for')))




# RB_BIND:0680
from web.routes.operations import install_0680 as _rb_install_0680
admin_qa_control = _rb_install_0680(_rb_bind(_rb_sys.modules[__name__], ('QARun', 'QATestResult', 'QA_N8N_RUN_WEBHOOK', 'app', 'render_template', 'require_perm')))


# RB_BIND:0681
from web.routes.operations import install_0681 as _rb_install_0681
admin_qa_control_run = _rb_install_0681(_rb_bind(_rb_sys.modules[__name__], ('QARun', 'QA_CALLBACK_TOKEN', 'QA_N8N_RUN_WEBHOOK', 'app', 'audit', 'current_user', 'datetime', 'db', 'flash', 'qa_n8n_call', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_BIND:0682
from web.routes.operations import install_0682 as _rb_install_0682
admin_qa_retry_failed = _rb_install_0682(_rb_bind(_rb_sys.modules[__name__], ('QARun', 'QATestResult', 'QA_CALLBACK_TOKEN', 'QA_N8N_RETRY_WEBHOOK', 'QA_N8N_RUN_WEBHOOK', 'app', 'flash', 'qa_n8n_call', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_BIND:0683
from web.routes.operations import install_0683 as _rb_install_0683
admin_qa_workflow_control = _rb_install_0683(_rb_bind(_rb_sys.modules[__name__], ('QA_N8N_DISABLE_WEBHOOK', 'QA_N8N_ENABLE_WEBHOOK', 'abort', 'app', 'audit', 'current_user', 'db', 'flash', 'qa_n8n_call', 'redirect', 'require_perm', 'url_for')))


# RB_BIND:0684
from web.routes.operations import install_0684 as _rb_install_0684
api_qa_callback = _rb_install_0684(_rb_bind(_rb_sys.modules[__name__], ('QARun', 'QATestResult', 'QA_CALLBACK_TOKEN', 'app', 'datetime', 'db', 'json', 'jsonify', 'qa_run_summary', 'request')))


# RB_BIND:0685
from web.routes.assignment import install_0685 as _rb_install_0685
admin_assignment_offers = _rb_install_0685(_rb_bind(_rb_sys.modules[__name__], ('AssignmentOfferPolicy', 'app', 'assignment_offer_policy', 'current_user', 'db', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))


# RB_BIND:0686
from application.components.shared import install_0686 as _rb_install_0686
_open_pool_visible_to_doctor = _rb_install_0686(_rb_bind(_rb_sys.modules[__name__], ('CaseAssignment', 'DoctorAssignmentProfile')))


# RB_BIND:0687
from application.components.registration import install_0687 as _rb_install_0687
order_is_open_pool = _rb_install_0687(_rb_bind(_rb_sys.modules[__name__], ('AssignmentEscalationEvent', 'assignment_strategy_for_order')))


# RB_BIND:0688
from web.routes.assignment import install_0688 as _rb_install_0688
open_pool_cases = _rb_install_0688(_rb_bind(_rb_sys.modules[__name__], ('Order', '_open_pool_visible_to_doctor', 'app', 'assignment_readiness_decision', 'current_user', 'login_required', 'order_is_open_pool', 'render_template')))


# RB_BIND:0689
from web.routes.assignment import install_0689 as _rb_install_0689
assign_case_to_me = _rb_install_0689(_rb_bind(_rb_sys.modules[__name__], ('AssignmentEscalationEvent', 'Order', '_open_pool_visible_to_doctor', 'abort', 'app', 'assignment_readiness_decision', 'audit', 'current_user', 'db', 'flash', 'login_required', 'manual_assign_order', 'order_is_open_pool', 'redirect', 'url_for')))


# RB_BIND:0690
from web.routes.assignment import install_0690 as _rb_install_0690
my_assignment_offers = _rb_install_0690(_rb_bind(_rb_sys.modules[__name__], ('AssignmentOffer', 'app', 'current_user', 'doctor_availability', 'expire_assignment_offers', 'login_required', 'render_template')))

# RB_BIND:0691
from web.routes.assignment import install_0691 as _rb_install_0691
assignment_availability_update = _rb_install_0691(_rb_bind(_rb_sys.modules[__name__], ('abort', 'app', 'current_user', 'datetime', 'db', 'doctor_availability', 'login_required', 'redirect', 'request', 'timedelta', 'url_for')))

# RB_BIND:0692
from web.routes.assignment import install_0692 as _rb_install_0692
assignment_offer_respond = _rb_install_0692(_rb_bind(_rb_sys.modules[__name__], ('AssignmentEscalationEvent', 'AssignmentOffer', 'CaseAssignment', 'Order', 'Report', 'User', 'abort', 'app', 'assignment_offer_policy', 'assignment_strategy_for_order', 'case_pacs_url', 'current_user', 'datetime', 'db', 'doctor_availability', 'doctor_workspace_user', 'expire_assignment_offers', 'flash', 'has_perm', 'jsonify', 'login_required', 'offer_next_doctor', 'ohif_url_for_order', 'redirect', 'request', 'timedelta', 'url_for')))

# RB_BIND:0693
from web.routes.assignment import install_0693 as _rb_install_0693
assignment_offer_start = _rb_install_0693(_rb_bind(_rb_sys.modules[__name__], ('AssignmentOffer', 'Order', 'app', 'current_user', 'flash', 'offer_next_doctor', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0694
from web.routes.assignment import install_0694 as _rb_install_0694
assignment_offers_expire_now = _rb_install_0694(_rb_bind(_rb_sys.modules[__name__], ('app', 'expire_assignment_offers', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))



# RB_BIND:0695
from web.routes.readiness import install_0695 as _rb_install_0695
case_readiness_dashboard = _rb_install_0695(_rb_bind(_rb_sys.modules[__name__], ('CaseReadiness', 'Order', 'app', 'case_delay_snapshot', 'render_template', 'require_perm')))


# RB_BIND:0696
from web.routes.shared import install_0696 as _rb_install_0696
case_operations_control_center = _rb_install_0696(_rb_bind(_rb_sys.modules[__name__], ('AssignmentEscalationEvent', 'AssignmentOffer', 'AssignmentStrategy', 'CaseAssignment', 'CaseReadiness', 'Order', 'PacsExamLink', 'PacsResolutionEvent', 'PacsStudyInbox', 'Report', 'app', 'db', 'has_perm', 'render_template', 'request', 'require_perm')))


# RB_BIND:0697
from web.routes.readiness import install_0697 as _rb_install_0697
case_exception_resolver = _rb_install_0697(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'CaseExceptionResolutionEvent', 'CaseReadiness', 'CaseReadinessOverride', 'Order', 'Patient', 'ReadinessSeriesAlias', 'Report', 'app', 'db', 'func', 'has_perm', 'json', 'quality_rule_for', 'render_template', 'request', 'require_perm')))


# RB_BIND:0698
from web.routes.readiness import install_0698 as _rb_install_0698
case_exception_series_map = _rb_install_0698(_rb_bind(_rb_sys.modules[__name__], ('CaseExceptionResolutionEvent', 'CaseReadiness', 'Order', 'ProcedureQualityRule', 'ReadinessSeriesAlias', 'Report', 'app', 'audit', 'current_user', 'db', 'evaluate_case_readiness', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_BIND:0699
from web.routes.readiness import install_0699 as _rb_install_0699
case_exception_series_map_deactivate = _rb_install_0699(_rb_bind(_rb_sys.modules[__name__], ('CaseExceptionResolutionEvent', 'Order', 'ReadinessSeriesAlias', 'Report', 'app', 'audit', 'current_user', 'db', 'evaluate_case_readiness', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_BIND:0700
from application.components.readiness import install_0700 as _rb_install_0700
case_exception_history_allowed = _rb_install_0700(_rb_bind(_rb_sys.modules[__name__], ('has_perm',)))


# RB_BIND:0701
from application.components.readiness import install_0701 as _rb_install_0701
case_exception_refresh_allowed = _rb_install_0701(_rb_bind(_rb_sys.modules[__name__], ('has_perm',)))


# RB_BIND:0702
from application.components.readiness import install_0702 as _rb_install_0702
case_exception_override_allowed = _rb_install_0702(_rb_bind(_rb_sys.modules[__name__], ('has_perm',)))


# RB_BIND:0703
from web.routes.readiness import install_0703 as _rb_install_0703
case_exception_history_add = _rb_install_0703(_rb_bind(_rb_sys.modules[__name__], ('CaseExceptionResolutionEvent', 'ClinicalNote', 'Order', 'Report', 'abort', 'app', 'audit', 'case_exception_history_allowed', 'current_user', 'db', 'evaluate_case_readiness', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_BIND:0704
from web.routes.readiness import install_0704 as _rb_install_0704
case_exception_image_minimum = _rb_install_0704(_rb_bind(_rb_sys.modules[__name__], ('CaseExceptionResolutionEvent', 'CaseReadiness', 'CaseReadinessOverride', 'Order', 'Report', 'abort', 'app', 'audit', 'case_exception_override_allowed', 'current_user', 'db', 'evaluate_case_readiness', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_BIND:0705
from web.routes.readiness import install_0705 as _rb_install_0705
case_exception_image_minimum_deactivate = _rb_install_0705(_rb_bind(_rb_sys.modules[__name__], ('CaseExceptionResolutionEvent', 'CaseReadinessOverride', 'Order', 'Report', 'abort', 'app', 'audit', 'case_exception_override_allowed', 'current_user', 'db', 'evaluate_case_readiness', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_BIND:0706
from web.routes.readiness import install_0706 as _rb_install_0706
case_exception_refresh = _rb_install_0706(_rb_bind(_rb_sys.modules[__name__], ('CaseExceptionResolutionEvent', 'CaseReadiness', 'Order', 'Report', 'abort', 'app', 'audit', 'case_exception_refresh_allowed', 'current_user', 'db', 'evaluate_case_readiness', 'flash', 'redirect', 'require_perm', 'url_for')))

# RB_BIND:0707
from web.routes.readiness import install_0707 as _rb_install_0707
case_readiness_check = _rb_install_0707(_rb_bind(_rb_sys.modules[__name__], ('Order', 'app', 'create_ris_notification', 'evaluate_case_readiness', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0708
from web.routes.readiness import install_0708 as _rb_install_0708
admin_procedure_quality_rules = _rb_install_0708(_rb_bind(_rb_sys.modules[__name__], ('ProcedureQualityRule', 'StudioTemplate', 'abort', 'app', 'audit', 'db', 'flash', 'has_perm', 'json', 'redirect', 'render_template', 'request', 'require_perm', 'technician_acquisition_notes', 'url_for')))

# RB_BIND:0709
from web.routes.reporting import install_0709 as _rb_install_0709
return_case_to_workflow = _rb_install_0709(_rb_bind(_rb_sys.modules[__name__], ('AssignmentOffer', 'CaseAssignment', 'CaseReturnReason', 'CaseWorkflowEvent', 'Order', '__name__', '_reporting_support_sys', 'abort', 'app', 'audit', 'create_ris_notification', 'current_user', 'datetime', 'db', 'flash', 'has_perm', 'login_required', 'redirect', 'request', 'url_for')))

# RB_BIND:0710
from web.routes.reporting import install_0710 as _rb_install_0710
return_case_form = _rb_install_0710(_rb_bind(_rb_sys.modules[__name__], ('CaseReturnReason', 'Order', '__name__', '_reporting_support_sys', 'app', 'login_required', 'render_template')))


# RB_BIND:0711
from web.routes.shared import install_0711 as _rb_install_0711
returned_cases_queue = _rb_install_0711(_rb_bind(_rb_sys.modules[__name__], ('CaseWorkflowEvent', 'Order', 'app', 'case_return_snapshot', 'current_user', 'db', 'login_required', 'render_template')))


# RB_BIND:0712
from web.routes.shared import install_0712 as _rb_install_0712
resolve_returned_case = _rb_install_0712(_rb_bind(_rb_sys.modules[__name__], ('AssignmentOffer', 'CaseAssignment', 'CaseWorkflowEvent', 'CorrectionRequest', 'Order', 'Report', 'User', '__name__', '_reporting_support_sys', 'abort', 'app', 'assign_case_atomically', 'assignment_readiness_decision', 'audit', 'case_return_snapshot', 'coordinate_post_exam_assignment', 'create_ris_notification', 'current_ui_language', 'current_user', 'db', 'evaluate_case_readiness', 'flash', 'login_required', 'pacs_lookup_order', 'redirect', 'request', 'technician_completion_evidence', 'url_for')))


# RB_BIND:0713
from web.routes.reporting import install_0713 as _rb_install_0713
report_portal_download = _rb_install_0713(_rb_bind(_rb_sys.modules[__name__], ('Order', 'REPORT_DIR', '_rb_final_revision_document', '_rb_revision_pdf_download', 'app', 'db', 'flash', 'make_pdf', 'os', 'redirect', 'require_perm', 'send_file', 'url_for')))


# RB_BIND:0714
from web.routes.reporting import install_0714 as _rb_install_0714
report_portal_download_pdf = _rb_install_0714(_rb_bind(_rb_sys.modules[__name__], ('Order', '_rb_final_revision_document', '_rb_revision_pdf_download', 'abort', 'app', 'os', 'require_perm', 'send_file')))


# RB_BIND:0715
from web.routes.reporting import install_0715 as _rb_install_0715
report_portal_download_docx = _rb_install_0715(_rb_bind(_rb_sys.modules[__name__], ('Order', '_rb_authoritative_docx_download', 'abort', 'app', 'build_report_filename', 'os', 'require_perm', 'send_file')))



# ============================================================
# REPORT SLA v1
# Reporting SLA is intentionally isolated from assignment
# offer timeouts / MWL / PACS / Order.status.
# ============================================================

REPORT_SLA_MINUTES = {
    'ROUTINE': 24 * 60,
    'VIP':       2 * 60,
    'URGENT':    4 * 60,
    'STAT':          60,
    'EMERGENCY':     60,
}

REPORT_SLA_NEAR_PERCENT = 80


# RB_BIND:0716
from application.components.reporting import install_0716 as _rb_install_0716
report_sla_snapshot = _rb_install_0716(_rb_bind(_rb_sys.modules[__name__], ('REPORT_SLA_MINUTES', 'REPORT_SLA_NEAR_PERCENT', 'datetime')))


# RB_BIND:0717
from application.components.readiness import install_0717 as _rb_install_0717
format_sla_minutes = _rb_install_0717(_rb_bind(_rb_sys.modules[__name__], ()))

# RB_BIND:0718
from web.routes.reporting import install_0718 as _rb_install_0718
reports_portal = _rb_install_0718(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'CaseWorkflowEvent', 'Order', 'Patient', 'Role', 'User', 'app', 'datetime', 'format_sla_minutes', 'or_', 'render_template', 'report_sla_snapshot', 'request', 'require_perm', 'timedelta')))


# RB_BIND:0719
from web.routes.reporting import install_0719 as _rb_install_0719
reports_portal_csv = _rb_install_0719(_rb_bind(_rb_sys.modules[__name__], ('CaseWorkflowEvent', 'Order', 'Patient', 'Response', 'app', 'csv', 'datetime', 'io', 'or_', 'request', 'require_perm', 'timedelta')))



# ============================================================
# RB_SLA_TAT_ADMIN_V1
# SLA / TAT policy administration.
#
# Existing TATPolicy schema only.
# No schema creation or migration is performed here.
# Resolution semantics:
#   priority + modality exact
#   -> fallback priority + "*"
# ============================================================

# RB_BIND:0720
from web.routes.assignment import install_0720 as _rb_install_0720
admin_tat_policies = _rb_install_0720(_rb_bind(_rb_sys.modules[__name__], ('Order', 'TATPolicy', 'abort', 'app', 'audit', 'db', 'flash', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))


# RB_BIND:0721
from web.routes.readiness import install_0721 as _rb_install_0721
tat_dashboard = _rb_install_0721(_rb_bind(_rb_sys.modules[__name__], ('Order', 'app', 'case_delay_snapshot', 'render_template', 'require_perm')))

# RB_BIND:0722
from web.routes.communications import install_0722 as _rb_install_0722
notification_center = _rb_install_0722(_rb_bind(_rb_sys.modules[__name__], ('RISNotification', 'app', 'current_user', 'login_required', 'or_', 'render_template')))

# RB_BIND:0723
from web.routes.communications import install_0723 as _rb_install_0723
notification_ack = _rb_install_0723(_rb_bind(_rb_sys.modules[__name__], ('RISNotification', 'app', 'current_user', 'datetime', 'db', 'login_required', 'redirect', 'request', 'url_for')))

# RB_BIND:0724
from web.routes.shared import install_0724 as _rb_install_0724
escalation_board = _rb_install_0724(_rb_bind(_rb_sys.modules[__name__], ('RISNotification', 'app', 'render_template', 'require_perm')))



# RB_BIND:0725
from web.routes.shared import install_0725 as _rb_install_0725
advanced_search = _rb_install_0725(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'CaseAssignment', 'CaseReadiness', 'Report', 'Role', 'SavedSearchFilter', 'User', 'advanced_order_query', 'app', 'current_user', 'login_required', 'render_template')))

# ===== RADBRIDGE SAVED VIEWS V1 =====

SAVED_VIEW_PAGES = {
    'NURSING': 'nursing_worklist',
    'TECHNICIAN': 'technician_worklist',
    'DOCTOR': 'doctor_workspace',
}

SAVED_VIEW_FILTER_KEYS = {
    'NURSING': __import__('staff_worklist_filters').KEYS - {'view'},
    'TECHNICIAN': __import__('staff_worklist_filters').KEYS,
    'DOCTOR': {'q','patient_id','accession','priority','modality','exam_status',
               'pacs_status','report_status','sla_status','date_from','date_to'},
}

# RB_BIND:0726
from web.components.shared import install_0726 as _rb_install_0726
_saved_view_page_or_404 = _rb_install_0726(_rb_bind(_rb_sys.modules[__name__], ('SAVED_VIEW_PAGES', '__name__', '_reporting_support_sys', 'abort')))

# RB_BIND:0727
from application.components.shared import install_0727 as _rb_install_0727
_saved_view_filtered_data = _rb_install_0727(_rb_bind(_rb_sys.modules[__name__], ('SAVED_VIEW_FILTER_KEYS',)))

# RB_BIND:0728
from application.components.shared import install_0728 as _rb_install_0728
_saved_views_for = _rb_install_0728(_rb_bind(_rb_sys.modules[__name__], ('SavedSearchFilter', 'current_user')))

# RB_BIND:0729
from web.routes.shared import install_0729 as _rb_install_0729
saved_view_save = _rb_install_0729(_rb_bind(_rb_sys.modules[__name__], ('SAVED_VIEW_PAGES', 'SavedSearchFilter', '_saved_view_filtered_data', '_saved_view_page_or_404', 'abort', 'app', 'current_user', 'db', 'json', 'login_required', 'redirect', 'request', 'url_for')))

# RB_BIND:0730
from web.routes.shared import install_0730 as _rb_install_0730
saved_view_update = _rb_install_0730(_rb_bind(_rb_sys.modules[__name__], ('SAVED_VIEW_PAGES', 'SavedSearchFilter', '_saved_view_filtered_data', '_saved_view_page_or_404', 'app', 'audit', 'current_user', 'db', 'flash', 'json', 'login_required', 'redirect', 'request', 'url_for')))

# RB_BIND:0731
from web.routes.shared import install_0731 as _rb_install_0731
saved_view_apply = _rb_install_0731(_rb_bind(_rb_sys.modules[__name__], ('SAVED_VIEW_PAGES', 'SavedSearchFilter', '_saved_view_filtered_data', '_saved_view_page_or_404', 'app', 'current_user', 'json', 'login_required', 'redirect', 'url_for')))

# RB_BIND:0732
from web.routes.shared import install_0732 as _rb_install_0732
saved_view_set_default = _rb_install_0732(_rb_bind(_rb_sys.modules[__name__], ('SAVED_VIEW_PAGES', 'SavedSearchFilter', '_saved_view_page_or_404', 'app', 'current_user', 'db', 'login_required', 'redirect', 'url_for')))

# RB_BIND:0733
from web.routes.shared import install_0733 as _rb_install_0733
saved_view_rename = _rb_install_0733(_rb_bind(_rb_sys.modules[__name__], ('SAVED_VIEW_PAGES', 'SavedSearchFilter', '_saved_view_page_or_404', 'abort', 'app', 'current_user', 'db', 'login_required', 'redirect', 'request', 'url_for')))

# RB_BIND:0734
from web.routes.shared import install_0734 as _rb_install_0734
saved_view_delete = _rb_install_0734(_rb_bind(_rb_sys.modules[__name__], ('SAVED_VIEW_PAGES', 'SavedSearchFilter', '_saved_view_page_or_404', 'app', 'current_user', 'db', 'login_required', 'redirect', 'url_for')))

# ===== END RADBRIDGE SAVED VIEWS V1 =====


# RB_BIND:0735
from web.routes.shared import install_0735 as _rb_install_0735
advanced_search_save = _rb_install_0735(_rb_bind(_rb_sys.modules[__name__], ('SavedSearchFilter', 'abort', 'app', 'current_user', 'db', 'json', 'login_required', 'redirect', 'request', 'url_for')))

# RB_BIND:0736
from web.routes.shared import install_0736 as _rb_install_0736
advanced_search_apply_saved = _rb_install_0736(_rb_bind(_rb_sys.modules[__name__], ('SavedSearchFilter', 'app', 'current_user', 'json', 'login_required', 'redirect', 'url_for')))

# RB_BIND:0737
from web.routes.registration import install_0737 as _rb_install_0737
order_actions = _rb_install_0737(_rb_bind(_rb_sys.modules[__name__], ('CASE_LIFECYCLE', 'Order', 'app', 'login_required', 'order_action_context', 'order_visual_state', 'render_template')))

# RB_BIND:0738
from web.routes.registration import install_0738 as _rb_install_0738
order_label_print = _rb_install_0738(_rb_bind(_rb_sys.modules[__name__], ('Order', 'app', 'default_label_template', 'json', 'login_required', 'render_template')))

# RB_BIND:0739
from web.routes.operations import install_0739 as _rb_install_0739
admin_label_templates = _rb_install_0739(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'LabelTemplate', 'app', 'db', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0740
from web.routes.reporting import install_0740 as _rb_install_0740
assignment_reports = _rb_install_0740(_rb_bind(_rb_sys.modules[__name__], ('app', 'assignment_report_rows', 'render_template', 'require_perm')))

# RB_BIND:0741
from web.routes.assignment import install_0741 as _rb_install_0741
assignment_dashboard = _rb_install_0741(_rb_bind(_rb_sys.modules[__name__], ('app', 'assignment_report_rows', 'render_template', 'request', 'require_perm')))

# RB_BIND:0742
from web.routes.reporting import install_0742 as _rb_install_0742
assignment_report_csv = _rb_install_0742(_rb_bind(_rb_sys.modules[__name__], ('app', 'assignment_report_rows', 'require_perm')))



# RB_BIND:0743
from web.routes.registration import install_0743 as _rb_install_0743
smart_intake_start = _rb_install_0743(_rb_bind(_rb_sys.modules[__name__], ('SmartIntakeSession', 'app', 'current_user', 'db', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0744
from web.routes.registration import install_0744 as _rb_install_0744
smart_intake_workspace = _rb_install_0744(_rb_bind(_rb_sys.modules[__name__], ('Contract', 'OCR_INTAKE_CORE_FIELDS', 'Patient', 'Procedure', 'SMART_DOCUMENT_TYPES', 'SmartIntakeDocument', 'SmartIntakeFieldReview', 'SmartIntakeSession', '_intake_final_identity', '_intake_match_context', 'app', 'enabled_ocr_provider', 'json', 'rebuild_intake_summary', 'render_template', 'require_perm', 'validate_intake_identity')))

# RB_BIND:0745
from web.routes.registration import install_0745 as _rb_install_0745
smart_intake_document_file = _rb_install_0745(_rb_bind(_rb_sys.modules[__name__], ('SmartIntakeDocument', 'abort', 'app', 'os', 'request', 'require_perm', 'send_file')))

# RB_BIND:0746
from web.routes.registration import install_0746 as _rb_install_0746
smart_intake_upload = _rb_install_0746(_rb_bind(_rb_sys.modules[__name__], ('SMART_DOCUMENT_TYPES', 'SmartIntakeDocument', 'SmartIntakeSession', '_apply_smart_intake_ocr_result', '_intake_file_sha256', '_intake_redirect', 'abort', 'app', 'audit', 'call_ocr_provider', 'db', 'enabled_ocr_provider', 'flash', 'ocr_configuration_hash', 'os', 're', 'rebuild_intake_summary', 'request', 'require_perm', 'smart_intake_dir', 'uuid')))


# RB_BIND:0747
from application.components.registration import install_0747 as _rb_install_0747
_apply_smart_intake_ocr_result = _rb_install_0747(_rb_bind(_rb_sys.modules[__name__], ('OCR_INTAKE_CORE_FIELDS', 'SmartIntakeFieldReview', 'classify_document_text', 'db', 'enabled_ocr_provider', 'json', 'normalize_intake_identity', 'ocr_configuration_hash', 'validate_intake_identity')))


# RB_BIND:0748
from web.routes.registration import install_0748 as _rb_install_0748
smart_intake_run_ocr = _rb_install_0748(_rb_bind(_rb_sys.modules[__name__], ('SmartIntakeDocument', 'SmartIntakeSession', '_apply_smart_intake_ocr_result', '_intake_redirect', 'abort', 'app', 'audit', 'call_ocr_provider', 'db', 'flash', 'rebuild_intake_summary', 'require_perm')))

# RB_BIND:0749
from web.routes.registration import install_0749 as _rb_install_0749
smart_intake_run_all_ocr = _rb_install_0749(_rb_bind(_rb_sys.modules[__name__], ('SmartIntakeDocument', 'SmartIntakeSession', '_apply_smart_intake_ocr_result', '_intake_redirect', 'abort', 'app', 'audit', 'call_ocr_provider', 'db', 'rebuild_intake_summary', 'require_perm')))

# RB_BIND:0750
from web.routes.registration import install_0750 as _rb_install_0750
smart_intake_document_review = _rb_install_0750(_rb_bind(_rb_sys.modules[__name__], ('SMART_DOCUMENT_TYPES', 'SmartIntakeDocument', 'SmartIntakeSession', '_intake_redirect', 'abort', 'app', 'audit', 'current_user', 'datetime', 'db', 'flash', 'rebuild_intake_summary', 'request', 'require_perm')))


# RB_BIND:0751
from application.components.registration import install_0751 as _rb_install_0751
_save_intake_review_fields = _rb_install_0751(_rb_bind(_rb_sys.modules[__name__], ('OCR_INTAKE_CORE_FIELDS', 'SmartIntakeFieldReview', 'current_user', 'datetime', 'db')))


# RB_BIND:0752
from application.components.registration import install_0752 as _rb_install_0752
_save_intake_visit_review = _rb_install_0752(_rb_bind(_rb_sys.modules[__name__], ('Contract', 'Procedure', 'current_user', 'datetime', 'json')))


# RB_BIND:0753
from application.components.registration import install_0753 as _rb_install_0753
_attach_intake_documents = _rb_install_0753(_rb_bind(_rb_sys.modules[__name__], ('PatientDocument', 'SmartIntakeDocument', 'db')))


# RB_BIND:0754
from web.components.registration import install_0754 as _rb_install_0754
_block_intake_registration = _rb_install_0754(_rb_bind(_rb_sys.modules[__name__], ('_intake_redirect', 'audit', 'db', 'flash')))


# RB_BIND:0755
from web.routes.registration import install_0755 as _rb_install_0755
smart_intake_review_save = _rb_install_0755(_rb_bind(_rb_sys.modules[__name__], ('OCR_INTAKE_CORE_FIELDS', 'Patient', 'SmartIntakeDocument', 'SmartIntakeSession', '_attach_intake_documents', '_block_intake_registration', '_intake_match_context', '_intake_redirect', '_save_intake_review_fields', '_save_intake_visit_review', 'abort', 'app', 'audit', 'current_user', 'datetime', 'db', 'flash', 'has_perm', 'intake_registration_fingerprint', 'next_patient_id', 'redirect', 'request', 'require_perm', 'text', 'url_for', 'validate_intake_identity')))

# RB_BIND:0756
from application.components.shared import install_0756 as _rb_install_0756
_ocr_config_from_input = _rb_install_0756(_rb_bind(_rb_sys.modules[__name__], ()))


# RB_BIND:0757
from application.components.shared import install_0757 as _rb_install_0757
_ocr_saved_validation = _rb_install_0757(_rb_bind(_rb_sys.modules[__name__], ('json', 'ocr_configuration_hash')))


# RB_BIND:0758
from application.components.shared import install_0758 as _rb_install_0758
_store_ocr_validation = _rb_install_0758(_rb_bind(_rb_sys.modules[__name__], ('current_user', 'datetime', 'json', 'ocr_configuration_hash')))


# RB_BIND:0759
from web.routes.operations import install_0759 as _rb_install_0759
admin_ocr_settings = _rb_install_0759(_rb_bind(_rb_sys.modules[__name__], ('OCRProviderSetting', 'SUPPORTED_PROVIDERS', '_ocr_config_from_input', '_ocr_saved_validation', '_store_ocr_validation', 'app', 'audit', 'current_user', 'db', 'flash', 'os', 'redirect', 'render_template', 'request', 'require_perm', 'url_for', 'urlparse', 'validate_ocr_local', 'validate_ocr_remote')))


# RB_BIND:0760
from web.routes.operations import install_0760 as _rb_install_0760
admin_ocr_settings_validate = _rb_install_0760(_rb_bind(_rb_sys.modules[__name__], ('OCRProviderSetting', '_ocr_config_from_input', '_store_ocr_validation', 'app', 'audit', 'db', 'jsonify', 'os', 'request', 'require_perm', 'validate_ocr_remote')))



# RB_BIND:0761
from web.routes.reporting import install_0761 as _rb_install_0761
doctor_clinical_workspace = _rb_install_0761(_rb_bind(_rb_sys.modules[__name__], ('Attachment', 'CaseAssignment', 'CaseReadiness', 'ClinicalAssessment', 'ClinicalDiagnosis', 'ClinicalLabItem', 'ClinicalMedication', 'Order', 'Report', 'VitalSign', 'app', 'audit', 'current_user', 'db', 'external_priors_for_order', 'redirect', 'related_prior_orders', 'render_template', 'request', 'require_perm', 'url_for')))



# RB_BIND:0762
from web.routes.operations import install_0762 as _rb_install_0762
admin_printer_profiles = _rb_install_0762(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'PrinterProfile', 'User', 'app', 'audit', 'db', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0763
from web.routes.operations import install_0763 as _rb_install_0763
admin_printer_profile_delete = _rb_install_0763(_rb_bind(_rb_sys.modules[__name__], ('PrinterProfile', 'app', 'audit', 'db', 'redirect', 'require_perm', 'url_for')))

# RB_BIND:0764
from web.routes.registration import install_0764 as _rb_install_0764
order_print_center = _rb_install_0764(_rb_bind(_rb_sys.modules[__name__], ('Order', 'app', 'login_required', 'order_action_context', 'render_template', 'request', 'resolve_printer_profile')))



# RB_BIND:0765
from web.routes.registration import install_0765 as _rb_install_0765
my_registration_defaults = _rb_install_0765(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Contract', 'Procedure', 'ReferringDoctor', 'RegistrationDefault', 'app', 'current_user', 'db', 'flash', 'login_required', 'redirect', 'render_template', 'request', 'url_for')))

# RB_BIND:0766
from web.routes.registration import install_0766 as _rb_install_0766
admin_registration_defaults = _rb_install_0766(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Contract', 'Procedure', 'ReferringDoctor', 'RegistrationDefault', 'User', 'app', 'audit', 'current_user', 'db', 'flash', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))



# RB_BIND:0767
from web.routes.shared import install_0767 as _rb_install_0767
studio_shortcuts = _rb_install_0767(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'STUDIO_ICON_LIBRARY', 'StudioShortcut', 'User', 'app', 'audit', 'current_user', 'db', 'flash', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0768
from web.routes.shared import install_0768 as _rb_install_0768
studio_shortcut_toggle = _rb_install_0768(_rb_bind(_rb_sys.modules[__name__], ('StudioShortcut', 'app', 'audit', 'db', 'redirect', 'require_perm', 'url_for')))

# RB_BIND:0769
from web.routes.shared import install_0769 as _rb_install_0769
studio_shortcut_delete = _rb_install_0769(_rb_bind(_rb_sys.modules[__name__], ('StudioShortcut', 'app', 'audit', 'db', 'redirect', 'require_perm', 'url_for')))

# RB_BIND:0770
from web.routes.shared import install_0770 as _rb_install_0770
studio_shortcut_reorder = _rb_install_0770(_rb_bind(_rb_sys.modules[__name__], ('StudioShortcut', 'app', 'db', 'redirect', 'request', 'require_perm', 'url_for')))



# RB_BIND:0771
from web.routes.finance import install_0771 as _rb_install_0771
inject_finance_helpers = _rb_install_0771(_rb_bind(_rb_sys.modules[__name__], ('app', 'effective_patient_due', 'total_order_discount')))

# RB_BIND:0772
from web.routes.shared import install_0772 as _rb_install_0772
inject_studio_shortcuts = _rb_install_0772(_rb_bind(_rb_sys.modules[__name__], ('STUDIO_ICON_MAP', 'app', 'studio_shortcuts_for')))



# RB_BIND:0773
from web.routes.shared import install_0773 as _rb_install_0773
management_dashboard = _rb_install_0773(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'ReferringDoctor', 'app', 'management_date_range', 'management_summary', 'render_template', 'require_perm')))

# RB_BIND:0774
from web.routes.shared import install_0774 as _rb_install_0774
referral_analytics = _rb_install_0774(_rb_bind(_rb_sys.modules[__name__], ('Order', 'ReferringDoctor', 'app', 'management_date_range', 'management_orders_query', 'referral_commission_for_order', 'render_template', 'require_perm')))

# RB_BIND:0775
from web.routes.finance import install_0775 as _rb_install_0775
admin_referral_commissions = _rb_install_0775(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Procedure', 'ReferralCommissionRule', 'ReferringDoctor', 'app', 'audit', 'current_user', 'db', 'flash', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0776
from web.routes.finance import install_0776 as _rb_install_0776
daily_closing = _rb_install_0776(_rb_bind(_rb_sys.modules[__name__], (
    'app', 'login_required', 'current_user', 'has_perm', 'request', 'abort', 'flash',
    'redirect', 'url_for', 'datetime', 'date', 'Branch', 'DailyClosing', 'Order',
    'Payment', 'PaymentRefund', 'db', 'audit', 'render_template',
)))


# RB_BIND:0777
from web.routes.finance import install_0777 as _rb_install_0777
admin_doctor_share_rules = _rb_install_0777(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Contract', 'PerformingDoctorShareRule', 'Procedure', 'User', 'app', 'audit', 'current_user', 'db', 'flash', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0778
from web.routes.finance import install_0778 as _rb_install_0778
doctor_share_rule_toggle = _rb_install_0778(_rb_bind(_rb_sys.modules[__name__], ('PerformingDoctorShareRule', 'app', 'audit', 'db', 'redirect', 'require_perm', 'url_for')))

# RB_BIND:0779
from web.routes.registration import install_0779 as _rb_install_0779
set_performing_doctor = _rb_install_0779(_rb_bind(_rb_sys.modules[__name__], ('Order', 'OrderPerformingDoctor', 'User', 'app', 'audit', 'current_user', 'datetime', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0780
from web.routes.finance import install_0780 as _rb_install_0780
admin_discount_policies = _rb_install_0780(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Contract', 'DiscountPolicy', 'Procedure', 'Role', 'User', 'app', 'audit', 'current_user', 'db', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0781
from web.routes.finance import install_0781 as _rb_install_0781
discount_policy_toggle = _rb_install_0781(_rb_bind(_rb_sys.modules[__name__], ('DiscountPolicy', 'app', 'audit', 'db', 'redirect', 'require_perm', 'url_for')))

# RB_BIND:0782
from web.routes.finance import install_0782 as _rb_install_0782
discount_approval_queue = _rb_install_0782(_rb_bind(_rb_sys.modules[__name__], ('DiscountApprovalRequest', 'app', 'db', 'render_template', 'require_perm')))

# RB_BIND:0783
from web.routes.finance import install_0783 as _rb_install_0783
discount_approval_decision = _rb_install_0783(_rb_bind(_rb_sys.modules[__name__], ('DiscountApprovalRequest', 'OrderDiscount', 'app', 'audit', 'base_patient_due', 'current_user', 'datetime', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0784
from web.routes.reporting import install_0784 as _rb_install_0784
doctor_earnings_report = _rb_install_0784(_rb_bind(_rb_sys.modules[__name__], ('Order', '__name__', '_reporting_support_sys', 'abort', 'app', 'effective_patient_due', 'has_perm', 'management_date_range', 'management_orders_query', 'performing_doctor_share_for_order', 'render_template', 'request', 'require_perm', 'total_order_discount')))



# RB_BIND:0785
from web.routes.inventory import install_0785 as _rb_install_0785
inventory_warehouses = _rb_install_0785(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Warehouse', 'app', 'audit', 'db', 'flash', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0786
from web.routes.inventory import install_0786 as _rb_install_0786
inventory_items = _rb_install_0786(_rb_bind(_rb_sys.modules[__name__], ('InventoryItem', 'app', 'audit', 'db', 'flash', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0787
from web.routes.inventory import install_0787 as _rb_install_0787
inventory_migrate_legacy_items = _rb_install_0787(_rb_bind(_rb_sys.modules[__name__], ('Consumable', 'InventoryItem', 'app', 'audit', 'db', 'flash', 'redirect', 'require_perm', 'url_for')))

# RB_BIND:0788
from web.routes.inventory import install_0788 as _rb_install_0788
inventory_stock = _rb_install_0788(_rb_bind(_rb_sys.modules[__name__], ('InventoryItem', 'InventoryLot', 'Warehouse', 'WarehouseStockPolicy', 'abort', 'app', 'audit', 'datetime', 'db', 'flash', 'has_perm', 'inventory_available_qty', 'inventory_move', 'inventory_on_hand_qty', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0789
from web.routes.inventory import install_0789 as _rb_install_0789
inventory_stock_policies = _rb_install_0789(_rb_bind(_rb_sys.modules[__name__], ('InventoryItem', 'Warehouse', 'WarehouseStockPolicy', 'app', 'db', 'flash', 'low_stock_rows', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0790
from web.routes.inventory import install_0790 as _rb_install_0790
inventory_procedure_defaults = _rb_install_0790(_rb_bind(_rb_sys.modules[__name__], ('InventoryItem', 'Procedure', 'ProcedureConsumableDefault', 'app', 'db', 'flash', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))



# RB_BIND:0791
from web.routes.inventory import install_0791 as _rb_install_0791
inventory_transfers = _rb_install_0791(_rb_bind(_rb_sys.modules[__name__], ('InventoryItem', 'StockTransfer', 'StockTransferLine', 'Warehouse', 'abort', 'app', 'audit', 'current_user', 'db', 'flash', 'has_perm', 'next_transfer_no', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0792
from web.routes.inventory import install_0792 as _rb_install_0792
inventory_transfer_detail = _rb_install_0792(_rb_bind(_rb_sys.modules[__name__], ('StockTransfer', 'StockTransferLine', 'app', 'render_template', 'require_perm')))

# RB_BIND:0793
from web.routes.inventory import install_0793 as _rb_install_0793
inventory_transfer_approve = _rb_install_0793(_rb_bind(_rb_sys.modules[__name__], ('StockTransfer', 'app', 'audit', 'current_user', 'datetime', 'db', 'flash', 'redirect', 'require_perm', 'url_for')))

# RB_BIND:0794
from web.routes.inventory import install_0794 as _rb_install_0794
inventory_transfer_reject = _rb_install_0794(_rb_bind(_rb_sys.modules[__name__], ('StockTransfer', 'app', 'audit', 'current_user', 'datetime', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0795
from web.routes.inventory import install_0795 as _rb_install_0795
inventory_transfer_dispatch = _rb_install_0795(_rb_bind(_rb_sys.modules[__name__], ('StockTransfer', 'StockTransferLine', 'app', 'audit', 'current_user', 'datetime', 'db', 'flash', 'inventory_available_qty', 'inventory_move', 'redirect', 'require_perm', 'stock_lots', 'url_for')))

# RB_BIND:0796
from web.routes.inventory import install_0796 as _rb_install_0796
inventory_transfer_receive = _rb_install_0796(_rb_bind(_rb_sys.modules[__name__], ('InventoryLot', 'StockTransfer', 'StockTransferLine', 'app', 'audit', 'current_user', 'datetime', 'db', 'flash', 'inventory_move', 'redirect', 'require_perm', 'url_for')))


# RB_BIND:0797
from web.routes.inventory import install_0797 as _rb_install_0797
order_consumables_available_stock = _rb_install_0797(_rb_bind(_rb_sys.modules[__name__], ('InventoryItem', 'Order', 'app', 'default_warehouse_for_order', 'inventory_available_qty', 'jsonify', 'require_perm')))

# RB_BIND:0798
from web.routes.inventory import install_0798 as _rb_install_0798
order_consumables = _rb_install_0798(_rb_bind(_rb_sys.modules[__name__], ('InventoryItem', 'Order', 'OrderConsumableUsage', 'abort', 'app', 'audit', 'consume_inventory', 'db', 'default_warehouse_for_order', 'flash', 'has_perm', 'order_consumable_cost', 'order_default_consumables', 'redirect', 'render_template', 'request', 'require_perm', 'return_inventory', 'url_for')))

# RB_BIND:0799
from web.routes.inventory import install_0799 as _rb_install_0799
order_consumables_apply_defaults = _rb_install_0799(_rb_bind(_rb_sys.modules[__name__], ('Order', 'app', 'audit', 'consume_inventory', 'db', 'default_warehouse_for_order', 'flash', 'order_default_consumables', 'redirect', 'require_perm', 'url_for')))

# RB_BIND:0800
from web.routes.reporting import install_0800 as _rb_install_0800
inventory_ledger = _rb_install_0800(_rb_bind(_rb_sys.modules[__name__], ('InventoryItem', 'Warehouse', 'WarehouseStockMovement', 'app', 'render_template', 'request', 'require_perm')))

# RB_BIND:0801
from web.routes.reporting import install_0801 as _rb_install_0801
inventory_consumption_report = _rb_install_0801(_rb_bind(_rb_sys.modules[__name__], ('OrderConsumableUsage', 'app', 'management_date_range', 'render_template', 'require_perm')))



# RB_BIND:0802
from web.routes.inventory import install_0802 as _rb_install_0802
inventory_warehouses_auto_create = _rb_install_0802(_rb_bind(_rb_sys.modules[__name__], ('app', 'audit', 'db', 'ensure_branch_warehouses', 'flash', 'redirect', 'require_perm', 'url_for')))

# RB_BIND:0803
from web.routes.inventory import install_0803 as _rb_install_0803
inventory_create_for_branch = _rb_install_0803(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Warehouse', '_warehouse_code_part', 'app', 'audit', 'db', 'ensure_main_warehouse', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))



# RB_BIND:0804
from web.routes.operations import install_0804 as _rb_install_0804
admin_data_management = _rb_install_0804(_rb_bind(_rb_sys.modules[__name__], ('Order', 'Patient', 'admin_backup_dir', 'admin_storage_paths', 'app', 'render_template', 'require_perm')))

# RB_BIND:0805
from web.routes.registration import install_0805 as _rb_install_0805
admin_export_patient = _rb_install_0805(_rb_bind(_rb_sys.modules[__name__], ('Attachment', 'Patient', 'Report', 'app', 'audit', 'csv', 'flash', 'io', 'json', 'os', 'patient_export_manifest', 'redirect', 'request', 'require_perm', 'safe_export_name', 'send_file', 'url_for', 'zipfile')))

# RB_BIND:0806
from web.routes.operations import install_0806 as _rb_install_0806
admin_create_manifest_backup = _rb_install_0806(_rb_bind(_rb_sys.modules[__name__], ('APP_VERSION', 'Order', 'Patient', 'admin_backup_dir', 'admin_storage_paths', 'app', 'audit', 'datetime', 'flash', 'hashlib', 'json', 'os', 'redirect', 'require_perm', 'url_for')))

# RB_BIND:0807
from web.routes.registration import install_0807 as _rb_install_0807
admin_anonymize_patient = _rb_install_0807(_rb_bind(_rb_sys.modules[__name__], ('Patient', 'app', 'audit', 'db', 'destructive_confirmation_token', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0808
from web.routes.registration import install_0808 as _rb_install_0808
admin_delete_patient = _rb_install_0808(_rb_bind(_rb_sys.modules[__name__], ('Order', 'Patient', 'app', 'audit', 'db', 'destructive_confirmation_token', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0809
from web.routes.operations import install_0809 as _rb_install_0809
admin_reset_module = _rb_install_0809(_rb_bind(_rb_sys.modules[__name__], ('abort', 'app', 'audit', 'db', 'destructive_confirmation_token', 'flash', 'NotificationEvent', 'redirect', 'request', 'require_perm', 'url_for')))



# RB_BIND:0810
from web.routes.registration import install_0810 as _rb_install_0810
admin_backup_all_patients = _rb_install_0810(_rb_bind(_rb_sys.modules[__name__], ('app', 'audit', 'create_all_patient_backup', 'current_user', 'db', 'flash', 'os', 'redirect', 'request', 'require_perm', 'url_for', 'verify_current_user_password')))

# RB_BIND:0811
from web.routes.registration import install_0811 as _rb_install_0811
admin_delete_all_patients = _rb_install_0811(_rb_bind(_rb_sys.modules[__name__], ('Order', 'Patient', 'app', 'audit', 'current_user', 'db', 'delete_all_patient_operational_data', 'flash', 'redirect', 'request', 'require_perm', 'url_for', 'verify_current_user_password')))



# RB_BIND:0812
from web.routes.shared import install_0812 as _rb_install_0812
executive_analytics = _rb_install_0812(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Contract', 'Device', 'Procedure', 'User', 'app', 'db', 'executive_analytics_filters', 'executive_analytics_payload', 'render_template', 'require_perm')))



# RB_BIND:0813
from web.routes.shared import install_0813 as _rb_install_0813
workflow_command_center = _rb_install_0813(_rb_bind(_rb_sys.modules[__name__], ('app', 'render_template', 'require_perm', 'workflow_command_snapshot')))

# RB_BIND:0814
from web.routes.imaging import install_0814 as _rb_install_0814
pacs_unregistered_cases = _rb_install_0814(_rb_bind(_rb_sys.modules[__name__], ('PacsStudyInbox', 'app', 'render_template', 'require_perm')))

# RB_BIND:0815
from web.routes.imaging import install_0815 as _rb_install_0815
pacs_unregistered_run = _rb_install_0815(_rb_bind(_rb_sys.modules[__name__], ('app', 'flash', 'pacs_auto_intake_once', 'redirect', 'require_perm', 'url_for')))

# RB_BIND:0816
from web.routes.imaging import install_0816 as _rb_install_0816
pacs_destinations = _rb_install_0816(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'DicomDestination', 'DiscoveredDicomNode', 'abort', 'app', 'audit', 'db', 'flash', 're', 'redirect', 'render_template', 'request', 'require_perm', 'sync_destination_to_orthanc', 'url_for')))

# RB_BIND:0817
from web.routes.imaging import install_0817 as _rb_install_0817
pacs_routing_rules = _rb_install_0817(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'DicomDestination', 'PacsRoutingRule', 'app', 'db', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))

# RB_BIND:0818
from web.routes.imaging import install_0818 as _rb_install_0818
pacs_case_actions = _rb_install_0818(_rb_bind(_rb_sys.modules[__name__], ('DicomDestination', 'Order', 'PacsTransferJob', 'app', 'current_user', 'or_', 'render_template', 'require_perm')))

# RB_BIND:0819
from web.routes.imaging import install_0819 as _rb_install_0819
pacs_case_send = _rb_install_0819(_rb_bind(_rb_sys.modules[__name__], ('DicomDestination', 'Order', 'abort', 'app', 'audit', 'current_user', 'db', 'enqueue_transfer', 'flash', 'pacs_lookup_order', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0820
from web.routes.imaging import install_0820 as _rb_install_0820
pacs_transfer_retry = _rb_install_0820(_rb_bind(_rb_sys.modules[__name__], ('PacsTransferJob', 'abort', 'app', 'audit', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0821
from web.routes.imaging import install_0821 as _rb_install_0821
pacs_transfer_cancel = _rb_install_0821(_rb_bind(_rb_sys.modules[__name__], ('PacsTransferJob', 'abort', 'app', 'audit', 'datetime', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0822
from web.routes.imaging import install_0822 as _rb_install_0822
pacs_destination_echo = _rb_install_0822(_rb_bind(_rb_sys.modules[__name__], ('DicomDestination', 'abort', 'app', 'audit', 'current_user', 'db', 'flash', 'redirect', 'request', 'require_perm', 'test_dicom_destination', 'url_for')))

# RB_BIND:0823
from web.routes.imaging import install_0823 as _rb_install_0823
pacs_transfer_process = _rb_install_0823(_rb_bind(_rb_sys.modules[__name__], ('app', 'flash', 'process_pacs_transfer_jobs', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0824
from web.routes.imaging import install_0824 as _rb_install_0824
pacs_archive_dashboard = _rb_install_0824(_rb_bind(_rb_sys.modules[__name__], ('DicomDestination', 'PacsTransferJob', 'app', 'render_template', 'require_perm')))

# RB_BIND:0825
from web.routes.reporting import install_0825 as _rb_install_0825
coordinator_queue = _rb_install_0825(_rb_bind(_rb_sys.modules[__name__], ('CoordinatorTask', 'app', 'current_user', 'or_', 'render_template', 'require_perm')))

# RB_BIND:0826
from web.routes.reporting import install_0826 as _rb_install_0826
coordinator_action = _rb_install_0826(_rb_bind(_rb_sys.modules[__name__], ('CoordinatorTask', 'abort', 'app', 'audit', 'current_user', 'datetime', 'db', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0827
from application.components.operations import install_0827 as _rb_install_0827
resolve_storage_target = _rb_install_0827(_rb_bind(_rb_sys.modules[__name__], ('StorageLocation',)))

# RB_STORAGE_BLUEPRINT_V1
from storage_routes import register_storage_routes

# Preserve the app module's callable and late lookup of the resolver.
api_storage_resolve = register_storage_routes(
    app,
    require_perm,
    lambda *args, **kwargs: resolve_storage_target(*args, **kwargs),
)

# RB_BIND:0828
from application.components.operations import install_0828 as _rb_install_0828
refresh_storage_health = _rb_install_0828(_rb_bind(_rb_sys.modules[__name__], ('datetime', 'os', 'shutil')))

# RB_BIND:0829
from web.routes.operations import install_0829 as _rb_install_0829
storage_dr_manager = _rb_install_0829(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'FailoverEvent', 'ReplicationNode', 'ReplicationPolicy', 'ReplicationState', 'StorageLocation', 'abort', 'app', 'audit', 'db', 'flash', 'has_perm', 'redirect', 'refresh_storage_health', 'render_template', 'request', 'require_perm', 'storage_defaults', 'url_for')))

# RB_BIND:0830
from web.routes.operations import install_0830 as _rb_install_0830
storage_target_health = _rb_install_0830(_rb_bind(_rb_sys.modules[__name__], ('StorageLocation', 'app', 'audit', 'db', 'redirect', 'refresh_storage_health', 'require_perm', 'url_for')))

# RB_BIND:0831
from web.routes.operations import install_0831 as _rb_install_0831
replication_state_update = _rb_install_0831(_rb_bind(_rb_sys.modules[__name__], ('ReplicationPolicy', 'ReplicationState', 'app', 'audit', 'datetime', 'db', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0832
from web.routes.operations import install_0832 as _rb_install_0832
dr_failover_request = _rb_install_0832(_rb_bind(_rb_sys.modules[__name__], ('FailoverEvent', 'ReplicationPolicy', 'abort', 'app', 'audit', 'current_user', 'db', 'flash', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0833
from web.routes.operations import install_0833 as _rb_install_0833
storage_cloud_manager = _rb_install_0833(_rb_bind(_rb_sys.modules[__name__], ('BackupRun', 'CloudBackupConfig', 'StorageLocation', 'abort', 'app', 'db', 'flash', 'has_perm', 'redirect', 'render_template', 'request', 'require_perm', 'storage_defaults', 'url_for')))

# RB_BIND:0834
from web.components.operations import install_0834 as _rb_install_0834
api_ops_health = _rb_install_0834(_rb_bind(_rb_sys.modules[__name__], ('PacsStudyInbox', 'PacsTransferJob', 'jsonify', 'login_required', 'workflow_command_snapshot')))



# RB_BIND:0835
from web.routes.imaging import install_0835 as _rb_install_0835
pacs_discovery = _rb_install_0835(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'DiscoveredDicomNode', 'PACS_ORTHANC_URL', '_orthanc_auth', 'app', 'dicom_ingress_policy_ready', 'modality_key', 'render_template', 'requests', 'require_perm')))


# RB_BIND:0836
from application.components.imaging import install_0836 as _rb_install_0836
_managed_dicom_modality = _rb_install_0836(_rb_bind(_rb_sys.modules[__name__], ('PACS_ORTHANC_URL', '_orthanc_auth', 'modality_key', 'requests')))


# RB_BIND:0837
from application.components.imaging import install_0837 as _rb_install_0837
_write_managed_dicom_modality = _rb_install_0837(_rb_bind(_rb_sys.modules[__name__], ('PACS_ORTHANC_URL', '_managed_dicom_modality', '_orthanc_auth', 'delete_modality', 'modality_key', 'requests')))


# RB_BIND:0838
from web.components.shared import install_0838 as _rb_install_0838
_apply_discovered_node_policy = _rb_install_0838(_rb_bind(_rb_sys.modules[__name__], ('PACS_ORTHANC_URL', '_managed_dicom_modality', '_orthanc_auth', '_write_managed_dicom_modality', 'abort', 'app', 'audit', 'current_user', 'datetime', 'db', 'dicom_ingress_policy_ready', 'modality_key', 'requests')))


# RB_BIND:0839
from application.components.reporting import install_0839 as _rb_install_0839
external_reporting_setting = _rb_install_0839(_rb_bind(_rb_sys.modules[__name__], ('ExternalReportingSetting', 'db')))


# RB_BIND:0840
from application.components.registration import install_0840 as _rb_install_0840
_external_intake_generated_accession = _rb_install_0840(_rb_bind(_rb_sys.modules[__name__], ('re',)))


# RB_BIND:0841
from application.components.registration import install_0841 as _rb_install_0841
_external_intake_local_patient_id = _rb_install_0841(_rb_bind(_rb_sys.modules[__name__], ('re',)))


# RB_BIND:0842
from application.components.registration import install_0842 as _rb_install_0842
external_intake_autofill = _rb_install_0842(_rb_bind(_rb_sys.modules[__name__], ('_external_intake_generated_accession', 'normalize_date')))


# RB_BIND:0843
from application.components.imaging import install_0843 as _rb_install_0843
capture_external_pacs_studies = _rb_install_0843(_rb_bind(_rb_sys.modules[__name__], ('ExternalIntakeCase', 'ExternalReportingCenter', 'PacsStudyInbox', '_orthanc_recent_studies', '_study_tags_extended', 'datetime', 'db', 'external_intake_autofill', 'external_reporting_setting')))


# RB_BIND:0844
from application.components.imaging import install_0844 as _rb_install_0844
external_intake_dicom_demographics = _rb_install_0844(_rb_bind(_rb_sys.modules[__name__], ('PACS_ORTHANC_URL', '_orthanc_auth', 'normalize_date', 'requests')))


# RB_BIND:0845
from application.components.registration import install_0845 as _rb_install_0845
external_intake_demographic_verification = _rb_install_0845(_rb_bind(_rb_sys.modules[__name__], ('datetime', 'external_intake_dicom_demographics', 'hashlib', 'json', 'normalize_date', 'normalize_person_name')))


# RB_BIND:0846
from web.components.registration import install_0846 as _rb_install_0846
queue_external_intake = _rb_install_0846(_rb_bind(_rb_sys.modules[__name__], ('Attachment', 'AuditLog', 'Branch', 'ExternalCenterDoctor', 'Order', 'Patient', '_external_intake_generated_accession', '_external_intake_local_patient_id', 'current_user', 'datetime', 'db', 'external_intake_demographic_verification', 'external_reporting_setting', 'normalize_date', 'offer_next_doctor', 'request')))


# RB_BIND:0847
from web.routes.assignment import install_0847 as _rb_install_0847
admin_external_intake = _rb_install_0847(_rb_bind(_rb_sys.modules[__name__], ('Attachment', 'Contract', 'ExternalCenterDoctor', 'ExternalIntakeCase', 'ExternalReportingCenter', 'Role', 'User', 'abort', 'app', 'audit', 'capture_external_pacs_studies', 'current_user', 'db', 'external_reporting_setting', 'json', 're', 'redirect', 'render_template', 'request', 'require_perm', 'url_for')))


# RB_BIND:0848
from web.routes.assignment import install_0848 as _rb_install_0848
admin_external_intake_attachment_upload = _rb_install_0848(_rb_bind(_rb_sys.modules[__name__], ('ATTACHMENT_DIR', 'Attachment', 'ExternalIntakeCase', 'abort', 'app', 'audit', 'current_user', 'db', 'flash', 'generated_attachment_filename', 'os', 'redirect', 'request', 'require_perm', 'url_for')))


# RB_BIND:0849
from web.routes.assignment import install_0849 as _rb_install_0849
admin_external_intake_decision = _rb_install_0849(_rb_bind(_rb_sys.modules[__name__], ('ExternalIntakeCase', 'abort', 'app', 'audit', 'current_user', 'datetime', 'db', 'external_intake_demographic_verification', 'flash', 'normalize_date', 'queue_external_intake', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0850
from web.routes.imaging import install_0850 as _rb_install_0850
pacs_discovery_approve = _rb_install_0850(_rb_bind(_rb_sys.modules[__name__], ('DiscoveredDicomNode', '_apply_discovered_node_policy', 'abort', 'app', 'flash', 'node_payload', 'redirect', 'request', 'require_perm', 'url_for')))

# RB_BIND:0851
from web.routes.imaging import install_0851 as _rb_install_0851
pacs_discovery_block = _rb_install_0851(_rb_bind(_rb_sys.modules[__name__], ('DiscoveredDicomNode', '_apply_discovered_node_policy', 'app', 'flash', 'redirect', 'require_perm', 'url_for')))

# RB_BIND:0852
from web.routes.imaging import install_0852 as _rb_install_0852
pacs_discovery_legacy_endpoint = _rb_install_0852(_rb_bind(_rb_sys.modules[__name__], ('DiscoveredDicomNode', 'PACS_ORTHANC_URL', '_orthanc_auth', 'abort', 'app', 'audit', 'db', 'flash', 'legacy_endpoint_matches', 'legacy_endpoint_payload', 'modality_key', 're', 'redirect', 'request', 'requests', 'require_perm', 'url_for')))

# RB_BIND:0853
from web.routes.imaging import install_0853 as _rb_install_0853
pacs_discovery_dismiss = _rb_install_0853(_rb_bind(_rb_sys.modules[__name__], ('DiscoveredDicomNode', '_managed_dicom_modality', 'abort', 'app', 'audit', 'db', 'flash', 'redirect', 'require_perm', 'url_for')))

# RB_BIND:0854
from web.routes.imaging import install_0854 as _rb_install_0854
pacs_discovery_manual = _rb_install_0854(_rb_bind(_rb_sys.modules[__name__], ('DiscoveredDicomNode', 'abort', 'app', 'datetime', 'db', 'node_payload', 'redirect', 'request', 'require_perm', 'url_for')))



# RB_BIND:0855
from application.components.assignment import install_0855 as _rb_install_0855
assignment_visit_summary = _rb_install_0855(_rb_bind(_rb_sys.modules[__name__], ()))


# RB_BIND:0856
from web.routes.assignment import install_0856 as _rb_install_0856
api_assignment_offer_pending = _rb_install_0856(_rb_bind(_rb_sys.modules[__name__], ('AssignmentOffer', 'app', 'assignment_visit_summary', 'case_pacs_url', 'current_user', 'expire_assignment_offers', 'jsonify', 'login_required', 'url_for')))

if os.getenv('RADBRIDGE_BOOTSTRAP_ON_IMPORT','0') == '1':
    with app.app_context():
        db.create_all()
        from runtime_feature_schema import ensure_runtime_feature_schema
        ensure_runtime_feature_schema(db, app.logger)
        seed_quality_defaults()
        seed_assignment_offer_policy()
        seed_communication_defaults()
        seed_v521_functional_defaults()
        migrate_existing_schema()
        seed_master_data()
        try:
            ensure_default_procedure_branch_mappings()
            seed_dictionary_and_ribbon()
            seed_full_ui_dictionary()
        except Exception as e:
            db.session.rollback()
            app.logger.warning(
                'v5.18 initialization skipped: %s',
                e
            )


# ============================================================
# OHIF Integrated Reporting API - TEST
# ============================================================


OHIF_REPORTING_TOKEN = os.getenv("OHIF_REPORTING_TOKEN", "").strip()

# RB_BIND:0857
from web.components.imaging import install_0857 as _rb_install_0857
require_ohif_token = _rb_install_0857(_rb_bind(_rb_sys.modules[__name__], ('OHIF_REPORTING_TOKEN', 'jsonify', 'request')))




# RB_BIND:0858
from application.components.shared import install_0858 as _rb_install_0858
_normalize_case_uid = _rb_install_0858(_rb_bind(_rb_sys.modules[__name__], ()))


# RB_BIND:0859
from web.components.imaging import install_0859 as _rb_install_0859
validate_ohif_case_context = _rb_install_0859(_rb_bind(_rb_sys.modules[__name__], ('_normalize_case_uid', 'ohif_study_uid', 'request')))


# RB_BIND:0860
from web.components.imaging import install_0860 as _rb_install_0860
require_valid_ohif_case_context = _rb_install_0860(_rb_bind(_rb_sys.modules[__name__], ('audit', 'db', 'jsonify', 'validate_ohif_case_context')))



# RB_BIND:0861
from application.components.reporting import install_0861 as _rb_install_0861
_ohif_report_json = _rb_install_0861(_rb_bind(_rb_sys.modules[__name__], ('allowed_report_transitions', 'ohif_study_uid')))


# RB_BIND:0862
from web.components.imaging import install_0862 as _rb_install_0862
_ohif_json_body = _rb_install_0862(_rb_bind(_rb_sys.modules[__name__], ('request',)))


# RB_BIND:0863
from web.routes.reporting import install_0863 as _rb_install_0863
api_ohif_report_get = _rb_install_0863(_rb_bind(_rb_sys.modules[__name__], ('AuditLog', 'Order', '_ohif_report_json', 'app', 'audit_as_user', 'datetime', 'db', 'get_report', 'jsonify', 'require_ohif_doctor_user', 'require_ohif_token', 'require_valid_ohif_case_context', 'timedelta')))


# RB_BIND:0864
from web.components.reporting import install_0864 as _rb_install_0864
api_ohif_reporting_agent_download = _rb_install_0864(_rb_bind(_rb_sys.modules[__name__], ('__file__', 'jsonify', 'os', 'require_ohif_token', 'send_file')))


# RB_BIND:0865
from web.routes.reporting import install_0865 as _rb_install_0865
api_ohif_report_word = _rb_install_0865(_rb_bind(_rb_sys.modules[__name__], ('Document', 'Inches', 'Order', '_next_report_version', '_rb_authoritative_docx_download', 'app', 'build_report_filename', 'get_report', 'io', 'require_ohif_doctor_user', 'require_ohif_token', 'require_valid_ohif_case_context', 'send_file')))


# RB_BIND:0866
from application.components.reporting import install_0866 as _rb_install_0866
_ohif_validate_word_upload = _rb_install_0866(_rb_bind(_rb_sys.modules[__name__], ('Document', 'audit', 'db', 'os', 're', 'tempfile')))


# RB_BIND:0867
from web.routes.reporting import install_0867 as _rb_install_0867
api_ohif_report_word_draft = _rb_install_0867(_rb_bind(_rb_sys.modules[__name__], ('Order', 'REPORT_DIR', 'Report', 'ReportEditControlError', 'ReportVersion', '_ohif_report_json', '_ohif_validate_word_upload', 'app', 'assert_lease_owner', 'assert_not_final', 'assert_role_allowed', 'audit', 'audit_as_user', 'db', 'extract_docx_sections', 'get_report', 'jsonify', 'request', 'require_ohif_doctor_user', 'require_ohif_token', 'require_valid_ohif_case_context')))



# RB_BIND:0868
from web.routes.reporting import install_0868 as _rb_install_0868
api_ohif_report_word_verify = _rb_install_0868(_rb_bind(_rb_sys.modules[__name__], ('Order', 'Report', 'ReportEditControlError', 'ReportVersion', '_rb_latest_report_version', 'app', 'assert_no_foreign_active_lease', 'assert_not_final', 'assert_role_allowed', 'audit', 'audit_as_user', 'db', 'get_report', 'jsonify', 'pacs_demographic_gate', 'request', 'require_ohif_doctor_user', 'require_ohif_token', 'require_valid_ohif_case_context')))


# RB_BIND:0869
from web.routes.reporting import install_0869 as _rb_install_0869
api_ohif_report_word_finalize = _rb_install_0869(_rb_bind(_rb_sys.modules[__name__], ('Order', 'Report', 'ReportEditControlError', 'ReportVersion', '_rb_latest_report_version', 'app', 'assert_no_foreign_active_lease', 'assert_not_final', 'assert_role_allowed', 'audit', 'audit_as_user', 'db', 'enqueue_event', 'get_report', 'jsonify', 'pacs_demographic_gate', 'queue_final_delivery', 'request', 'require_ohif_doctor_user', 'require_ohif_token', 'require_valid_ohif_case_context')))



# RB_BIND:0870
from web.routes.reporting import install_0870 as _rb_install_0870
api_ohif_report_lease_acquire = _rb_install_0870(_rb_bind(_rb_sys.modules[__name__], ('Order', 'ReportEditControlError', 'acquire_or_renew_lease', 'app', 'audit_as_user', 'db', 'get_report', 'jsonify', 'request', 'require_ohif_doctor_user', 'require_ohif_token', 'require_valid_ohif_case_context')))



# RB_BIND:0871
from web.routes.reporting import install_0871 as _rb_install_0871
api_ohif_report_lease_release = _rb_install_0871(_rb_bind(_rb_sys.modules[__name__], ('Order', 'app', 'audit_as_user', 'db', 'jsonify', 'release_lease', 'request', 'require_ohif_doctor_user', 'require_ohif_token', 'require_valid_ohif_case_context')))


# RB_BIND:0872
from web.routes.reporting import install_0872 as _rb_install_0872
api_ohif_report_draft = _rb_install_0872(_rb_bind(_rb_sys.modules[__name__], ('Order', 'REPORT_DIR', 'Report', 'ReportEditControlError', 'ReportVersion', '_ohif_report_json', 'app', 'assert_lease_owner', 'assert_not_final', 'assert_role_allowed', 'audit', 'db', 'get_report', 'jsonify', 'make_docx', 'request', 'require_ohif_doctor_user', 'require_ohif_token', 'require_valid_ohif_case_context')))



# RB_BIND:0873
from web.routes.reporting import install_0873 as _rb_install_0873
api_ohif_report_verify = _rb_install_0873(_rb_bind(_rb_sys.modules[__name__], ('Order', 'Report', 'ReportEditControlError', 'ReportVersion', '_rb_latest_report_version', 'app', 'assert_no_foreign_active_lease', 'assert_not_final', 'assert_role_allowed', 'audit', 'audit_as_user', 'db', 'get_report', 'jsonify', 'pacs_demographic_gate', 'request', 'require_ohif_doctor_user', 'require_ohif_token', 'require_valid_ohif_case_context')))


# RB_BIND:0874
from web.routes.reporting import install_0874 as _rb_install_0874
api_ohif_report_finalize = _rb_install_0874(_rb_bind(_rb_sys.modules[__name__], ('Order', 'Report', 'ReportEditControlError', 'ReportVersion', '_rb_latest_report_version', 'app', 'assert_no_foreign_active_lease', 'assert_not_final', 'assert_role_allowed', 'audit', 'audit_as_user', 'db', 'enqueue_event', 'get_report', 'jsonify', 'pacs_demographic_gate', 'queue_final_delivery', 'request', 'require_ohif_doctor_user', 'require_ohif_token', 'require_valid_ohif_case_context')))


# RB_BIND:0875
from web.routes.reporting import install_0875 as _rb_install_0875
api_ohif_report_return = _rb_install_0875(_rb_bind(_rb_sys.modules[__name__], ('CaseAssignment', 'CaseReturnReason', 'CaseWorkflowEvent', 'Order', '__name__', '_ohif_json_body', '_reporting_support_sys', 'abort', 'app', 'audit_as_user', 'create_ris_notification', 'datetime', 'db', 'jsonify', 'require_ohif_doctor_user', 'require_ohif_token', 'require_valid_ohif_case_context')))


# RB_BIND:0876
from web.components.reporting import install_0876 as _rb_install_0876
api_ohif_reporting_health = _rb_install_0876(_rb_bind(_rb_sys.modules[__name__], ('jsonify',)))



# ============================================================
# RADBRIDGE ONLYOFFICE INTEGRATION - PHASE 1
# ============================================================

ONLYOFFICE_INTERNAL_URL = (
    os.getenv('ONLYOFFICE_INTERNAL_URL', 'http://radbridge-onlyoffice')
    .rstrip('/')
)

ONLYOFFICE_RIS_INTERNAL_URL = (
    os.getenv('ONLYOFFICE_RIS_INTERNAL_URL', 'http://ris-web:8000')
    .rstrip('/')
)

ONLYOFFICE_JWT_SECRET = (
    os.getenv('ONLYOFFICE_JWT_SECRET', '').strip()
)

ONLYOFFICE_LINK_TTL = int(
    os.getenv('ONLYOFFICE_LINK_TTL', '7200') or 7200
)


ONLYOFFICE_KEY_REV = (
    os.getenv('ONLYOFFICE_KEY_REV', '1')
    .strip()
    or '1'
)


# RB_BIND:0877
from application.components.reporting import install_0877 as _rb_install_0877
_onlyoffice_serializer = _rb_install_0877(_rb_bind(_rb_sys.modules[__name__], ('URLSafeTimedSerializer', 'app')))


# RB_BIND:0878
from application.components.reporting import install_0878 as _rb_install_0878
_onlyoffice_link_token = _rb_install_0878(_rb_bind(_rb_sys.modules[__name__], ('_onlyoffice_serializer',)))


# RB_BIND:0879
from application.components.reporting import install_0879 as _rb_install_0879
_onlyoffice_check_link_token = _rb_install_0879(_rb_bind(_rb_sys.modules[__name__], ('ONLYOFFICE_LINK_TTL', '_onlyoffice_serializer')))


# RB_BIND:0880
from application.components.reporting import install_0880 as _rb_install_0880
_onlyoffice_b64url = _rb_install_0880(_rb_bind(_rb_sys.modules[__name__], ()))


# RB_BIND:0881
from application.components.reporting import install_0881 as _rb_install_0881
_onlyoffice_jwt_encode = _rb_install_0881(_rb_bind(_rb_sys.modules[__name__], ('ONLYOFFICE_JWT_SECRET', '_onlyoffice_b64url', 'hmac')))


# RB_BIND:0882
from application.components.reporting import install_0882 as _rb_install_0882
_onlyoffice_jwt_valid = _rb_install_0882(_rb_bind(_rb_sys.modules[__name__], ('ONLYOFFICE_JWT_SECRET', 'hmac')))


# RB_BIND:0883
from application.components.reporting import install_0883 as _rb_install_0883
_onlyoffice_working_dir = _rb_install_0883(_rb_bind(_rb_sys.modules[__name__], ('REPORT_DIR', 'os')))


# RB_BIND:0884
from application.components.reporting import install_0884 as _rb_install_0884
_onlyoffice_working_path = _rb_install_0884(_rb_bind(_rb_sys.modules[__name__], ('_onlyoffice_working_dir', 'os', 're')))


# RB_BIND:0885
from application.components.reporting import install_0885 as _rb_install_0885
_onlyoffice_prepare_document = _rb_install_0885(_rb_bind(_rb_sys.modules[__name__], ('ReportVersion', '_onlyoffice_working_path', 'make_docx', 'os')))


# RB_BIND:0886
from web.routes.reporting import install_0886 as _rb_install_0886
api_ohif_reporting_preferences = _rb_install_0886(_rb_bind(_rb_sys.modules[__name__], ('Order', 'app', 'audit_as_user', 'db', 'get_user_profile', 'jsonify', 'request', 'require_ohif_doctor_user', 'require_ohif_token', 'require_valid_ohif_case_context', 'set_user_reporting_preferences', 'user_reporting_preferences')))



# RB_BIND:0887
from web.components.reporting import install_0887 as _rb_install_0887
api_ohif_onlyoffice_health = _rb_install_0887(_rb_bind(_rb_sys.modules[__name__], ('ONLYOFFICE_INTERNAL_URL', 'ONLYOFFICE_JWT_SECRET', 'jsonify', 'requests')))


# RB_BIND:0888
from web.routes.reporting import install_0888 as _rb_install_0888
api_ohif_onlyoffice_config = _rb_install_0888(_rb_bind(_rb_sys.modules[__name__], ('ONLYOFFICE_KEY_REV', 'ONLYOFFICE_RIS_INTERNAL_URL', 'Order', 'ReportEditControlError', '__name__', '_next_report_version', '_onlyoffice_jwt_encode', '_onlyoffice_link_token', '_rb_authoritative_docx_download', '_rb_latest_report_version', 'app', 'assert_lease_owner', 'assert_role_allowed', 'audit_as_user', 'build_report_filename', 'db', 'get_report', 'jsonify', 'os', 'quote', 'request', 'require_ohif_doctor_user', 'require_ohif_token', 'require_valid_ohif_case_context')))



# RB_BIND:0889
from application.components.reporting import install_0889 as _rb_install_0889
_onlyoffice_forcesave_ack_path = _rb_install_0889(_rb_bind(_rb_sys.modules[__name__], ('_onlyoffice_working_dir', 'os')))


# RB_BIND:0890
from application.components.reporting import install_0890 as _rb_install_0890
_onlyoffice_write_forcesave_ack = _rb_install_0890(_rb_bind(_rb_sys.modules[__name__], ('_onlyoffice_forcesave_ack_path', 'os')))


# RB_BIND:0891
from application.components.reporting import install_0891 as _rb_install_0891
_onlyoffice_read_forcesave_ack = _rb_install_0891(_rb_bind(_rb_sys.modules[__name__], ('_onlyoffice_forcesave_ack_path',)))


# RB_BIND:0892
from web.routes.reporting import install_0892 as _rb_install_0892
api_ohif_onlyoffice_forcesave = _rb_install_0892(_rb_bind(_rb_sys.modules[__name__], ('ONLYOFFICE_INTERNAL_URL', 'Order', 'ReportEditControlError', '_onlyoffice_jwt_encode', '_onlyoffice_read_forcesave_ack', '_onlyoffice_working_path', 'app', 'assert_lease_owner', 'assert_not_final', 'assert_role_allowed', 'db', 'get_report', 'jsonify', 'os', 'request', 'requests', 'require_ohif_doctor_user', 'require_ohif_token', 'require_valid_ohif_case_context')))


# RB_BIND:0893
from web.routes.reporting import install_0893 as _rb_install_0893
api_ohif_onlyoffice_document = _rb_install_0893(_rb_bind(_rb_sys.modules[__name__], ('Order', '_onlyoffice_check_link_token', '_onlyoffice_prepare_document', '_rb_authoritative_docx_download', 'app', 'get_report', 'jsonify', 'os', 'request', 'send_file')))




# ============================================================
# RadBridge Unified Clinical Report Versioning
# ONLYOFFICE explicit commit layer.
#
# IMPORTANT:
# ONLYOFFICE callback remains WORKING COPY ONLY.
# Clinical versions are created only by explicit doctor action.
# ============================================================

# RB_BIND:0894
from application.components.reporting import install_0894 as _rb_install_0894
_rb_latest_report_version = _rb_install_0894(_rb_bind(_rb_sys.modules[__name__], ('ReportVersion',)))


# RB_BIND:0895
from application.components.shared import install_0895 as _rb_install_0895
_rb_authoritative_docx_download = _rb_install_0895(_rb_bind(_rb_sys.modules[__name__], ('_rb_final_revision_document', '_rb_latest_report_version', 'os')))


# RB_BIND:0896
from application.components.shared import install_0896 as _rb_install_0896
_rb_latest_authoritative_version = _rb_install_0896(_rb_bind(_rb_sys.modules[__name__], ('ReportVersion',)))


# RB_BIND:0897
from application.components.reporting import install_0897 as _rb_install_0897
_rb_onlyoffice_working_bytes = _rb_install_0897(_rb_bind(_rb_sys.modules[__name__], ('_onlyoffice_working_path', 'os')))


# RB_BIND:0898
from application.components.reporting import install_0898 as _rb_install_0898
_rb_commit_onlyoffice_version = _rb_install_0898(_rb_bind(_rb_sys.modules[__name__], ('REPORT_DIR', 'Report', 'ReportVersion', '_rb_latest_report_version', '_rb_onlyoffice_working_bytes', 'audit_as_user', 'db', 'extract_docx_sections', 'get_report', 'os', 'validate_report_patient_file')))


# RB_BIND:0899
from web.routes.reporting import install_0899 as _rb_install_0899
api_ohif_onlyoffice_commit_draft = _rb_install_0899(_rb_bind(_rb_sys.modules[__name__], ('Order', 'ReportEditControlError', '_rb_commit_onlyoffice_version', 'app', 'assert_lease_owner', 'assert_not_final', 'assert_role_allowed', 'audit', 'db', 'get_report', 'jsonify', 'request', 'require_ohif_doctor_user', 'require_ohif_token', 'require_valid_ohif_case_context')))


# RB_BIND:0900
from web.routes.reporting import install_0900 as _rb_install_0900
api_ohif_onlyoffice_verify = _rb_install_0900(_rb_bind(_rb_sys.modules[__name__], ('Order', 'ReportEditControlError', '_rb_commit_onlyoffice_version', 'app', 'assert_no_foreign_active_lease', 'assert_not_final', 'assert_role_allowed', 'audit', 'db', 'get_report', 'jsonify', 'pacs_demographic_gate', 'request', 'require_ohif_doctor_user', 'require_ohif_token', 'require_valid_ohif_case_context')))


# RB_BIND:0901
from web.routes.reporting import install_0901 as _rb_install_0901
api_ohif_onlyoffice_finalize = _rb_install_0901(_rb_bind(_rb_sys.modules[__name__], ('Order', 'Report', 'ReportEditControlError', 'ReportVersion', 'app', 'assert_no_foreign_active_lease', 'assert_not_final', 'assert_role_allowed', 'audit', 'db', 'get_report', 'jsonify', 'pacs_demographic_gate', 'require_ohif_doctor_user', 'require_ohif_token', 'require_valid_ohif_case_context')))

# RB_BIND:0902
from web.routes.reporting import install_0902 as _rb_install_0902
api_ohif_onlyoffice_callback = _rb_install_0902(_rb_bind(_rb_sys.modules[__name__], ('ONLYOFFICE_INTERNAL_URL', 'Order', '_ohif_validate_word_upload', '_onlyoffice_check_link_token', '_onlyoffice_jwt_valid', '_onlyoffice_working_path', '_onlyoffice_write_forcesave_ack', 'app', 'audit', 'db', 'jsonify', 'os', 'request', 'requests', 'shutil')))


# ============================================================
# END RADBRIDGE ONLYOFFICE INTEGRATION - PHASE 1
# ============================================================


# ============================================================
# End OHIF Integrated Reporting API - TEST
# ============================================================



# ==========================================================
# RadBridge v5.37 - AI Voice Reception API
# ==========================================================
try:
    import sys as _sys
    from voice_ai import register_voice_ai

    register_voice_ai(_sys.modules[__name__])

    app.logger.info(
        "RadBridge Voice AI API registered"
    )

except Exception as _voice_error:
    app.logger.exception(
        "VOICE_AI_REGISTER_ERROR: %s",
        _voice_error
    )


# ===== RADBRIDGE VOICE BOOKING API V5.37 =====

# RB_BIND:0903
from web.components.communications import install_0903 as _rb_install_0903
_voice_api_authorized = _rb_install_0903(_rb_bind(_rb_sys.modules[__name__], ('os', 'request')))


# RB_BIND:0904
from web.components.communications import install_0904 as _rb_install_0904
_voice_api_unauthorized = _rb_install_0904(_rb_bind(_rb_sys.modules[__name__], ('jsonify',)))


# RB_BIND:0905
from web.components.communications import install_0905 as _rb_install_0905
voice_api_branches = _rb_install_0905(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Device', '_voice_api_authorized', '_voice_api_unauthorized', 'jsonify')))


# RB_BIND:0906
from web.components.communications import install_0906 as _rb_install_0906
voice_api_procedures = _rb_install_0906(_rb_bind(_rb_sys.modules[__name__], ('Procedure', '_voice_api_authorized', '_voice_api_unauthorized', 'jsonify', 'procedure_available_in_branch', 'request')))


# RB_BIND:0907
from web.routes.communications import install_0907 as _rb_install_0907
voice_api_patient_search = _rb_install_0907(_rb_bind(_rb_sys.modules[__name__], ('_voice_api_authorized', '_voice_api_unauthorized', 'app', 'jsonify', 'patient_by_phone', 'request')))


# RB_BIND:0908
from web.routes.communications import install_0908 as _rb_install_0908
voice_api_patient_create = _rb_install_0908(_rb_bind(_rb_sys.modules[__name__], ('Patient', '_voice_api_authorized', '_voice_api_unauthorized', 'app', 'arabic_to_latin_name', 'audit', 'db', 'generate_patient_id', 'jsonify', 'patient_by_phone', 'request', 'secrets')))



# ============================================================
# Voice Patient / Exam / Report Inquiry - v5.40.3
# READ ONLY
# ============================================================

# RB_BIND:0909
from web.routes.communications import install_0909 as _rb_install_0909
voice_api_inquiry = _rb_install_0909(_rb_bind(_rb_sys.modules[__name__], ('Order', 'Report', '_voice_api_authorized', '_voice_api_unauthorized', 'app', 'jsonify', 'patient_by_phone', 'request')))


# RB_BIND:0910
from web.components.communications import install_0910 as _rb_install_0910
voice_api_slots = _rb_install_0910(_rb_bind(_rb_sys.modules[__name__], ('Branch', 'Device', 'Procedure', '_voice_api_authorized', '_voice_api_unauthorized', 'booking_restriction_reason', 'booking_v2_slots', 'datetime', 'db', 'jsonify', 'procedure_available_in_branch', 'request')))


# RB_BIND:0911
from web.routes.communications import install_0911 as _rb_install_0911
voice_api_book = _rb_install_0911(_rb_bind(_rb_sys.modules[__name__], ('Appointment', 'Branch', 'Device', 'Patient', 'Procedure', '_voice_api_authorized', '_voice_api_unauthorized', 'app', 'audit', 'booking_restriction_reason', 'datetime', 'db', 'jsonify', 'request')))


# ===== END RADBRIDGE VOICE BOOKING API V5.37 =====



# ============================================================
# RadBridge v5.40 - Reception Voice Intake
# ============================================================

VOICE_INTAKE_STORAGE = os.environ.get(
    "VOICE_INTAKE_STORAGE",
    "/data/voice-intake"
)


# RB_BIND:0912
from web.routes.communications import install_0912 as _rb_install_0912
voice_intake_review = _rb_install_0912(_rb_bind(_rb_sys.modules[__name__], ('app', 'db', 'login_required', 'render_template', 'request')))


# RB_BIND:0913
from web.routes.communications import install_0913 as _rb_install_0913
voice_intake_audio = _rb_install_0913(_rb_bind(_rb_sys.modules[__name__], ('VOICE_INTAKE_STORAGE', 'abort', 'app', 'db', 'login_required', 'os', 'send_file')))


# RB_BIND:0914
from web.routes.communications import install_0914 as _rb_install_0914
voice_intake_status = _rb_install_0914(_rb_bind(_rb_sys.modules[__name__], ('abort', 'app', 'db', 'flash', 'login_required', 'redirect', 'request', 'url_for')))


# ============================================================
# End RadBridge v5.40 - Reception Voice Intake
# ============================================================


# ============================================================
# RadBridge v5.40 - Voice Caller Controls
# ============================================================

# RB_BIND:0915
from web.routes.communications import install_0915 as _rb_install_0915
voice_intake_caller_control = _rb_install_0915(_rb_bind(_rb_sys.modules[__name__], ('abort', 'app', 'audit', 'current_user', 'db', 'flash', 'login_required', 'redirect', 'request', 'url_for')))


# RB_BIND:0916
from web.routes.communications import install_0916 as _rb_install_0916
voice_intake_caller_history = _rb_install_0916(_rb_bind(_rb_sys.modules[__name__], ('abort', 'app', 'db', 'login_required', 'render_template')))

# ============================================================
# End RadBridge v5.40 - Voice Caller Controls
# ============================================================


# ============================================================
# RadBridge v5.41.1 - Authenticated Staff Telephony
# ============================================================

# RB_BIND:0917
from web.components.communications import install_0917 as _rb_install_0917
telephony_me = _rb_install_0917(_rb_bind(_rb_sys.modules[__name__], ('TelephonyExtension', 'current_user', 'jsonify', 'login_required')))


# RB_BIND:0918
from web.components.communications import install_0918 as _rb_install_0918
telephony_directory = _rb_install_0918(_rb_bind(_rb_sys.modules[__name__], ('TelephonyExtension', 'jsonify', 'login_required')))


# RB_BIND:0919
from web.routes.communications import install_0919 as _rb_install_0919
telephony_console = _rb_install_0919(_rb_bind(_rb_sys.modules[__name__], ('TelephonyExtension', 'app', 'current_user', 'login_required', 'render_template')))


# RADBRIDGE V5.41.2 WEBRTC CREDENTIAL API

# RB_BIND:0920
from application.components.shared import install_0920 as _rb_install_0920
_rb_sip_secret_map = _rb_install_0920(_rb_bind(_rb_sys.modules[__name__], ('os',)))


# RB_BIND:0921
from web.routes.communications import install_0921 as _rb_install_0921
telephony_credentials = _rb_install_0921(_rb_bind(_rb_sys.modules[__name__], ('TelephonyExtension', '_rb_sip_secret_map', 'app', 'current_user', 'jsonify', 'login_required')))


# ============================================================
# RADBRIDGE V5.41.7 PUBLIC CLICK TO CALL
# ============================================================

# RB_BIND:0922
from web.routes.shared import install_0922 as _rb_install_0922
public_click_to_call = _rb_install_0922(_rb_bind(_rb_sys.modules[__name__], ('app', 'render_template')))


# RB_BIND:0923
from web.routes.communications import install_0923 as _rb_install_0923
public_call_config = _rb_install_0923(_rb_bind(_rb_sys.modules[__name__], ('_rb_sip_secret_map', 'app', 'jsonify')))


# ============================================================
# RADBRIDGE V5.42.0 PATIENT COMMUNICATIONS
# ============================================================

import secrets as _rb_secrets
import re as _rb_re


# RB_BIND:0924
from application.components.shared import install_0924 as _rb_install_0924
_rb_public_reference = _rb_install_0924(_rb_bind(_rb_sys.modules[__name__], ('_rb_secrets',)))


# RB_BIND:0925
from application.components.shared import install_0925 as _rb_install_0925
_rb_normalize_public_text = _rb_install_0925(_rb_bind(_rb_sys.modules[__name__], ()))


# RB_BIND:0926
from application.components.shared import install_0926 as _rb_install_0926
_rb_human_requested = _rb_install_0926(_rb_bind(_rb_sys.modules[__name__], ()))


# RB_BIND:0927
from application.components.shared import install_0927 as _rb_install_0927
_rb_public_bot_answer = _rb_install_0927(_rb_bind(_rb_sys.modules[__name__], ()))


# RB_BIND:0928
from application.components.shared import install_0928 as _rb_install_0928
_rb_staff_role_allowed = _rb_install_0928(_rb_bind(_rb_sys.modules[__name__], ('current_user',)))


# RB_BIND:0929
from web.routes.communications import install_0929 as _rb_install_0929
public_contact_request_create = _rb_install_0929(_rb_bind(_rb_sys.modules[__name__], ('_rb_normalize_public_text', '_rb_public_reference', 'app', 'db', 'jsonify', 'request')))


# RB_BIND:0930
from web.routes.communications import install_0930 as _rb_install_0930
public_chat_start = _rb_install_0930(_rb_bind(_rb_sys.modules[__name__], ('_rb_secrets', 'app', 'db', 'jsonify', 'request')))


# RB_BIND:0931
from web.routes.communications import install_0931 as _rb_install_0931
public_chat_messages = _rb_install_0931(_rb_bind(_rb_sys.modules[__name__], ('app', 'db', 'jsonify')))


# RB_BIND:0932
from web.routes.communications import install_0932 as _rb_install_0932
public_chat_message = _rb_install_0932(_rb_bind(_rb_sys.modules[__name__], ('_rb_human_requested', '_rb_normalize_public_text', '_rb_public_bot_answer', 'app', 'db', 'jsonify', 'request')))


# RB_BIND:0933
from web.routes.communications import install_0933 as _rb_install_0933
communications_dashboard = _rb_install_0933(_rb_bind(_rb_sys.modules[__name__], ('_rb_staff_role_allowed', 'abort', 'app', 'login_required', 'render_template')))


# RB_BIND:0934
from web.components.communications import install_0934 as _rb_install_0934
staff_communications_queue = _rb_install_0934(_rb_bind(_rb_sys.modules[__name__], ('_rb_staff_role_allowed', 'abort', 'db', 'jsonify', 'login_required')))


# RB_BIND:0935
from web.routes.communications import install_0935 as _rb_install_0935
staff_communications_claim = _rb_install_0935(_rb_bind(_rb_sys.modules[__name__], ('_rb_staff_role_allowed', 'abort', 'app', 'current_user', 'db', 'jsonify', 'login_required')))


# RB_BIND:0936
from web.routes.communications import install_0936 as _rb_install_0936
staff_communications_confirm = _rb_install_0936(_rb_bind(_rb_sys.modules[__name__], ('_rb_normalize_public_text', '_rb_staff_role_allowed', 'abort', 'app', 'current_user', 'db', 'jsonify', 'login_required', 'request')))


# RB_BIND:0937
from web.routes.communications import install_0937 as _rb_install_0937
staff_communications_reject = _rb_install_0937(_rb_bind(_rb_sys.modules[__name__], ('_rb_staff_role_allowed', 'abort', 'app', 'current_user', 'db', 'jsonify', 'login_required')))


# RB_BIND:0938
from web.components.communications import install_0938 as _rb_install_0938
staff_chat_messages = _rb_install_0938(_rb_bind(_rb_sys.modules[__name__], ('_rb_staff_role_allowed', 'abort', 'db', 'jsonify', 'login_required')))


# RB_BIND:0939
from web.routes.communications import install_0939 as _rb_install_0939
staff_chat_message = _rb_install_0939(_rb_bind(_rb_sys.modules[__name__], ('_rb_normalize_public_text', '_rb_staff_role_allowed', 'abort', 'app', 'current_user', 'db', 'jsonify', 'login_required', 'request')))




# ============================================================
# RADBRIDGE V5.42.1 PATIENT DOCUMENTS
# ============================================================

import os as _rb_os
import uuid as _rb_uuid
from datetime import datetime as _rb_datetime


_RB_PUBLIC_ATTACHMENT_TYPES = {
    "INSURANCE_CARD",
    "MEDICAL_APPROVAL",
    "OLD_REPORT",
    "OLD_IMAGE",
    "PRESCRIPTION",
    "LAB_RESULT",
    "REFERRAL",
    "OTHER",
}


_RB_ALLOWED_UPLOAD_EXTENSIONS = {
    ".pdf",
    ".jpg",
    ".jpeg",
    ".png",
}


_RB_ALLOWED_UPLOAD_MIMES = {
    "application/pdf",
    "image/jpeg",
    "image/png",
}


# RB_BIND:0940
from application.components.shared import install_0940 as _rb_install_0940
_rb_public_attachment_root = _rb_install_0940(_rb_bind(_rb_sys.modules[__name__], ('_rb_os', 'os')))


# RB_BIND:0941
from application.components.shared import install_0941 as _rb_install_0941
_rb_safe_upload_extension = _rb_install_0941(_rb_bind(_rb_sys.modules[__name__], ('_rb_os',)))


# RB_BIND:0942
from web.routes.communications import install_0942 as _rb_install_0942
public_contact_request_details = _rb_install_0942(_rb_bind(_rb_sys.modules[__name__], ('_rb_normalize_public_text', 'app', 'db', 'jsonify', 'request')))


# RB_BIND:0943
from web.routes.communications import install_0943 as _rb_install_0943
public_contact_attachment_upload = _rb_install_0943(_rb_bind(_rb_sys.modules[__name__], ('_RB_ALLOWED_UPLOAD_EXTENSIONS', '_RB_ALLOWED_UPLOAD_MIMES', '_RB_PUBLIC_ATTACHMENT_TYPES', '_rb_normalize_public_text', '_rb_os', '_rb_public_attachment_root', '_rb_safe_upload_extension', '_rb_uuid', 'app', 'db', 'jsonify', 'request')))


# RB_BIND:0944
from web.components.communications import install_0944 as _rb_install_0944
staff_contact_attachments = _rb_install_0944(_rb_bind(_rb_sys.modules[__name__], ('_rb_staff_role_allowed', 'abort', 'db', 'jsonify', 'login_required')))


# RB_BIND:0945
from web.routes.shared import install_0945 as _rb_install_0945
staff_public_attachment_view = _rb_install_0945(_rb_bind(_rb_sys.modules[__name__], ('_rb_os', '_rb_public_attachment_root', '_rb_staff_role_allowed', 'abort', 'app', 'db', 'login_required', 'send_file')))


# RB_BIND:0946
from web.routes.shared import install_0946 as _rb_install_0946
staff_attachment_verify = _rb_install_0946(_rb_bind(_rb_sys.modules[__name__], ('_rb_staff_role_allowed', 'abort', 'app', 'current_user', 'db', 'jsonify', 'login_required')))


# RB_BIND:0947
from web.routes.shared import install_0947 as _rb_install_0947
staff_attachment_reject = _rb_install_0947(_rb_bind(_rb_sys.modules[__name__], ('_rb_staff_role_allowed', 'abort', 'app', 'current_user', 'db', 'jsonify', 'login_required')))


# ============================================================
# RADBRIDGE V5.48.0 IVR CONTROL PLANE
# ============================================================

# RB_BIND:0948
from web.routes.operations import install_0948 as _rb_install_0948
admin_ivr = _rb_install_0948(_rb_bind(_rb_sys.modules[__name__], ('TelephonyExtension', 'User', 'app', 'db', 'login_required', 'redirect', 'render_template', 'session', 'url_for')))


# RB_REPORTING_SUPPORT_BLUEPRINT_V1
# Register after every retained view implementation is defined.  This keeps
# public URLs stable while moving their ownership out of the monolithic app.
import sys as _reporting_support_sys
from reporting_support_routes import register_reporting_support_routes

reporting_support_blueprint = register_reporting_support_routes(
    _reporting_support_sys.modules[__name__]
)


# RB_SCHEDULING_REFERENCE_BLUEPRINT_V1
# Read-only scheduling lookup APIs keep their public URLs and authorization
# wrappers while their HTTP ownership moves out of the monolithic application.
from scheduling_reference_routes import register_scheduling_reference_routes

scheduling_reference_blueprint = register_scheduling_reference_routes(
    _reporting_support_sys.modules[__name__]
)


# RB_TELEPHONY_REFERENCE_BLUEPRINT_V1
# Authenticated, read-only staff identity and directory APIs retain their
# public URLs while their HTTP ownership moves out of the monolithic app.
from telephony_reference_routes import register_telephony_reference_routes

telephony_reference_blueprint = register_telephony_reference_routes(
    _reporting_support_sys.modules[__name__]
)


# RB_OPERATIONAL_HEALTH_BLUEPRINT_V1
# Read-only operational and ONLYOFFICE health endpoints retain their public
# URLs and existing authorization behavior outside the monolithic app router.
from operational_health_routes import register_operational_health_routes

operational_health_blueprint = register_operational_health_routes(
    _reporting_support_sys.modules[__name__]
)


# RB_STAFF_COMMUNICATIONS_REFERENCE_BLUEPRINT_V1
# Authenticated, role-gated staff communication reads retain their public URLs
# while all claim, confirm, reject, message and upload writes remain on app.
from staff_communications_reference_routes import (
    register_staff_communications_reference_routes,
)

staff_communications_reference_blueprint = (
    register_staff_communications_reference_routes(
        _reporting_support_sys.modules[__name__]
    )
)


# RB_VOICE_SCHEDULING_REFERENCE_BLUEPRINT_V1
# Token-protected, read-only branch, procedure and availability references keep
# their public URLs while patient-data reads and all voice writes remain on app.
from voice_scheduling_reference_routes import (
    register_voice_scheduling_reference_routes,
)

voice_scheduling_reference_blueprint = (
    register_voice_scheduling_reference_routes(
        _reporting_support_sys.modules[__name__]
    )
)


# RB_ANALYTICS_REFERENCE_BLUEPRINT_V1
# Permission-protected, read-only analytics views retain their public URLs
# while financial writes, imports and patient-level files remain on app.
from analytics_reference_routes import register_analytics_reference_routes

analytics_reference_blueprint = register_analytics_reference_routes(
    _reporting_support_sys.modules[__name__]
)


from report_template_routes import register as register_report_template_routes
register_report_template_routes(_reporting_support_sys.modules[__name__])

from readiness_operations import register as register_readiness_operations
register_readiness_operations(_reporting_support_sys.modules[__name__])

from web.extensions.branch_scoped_reads import register as register_branch_scoped_reads
register_branch_scoped_reads(_reporting_support_sys.modules[__name__])

from web.compatibility.returned_documents import correction_documents_ready as _rb_returned_documents_ready
_rb_sys.modules[__name__].returned_documents_ready = _rb_returned_documents_ready
from infrastructure.adapters.clinical.technician_queue_queries import SqlAlchemyTechnicianQueueQueries
_rb_sys.modules[__name__].technician_queue_queries = SqlAlchemyTechnicianQueueQueries(
    Order=Order, Patient=Patient, CaseWorkflowEvent=CaseWorkflowEvent)
from infrastructure.adapters.assignment.coordinator_case_search import SqlAlchemyCoordinatorCaseSearch
_rb_sys.modules[__name__].coordinator_case_search = SqlAlchemyCoordinatorCaseSearch(
    Patient=Patient, Order=Order)

from technician_queue import register as register_technician_queue
register_technician_queue(_reporting_support_sys.modules[__name__])

from coordinator_bulk_assignment import register as register_coordinator_bulk_assignment
register_coordinator_bulk_assignment(_reporting_support_sys.modules[__name__])

from report_identity_routes import register as register_report_identity_routes
register_report_identity_routes(_reporting_support_sys.modules[__name__])

from dicom_export_routes import register as register_dicom_export_routes
register_dicom_export_routes(_reporting_support_sys.modules[__name__])

from report_draft_warnings import register as register_report_draft_warnings
register_report_draft_warnings(_reporting_support_sys.modules[__name__])

from report_revisions import register as register_report_revisions
register_report_revisions(_reporting_support_sys.modules[__name__])
from report_revision_editor import register as register_report_revision_editor
register_report_revision_editor(_reporting_support_sys.modules[__name__])

from physician_final_counts import register as register_physician_final_counts
register_physician_final_counts(_reporting_support_sys.modules[__name__])

from replacement_settlements import register as register_replacement_settlements
register_replacement_settlements(_reporting_support_sys.modules[__name__])

from physician_report_workload import register as register_physician_report_workload
register_physician_report_workload(_reporting_support_sys.modules[__name__])

from report_exports import register as register_report_exports
register_report_exports(_reporting_support_sys.modules[__name__])

from viewer_choice import register as register_viewer_choice
register_viewer_choice(_reporting_support_sys.modules[__name__])

from user_csv_transfer import register as register_user_csv_transfer
register_user_csv_transfer(_reporting_support_sys.modules[__name__])

from clinical_scope import register as register_clinical_scope
register_clinical_scope(_reporting_support_sys.modules[__name__])

from inventory_transfer_policy import register as register_inventory_transfer_policy
register_inventory_transfer_policy(_reporting_support_sys.modules[__name__])

from security_dashboard import register as register_security_dashboard
register_security_dashboard(_reporting_support_sys.modules[__name__])

from procedure_staff_workflow import register as register_procedure_staff_workflow
register_procedure_staff_workflow(_reporting_support_sys.modules[__name__])

from radiotherapy import register as register_radiotherapy
register_radiotherapy(_reporting_support_sys.modules[__name__])
