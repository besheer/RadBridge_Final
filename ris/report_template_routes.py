"""Report template administration and case-scoped ONLYOFFICE choices."""
import hashlib
import io
import json
import uuid
import secrets
import hmac
from datetime import datetime
from pathlib import Path
from flask import abort, flash, jsonify, redirect, render_template, request, send_file, url_for, session
from docx import Document
import report_template_library as lib

MIME='application/vnd.openxmlformats-officedocument.wordprocessingml.document'

def register(m):
    from report_template_conversion import convert_doc,register_source
    register_source(m)
    def csrf():
        token=session.get('report_template_csrf')
        if not token:token=secrets.token_urlsafe(32);session['report_template_csrf']=token
        if request.method=='POST' and not hmac.compare_digest(str(request.form.get('csrf_token') or ''),token):abort(400)
        return token

    def user_access():
        user=m.current_user()
        if not lib.clinician(user) or not (lib.manager(user) or m.has_perm('reports.edit')):abort(403)
        return user

    @m.require_perm('reports.edit')
    def library():
        user=user_access();csrf_token=csrf()
        if request.method=='POST':
            path=None
            try:
                f=request.files.get('file')
                if not f or not f.filename.lower().endswith(('.doc','.docx')):raise lib.TemplateError('Select a DOC or DOCX template.')
                blob=f.stream.read(5*1024*1024+1)
                source_sha=hashlib.sha256(blob).hexdigest()
                if f.filename.lower().endswith('.doc'):blob=convert_doc(m,blob)
                sha=lib.validate_docx(blob)
                visibility=str(request.form.get('visibility') or 'PERSONAL').upper()
                if visibility not in ('PERSONAL','DEPARTMENT','SHARED'):abort(400)
                if visibility!='PERSONAL' and not lib.manager(user):abort(403)
                owner_id=request.form.get('owner_user_id',type=int) or user.id
                if owner_id!=user.id and not lib.manager(user):abort(403)
                owner=m.db.session.get(m.User,owner_id)
                if not lib.clinician(owner):raise lib.TemplateError('Choose an active reporting doctor.')
                proc=request.form.get('procedure_id',type=int)
                if proc and not m.db.session.get(m.Procedure,proc):abort(400)
                modality=str(request.form.get('modality') or '').strip().upper()
                if len(modality)>30:abort(400)
                folder=Path(m.REPORT_DIR)/'doctor-templates'/str(owner_id);folder.mkdir(parents=True,exist_ok=True)
                path=folder/(uuid.uuid4().hex+'.docx')
                with path.open('xb') as out:out.write(blob)
                row=m.DoctorWordTemplate(owner_user_id=owner_id,name=(request.form.get('name') or f.filename).strip()[:180],modality=modality,procedure_id=proc,visibility=visibility,file_path=str(path),original_name=Path(f.filename).name[:255])
                m.db.session.add(row);m.db.session.flush()
                m.audit_as_user(user,'REPORT_TEMPLATE_UPLOADED','doctor_word_template',row.id,json.dumps({'sha256':sha,'source_sha256':source_sha,'source_format':Path(f.filename).suffix.lower(),'owner_user_id':owner_id,'visibility':visibility,'procedure_id':proc,'modality':modality}))
                m.db.session.commit();flash('Template uploaded. Select it before starting a new ONLYOFFICE report.','success')
                return redirect(url_for('doctor_templates'))
            except lib.TemplateError as exc:
                m.db.session.rollback()
                if path and path.exists():path.unlink()
                flash(str(exc),'danger')
            except Exception:
                m.db.session.rollback()
                if path and path.exists():path.unlink()
                raise
        doctors=[u for u in m.User.query.order_by(m.User.username).all() if lib.clinician(u)] if lib.manager(user) else [user]
        return render_template('doctor_templates.html',rows=lib.available(m,user),procedures=m.Procedure.query.filter_by(active=True).order_by(m.Procedure.name).all(),doctors=doctors,can_manage=lib.manager(user),placeholders=lib.PLACEHOLDERS,user=user,template_csrf=csrf_token)

    m.app.view_functions['doctor_templates']=library

    @m.app.post('/doctor/templates/<int:template_id>/deactivate')
    @m.require_perm('reports.edit')
    def report_template_deactivate(template_id):
        user=user_access();csrf();row=m.DoctorWordTemplate.query.get_or_404(template_id)
        if not lib.manager(user) and (row.owner_user_id!=user.id or row.visibility!='PERSONAL'):abort(403)
        row.active=False;m.audit_as_user(user,'REPORT_TEMPLATE_DEACTIVATED','doctor_word_template',row.id,'File retained; existing reports unchanged.');m.db.session.commit()
        return redirect(url_for('doctor_templates'))

    @m.app.get('/doctor/templates/sample.docx')
    @m.require_perm('reports.edit')
    def report_template_sample():
        user_access();doc=Document()
        doc.sections[0].header.add_paragraph('RADIOLOGY REPORT')
        doc.sections[0].header.add_paragraph('Patient: {{PATIENT_NAME}} | ID: {{PATIENT_ID}} | Accession: {{ACCESSION}}')
        doc.add_heading('{{EXAM}}',0)
        for title,placeholder in [('FINDINGS','FINDINGS'),('IMPRESSION','IMPRESSION')]:doc.add_heading(title,1);doc.add_paragraph('{{'+placeholder+'}}')
        doc.add_paragraph('Radiologist: {{RADIOLOGIST}}')
        out=io.BytesIO();doc.save(out);out.seek(0)
        return send_file(out,download_name='RadBridge-report-template.docx',as_attachment=True,mimetype=MIME)

    @m.app.route('/admin/report-header',methods=['GET','POST'])
    @m.require_perm('users.manage')
    def report_header_settings():
        user=m.current_user();csrf_token=csrf()
        if not lib.manager(user):abort(403)
        branch_id=request.values.get('branch_id',type=int)
        if branch_id and not m.db.session.get(m.Branch,branch_id):abort(400)
        layout,profile=lib.header_profile(m,branch_id=branch_id)
        logo_path=None
        if request.method=='POST':
            try:
                layout=dict(layout)
                for key in ('title','subtitle','address','phone','footer'):layout[key]=str(request.form.get(key) or '').strip()[:500]
                for key in ('preserve_template_header','show_dob','show_sex','show_study_date','show_referrer'):layout[key]=request.form.get(key)=='1'
                if request.form.get('remove_logo')=='1':layout['logo_path']=''
                logo=request.files.get('logo')
                if logo and logo.filename:
                    from PIL import Image
                    blob=logo.stream.read(2*1024*1024+1)
                    if len(blob)>2*1024*1024:raise lib.TemplateError('Logo must be at most 2 MB.')
                    try:
                        img=Image.open(io.BytesIO(blob))
                        if img.format not in ('PNG','JPEG') or max(img.size)>4000:raise ValueError()
                        extension='.png' if img.format=='PNG' else '.jpg';img.verify()
                    except Exception:raise lib.TemplateError('Use a valid PNG or JPEG logo up to 4000 pixels.') from None
                    folder=Path(m.REPORT_DIR)/'report-header-assets';folder.mkdir(parents=True,exist_ok=True)
                    logo_path=folder/(uuid.uuid4().hex+extension)
                    with logo_path.open('xb') as out:out.write(blob)
                    layout['logo_path']=str(logo_path)
                prior=m.StudioTemplate.query.filter_by(template_type='REPORT_HEADER',branch_id=branch_id).order_by(m.StudioTemplate.version_no.desc()).first()
                version=(prior.version_no if prior else 0)+1
                row=m.StudioTemplate(template_type='REPORT_HEADER',name='Report header',scope='BRANCH' if branch_id else 'GLOBAL',branch_id=branch_id,version_no=version,status='PUBLISHED',layout_json=json.dumps(layout),created_by=user.id,published_at=datetime.utcnow())
                m.db.session.add(row);m.db.session.flush()
                m.db.session.add(m.StudioTemplateVersion(template_id=row.id,version_no=version,snapshot_json=json.dumps(layout),created_by=user.id))
                m.audit_as_user(user,'REPORT_HEADER_PUBLISHED','studio_template',row.id,json.dumps({'branch_id':branch_id,'version':version,'old_profile_id':profile.id if profile else None,'layout':layout}))
                m.db.session.commit();flash('Header published for new reports. Existing report documents remain unchanged.','success')
                return redirect(url_for('report_header_settings',branch_id=branch_id) if branch_id else url_for('report_header_settings'))
            except lib.TemplateError as exc:
                m.db.session.rollback()
                if logo_path and logo_path.exists():logo_path.unlink()
                flash(str(exc),'danger')
            except Exception:
                m.db.session.rollback()
                if logo_path and logo_path.exists():logo_path.unlink()
                raise
        return render_template('report_header_settings.html',layout=layout,profile=profile,branch_id=branch_id,branches=m.Branch.query.order_by(m.Branch.name).all(),template_csrf=csrf_token)

    @m.app.get('/api/ohif/report/<int:order_id>/onlyoffice/template-options')
    @m.require_ohif_token
    def onlyoffice_template_options(order_id):
        order=m.Order.query.get_or_404(order_id)
        block=m.require_valid_ohif_case_context(order)
        if block:return block
        user,block=m.require_ohif_doctor_user(order.id)
        if block:return block
        report=m.get_report(order)
        layout,profile=lib.header_profile(m,order)
        return jsonify(ok=True,can_initialize=lib.can_initialize(m,order,report),templates=[{'id':t.id,'name':t.name,'modality':t.modality,'visibility':t.visibility} for t in lib.available(m,user,order)],header_version=profile.version_no if profile else 0,report_status=report.status)

    @m.require_perm('reports.view')
    def download(order_id,template_id):
        order=m.Order.query.get_or_404(order_id);user=user_access()
        if not m.doctor_case_access_allowed(user,order):abort(403)
        tpl=m.DoctorWordTemplate.query.get_or_404(template_id)
        if not lib.visible(m,tpl,user,order):abort(403)
        report=m.get_report(order);artifact=m._rb_authoritative_docx_download(order,report)
        if artifact:return send_file(artifact['path'],as_attachment=True,download_name=m.build_report_filename(order,artifact['version_no'],artifact['status'],'docx'),mimetype=MIME)
        try:blob,meta=lib.build(m,order,report,user,template_id)
        except lib.TemplateError as exc:return jsonify(ok=False,error=str(exc)),400
        m.audit_as_user(user,'REPORT_TEMPLATE_DOWNLOADED','order',order.id,json.dumps(meta));m.db.session.commit()
        return send_file(io.BytesIO(blob),as_attachment=True,download_name=m.build_report_filename(order,1,'DRAFT','docx'),mimetype=MIME)
    m.app.view_functions['report_word_from_doctor_template']=download
