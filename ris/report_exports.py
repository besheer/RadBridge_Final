"""Downloads of authorized report views: no clinical writes or separate queries.

Tables and metrics come from the same server-rendered view, never posted HTML.
Legacy downloads delegate to the corresponding protected report view.
"""
import csv
from datetime import datetime, timezone
from functools import lru_cache
from html import escape
from html.parser import HTMLParser
from io import BytesIO, StringIO
from pathlib import Path
import re
from urllib.parse import urlencode

from flask import g, request, send_file, template_rendered, url_for, make_response

REPORTS = frozenset({
    'reports', 'reports_portal', 'analytics_reference.analytics',
    'financial_center', 'receivables', 'accounting_reports',
    'assignment_reports', 'tat_dashboard', 'management_dashboard',
    'operations_dashboard', 'referral_analytics', 'doctor_earnings_report',
    'executive_analytics', 'inventory_ledger', 'inventory_consumption_report',
    'inventory_stock', 'financial_entries', 'daily_closing',
    'physician_final_counts', 'physician_report_workload',
    'replacement_settlements', 'admin_audit', 'pacs_reconciliation_audit',
})
LEGACY = {
    'analytics_reference.analytics_export_csv': ('analytics_reference.analytics', 'csv'),
    'analytics_export_csv': ('analytics_reference.analytics', 'csv'),
    'accounting_reports_csv': ('accounting_reports', 'csv'),
    'receivables_csv': ('receivables', 'csv'),
    'assignment_report_csv': ('assignment_reports', 'csv'),
    'reports_portal_csv': ('reports_portal', 'csv'),
    'financial_center_export': ('financial_center', None),
}
EXPORT_PERMISSION = frozenset({'analytics_reference.analytics', 'accounting_reports', 'receivables', 'financial_center'})
FILTERS = frozenset({'start','end','date_from','date_to','branch_id','doctor_id',
    'modality','contract_id','device_id','warehouse_id','item_id','status',
    'report_status','view','q','report','page','action','username','ip','entity','date'})
FORMATS = {'csv':'text/csv; charset=utf-8', 'xlsx':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'pdf':'application/pdf'}
VOID = frozenset({'area','base','br','col','embed','hr','img','input','link','meta','param','source','track','wbr'})
SKIP = frozenset({'form','button','input','select','textarea','script','style','nav','svg'})


class Node:
    def __init__(self, tag='', attrs=(), parent=None):
        self.tag, self.attrs, self.parent, self.children = tag, dict(attrs), parent, []

    @property
    def classes(self):
        return set((self.attrs.get('class') or '').split())

    def hidden(self):
        style=(self.attrs.get('style') or '').replace(' ','').lower()
        return ('hidden' in self.attrs or self.attrs.get('aria-hidden')=='true'
            or 'display:none' in style or 'visibility:hidden' in style)

    def nodes(self):
        for child in self.children:
            if isinstance(child,Node):
                if child.hidden() or child.tag in SKIP:
                    continue
                yield child
                yield from child.nodes()

    def text(self):
        if self.hidden() or self.tag in SKIP:
            return ''
        return re.sub(r'\s+',' ',' '.join(c.text() if isinstance(c,Node) else c for c in self.children)).strip()


class Document(HTMLParser):
    def __init__(self, html):
        super().__init__(convert_charrefs=True)
        self.root=Node();self.current=self.root;self.feed(html)

    def handle_starttag(self,tag,attrs):
        n=Node(tag,attrs,self.current);self.current.children.append(n)
        if tag not in VOID:self.current=n

    def handle_startendtag(self,tag,attrs):
        self.handle_starttag(tag,attrs)
        if tag not in VOID:self.handle_endtag(tag)

    def handle_endtag(self,tag):
        n=self.current
        while n.parent:
            if n.tag==tag:
                self.current=n.parent;return
            n=n.parent

    def handle_data(self,data):
        self.current.children.append(data)


def section_name(n, fallback):
    current=n;parent=n.parent
    while parent:
        # A section heading often precedes a table-wrap rather than containing it.
        siblings=parent.children[:parent.children.index(current)]
        for sibling in reversed(siblings):
            if isinstance(sibling,Node) and sibling.tag in ('h2','h3','h4'):
                return sibling.text()
        if 'page-content' in parent.classes:break
        headings=[x.text() for x in parent.nodes() if x.tag in ('h2','h3','h4')]
        if headings:return headings[0]
        current=parent;parent=parent.parent
    return fallback


def extract(html, context=None, endpoint=None):
    """Export the report body only, excluding navigation, forms and action cells."""
    doc=Document(html);nodes=list(doc.root.nodes())
    title=next((x.text() for x in nodes if x.tag=='h1'),'RadBridge Report')
    body=next((x for x in nodes if 'page-content' in x.classes),None)
    if body is None:raise ValueError('Report content missing')
    tables=[];metrics=[]
    for table in (x for x in body.nodes() if x.tag=='table'):
        headers=[];rows=[];keep=[]
        for tr in (x for x in table.nodes() if x.tag=='tr'):
            cells=[x for x in tr.children if isinstance(x,Node) and x.tag in ('td','th')]
            if not cells:continue
            if any(c.tag=='th' for c in cells):
                headers=[c.text() for c in cells]
                # Empty/action columns contain workflow controls, never report data.
                keep=[i for i,h in enumerate(headers) if h and h.lower() not in ('action','actions','إجراءات','الإجراءات')]
                headers=[headers[i] for i in keep]
            elif headers and not any(int(c.attrs.get('colspan') or 1)>1 for c in cells):
                rows.append([cells[i].text() if i<len(cells) else '' for i in keep])
        if headers:tables.append({'title':section_name(table,f'Table {len(tables)+1}'),'headers':headers,'rows':rows})
    for n in body.nodes():
        is_metric=bool(n.classes & {'kpi-card','metric'}) or (n.parent and n.parent.classes & {'management-kpis','exec-kpis'})
        if is_metric:
            labels=[x.text() for x in n.nodes() if x.tag in ('small','span')]
            values=[x.text() for x in n.nodes() if x.tag in ('b','strong')]
            if labels and values:metrics.append([labels[0],values[0]])
    grouped={}
    for n in body.nodes():
        if 'bi-row' in n.classes:
            grouped.setdefault(section_name(n,'Summary'),[]).append([x.text() for x in n.children if isinstance(x,Node)])
    for name,rows in grouped.items():
        tables.append({'title':name,'headers':['Name','Count'],'rows':[r[:2] for r in rows if len(r)>=2]})
    if endpoint=='executive_analytics':
        data=(context or {}).get('data',{})
        for key in ('daily_exams','daily_revenue','modalities','statuses','branches','doctors','devices'):
            tables.append({'title':key.replace('_',' ').title(),'headers':['Name / Date','Value'],
                'rows':[[r['label'],r['value']] for r in data.get(key,[])]})
    if metrics:tables.insert(0,{'title':'Summary','headers':['Metric','Value'],'rows':metrics})
    return {'title':title,'tables':tables}


def safe_csv(value):
    text=str(value if value is not None else '')
    # Spreadsheet programs execute formulas even when CSV fields are quoted.
    if text.lstrip(' \t\r\n\ufeff').startswith(('=','+','-','@')) or text.startswith(('\t','\r','\n')):
        if not isinstance(value,(int,float)):text="'"+text
    return text


def csv_bytes(data, metadata):
    buf=StringIO(newline='');w=csv.writer(buf)
    w.writerow([safe_csv(data['title'])]);w.writerows([[safe_csv(k),safe_csv(v)] for k,v in metadata])
    for section in data['tables']:
        w.writerow([]);w.writerow([safe_csv(section['title'])]);w.writerow([safe_csv(v) for v in section['headers']])
        w.writerows([[safe_csv(v) for v in row] for row in section['rows']])
    return buf.getvalue().encode('utf-8-sig')


def xlsx_bytes(data,metadata,rtl=False):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment,Font,PatternFill
    from openpyxl.utils import get_column_letter
    wb=Workbook();wb.remove(wb.active)
    sections=[{'title':'Report Information','headers':['Field','Value'],'rows':[['Report',data['title']]]+metadata}]+data['tables']
    used=set()
    for index,section in enumerate(sections):
        base=re.sub(r'[\\/*?:\[\]]',' ',section['title']).strip()[:26] or 'Report'
        name=base
        if name.lower() in used:name=f'{base} {index}'
        used.add(name.lower());ws=wb.create_sheet(name);ws.sheet_view.rightToLeft=rtl
        for r,values in enumerate([section['headers']]+section['rows'],1):
            for c,value in enumerate(values,1):
                text=str(value if value is not None else '')
                if len(text)>32767:raise ValueError('A cell exceeds the Excel text limit; use CSV or PDF')
                cell=ws.cell(r,c,text);cell.data_type='s'
                header=str(section['headers'][c-1]).lower() if c<=len(section['headers']) else ''
                identifier=any(k in header for k in ('accession',' id','patient','case','رقم','كود','الحالة','المريض'))
                if r>1 and index>0 and not identifier and re.fullmatch(r'-?(?:0|[1-9]\d*)(?:\.\d+)?',text):
                    cell.value=float(text) if '.' in text else int(text)
                cell.alignment=Alignment(horizontal='right' if rtl else 'left',vertical='top',wrap_text=True)
                if r==1:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='18486B')
            if r>1:ws.row_dimensions[r].height=min(300,max(32,16*(1+max((len(str(v))//24 for v in values),default=0))))
        for c in range(1,len(section['headers'])+1):ws.column_dimensions[get_column_letter(c)].width=26
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        ws.print_title_rows='1:1';ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    out=BytesIO();wb.save(out);return out.getvalue()


def pdf_bytes(data,metadata,rtl=False):
    import arabic_reshaper
    from bidi.algorithm import get_display
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT,TA_RIGHT
    from reportlab.lib.pagesizes import A4,landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import LongTable,Paragraph,SimpleDocTemplate,Spacer
    font='RadBridgeAmiri'
    if font not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font,str(Path(__file__).parent/'assets/fonts/Amiri-Regular.ttf')))
    reshaper=arabic_reshaper.ArabicReshaper({'delete_harakat':False,'shift_harakat_position':True})
    @lru_cache(maxsize=4096)
    def visual(text):
        if not re.search(r'[\u0600-\u06ff]',text):return text
        return get_display(reshaper.reshape(text))
    def paragraph(value,width,size=11,header=False):
        text=str(value if value is not None else '')
        # Wrap logical text before shaping each line; RTL order stays correct.
        lines=[];line=''
        for word in text.split():
            candidate=(line+' '+word).strip()
            if line and pdfmetrics.stringWidth(visual(candidate),font,size)>width-14:
                lines.append(line);line=word
            else:line=candidate
            # Break long identifiers without dropping any characters.
            if pdfmetrics.stringWidth(visual(line),font,size)>width-14:
                chunk=''
                for char in line:
                    if chunk and pdfmetrics.stringWidth(visual(chunk+char),font,size)>width-14:
                        lines.append(chunk);chunk=''
                    chunk+=char
                line=chunk
        if line:lines.append(line)
        style=ParagraphStyle('cell',fontName=font,fontSize=size,leading=size*1.55,
            keepWithNext=size>=14,spaceAfter=6 if size>=14 else 0,
            alignment=TA_RIGHT if rtl or re.search(r'[\u0600-\u06ff]',text) else TA_LEFT,
            textColor=colors.white if header else colors.HexColor('#18334a'))
        return Paragraph('<br/>'.join(escape(visual(line)) for line in lines),style)
    out=BytesIO();width,height=landscape(A4);available=width-64
    doc=SimpleDocTemplate(out,pagesize=(width,height),leftMargin=32,rightMargin=32,topMargin=32,bottomMargin=36,
        title=data['title'],author='RadBridge')
    story=[paragraph(data['title'],available,18),Spacer(1,12)]
    for k,v in metadata:story.extend([paragraph(f'{k}: {v}',available,10),Spacer(1,2)])
    story.append(Spacer(1,12))
    for section in data['tables']:
        headers=section['headers'];rows=section['rows']
        # Wide reports use panels with the first identifier repeated on each.
        panels=[list(range(len(headers)))] if len(headers)<=6 else [([0]+list(range(i,min(i+5,len(headers))))) for i in range(1,len(headers),5)]
        for number,columns in enumerate(panels,1):
            title=section['title']+(f' ({number}/{len(panels)})' if len(panels)>1 else '')
            story.append(paragraph(title,available,14))
            col_width=available/len(columns)
            content=[[paragraph(headers[c],col_width,11,True) for c in columns]]
            content.extend([[paragraph(row[c] if c<len(row) else '',col_width) for c in columns] for row in rows])
            if not rows:content.append([paragraph('—',col_width) for c in columns])
            table=LongTable(content,colWidths=[col_width]*len(columns),repeatRows=1,splitByRow=1,splitInRow=1,hAlign='RIGHT' if rtl else 'LEFT')
            table.setStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#18486b')),
                ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f0f5f9')]),
                ('GRID',(0,0),(-1,-1),0.4,colors.HexColor('#d5e1eb')),
                ('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),7),
                ('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)])
            story.extend([table,Spacer(1,14)])
    def footer(canvas,document):
        canvas.setFont(font,10);canvas.drawString(32,18,'RadBridge');canvas.drawRightString(width-32,18,str(document.page))
    doc.build(story,onFirstPage=footer,onLaterPages=footer)
    return out.getvalue()


def register(m):
    def selected_endpoint():return getattr(g,'report_export_endpoint',request.endpoint)

    @m.app.before_request
    def prepare_report_export():
        for key in ('report_export_endpoint','report_export_format','report_export_context'):
            g.pop(key,None)
        endpoint=request.endpoint
        if endpoint not in REPORTS and endpoint not in LEGACY:return
        if request.method!='GET':return
        fmt=request.args.get('export')
        if endpoint in LEGACY:
            target,fixed=LEGACY[endpoint];g.report_export_endpoint=target
            fmt=fixed or (request.view_args or {}).get('fmt')
        if not fmt:return
        user=m.current_user()
        if not user:return  # Original login/permission decorators handle the denial.
        if not user.active:m.abort(403)
        if fmt not in FORMATS:m.abort(400)
        if selected_endpoint() in EXPORT_PERMISSION and not m.has_perm('export.view'):m.abort(403)
        g.report_export_format=fmt

    def remember_context(sender,template,context,**extra):
        if getattr(g,'report_export_format',None):g.report_export_context=context
    template_rendered.connect(remember_context,m.app,weak=False)

    @m.app.context_processor
    def export_controls():
        endpoint=selected_endpoint();allowed=endpoint in REPORTS and request.method=='GET'
        allowed=allowed and bool(m.current_user()) and m.current_user().active
        if endpoint in EXPORT_PERMISSION:allowed=allowed and m.has_perm('export.view')
        links={}
        if allowed:
            params=[(k,v) for k,v in request.args.items(multi=True) if k in FILTERS]
            if endpoint=='physician_report_workload':params=[(k,v) for k,v in params if k!='page']
            for fmt in FORMATS:links[fmt]=url_for(endpoint)+'?'+urlencode(params+[('export',fmt)])
        return {'report_export_links':links,'report_export_all_rows':endpoint=='physician_report_workload'}

    @m.app.after_request
    def download_report(response):
        fmt=getattr(g,'report_export_format',None)
        if not fmt or response.status_code!=200 or response.mimetype!='text/html':return response
        data=extract(response.get_data(as_text=True),getattr(g,'report_export_context',{}),selected_endpoint())
        context=getattr(g,'report_export_context',{})
        metadata=[['Generated UTC',datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S')],
            ['Scope','All matching cases' if selected_endpoint()=='physician_report_workload' else 'Displayed report data (view limits apply)']]
        params={k:v for k,v in request.args.items() if k in FILTERS and v}
        for key in ('start','end','branch_id','doctor_id','status'):
            if key in context and context[key] is not None:params[key]=str(context[key])
        metadata.extend(sorted(params.items()))
        rtl=(context.get('ui_lang') or m.session.get('ui_lang'))=='ar'
        try:
            raw=csv_bytes(data,metadata) if fmt=='csv' else xlsx_bytes(data,metadata,rtl) if fmt=='xlsx' else pdf_bytes(data,metadata,rtl)
        except ValueError as error:
            return make_response(str(error),400)
        filename=re.sub(r'[^a-zA-Z0-9_-]','-',selected_endpoint())+'-'+datetime.now(timezone.utc).strftime('%Y%m%d-%H%M%S')+'.'+fmt
        result=send_file(BytesIO(raw),mimetype=FORMATS[fmt],as_attachment=True,download_name=filename,max_age=0)
        result.headers['Cache-Control']='private, no-store';result.headers['X-Content-Type-Options']='nosniff'
        return result

    for alias,(target,fmt) in LEGACY.items():
        if alias in m.app.view_functions and target in m.app.view_functions:
            def delegate(target=target,**kwargs):
                return m.app.view_functions[target]()
            m.app.view_functions[alias]=delegate
