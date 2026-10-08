"""ONLYOFFICE editing for approved corrections, using the shared clinical writer."""
import hashlib
import base64
import hmac
import io
import json
import os
import secrets
import shutil
import tempfile
import time
import uuid
from datetime import datetime,timedelta
from pathlib import Path
from urllib.parse import urlparse,quote,unquote

from docx import Document
from docx.oxml.ns import qn
from flask import request,jsonify,render_template,send_file
from itsdangerous import URLSafeTimedSerializer,BadSignature,SignatureExpired

import report_revisions as service

def root(m,row):
    return Path(m.REPORT_DIR).resolve()/'_report_revisions'/str(row.id)

def token_hash(token):return hashlib.sha256(str(token or '').encode()).hexdigest()

def access(m,rid,write=False,lock=False):
    row=m.ReportRevision.query.get_or_404(rid)
    order=m.Order.query.get_or_404(row.order_id)
    user=m.current_user()
    if not service.in_scope(m,user,order):m.abort(403)
    role=user.role.code if user and user.role else ''
    if not user.active or role not in service.CLINICAL_ROLES|service.COORDINATOR_ROLES:m.abort(403)
    can_coordinate=role in service.COORDINATOR_ROLES and (m.has_perm('reports.coordinate') or m.has_perm('assignment.manage'))
    can_read=(m.has_perm('reports.view') and (user.id in (row.requested_by,row.doctor_id) or m.doctor_case_access_allowed(user,order)))
    if not (can_coordinate or can_read):m.abort(403)
    if write:
        if row.doctor_id!=user.id or not m.has_perm('reports.edit') or not service.eligible(m,order,user):m.abort(403)
        if row.status not in ('DRAFT','VERIFIED'):m.abort(409,description='CORRECTION_NOT_EDITABLE')
    if lock:
        order=m.Order.query.filter_by(id=order.id).populate_existing().with_for_update().one()
        row=m.ReportRevision.query.filter_by(id=rid).populate_existing().with_for_update().one()
        # Permission is re-evaluated after the row lock, closing review races.
        if write and (row.doctor_id!=user.id or row.status not in ('DRAFT','VERIFIED')):m.abort(409)
        if not service.in_scope(m,user,order):m.abort(403)
        if write and (not m.has_perm('reports.edit') or not service.eligible(m,order,user)):m.abort(403)
    return row,order,user

def lease(row,token):
    if not token or not row.lease_hash or not hmac.compare_digest(row.lease_hash,token_hash(token)):
        raise service.RevisionError('REVISION_EDIT_LOCKED')
    if not row.lease_expires_at or row.lease_expires_at<=datetime.utcnow():raise service.RevisionError('REVISION_LEASE_EXPIRED')

def signer(m):return URLSafeTimedSerializer(m.app.config['SECRET_KEY'],salt='report-revision-document-v1')

def link(m,row,purpose,version_id=None):
    return signer(m).dumps({'revision_id':row.id,'purpose':purpose,'key':row.editor_key,
                            'lease_hash':row.lease_hash,'doctor_id':row.doctor_id,'version_id':version_id})

def check_link(m,row,purpose,value):
    try:claims=signer(m).loads(value,max_age=14400)
    except (BadSignature,SignatureExpired):return False
    if claims.get('revision_id')!=row.id or claims.get('purpose')!=purpose:return False
    if purpose=='callback':
        return (row.status=='DRAFT' and claims.get('key')==row.editor_key and claims.get('lease_hash')==row.lease_hash
            and claims.get('doctor_id')==row.doctor_id and row.lease_expires_at and row.lease_expires_at>datetime.utcnow())
    return claims.get('key')==row.editor_key

def body_text(path):
    doc=Document(path)
    return ' '.join(' '.join(node.text or '' for node in doc.element.body.iter(qn('w:t'))).split())

def prepare(m,row,order):
    base=service.check_reference(m,order,row)
    target=root(m,row)/'working.docx'
    target.parent.mkdir(parents=True,exist_ok=True,mode=0o700)
    if not target.exists():
        current=service.latest(m,row)
        if current:
            if not current.docx_path or service.digest(current.docx_path)!=current.sha256:raise service.RevisionError('DOCUMENT')
            shutil.copyfile(current.docx_path,target)
        else:
            doc=Document(base)
            # Addendum is append-only. Its original body is compared on every
            # explicit clinical commit, independently from Word layout changes.
            if row.kind=='ADDENDUM':doc.add_page_break()
            doc.add_paragraph(f'RADBRIDGE {row.kind} #{row.revision_no} / correction request #{row.id}')
            doc.add_paragraph('Reference FINAL SHA256: '+row.base_sha256)
            doc.add_paragraph('Correction reason: '+row.reason)
            if row.kind=='ADDENDUM':doc.add_heading('ADDITIONAL FINDINGS / IMPRESSION',level=1)
            fd,name=tempfile.mkstemp(dir=target.parent,suffix='.docx');os.close(fd)
            try:doc.save(name);os.link(name,target)
            finally:os.unlink(name)
    row.working_path=str(target)
    return target

def validate_content(m,row,order,path):
    from report_version_safety import validate_docx_bytes
    raw=Path(path).read_bytes()
    if len(raw)>32*1024*1024:raise service.RevisionError('DOCUMENT_TOO_LARGE')
    meta=validate_docx_bytes(raw)
    identity=m.validate_report_patient_file(order,str(path),'.docx')
    if identity.get('result')!='MATCH':
        if identity.get('result')=='BLOCKED':code='DOCUMENT_IDENTITY_MISMATCH'
        elif not identity.get('found_patient_id') or not identity.get('found_accession'):code='DOCUMENT_IDENTIFIERS_MISSING'
        else:code='DOCUMENT_IDENTITY_REVIEW_REQUIRED'
        raise service.RevisionError(code)
    base=service.check_reference(m,order,row)
    content=body_text(path)
    if row.kind=='ADDENDUM':
        if not content.startswith(body_text(base)):raise service.RevisionError('ADDENDUM_ORIGINAL_CONTENT_CHANGED')
        marker=f'RADBRIDGE ADDENDUM #{row.revision_no} / correction request #{row.id}'
        if marker not in content:raise service.RevisionError('ADDENDUM_REFERENCE_MISSING')
    return meta

def callback_bytes(m,url):
    # The URL is trusted only when JWT validation also bound the signed payload.
    # Redirects and arbitrary network destinations are never fetched.
    actual=urlparse(url)
    artifact_path=actual.path
    # Browser-facing Document Server links include our fixed reverse-proxy
    # prefix. Remove only this known prefix, then apply the same cache guard.
    if artifact_path.startswith('/onlyoffice/cache/files/'):
        artifact_path=artifact_path[len('/onlyoffice'):]
    decoded=unquote(unquote(artifact_path))
    if not decoded.startswith('/cache/files/') or '..' in decoded.split('/') or '\\' in decoded:
        raise service.RevisionError('CALLBACK_DESTINATION_DENIED')
    # ONLYOFFICE can announce its browser-facing host. Fetch only the signed
    # cache artifact path from our configured internal server, never that host.
    internal=m.ONLYOFFICE_INTERNAL_URL.rstrip('/')+artifact_path+('?' + actual.query if actual.query else '')
    r=m.requests.get(internal,timeout=(5,15),stream=True,allow_redirects=False)
    try:
        if r.status_code!=200:raise service.RevisionError('CALLBACK_DOWNLOAD_FAILED')
        data=bytearray()
        for chunk in r.iter_content(65536):
            data.extend(chunk)
            if len(data)>32*1024*1024:raise service.RevisionError('DOCUMENT_TOO_LARGE')
        return bytes(data)
    finally:r.close()

def callback_failure(m,row,nonce,code):
    """Record only a signed force-save attempt's technical rejection reason."""
    try:
        nonce=str(uuid.UUID(str(nonce)))
    except (ValueError,TypeError,AttributeError):return
    code=str(code or 'CORRECTION_CALLBACK_REJECTED')
    m.audit_as_user(m.db.session.get(m.User,row.doctor_id),
        'REPORT_REVISION_CALLBACK_REJECTED','report_revision',row.id,
        json.dumps({'forcesave_nonce':nonce,'error':code}))
    m.db.session.commit()

def callback_failure_code(m,revision_id,nonce):
    events=(m.AuditLog.query.filter_by(action='REPORT_REVISION_CALLBACK_REJECTED',
        entity='report_revision',entity_id=str(revision_id))
        .order_by(m.AuditLog.id.desc()).limit(20).all())
    for event in events:
        try:data=json.loads(event.details or '{}')
        except (TypeError,ValueError):continue
        if isinstance(data,dict) and data.get('forcesave_nonce')==nonce:
            return data.get('error') or 'CORRECTION_CALLBACK_REJECTED'
    return None

def final_pdf(m,document):
    """Render the exact finalized DOCX. A stale original PDF is never a fallback."""
    row=m.db.session.get(m.ReportRevision,document['revision_id'])
    version=service.latest(m,row)
    if row.status!='FINAL' or not version or version.sha256!=document['sha256']:
        raise service.RevisionError('DOCUMENT')
    target=root(m,row)/('final-'+version.sha256+'.pdf')
    if not target.is_file():
        payload={'async':False,'filetype':'docx','outputtype':'pdf',
            'key':f'rbr-pdf-{row.id}-{version.sha256[:32]}',
            'url':m.ONLYOFFICE_RIS_INTERNAL_URL+f'/report-revisions/{row.id}/document?token='+quote(link(m,row,'document',version.id))}
        signed=m._onlyoffice_jwt_encode(payload)
        response=m.requests.post(m.ONLYOFFICE_INTERNAL_URL.rstrip('/')+'/converter',
            json={**payload,'token':signed},headers={'Authorization':'Bearer '+signed},
            timeout=(5,45),allow_redirects=False)
        try:
            response.raise_for_status();result=response.json()
        finally:response.close()
        if not result.get('endConvert') or result.get('error') or not result.get('fileUrl'):
            raise service.RevisionError('CORRECTION_PDF_NOT_READY')
        data=callback_bytes(m,result['fileUrl'])
        if not data.startswith(b'%PDF-') or b'%%EOF' not in data[-2048:]:raise service.RevisionError('DOCUMENT')
        fd,name=tempfile.mkstemp(dir=target.parent,suffix='.pdf')
        try:
            with os.fdopen(fd,'wb') as stream:stream.write(data)
            try:os.link(name,target)
            except FileExistsError:pass
        finally:os.unlink(name)
    data=target.read_bytes()
    if not data.startswith(b'%PDF-') or b'%%EOF' not in data[-2048:]:raise service.RevisionError('DOCUMENT')
    return target

def commit(m,row,order,user,data,artifacts=None):
    from report_commit_core import validate_base_version
    from report_transaction_writer import commit_report_version,ReportCommitDuplicate
    from onlyoffice_report_bridge import onlyoffice_verify_latest,onlyoffice_finalize_latest
    action=data.get('action');base=data.get('base_version_no');path=prepare(m,row,order)
    lease(row,request.headers.get('X-RadBridge-Revision-Lease'))
    if action not in ('draft','verify','finalize'):raise service.RevisionError('STATE')
    m.assert_role_allowed(user,{'draft':'DRAFT','verify':'VERIFY','finalize':'FINALIZE'}[action])
    meta=validate_content(m,row,order,path)
    current=service.latest(m,row);current_no=current.version_no if current else 0
    if action!='draft':
        base_document=service.check_reference(m,order,row)
        initial=Document(base_document)
        initial.add_paragraph(f'RADBRIDGE {row.kind} #{row.revision_no} / correction request #{row.id}')
        initial.add_paragraph('Reference FINAL SHA256: '+row.base_sha256)
        initial.add_paragraph('Correction reason: '+row.reason)
        if row.kind=='ADDENDUM':initial.add_heading('ADDITIONAL FINDINGS / IMPRESSION',level=1)
        initial_text=' '.join(' '.join(n.text or '' for n in initial.element.body.iter(qn('w:t'))).split())
        if body_text(path)==initial_text:raise service.RevisionError('CORRECTION_CONTENT_REQUIRED')
    adapter=service.order_adapter(order,row)
    if action=='draft':
        if row.status!='DRAFT':raise service.RevisionError('STATE')
        previous=m.ReportRevisionVersion.query.filter_by(report_id=row.id,commit_id=data.get('commit_id')).first() if data.get('commit_id') else None
        if previous and previous.sha256!=meta['sha256']:raise service.RevisionError('IDEMPOTENT_CONTENT_CHANGED')
        if not previous and current and current.sha256==meta['sha256']:
            validate_base_version(current_version_no=current_no,base_version_no=base)
            return {'ok':True,'version_no':current_no,'status':row.status,'no_changes':True}
        result=commit_report_version(db=m.db,Report=m.ReportRevision,ReportVersion=m.ReportRevisionVersion,
            report_root=str(root(m,row)/'versions'),order=adapter,patient=order.patient,user=user,
            editor_type='ONLYOFFICE',requested_status='DRAFT',source_file=str(path),base_version_no=base,
            commit_id=data.get('commit_id'),change_summary=f'{row.kind} #{row.revision_no}; original preserved')
        version=result['version']
        if not result['retry'] and artifacts is not None:artifacts.append(version.source_path)
        if not result['retry']:service.audit(m,user,'REPORT_REVISION_DRAFT_SAVED',row,{'version_no':current_no},{'version_no':version.version_no,'sha256':version.sha256})
        return {'ok':True,'version_no':version.version_no,'status':row.status,'retry':result['retry']}
    validate_base_version(current_version_no=current_no,base_version_no=base)
    if not current or current.sha256!=meta['sha256']:raise service.RevisionError('SAVE_CORRECTION_DRAFT_FIRST')
    # The same PACS patient safety gate applies to revision approval as to the
    # original report. It can never be bypassed by a coordinator's request.
    allowed,reason=m.pacs_demographic_gate(order)
    if not allowed:raise service.RevisionError('PACS_IDENTITY_NOT_VERIFIED')
    before={'status':row.status,'version_no':current_no}
    args=dict(db=m.db,Report=m.ReportRevision,ReportVersion=m.ReportRevisionVersion,order=adapter,user=user)
    if action=='verify':
        version=onlyoffice_verify_latest(**args)
    else:
        version=onlyoffice_finalize_latest(**args)
        row.active_order_id=None
        service.notify(m,order,'REPORT_REVISION_FINAL','تصحيح تقرير معتمد / FINAL report correction',row.reason,'CASE_COORDINATOR')
        service.notify(m,order,'REPORT_REVISION_FINAL','تصحيح تقرير معتمد / FINAL report correction',row.reason,'RADIOLOGIST',row.requested_by)
    row.lease_hash='';row.lease_expires_at=None;row.editor_key=''
    service.audit(m,user,'REPORT_REVISION_'+action.upper(),row,before,{'status':row.status,'version_no':version.version_no,'sha256':version.sha256})
    return {'ok':True,'version_no':version.version_no,'status':row.status}

def register(m):
    def corrected_pdf_download(order,document):
        try:path=final_pdf(m,document)
        except Exception:
            m.abort(503,description='CORRECTION_PDF_NOT_READY; download the corrected DOCX or retry.')
        return send_file(path,as_attachment=True,download_name=f'{order.accession}-CORRECTION-{document["revision_no"]}-FINAL.pdf')
    m._rb_revision_pdf_download=corrected_pdf_download
    @m.app.get('/report-revisions/<int:revision_id>')
    @m.login_required
    def report_revision_view(revision_id):
        row,order,user=access(m,revision_id)
        return render_template('report_revision_editor.html',revision=row,order=order,
            can_edit=(row.doctor_id==user.id and m.has_perm('reports.edit') and service.eligible(m,order,user)),
            current=service.latest(m,row))

    @m.app.post('/report-revisions/<int:revision_id>/lease')
    @m.login_required
    def report_revision_lease(revision_id):
        service.csrf(m)
        row,order,user=access(m,revision_id,write=True,lock=True)
        data=request.get_json(silent=True) or {};supplied=str(data.get('token') or '')
        if row.lease_hash and row.lease_expires_at and row.lease_expires_at>datetime.utcnow():
            if not supplied or not hmac.compare_digest(row.lease_hash,token_hash(supplied)):
                m.db.session.rollback();return jsonify(ok=False,error='REVISION_EDIT_LOCKED'),409
            token=supplied
        else:
            token=secrets.token_urlsafe(32)
            row.lease_hash=token_hash(token);row.editor_key=f'rbr-{row.id}-{row.lease_hash[:24]}'
        row.lease_expires_at=datetime.utcnow()+timedelta(seconds=120)
        m.db.session.commit()
        return jsonify(ok=True,token=token)

    @m.app.get('/report-revisions/<int:revision_id>/config')
    @m.login_required
    def report_revision_config(revision_id):
        row,order,user=access(m,revision_id,lock=True)
        try:
            locked=row.status in ('VERIFIED','FINAL') or not (
                row.doctor_id==user.id and m.has_perm('reports.edit') and service.eligible(m,order,user))
            if not locked:
                row,order,user=access(m,revision_id,write=True,lock=True)
                lease(row,request.headers.get('X-RadBridge-Revision-Lease'))
                if row.status!='DRAFT':raise service.RevisionError('STATE')
                prepare(m,row,order)
            current=service.latest(m,row)
            if locked and not current:raise service.RevisionError('DOCUMENT')
            base=m.ONLYOFFICE_RIS_INTERNAL_URL+f'/report-revisions/{row.id}'
            key=f'rbr-view-{row.id}-{current.sha256[:24]}' if locked else row.editor_key
            config={'document':{'fileType':'docx','key':key,
                'title':f'{order.accession}-{row.kind}-{row.revision_no}.docx',
                'url':base+'/document?token='+quote(link(m,row,'document',current.id if locked else None)),
                'permissions':{'edit':not locked,'download':True,'print':True}},'documentType':'word',
                'editorConfig':{'mode':'view' if locked else 'edit','lang':m.current_ui_language(),
                    'user':{'id':str(user.id),'name':user.full_name or user.username},
                    'callbackUrl':base+'/callback?token='+quote(link(m,row,'callback')),
                    'customization':{'autosave':not locked,'forcesave':not locked}}}
            if locked:config['editorConfig'].pop('callbackUrl',None)
            config['token']=m._onlyoffice_jwt_encode(config)
            service.audit(m,user,'REPORT_REVISION_EDITOR_OPENED',row,{}, {'status':row.status})
            m.db.session.commit()
            return jsonify(ok=True,config=config,version_no=current.version_no if current else 0)
        except service.RevisionError as exc:
            m.db.session.rollback();return jsonify(ok=False,error=exc.code),409

    @m.app.get('/report-revisions/<int:revision_id>/document')
    def report_revision_document(revision_id):
        row=m.ReportRevision.query.get_or_404(revision_id)
        if not check_link(m,row,'document',str(request.args.get('token') or '')):m.abort(403)
        claims=signer(m).loads(str(request.args.get('token') or ''),max_age=14400)
        path=row.working_path
        if claims.get('version_id') is not None:
            current=m.ReportRevisionVersion.query.filter_by(id=claims['version_id'],report_id=row.id).first()
            path=current.docx_path if current else None
            if not path or not Path(path).is_file() or service.digest(path)!=current.sha256:m.abort(404)
        elif not path or Path(path).resolve()!=root(m,row)/'working.docx' or not Path(path).is_file():m.abort(404)
        return send_file(path,as_attachment=False,mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

    @m.app.post('/report-revisions/<int:revision_id>/callback')
    def report_revision_callback(revision_id):
        row=m.ReportRevision.query.get_or_404(revision_id)
        if not check_link(m,row,'callback',str(request.args.get('token') or '')):return jsonify(error=1),403
        data=request.get_json(silent=True) or {}
        supplied=str(data.get('token') or request.headers.get('Authorization','').removeprefix('Bearer '))
        # Validate signed payload, not just the signature of an unrelated JWT.
        if not m._onlyoffice_jwt_valid(supplied):return jsonify(error=1),403
        try:
            parts=supplied.split('.')
            header=json.loads(base64.urlsafe_b64decode(parts[0]+'='*(-len(parts[0])%4)))
            claims=json.loads(base64.urlsafe_b64decode(parts[1]+'='*(-len(parts[1])%4)))
            if header.get('alg')!='HS256' or not isinstance(claims,dict):raise ValueError()
            if claims.get('exp') and float(claims['exp'])<time.time():raise ValueError()
        except Exception:return jsonify(error=1),403
        claims=claims.get('payload',claims)
        if any(claims.get(k)!=data.get(k) for k in ('key','status','url','userdata')):return jsonify(error=1),403
        if data.get('key')!=row.editor_key:return jsonify(error=1),403
        if data.get('status') not in (2,6):return jsonify(error=0)
        try:
            raw=callback_bytes(m,str(data.get('url') or ''))
            from report_version_safety import validate_docx_bytes
            validate_docx_bytes(raw)
            order=m.Order.query.filter_by(id=row.order_id).with_for_update().one()
            row=m.ReportRevision.query.filter_by(id=row.id).populate_existing().with_for_update().one()
            if not check_link(m,row,'callback',str(request.args.get('token') or '')):raise service.RevisionError('STATE')
            if not service.eligible(m,order,m.db.session.get(m.User,row.doctor_id)):raise service.RevisionError('DOCTOR')
            path=root(m,row)/'working.docx'
            if not row.working_path or Path(row.working_path).resolve()!=path:raise service.RevisionError('DOCUMENT')
            fd,name=tempfile.mkstemp(dir=path.parent,suffix='.docx');os.close(fd)
            try:
                Path(name).write_bytes(raw)
                validate_content(m,row,order,name)
                os.replace(name,path)
            finally:
                if Path(name).exists():os.unlink(name)
            row.forcesave_nonce=str(data.get('userdata') or '')[:36]
            m.db.session.commit();return jsonify(error=0)
        except Exception as exc:
            m.db.session.rollback()
            # Authentication, payload binding, key and lease checks succeeded
            # before this try block. A rejected artifact cannot become a save.
            callback_failure(m,row,data.get('userdata'),getattr(exc,'code','CORRECTION_CALLBACK_REJECTED'))
            m.app.logger.warning('Revision callback rejected revision=%s type=%s code=%s',revision_id,type(exc).__name__,getattr(exc,'code','UNEXPECTED'))
            return jsonify(error=1),409

    @m.app.post('/report-revisions/<int:revision_id>/forcesave')
    @m.login_required
    def report_revision_forcesave(revision_id):
        service.csrf(m)
        row,order,user=access(m,revision_id,write=True)
        try:
            lease(row,request.headers.get('X-RadBridge-Revision-Lease'))
            if row.status!='DRAFT':raise service.RevisionError('STATE')
            data=request.get_json(silent=True) or {};nonce=str(uuid.uuid4())
            command={'c':'forcesave','key':row.editor_key,'userdata':nonce}
            token=m._onlyoffice_jwt_encode(command)
            r=m.requests.post(m.ONLYOFFICE_INTERNAL_URL+'/coauthoring/CommandService.ashx',json=command,
                headers={'Authorization':'Bearer '+token},timeout=10)
            r.raise_for_status();code=int(r.json().get('error') or 0)
            if code==4 and not data.get('dirty'):return jsonify(ok=True,no_changes=True)
            if code!=0:raise service.RevisionError('ONLYOFFICE_CHANGES_NOT_SYNCHRONIZED')
            deadline=time.monotonic()+12
            while time.monotonic()<deadline:
                m.db.session.expire_all();fresh=m.db.session.get(m.ReportRevision,row.id)
                if fresh.forcesave_nonce==nonce:return jsonify(ok=True)
                failure=callback_failure_code(m,revision_id,nonce)
                if failure:return jsonify(ok=False,error=failure),409
                time.sleep(.25)
            return jsonify(ok=False,error='ONLYOFFICE_CALLBACK_TIMEOUT'),504
        except service.RevisionError as exc:return jsonify(ok=False,error=exc.code),409
        except Exception:return jsonify(ok=False,error='ONLYOFFICE_FORCESAVE_FAILED'),502

    @m.app.post('/report-revisions/<int:revision_id>/commit')
    @m.login_required
    def report_revision_commit(revision_id):
        service.csrf(m)
        row,order,user=access(m,revision_id,write=True,lock=True)
        artifacts=[]
        try:
            result=commit(m,row,order,user,request.get_json(silent=True) or {},artifacts)
            m.db.session.commit();return jsonify(result)
        except Exception as exc:
            m.db.session.rollback()
            for artifact in artifacts:
                path=Path(artifact).resolve()
                if path.is_relative_to(root(m,row)/'versions') and path.is_file():path.unlink()
            return jsonify(ok=False,error=getattr(exc,'code','CORRECTION_COMMIT_REJECTED')),409

    @m.app.get('/report-revisions/<int:revision_id>/state')
    @m.login_required
    def report_revision_state(revision_id):
        row,order,user=access(m,revision_id)
        current=service.latest(m,row)
        return jsonify(ok=True,status=row.status,version_no=current.version_no if current else 0)

    @m.app.get('/report-revisions/<int:revision_id>/download')
    @m.login_required
    def report_revision_download(revision_id):
        row,order,user=access(m,revision_id)
        if row.status not in ('VERIFIED','FINAL'):m.abort(409)
        current=service.latest(m,row)
        if not current or not current.docx_path or service.digest(current.docx_path)!=current.sha256:m.abort(404)
        return send_file(current.docx_path,as_attachment=True,
            download_name=f'{order.accession}-{row.kind}-{row.revision_no}-{row.status}.docx')

    @m.app.get('/report-revisions/<int:revision_id>/original')
    @m.login_required
    def report_revision_original(revision_id):
        row,order,user=access(m,revision_id)
        report,v,path,sha=service.original(m,order)
        if report.id!=row.original_report_id or v.id!=row.original_version_id or sha!=row.original_sha256:m.abort(409)
        return send_file(path,as_attachment=True,download_name=f'{order.accession}-ORIGINAL-FINAL.docx')
