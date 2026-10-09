import pytest
import app as m
import delivery_policy as d
from test_assignment_distribution import doctor,order

@pytest.fixture(autouse=True)
def isolated(app_ctx,monkeypatch):
    monkeypatch.setattr(m.db.session,'commit',m.db.session.flush)
    yield
    m.db.session.rollback();m.db.session.remove()

def client(monkeypatch,role='SUPER_ADMIN'):
    user=doctor('CODEX-UAT-delivery-'+role,role)
    monkeypatch.setattr(m,'current_user',lambda:user)
    c=m.app.test_client()
    with c.session_transaction() as session:session['uid']=user.id
    return c

def final_report(o,tmp_path,monkeypatch,status='FINAL',valid=True):
    from docx import Document
    monkeypatch.setattr(m,'REPORT_DIR',str(tmp_path))
    file=tmp_path/'synthetic-final.docx'
    doc=Document();doc.add_paragraph('Patient ID: '+(o.patient.patient_id if valid else 'OTHER-PATIENT'))
    doc.add_paragraph('Accession: '+o.accession);doc.save(file)
    report=m.Report(order_id=o.id,status=status,docx_path=str(file))
    m.db.session.add(report);m.db.session.flush()
    return report,file

def linked_images(o,monkeypatch,count=1,offline=False,mismatch=False):
    o.pacs_study_id='CODEX-UAT-STUDY';o.pacs_status='IMAGES_COMPLETE'
    m.db.session.add(m.PacsStudyMatch(order_id=o.id,orthanc_study_id=o.pacs_study_id,is_current=True,match_method='DEMOGRAPHICS_EXACT',confidence='AUTO_SAFE',pacs_patient_id=o.patient.patient_id,pacs_accession=o.accession))
    m.db.session.flush()
    class Response:
        def __init__(self,data):self.data=data
        def raise_for_status(self):pass
        def json(self):return self.data
    def get(url,**kwargs):
        if offline:raise RuntimeError('Synthetic PACS unavailable')
        if url.endswith('/statistics'):return Response(dict(CountInstances=count))
        return Response(dict(ID=o.pacs_study_id,PatientMainDicomTags=dict(PatientID='OTHER' if mismatch else o.patient.patient_id),MainDicomTags=dict(AccessionNumber=o.accession)))
    monkeypatch.setattr(m.requests,'get',get)

@pytest.mark.parametrize('status',['UNREPORTED','DRAFT','VERIFIED'])
def test_nonfinal_report_cannot_be_delivered(monkeypatch,tmp_path,status):
    c=client(monkeypatch);o=order('CODEX-UAT-REPORT-'+status)
    final_report(o,tmp_path,monkeypatch,status=status)
    response=c.post(f'/orders/{o.id}/delivery/REPORT',data=dict(delivery_method='IN_PERSON',recipient_name='CODEX UAT Recipient'))
    assert response.status_code==303 and not o.report_delivered
    assert m.DeliveryEvent.query.filter_by(order_id=o.id).count()==0

@pytest.mark.parametrize('missing,valid',[(True,True),(False,False)])
def test_missing_or_wrong_patient_final_file_blocks_delivery(monkeypatch,tmp_path,missing,valid):
    c=client(monkeypatch);o=order('CODEX-UAT-REPORT-FILE')
    report,file=final_report(o,tmp_path,monkeypatch,valid=valid)
    if missing:file.unlink()
    c.post(f'/orders/{o.id}/delivery/REPORT',data=dict(delivery_method='IN_PERSON',recipient_name='CODEX UAT Recipient'))
    assert not o.report_delivered and m.DeliveryEvent.query.filter_by(order_id=o.id).count()==0

def test_valid_final_report_delivered_once(monkeypatch,tmp_path):
    c=client(monkeypatch);o=order('CODEX-UAT-FINAL-DELIVERY')
    final_report(o,tmp_path,monkeypatch)
    data=dict(delivery_method='IN_PERSON',recipient_name='CODEX UAT Recipient')
    assert c.post(f'/orders/{o.id}/delivery/REPORT',data=data).status_code==303
    assert o.report_delivered and o.report_delivered_at
    timestamp=o.report_delivered_at
    c.post(f'/orders/{o.id}/delivery/REPORT',data=data)
    assert m.DeliveryEvent.query.filter_by(order_id=o.id,item_type='REPORT').count()==1
    assert o.report_delivered_at==timestamp

@pytest.mark.parametrize('count,offline,mismatch',[(0,False,False),(1,True,False),(1,False,True)])
def test_cached_image_count_does_not_bypass_live_pacs_check(monkeypatch,count,offline,mismatch):
    c=client(monkeypatch);o=order('CODEX-UAT-IMAGE-BLOCK')
    linked_images(o,monkeypatch,count=count,offline=offline,mismatch=mismatch)
    c.post(f'/orders/{o.id}/delivery/IMAGES',data=dict(delivery_method='CD',recipient_name='CODEX UAT Recipient'))
    assert not o.images_delivered and m.DeliveryEvent.query.filter_by(order_id=o.id).count()==0

def test_no_verified_study_link_blocks_images(monkeypatch):
    c=client(monkeypatch);o=order('CODEX-UAT-NO-LINK')
    c.post(f'/orders/{o.id}/delivery/IMAGES',data=dict(delivery_method='CD',recipient_name='CODEX UAT Recipient'))
    assert not o.images_delivered

def test_live_verified_images_can_be_delivered_without_report(monkeypatch):
    c=client(monkeypatch);o=order('CODEX-UAT-IMAGE-DELIVERY')
    linked_images(o,monkeypatch)
    assert c.post(f'/orders/{o.id}/delivery/IMAGES',data=dict(delivery_method='CD',recipient_name='CODEX UAT Recipient')).status_code==303
    assert o.images_delivered and not o.report_delivered
    assert m.DeliveryEvent.query.filter_by(order_id=o.id,item_type='IMAGES').count()==1

@pytest.mark.parametrize('role',['NURSING','TECHNICIAN','RADIOLOGIST'])
def test_delivery_permission_cannot_be_bypassed(monkeypatch,role):
    c=client(monkeypatch,role);o=order('CODEX-UAT-ROLE-'+role)
    assert c.post(f'/orders/{o.id}/delivery/REPORT',data=dict(delivery_method='IN_PERSON',recipient_name='CODEX UAT Recipient')).status_code==403
    assert not o.report_delivered

def test_invalid_method_and_empty_recipient_do_not_record_delivery(monkeypatch):
    c=client(monkeypatch);o=order('CODEX-UAT-INVALID-METHOD')
    assert c.post(f'/orders/{o.id}/delivery/REPORT',data=dict(delivery_method='CD',recipient_name='CODEX UAT Recipient')).status_code==400
    c.post(f'/orders/{o.id}/delivery/IMAGES',data=dict(delivery_method='CD',recipient_name=''))
    assert m.DeliveryEvent.query.filter_by(order_id=o.id).count()==0

def test_cancelled_case_delivery_blocked_even_with_final_report(monkeypatch,tmp_path):
    c=client(monkeypatch);o=order('CODEX-UAT-CANCELLED');o.status='CANCELLED'
    final_report(o,tmp_path,monkeypatch)
    c.post(f'/orders/{o.id}/delivery/REPORT',data=dict(delivery_method='IN_PERSON',recipient_name='CODEX UAT Recipient'))
    assert not o.report_delivered
