"""Read-only export boundaries, filtered downloads and actual file semantics."""
from io import BytesIO
import json
import pytest
import report_exports as e
from test_physician_report_workload import workload
from test_physician_final_counts import final_counts

HTML='''<aside><table><tr><th>Private</th></tr><tr><td>DO NOT EXPORT</td></tr></table></aside>
<h1>تقرير الأطباء</h1><div class="page-content"><form><input value="SECRET"><button>Confirm</button></form>
<div class="kpi-card"><small>Total</small><b>2</b></div><h2>Cases</h2><table><tr><th>Accession</th><th>Patient</th><th>Amount</th><th>Actions</th></tr>
<tr><td>00123</td><td>محمد أحمد</td><td>12.50</td><td><button>DELETE</button></td></tr>
<tr><td>=1+1</td><td><span>Sample</span><input value="TOKEN"></td><td>8</td><td>Approve</td></tr></table>
<div hidden><table><tr><th>Hidden</th></tr><tr><td>HIDDEN SECRET</td></tr></table></div>
<script>console.log('SCRIPT SECRET')</script></div>'''


def test_report_body_excludes_actions_secrets_and_navigation():
    d=e.extract(HTML)
    assert d['title']=='تقرير الأطباء'
    assert d['tables'][0]['rows']==[['Total','2']]
    assert d['tables'][1]['headers']==['Accession','Patient','Amount']
    assert d['tables'][1]['rows'][0]==['00123','محمد أحمد','12.50']
    assert not any(word in json.dumps(d) for word in ('SECRET','Approve','DELETE','DO NOT EXPORT'))


@pytest.mark.parametrize('value',['=SUM(1,2)',' +cmd','-1+1','@cmd','\t=1','\r1','\n1','\ufeff=1'])
def test_csv_blocks_formula_execution(value):assert e.safe_csv(value).startswith("'")


def test_csv_bom_arabic_identifiers_and_sections():
    raw=e.csv_bytes(e.extract(HTML),[['start','2038-02-01']])
    assert raw.startswith(b'\xef\xbb\xbf')
    text=raw.decode('utf-8-sig')
    assert 'محمد أحمد' in text and '00123' in text and "'=1+1" in text and '2038-02-01' in text


def test_excel_is_a_workbook_with_typed_money_and_safe_text():
    from openpyxl import load_workbook
    wb=load_workbook(BytesIO(e.xlsx_bytes(e.extract(HTML),[['start','2038-02-01']],True)))
    ws=wb.worksheets[2]
    assert ws['A2'].value=='00123' and ws['A3'].value=='=1+1' and ws['A3'].data_type=='s'
    assert ws['C2'].value==12.5 and ws['B2'].value=='محمد أحمد'
    assert ws.freeze_panes=='A2' and ws.sheet_view.rightToLeft and ws.auto_filter.ref


def test_pdf_has_embedded_arabic_font_and_multiple_wide_panels():
    from pypdf import PdfReader
    d={'title':'تقرير الحالات', 'tables':[{'title':'بيانات تجريبية','headers':['رقم الحالة']+[f'Column {i}' for i in range(8)],
        'rows':[[f'CASE-{i}','محمد أحمد']+['تفاصيل الفحص طويلة للاختبار']*7 for i in range(100)]}]}
    raw=e.pdf_bytes(d,[['Period','2038-02-01 / 2038-02-28']],True)
    reader=PdfReader(BytesIO(raw))
    assert len(reader.pages)>3 and raw.startswith(b'%PDF-')
    text='\n'.join(p.extract_text() for p in reader.pages)
    assert 'CASE-99' in text and 'Column 7' in text
    assert text.count('CASE-99')==2
    assert b'FontFile2' in raw  # Font embedded for Arabic on machines without fonts.


@pytest.mark.parametrize('fmt',['csv','xlsx','pdf'])
def test_workload_formats_preserve_branch_scope_and_filters(workload,fmt):
    c=workload
    with c.client.session_transaction() as s:s['uid']=c.actor.id
    r=c.client.get('/coordinator/physician-reports?start=2038-02-01&end=2038-02-28&status=FINAL&export='+fmt)
    assert r.status_code==200 and 'attachment;' in r.headers['Content-Disposition']
    assert r.headers['Cache-Control']=='private, no-store'
    if fmt=='csv':
        text=r.data.decode('utf-8-sig');assert c.orders[0].accession in text and c.orders[1].accession not in text
        assert c.verified_order.accession not in text and 'status,FINAL' in text
    denied=c.client.get(f'/coordinator/physician-reports?branch_id={c.branches[1].id}&export={fmt}')
    assert denied.status_code==403


def test_workload_export_includes_beyond_first_page(workload,monkeypatch):
    import physician_report_workload as w
    c=workload;monkeypatch.setattr(w,'PAGE_SIZE',1)
    with c.client.session_transaction() as s:s['uid']=c.actor.id
    r=c.client.get('/coordinator/physician-reports?start=2038-02-01&end=2038-02-28&page=2&export=csv')
    assert r.status_code==200
    text=r.data.decode('utf-8-sig');assert c.verified_order.accession in text and c.orders[0].accession in text
    assert text.count('REPLACEMENT #')==2 and 'All matching cases' in text


@pytest.mark.parametrize('role',['RECEPTION','NURSING','TECHNICIAN','RADIOLOGIST','CASHIER'])
@pytest.mark.parametrize('fmt',['csv','xlsx','pdf'])
def test_workload_denies_unrelated_roles(workload,role,fmt):
    c=workload;r=c.m.Role.query.filter_by(code=role).first()
    if not r:
        r=c.m.Role(code=role,name=role,permissions_json='[]');c.m.db.session.add(r);c.m.db.session.flush()
    c.actor.role=r;c.m.db.session.flush()
    with c.client.session_transaction() as s:s['uid']=c.actor.id
    assert c.client.get('/coordinator/physician-reports?export='+fmt).status_code==403


def test_invalid_format_and_inactive_actor_fail_closed(workload):
    c=workload
    with c.client.session_transaction() as s:s['uid']=c.actor.id
    assert c.client.get('/coordinator/physician-reports?export=exe').status_code==400
    c.actor.active=False
    assert c.client.get('/coordinator/physician-reports?export=csv').status_code==403


def test_existing_csv_delegates_to_report_permission(workload):
    c=workload;c.actor.role.permissions_json=json.dumps(['export.view'])
    with c.client.session_transaction() as s:s['uid']=c.actor.id
    assert c.client.get('/accounting/reports.csv').status_code==403


def test_export_permission_retained_for_existing_finance_downloads(workload):
    c=workload;c.actor.role.permissions_json=json.dumps(['financial.view'])
    with c.client.session_transaction() as s:s['uid']=c.actor.id
    assert c.client.get('/accounting/reports?export=csv').status_code==403


def test_toolbar_keeps_filters_and_excludes_exported_secrets(workload):
    c=workload
    with c.client.session_transaction() as s:s['uid']=c.actor.id
    r=c.client.get('/coordinator/physician-reports?start=2038-02-01&end=2038-02-28&status=FINAL&token=SECRET')
    from html import unescape
    html=unescape(r.data.decode())
    assert 'status=FINAL&export=xlsx' in html and 'token=SECRET&export' not in html


def test_registered_report_views_have_all_three_downloads(app_ctx):
    m=app_ctx
    assert e.REPORTS<=m.app.view_functions.keys()
    assert set(e.FORMATS)=={'csv','xlsx','pdf'}


def test_executive_charts_export_their_actual_underlying_values():
    context={'data':{'daily_exams':[{'label':'2038-01-01','value':12}],'doctors':[{'label':'Synthetic Doctor','value':3}]}}
    d=e.extract(HTML,context,'executive_analytics')
    assert next(t for t in d['tables'] if t['title']=='Daily Exams')['rows']==[['2038-01-01',12]]
    assert next(t for t in d['tables'] if t['title']=='Doctors')['rows']==[['Synthetic Doctor',3]]
