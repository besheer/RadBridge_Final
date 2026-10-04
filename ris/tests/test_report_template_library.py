import io
import zipfile
from types import SimpleNamespace as NS
from pathlib import Path
import pytest
from docx import Document
import report_template_library as lib

def doc_bytes(doc=None):
    doc=doc or Document()
    if not doc.paragraphs:
        doc.add_paragraph('{{PATIENT_NAME}} {{PATIENT_ID}} {{ACCESSION}}')
        doc.add_heading('FINDINGS',1);doc.add_paragraph('{{FINDINGS}}')
        doc.add_heading('IMPRESSION',1);doc.add_paragraph('{{IMPRESSION}}')
    out=io.BytesIO();doc.save(out);return out.getvalue()

def modify_zip(blob,name,value):
    out=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(blob)) as source,zipfile.ZipFile(out,'w') as target:
        for entry in source.infolist():target.writestr(entry,value if entry.filename==name else source.read(entry))
        if name not in source.namelist():target.writestr(name,value)
    return out.getvalue()

def test_identity_is_required_before_upload():
    doc=Document();doc.add_paragraph('No patient fields')
    with pytest.raises(lib.TemplateError,match='placeholders'):lib.validate_docx(doc_bytes(doc))

@pytest.mark.parametrize('name,value',[('word/vbaProject.bin',b'x'),('word/embeddings/object.bin',b'x'),('../escape.xml',b'<x/>'),('word/_rels/custom.rels',b'<Relationships><Relationship TargetMode="External" Target="file:///tmp/a"/></Relationships>')])
def test_unsafe_packages_rejected(name,value):
    with pytest.raises(lib.TemplateError):lib.validate_docx(modify_zip(doc_bytes(),name,value))

def test_missing_report_sections_rejected():
    doc=Document();doc.add_paragraph('{{PATIENT_NAME}} {{PATIENT_ID}} {{ACCESSION}}')
    with pytest.raises(lib.TemplateError,match='headings'):lib.validate_docx(doc_bytes(doc))

def test_plain_normal_template_is_prepared_without_losing_text():
    doc=Document();doc.add_paragraph('Synthetic reusable normal findings.')
    normalized=lib.normalize_upload(doc_bytes(doc))
    assert lib.validate_docx(normalized)
    result=Document(io.BytesIO(normalized))
    assert [p.text for p in result.paragraphs][:2]==['FINDINGS','Synthetic reusable normal findings.']
    assert 'IMPRESSION' in [p.text for p in result.paragraphs]

def test_normalization_rejects_an_existing_patient_report_marker():
    doc=Document();doc.add_paragraph('RIS_ACCESSION: REAL-EXISTING-REPORT')
    with pytest.raises(lib.TemplateError,match='fixed'):lib.normalize_upload(doc_bytes(doc))

def test_split_placeholder_preserves_surrounding_run_formatting():
    doc=Document();p=doc.add_paragraph();p.add_run('Before ').italic=True
    p.add_run('{{PATIENT_').bold=True;p.add_run('NAME}} / {{PATIENT_ID}}');p.add_run(' after').italic=True
    cell=doc.add_table(rows=1,cols=1).cell(0,0)
    cell.add_table(rows=1,cols=1).cell(0,0).text='{{ACCESSION}}'
    doc.sections[0].header.add_paragraph('{{EXAM}}')
    lib.substitute(doc,{'PATIENT_NAME':'Synthetic Person','PATIENT_ID':'SYN-1','ACCESSION':'SYN-ACC','EXAM':'Synthetic Exam'})
    assert p.text=='Before Synthetic Person / SYN-1 after'
    assert p.runs[0].italic and p.runs[1].bold and p.runs[-1].italic
    assert cell.tables[0].cell(0,0).text=='SYN-ACC'
    assert doc.sections[0].header.paragraphs[-1].text=='Synthetic Exam'

def test_unrecognized_field_fails_closed():
    doc=Document();doc.add_paragraph('{{UNKNOWN}}')
    with pytest.raises(lib.TemplateError):lib.substitute(doc,{})

def test_template_cannot_escape_report_storage(tmp_path):
    outsider=tmp_path/'outside.docx';outsider.write_bytes(doc_bytes())
    m=NS(REPORT_DIR=str(tmp_path/'reports'))
    with pytest.raises(lib.TemplateError):lib.template_bytes(m,NS(file_path=str(outsider)))

def test_department_scope_and_exam_compatibility():
    owner=NS(branch_id=1);m=NS(db=NS(session=NS(get=lambda *args:owner)),User=object)
    user=NS(id=2,active=True,branch_id=2,role=NS(code='RADIOLOGIST',permissions=lambda:[]))
    template=NS(active=True,owner_user_id=1,visibility='DEPARTMENT',modality='CT',procedure_id=None)
    assert not lib.visible(m,template,user,NS(branch_id=2,modality='CT',procedure_id=None))
    user.branch_id=1
    assert lib.visible(m,template,user,NS(branch_id=1,modality='CT',procedure_id=None))
    assert not lib.visible(m,template,user,NS(branch_id=1,modality='MR',procedure_id=None))

def test_header_always_contains_case_identity(tmp_path):
    doc=Document();doc.sections[0].header.add_paragraph('Custom facility header')
    layout=dict(lib.DEFAULT_HEADER,show_dob=False,show_sex=False,show_study_date=False)
    values={'PATIENT_NAME':'Synthetic Person','PATIENT_ID':'SYN-1','ACCESSION':'SYN-ACC'}
    lib.apply_header(doc,NS(REPORT_DIR=str(tmp_path)),layout,values,True)
    assert 'Custom facility header' in '\n'.join(p.text for p in doc.sections[0].header.paragraphs)
    assert [row.cells[1].text for row in doc.sections[0].header.tables[0].rows]==list(values.values())

@pytest.mark.parametrize('status,findings,version,existing',[('FINAL','',False,False),('VERIFIED','',False,False),('DRAFT','Started',False,False),('DRAFT','',True,False),('DRAFT','',False,True),('DRAFT','',False,False)])
def test_template_initialization_never_replaces_existing_work(tmp_path,status,findings,version,existing):
    path=tmp_path/'report.docx'
    if existing:path.write_bytes(b'existing document')
    m=NS(ReportVersion=NS(query=NS(filter_by=lambda **kw:NS(first=lambda:object() if version else None))),_onlyoffice_working_path=lambda order:str(path))
    allowed=lib.can_initialize(m,NS(),NS(id=1,status=status,findings=findings,impression=''))
    assert allowed==(status=='DRAFT' and not findings and not version and not existing)
    if existing:assert path.read_bytes()==b'existing document'
