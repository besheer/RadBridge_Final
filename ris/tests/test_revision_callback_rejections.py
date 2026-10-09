import io,json
from uuid import uuid4
from unittest.mock import Mock
from docx import Document
import pytest
from test_report_revisions import revision_case,request,approve,token
import report_revision_editor as e
import report_revisions as s

def prepared(c):
 row=request(c);approve(c,row);lease=token(row)
 path=e.prepare(c.m,row,c.order);c.m.db.session.flush()
 return row,lease,path

def test_signed_rejected_callback_keeps_original_and_working_copy(revision_case,monkeypatch):
 c=revision_case;row,lease,path=prepared(c);before=s.digest(path)
 c.m.db.session().commit()
 doc=Document();doc.add_paragraph('SYNTHETIC findings. Unlinked document.')
 stream=io.BytesIO();doc.save(stream)
 monkeypatch.setattr(e,'callback_bytes',lambda *a:stream.getvalue())
 nonce=str(uuid4());payload={'key':row.editor_key,'status':6,'url':'http://configured/cache/files/synthetic.docx','userdata':nonce}
 payload['token']=c.m._onlyoffice_jwt_encode(payload)
 result=c.m.app.test_client().post(f'/report-revisions/{row.id}/callback?token='+e.link(c.m,row,'callback'),json=payload)
 assert result.status_code==409
 assert e.callback_failure_code(c.m,row.id,nonce)=='DOCUMENT_IDENTIFIERS_MISSING'
 assert e.callback_failure_code(c.m,row.id,str(uuid4())) is None
 assert s.digest(path)==before and s.digest(c.path)==c.sha
 assert c.report.status==c.order.status=='FINAL' and row.status=='DRAFT'
 assert c.m.ReportRevisionVersion.query.filter_by(report_id=row.id).count()==0

def test_forcesave_returns_rejection_instead_of_timeout(revision_case,monkeypatch):
 c=revision_case;row,lease,path=prepared(c)
 def command(*a,**kw):
  e.callback_failure(c.m,row,kw['json']['userdata'],'DOCUMENT_IDENTIFIERS_MISSING')
  response=Mock();response.json.return_value={'error':0};return response
 monkeypatch.setattr(c.m.requests,'post',command)
 client=c.m.app.test_client()
 with client.session_transaction() as session:session['report_revision_csrf']='SYNTHETIC-CSRF'
 result=client.post(f'/report-revisions/{row.id}/forcesave',json={'dirty':True},headers={'X-Report-Revision-CSRF':'SYNTHETIC-CSRF','X-RadBridge-Revision-Lease':lease})
 assert result.status_code==409 and result.json=={'ok':False,'error':'DOCUMENT_IDENTIFIERS_MISSING'}
 assert row.status=='DRAFT' and s.digest(c.path)==c.sha

def test_missing_confirmation_still_times_out_without_saving(revision_case,monkeypatch):
 c=revision_case;row,lease,path=prepared(c)
 response=Mock();response.json.return_value={'error':0}
 monkeypatch.setattr(c.m.requests,'post',lambda *a,**kw:response)
 times=iter([0,13]);monkeypatch.setattr(e.time,'monotonic',lambda:next(times))
 client=c.m.app.test_client()
 with client.session_transaction() as session:session['report_revision_csrf']='SYNTHETIC-CSRF'
 result=client.post(f'/report-revisions/{row.id}/forcesave',json={'dirty':True},headers={'X-Report-Revision-CSRF':'SYNTHETIC-CSRF','X-RadBridge-Revision-Lease':lease})
 assert result.status_code==504 and result.json['error']=='ONLYOFFICE_CALLBACK_TIMEOUT'
 assert c.m.ReportRevisionVersion.query.filter_by(report_id=row.id).count()==0

def test_wrong_identifiers_remain_blocked(revision_case):
 c=revision_case;row,lease,path=prepared(c);d=Document(path)
 d.paragraphs[1].text='Patient ID: SYNTHETIC-WRONG';d.save(path)
 with pytest.raises(s.RevisionError,match='DOCUMENT_IDENTITY_MISMATCH'):e.validate_content(c.m,row,c.order,path)
 assert s.digest(c.path)==c.sha
